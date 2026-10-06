"""Documented adjustments between AER and distributor prices: the explanation classes of scripts/reconcile.py and the
price_adjustment rows of the tariff database (scripts/tariffdb/build.py).

A value difference is explained only when all of these hold; anything else stays `unexplained`:
  * the distributor-year and the tariff are inside the scope below, which is exactly what the aer-rules review
    verified against the source documents (with its exceptions);
  * the amount is read from a source document (the AER file's own metering price, the distributor's Metering block,
    or the LFiT figure Evoenergy prints in the very document being compared);
  * distributor value - AER value reproduces that amount within the published rounding of both values.

metering_adder  Fixed charge only (c/day). Distributor daily charge = AER daily charge + metering. The AER prices
                metering separately: the 'Metering' worksheet of the consolidated reports (2025-26 on) and 'Tariff
                schedule 1' of the 2024-25 per-distributor reports, in $/yr, so the adder is 100 x $/yr / 365 c/day.
                Energex and Ergon print a per-tariff Metering block in $/day next to the DUoS/TUoS/JS blocks
                (NUoS = DUoS + TUoS + JS + Metering), which is the amount used for them.
lfit_adder      Evoenergy energy charges only (c/kWh). Distributor c/kWh = AER c/kWh + the ACT Large-scale Feed-in
                Tariff (LFiT) amount Evoenergy states in its LFiT-inclusive schedule; fixed and demand charges are
                unchanged. The 2023-24 LFiT *rebate* (2.27 c/kWh "on average", not per component) is not an adder
                and stays unexplained.
"""
import os
import re
from functools import lru_cache

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ENDEAVOUR_METERED = {"N50", "N54", "N70", "N71", "N72", "N73", "N90", "N91", "N92", "N93", "N95", "N62", "N63", "N96",
                     "N97"}
ESSENTIAL_NOT_METERED = {"BLNC1AU", "BLNC2AU", "BLNP1AO"}  # controlled load and unmetered LV tariffs
ESSENTIAL_METERING_NOTE = ("Please note this metering cost is included in the relevant LV tariffs daily charge, however "
                           "the published AER standard control services pricing model lists the metering charge "
                           "separately on the Tariff Schedule tab.")
ESSENTIAL_NOTE_DOCS = {"2025-26": "sources/dnsp/essential/Essential_Price_List_and_Explanatory_Notes_2025-26.pdf",
                       "2026-27": "sources/dnsp/essential/Essential_Price_List_and_Explanatory_Notes_2026-27.pdf"}

# distributor -> (years, scope description, scope test on the distributor's tariff code, amount source)
METERING = {
    "Endeavour Energy": (("2024-25", "2025-26", "2026-27"),
                         "small-customer tariffs the AER Metering sheet lists beside the metering price (N50, N54, "
                         "N62, N63, N70-N73, N90-N93, N95-N97); not N19/N20/N29/N39/N89",
                         lambda code: code in ENDEAVOUR_METERED, "aer"),
    "Essential Energy": (("2024-25", "2025-26", "2026-27"),
                         "LV tariffs (BLN*) except controlled load BLNC1AU/BLNC2AU and unmetered BLNP1AO; HV and "
                         "sub-transmission tariffs (BHN*, BSS*) carry no metering",
                         lambda code: code.startswith("BLN") and code not in ESSENTIAL_NOT_METERED, "aer"),
    "Energex": (("2025-26", "2026-27"),
                "SAC tariffs whose row in the distributor's price list carries a Metering block value; CAC and large "
                "tariffs carry none; 2024-25 has no adder",
                lambda code: True, "distributor_block"),
    "Ergon Energy": (("2025-26", "2026-27"),
                     "SAC tariffs whose row in the distributor's price list carries a Metering block value; CAC "
                     "(EC*/WC*) tariffs carry none; 2024-25 has no adder",
                     lambda code: True, "distributor_block"),
}

# Evoenergy LFiT adder: fin_year -> (c/kWh, document that states it, locator, verbatim statement)
LFIT_ADDERS = {
    "2024-25": ("0.258", "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2024-25_LFiT_adjusted.pdf", "pdf:p3",
                "The 2024-25 LFiT amount has been applied as an adjustment to the AER's approved charges for 2024-25, "
                "and is equivalent to an additional 0.258 cents per kilowatt-hour (kWh). This adjustment has been "
                "applied uniformly to the consumption charges (c/kWh) in Evoenergy's tariffs."),
    "2025-26": ("1.593", "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2025-26_incl_LFiT_May2025.xlsx",
                "xlsx:Network tariffs!B6",
                "The 2025-26 LFiT cost has been applied as an adjustment to the AER's approved charges for 2025-26, "
                "equivalent to an additional 1.593 cents per kilowatt-hour (kWh). The LFiT cost has been applied "
                "uniformly to the consumption charges (c/kWh) in Evoenergy's tariffs."),
    "2026-27": ("3.035", "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2026-27_incl_LFiT_June2026.xlsx",
                "xlsx:Network tariffs!B6",
                "The 2026-27 LFiT cost has been applied as an adjustment to the AER's approved charges for 2026-27, "
                "equivalent to an additional 3.035 cents per kilowatt-hour (kWh). The LFiT cost has been applied "
                "uniformly to the consumption charges (c/kWh) in Evoenergy's tariffs."),
}


def code_of(code):
    return re.sub(r"\s+", "", (code or "").upper()).rstrip("*")


def metering_scope(distributor, fin_year, code):
    """(scope description, amount source) when the tariff is inside the verified metering scope, else None."""
    m = METERING.get(distributor)
    if not m or fin_year not in m[0] or not m[2](code_of(code)):
        return None
    return m[1], m[3]


@lru_cache(maxsize=None)
def aer_metering_price(path, distributor, fin_year):
    """The distributor's annual metering price in an AER workbook: (value $/yr, published, locator, basis note), or
    None when the workbook has no single non-zero price for it. Consolidated reports: the 'Metering' worksheet block
    '<distributor> <year> Metering prices'; 2024-25 per-distributor reports: 'Tariff schedule 1'."""
    import openpyxl
    from published import cell_value
    from tariffdb import locators
    if not path.endswith(".xlsx"):
        return None
    wb = openpyxl.load_workbook(os.path.join(ROOT, path), data_only=True)
    found = []
    if "Metering" in wb.sheetnames:
        ws, cur = wb["Metering"], False
        for r in range(1, ws.max_row + 1):
            b = ws.cell(r, 2).value
            if isinstance(b, str) and "Metering prices" in b:
                cur = b.startswith(distributor + " ")
                continue
            if not cur or not isinstance(b, (int, float)) or isinstance(b, bool):
                continue
            label, code, per, val = (ws.cell(r, c) for c in (3, 8, 10, 12))
            if not isinstance(val.value, (int, float)) or not val.value or code.value == "Exit fee" \
                    or per.value == "per meter":
                continue
            found.append((val, "per year" if per.value == "per year" else "period not printed"))
    elif "Tariff schedule" in wb.sheetnames:
        ws = wb["Tariff schedule"]
        start = next((r for r in range(1, 40) if str(ws.cell(r, 2).value or "").startswith("Tariff schedule 1")), None)
        ycol = start and next((c for c in range(3, 20)
                               if str(ws.cell(start, c).value or "").strip() == fin_year.replace("-", "–")), None)
        if start and ycol:
            for r in range(start + 1, start + 40):
                if str(ws.cell(r, 2).value or "").startswith("Tariff schedule 2"):
                    break
                val = ws.cell(r, ycol)
                if isinstance(val.value, (int, float)) and val.value:
                    found.append((val, "period not printed; the AER Metering sheets of later years state 'per year'"))
    values = {c.value for c, _ in found}
    if len(values) != 1:
        return None
    cell, basis = found[0]
    return cell.value, cell_value(cell), locators.xlsx(cell.parent, cell), basis


@lru_cache(maxsize=None)
def _distributor_blocks(path):
    """Energex/Ergon price list (scripts/dnsp/energex_ergon.py): {locator of a NUoS price cell: the Metering block
    cell of the same tariff (code, name, zone) and component}."""
    import openpyxl
    from dnsp import energex_ergon as ee
    wb = openpyxl.load_workbook(os.path.join(ROOT, path), read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        if ws.title.strip() not in ee.IN_SCOPE_SHEETS:
            continue
        sheet = ee.parse_sheet(ws)
        blocks = {}
        for m in sheet.metering_cells:
            blocks.setdefault((ee.tid(m), ee.comp_key(m["comp"])), []).append(m)
        for r in sheet.records:
            cells = blocks.get((ee.tid(r), ee.comp_key(r["comp"])), [])
            if r["basis"] == "NUoS" and len(cells) == 1:
                out[r["locator"]] = cells[0]
    wb.close()
    return out


def distributor_block_price(path, locator):
    """Metering block value the distributor's price list prints for the tariff and component of the NUoS charge at
    `locator`, in $/day: (value, published, locator) or None."""
    if not path.endswith(".xlsx"):
        return None
    m = _distributor_blocks(path).get(locator)
    if not m or not m["value"] or (m["unit"] or "").replace(" ", "").lower() != "$/day":
        return None
    return m["value"], m["published"], m["locator"]


def explain(a, d, tolerance):
    """(explanation class, detail) when a documented adjustment reproduces distributor - AER for this compared pair,
    else None. `a`/`d` are reconcile rows (value_std_f, unit_std, charge_type, tariff_code, source_file, locator);
    `tolerance` is the half-unit published rounding of both values in standard units."""
    dist, fy = d["distributor"], d["fin_year"]
    diff = d["value_std_f"] - a["value_std_f"]
    if a["charge_type"] == d["charge_type"] == "fixed" and a["unit_std"] == d["unit_std"] == "c/day":
        scope = metering_scope(dist, fy, d["tariff_code"])
        if scope:
            text, source = scope
            if source == "aer":
                price = aer_metering_price(a["source_file"], dist, fy)
                if price is None:
                    return None
                value, shown, loc, basis = price
                amount = value * 100 / 365
                how = (f"100 x {value:.15g} $/yr ({os.path.basename(a['source_file'])} {loc}, shown {shown}; "
                       f"{basis}) / 365")
            else:
                price = distributor_block_price(d["source_file"], d.get("locator"))
                if price is None:
                    return None
                value, _, loc = price
                amount = value * 100
                how = f"Metering block {value:.15g} $/day for this tariff in the distributor's price list ({loc})"
            if abs(diff - amount) <= tolerance:
                evidence = (f"; Essential's price list: \"{ESSENTIAL_METERING_NOTE}\"" if dist == "Essential Energy" else "")
                return "metering_adder", (f"distributor daily charge = AER daily charge + metering {amount:.4f} c/day "
                                          f"[{how}]; residual {diff - amount:+.4f} within published rounding "
                                          f"({tolerance:.4f}); scope: {text}{evidence}")
    if dist == "Evoenergy" and fy in LFIT_ADDERS and a["charge_type"] == d["charge_type"] == "energy" \
            and a["unit_std"] == d["unit_std"] == "c/kWh":
        amount, doc, loc, quote = LFIT_ADDERS[fy]
        if d["source_file"] == doc and abs(diff - float(amount)) <= tolerance:
            return "lfit_adder", (f"distributor c/kWh = AER c/kWh + ACT LFiT {amount} c/kWh; residual "
                                  f"{diff - float(amount):+.4f} within published rounding ({tolerance:.4f}); "
                                  f"{os.path.basename(doc)} {loc}: \"{quote}\"")
    return None
