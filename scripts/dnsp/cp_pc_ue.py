"""CitiPower / Powercor / United Energy network tariff parser.

Parses the DNSP "Tariff Summary" workbooks and the "Pricing (Proposal)" PDFs listed in FILES and
writes out/dnsp/cp_pc_ue.csv (schema.COLUMNS).

Run from the repo root:  .venv/bin/python scripts/dnsp/cp_pc_ue.py

Workbooks: every sheet whose title row carries the file's financial year is parsed (NUOS / DUOS /
TUOS / JUOS|JSA / TRIAL sheets).  Header detection is label based: the row holding "Code" is the
group-header row, the next row holding unit strings (c/day, c/kWh, $/kVA/month ...) is the unit
row, rows in between are time-band rows.  Group headers live in merged cells, so they are
forward-filled across the price columns.  Indicative / prior-year sheets (UE_IND_*, *INDICATIVE*)
are skipped.

PDFs: the price tables are Excel pastes whose characters pdfminer flags as non-upright (tiny
negative matrix terms), which makes pdfplumber read them as vertical text.  We force upright=True,
extract words, and rebuild the table geometrically: unit tokens define the price columns, header
words are attached to the nearest column, group headers (Fixed / Demand Charges / Usage ...) are
assigned to contiguous column ranges, and each data token is placed in the nearest column so that
split digits ("2 4.66"), spaced minus signs ("- 1.50") and parenthesised rebates ("(1.50)") are
rebuilt correctly.
"""
from __future__ import annotations

import bisect
import csv
import re
import sys
from pathlib import Path

import openpyxl
import pdfplumber

sys.path.insert(0, "scripts")
from published import cell_value
import schema  # noqa: E402
import units  # noqa: E402

ROOT = Path(".")
OUT = ROOT / "out" / "dnsp" / "cp_pc_ue.csv"
INVENTORY = ROOT / "sources" / "inventory.csv"

# (distributor, fin_year, repo path, version note)
FILES = [
    ("CitiPower", "2023-24", "sources/aer/2023-24_price_lists/CitiPower_2023-24_Tariff_Summary_31Mar2023.xlsx", "Tariff Summary 31 Mar 2023 (annual pricing proposal submission)"),
    ("CitiPower", "2023-24", "sources/dnsp/citipower/CitiPower_Pricing_Proposal_2023-24_wayback.pdf", "2023/24 Pricing Proposal PDF (Wayback copy, file dated May 2023)"),
    ("CitiPower", "2024-25", "sources/aer/dnsp_copies/CitiPower_2024-25_Tariff_Summary_28Mar2024.xlsx", "Tariff Summary 28 Mar 2024 (annual pricing proposal submission)"),
    ("CitiPower", "2024-25", "sources/dnsp/citipower/CitiPower_Pricing_Proposal_2024-25_28Mar2024_wayback.pdf", "2024/25 Pricing PDF 28 Mar 2024 (Wayback copy)"),
    ("CitiPower", "2025-26", "sources/dnsp/citipower/CitiPower_Pricing_Proposal_2025-26_31Mar2025_wayback.pdf", "2025/26 Final Pricing PDF 31 Mar 2025 (Wayback copy)"),
    ("CitiPower", "2026-27", "sources/dnsp/citipower/CitiPower_Tariff_Summary_2026-27_07May2026.xlsx", "FINAL 2026-27 Tariff Summary 7 May 2026"),
    ("CitiPower", "2026-27", "sources/dnsp/citipower/CitiPower_Pricing_Proposal_2026-27_07May2026.pdf", "2026/27 Pricing PDF 7 May 2026"),
    ("Powercor", "2023-24", "sources/aer/2023-24_price_lists/Powercor_2023-24_Tariff_Summary_31Mar2023.xlsx", "Tariff Summary 31 Mar 2023 (annual pricing proposal submission)"),
    ("Powercor", "2024-25", "sources/aer/dnsp_copies/Powercor_2024-25_Tariff_Summary_28Mar2024.xlsx", "Tariff Summary 28 Mar 2024 (annual pricing proposal submission)"),
    ("Powercor", "2024-25", "sources/dnsp/powercor/Powercor_Pricing_Proposal_2024-25_28Mar2024_wayback.pdf", "2024/25 Pricing PDF 28 Mar 2024 (Wayback copy)"),
    ("Powercor", "2025-26", "sources/dnsp/powercor/Powercor_Tariff_Summary_2025-26_31Mar2025_wayback.xlsx", "FINAL 2025-26 Tariff Summary 31 Mar 2025 (Wayback copy)"),
    ("Powercor", "2025-26", "sources/dnsp/powercor/Powercor_Pricing_Proposal_2025-26_31Mar2025_wayback.pdf", "2025/26 Final Pricing PDF 31 Mar 2025 (Wayback copy)"),
    ("Powercor", "2026-27", "sources/dnsp/powercor/Powercor_Tariff_Summary_2026-27_07May2026.xlsx", "FINAL 2026-27 Tariff Summary 7 May 2026"),
    ("Powercor", "2026-27", "sources/dnsp/powercor/Powercor_Pricing_Proposal_2026-27_07May2026.pdf", "2026/27 Pricing PDF 7 May 2026"),
    ("United Energy", "2023-24", "sources/aer/2023-24_price_lists/UnitedEnergy_2023-24_Tariff_Summary_31Mar2023.xlsx", "Tariff Summary 31 Mar 2023 (annual pricing proposal submission)"),
    ("United Energy", "2023-24", "sources/dnsp/unitedenergy/UE_Pricing_Proposal_2023-24_wayback.pdf", "2023/24 Pricing Proposal PDF (Wayback copy, Aug 2023)"),
    ("United Energy", "2024-25", "sources/aer/dnsp_copies/UnitedEnergy_2024-25_Tariff_Summary_28Mar2024.xlsx", "Tariff Summary 28 Mar 2024 (annual pricing proposal submission)"),
    ("United Energy", "2025-26", "sources/aer/dnsp_copies/UnitedEnergy_2025-26_Final_Pricing_31Mar2025.pdf", "2025/26 Final Pricing PDF 31 Mar 2025 (AER-hosted copy)"),
    ("United Energy", "2026-27", "sources/dnsp/unitedenergy/UE_Tariff_Summary_2026-27_07May2026.xlsx", "FINAL 2026-27 Tariff Summary 7 May 2026"),
    ("United Energy", "2026-27", "sources/dnsp/unitedenergy/UE_Pricing_Proposal_2026-27_07May2026.pdf", "2026/27 Pricing PDF 7 May 2026"),
]

# PDF price tables to parse: (title regex, basis, extra note).  Indicative tables are deliberately absent.
PDF_TABLES = [
    (re.compile(r"^Table A\.\s*1\s+Network", re.I), "NUoS", ""),
    (re.compile(r"^Table A\.\s*2\s+Distribution", re.I), "DUoS", ""),
    (re.compile(r"^Table A\.\s*3\s+Transmission", re.I), "TUoS", ""),
    (re.compile(r"^Table A\.\s*4\s+Jurisdictional", re.I), "JSA", ""),
    (re.compile(r"^Table 6\s+Trial tariffs", re.I), "NUoS", "trial tariff"),
    (re.compile(r"^TABLE 1 NETWORK TARIFF", re.I), "NUoS", ""),
    (re.compile(r"^TABLE 2 TRIAL TARIFFS", re.I), "NUoS", "trial tariff"),
]

UNIT_RE = re.compile(r"^\(?(c|\$)\s*/\s*(day|kWh|kVA|kW|MWh)(\s*/\s*(day|month|year|mth))?\)?$", re.I)
UNIT_FRAG_RE = re.compile(r"^\$/k(VA|W)/$", re.I)  # "$/kVA/" with "month" on another line
PERIOD_FRAG_RE = re.compile(r"^(month|day|year)$", re.I)
GROUP_RE = re.compile(r"^(Fixed\*?|Demand( Charges)?|Usage( Charges)?( - (IMPORT|EXPORT))?|Capacity charge)$", re.I)
NUM_RE = re.compile(r"^\(?-?\d+(\.\d+)?\)?$")

# ---------------------------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------------------------

def load_inventory() -> dict[str, str]:
    urls = {}
    with INVENTORY.open(newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row.get("local_path"):
                urls[row["local_path"].strip()] = row["source_url"].strip()
    return urls


def side_for(path: str) -> str:
    return "DNSP" if path.startswith("sources/dnsp/") else "AER_HOSTED"


def norm_unit(u: str) -> str:
    """'$/kW/ month' (line-wrapped cell) -> '$/kW/month'."""
    return re.sub(r"\s*/\s*", "/", norm_label(u))


def norm_label(s: str) -> str:
    s = re.sub(r"\s+", " ", str(s)).strip()
    s = s.replace("Non- Summer", "Non-Summer").replace("Non- summer", "Non-summer")
    return s


def band_overrides(label: str) -> tuple[str, str]:
    """(time_band, season) with distributor-specific vocabulary applied before the generic helpers."""
    l = label.lower()
    tb = schema.time_band_from_label(label)
    season = schema.season_from_label(label)
    if "saver" in l:
        tb = "solar_soak"  # CP/PAL/UE 'Saver' window is 10am-3pm / 11am-4pm daytime
    if "jul-jun" in l or "rolling" in l:
        season = ""  # 12-month / annual measurement
    if "dec-mar" in l:
        season = "summer"  # documents: "summer period covers December to March"
    if "apr-nov" in l:
        season = "non_summer"  # documents: "non-summer is April to November"
    return tb, season


def make_row(*, dist, fin_year, code, name, component, unit, value, gst, basis, source_file, url, note) -> dict:
    tb, season = band_overrides(component)
    vstd, ustd = units.to_std(value, unit, component)
    return {
        "side": side_for(source_file),
        "distributor": dist,
        "fin_year": fin_year,
        "tariff_code": code.strip(),
        "tariff_name": norm_label(name),
        "customer_class": "",
        "component": component,
        "charge_type": schema.charge_type_from_label(component, unit),
        "time_band": tb,
        "season": season,
        "unit": unit,
        "value": str(value),
        "value_std": str(vstd) if vstd is not None else "",
        "unit_std": ustd,
        "gst": gst,
        "basis": basis,
        "source_file": source_file,
        "source_url": url,
        "note": note,
    }


def dedupe_components(cols: list[dict]) -> None:
    """Two columns may share a label (e.g. 'Demand Charges - Dec-Mar' in $/kVA and $/kW); append the unit."""
    counts = {}
    for c in cols:
        counts[c["component"]] = counts.get(c["component"], 0) + 1
    for c in cols:
        if counts[c["component"]] > 1:
            c["component"] = f"{c['component']} ({c['unit']})"


def basis_from_name(name: str) -> str:
    n = name.upper()
    if "NUOS" in n or "NETWORK" in n:
        return "NUoS"
    if "DUOS" in n or "DISTRIBUTION" in n:
        return "DUoS"
    if "TUOS" in n or "TRANSMISSION" in n:
        return "TUoS"
    if "JUOS" in n or "JSA" in n or "JURISDICTIONAL" in n:
        return "JSA"
    return "unknown"


def fin_year_in(text: str) -> str | None:
    m = re.search(r"(20\d\d)\s*[-/_]\s*(\d\d)\b", text or "")
    return f"{m.group(1)}-{m.group(2)}" if m else None


# ---------------------------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------------------------

def parse_summary_xlsx(path: str, dist: str, fin_year: str, url: str, version_note: str, log: list[str]) -> list[dict]:
    rows_out: list[dict] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [[cell_value(c) if c.data_type == "n" else c.value for c in r] for r in ws.iter_rows()]
        title = " ".join(str(c) for c in rows[0] if c) if rows else ""
        sheet_fy = fin_year_in(ws.title) or fin_year_in(title)
        if re.search(r"_IND_|INDICATIVE", ws.title.upper()) or "indicative" in title.lower():
            log.append(f"  skip sheet {ws.title!r} (indicative / prior-year schedule)")
            continue
        if sheet_fy != fin_year:
            log.append(f"  skip sheet {ws.title!r} (year {sheet_fy} != {fin_year})")
            continue
        parsed = parse_summary_sheet(ws.title, rows, dist, fin_year, path, url, version_note, log)
        rows_out.extend(parsed)
    return rows_out


def parse_summary_sheet(sheet, rows, dist, fin_year, path, url, version_note, log) -> list[dict]:
    head_text = " ".join(str(c) for r in rows[:3] for c in r if c)
    gst = "excl" if "EXCLUSIVE OF GST" in head_text.upper() else "unknown"
    gst_note = "" if gst == "excl" else "GST treatment not stated; assumed excl"
    if gst == "unknown":
        gst = "excl"
    basis = basis_from_name(sheet)
    if basis == "unknown":
        basis = basis_from_name(head_text)  # e.g. 'Trial tariffs' sheet titled 'NETWORK TARIFF SCHEDULE - TRIAL TARIFFS'
    # header rows
    h = next(i for i, r in enumerate(rows) if any(str(c).strip() == "Code" for c in r if c))
    code_col = next(j for j, c in enumerate(rows[h]) if c and str(c).strip() == "Code")
    u = next(i for i in range(h + 1, h + 5) if any(c and UNIT_RE.match(str(c).strip()) for c in rows[i]))
    group_row, band_rows, unit_row = rows[h], rows[h + 1:u], rows[u]
    name_col = next(j for j, c in enumerate(group_row) if c)
    is_trial = "trial" in str(group_row[name_col]).lower() or "TRIAL" in sheet.upper()
    sheet_basis_text = str(group_row[name_col]).strip()
    # forward-fill merged group headers
    ff, cur = [], None
    for j, c in enumerate(group_row):
        if c not in (None, ""):
            cur = str(c).strip()
        ff.append(cur if j > code_col else None)
    price_cols = [j for j, c in enumerate(unit_row) if c and UNIT_RE.match(str(c).strip())]
    extra_cols = [j for j in range(code_col + 1, price_cols[0]) if group_row[j] not in (None, "")]
    # footnotes (e.g. "* Includes JSA charges", "• Peak export credit applies in ...")
    notes_text = [norm_label(c) for r in rows[u + 1:] for c in r if isinstance(c, str) and c.strip()]
    cols = []
    for j in price_cols:
        group = norm_label(ff[j] or "")
        band = " ".join(norm_label(r[j]) for r in band_rows if r[j] not in (None, "")).strip()
        comp = f"{group.rstrip('*')} - {band.rstrip('*')}" if band else group.rstrip("*")
        note = ""
        if "*" in group + band:
            lab = (band or group).rstrip("*").lower()
            hit = [n for n in notes_text if lab in n.lower()] or [n for n in notes_text if n.startswith("*")]
            if hit:
                note = re.sub(r"^[•*\s]+", "", hit[0])
        cols.append({"j": j, "component": comp, "unit": norm_unit(unit_row[j]), "note": note})
    dedupe_components(cols)
    hdr_fy = fin_year_in(sheet_basis_text)
    year_note = f"header row labelled {sheet_basis_text!r} (sheet/title say {fin_year})" if hdr_fy and hdr_fy != fin_year else ""
    out = []
    n_codes = 0
    for r in rows[u + 1:]:
        code = r[code_col]
        if code in (None, ""):
            if r[name_col] and str(r[name_col]).strip().lower().startswith("notes"):
                break
            continue
        code = str(code).strip()
        name = r[name_col] or ""
        extras = []
        for j in extra_cols:
            v = r[j]
            if v not in (None, "") and str(v).strip() not in ("-", "na"):
                extras.append(f"{norm_label(group_row[j])}: {norm_label(v)}")
        base_note = [f"sheet '{sheet}'", version_note, year_note]
        if is_trial:
            base_note.append("trial tariff")
        base_note += extras
        if gst_note:
            base_note.append(gst_note)
        emitted = 0
        for c in cols:
            v = r[c["j"]]
            if v in (None, ""):
                continue
            if isinstance(v, str):
                s = v.strip().replace(",", "")
                if not NUM_RE.match(s):
                    continue
                v = ("-" if s.startswith("(") else "") + s.strip("()")
            if float(v) == 0:
                continue  # 0 in these matrices means 'component not applicable'
            note = "; ".join(x for x in base_note + [c["note"]] if x)
            out.append(make_row(dist=dist, fin_year=fin_year, code=code, name=name, component=c["component"],
                                unit=c["unit"], value=v, gst=gst, basis=basis, source_file=path,
                                url=url, note=note))
            emitted += 1
        if emitted:
            n_codes += 1
        else:
            log.append(f"  sheet {sheet!r}: code {code!r} ({norm_label(name)}) has no numeric prices - not emitted")
    log.append(f"  sheet {sheet!r} [{sheet_basis_text}] -> basis {basis}, gst {gst}, {len(out)} rows, {n_codes} codes")
    return out


# ---------------------------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------------------------

def page_vlines(page):
    """Vertical rulings (x, top, bottom) of the page's tables; used to split name/code/extra cells."""
    out = []
    for l in page.lines:
        if abs(l["x0"] - l["x1"]) < 1:
            out.append((l["x0"], l["top"], l["bottom"]))
    for r in page.rects:
        if r["width"] < 2:
            out.append((r["x0"], r["top"], r["bottom"]))
    return out


def page_lines(page, x_tol=1.5, y_tol=2.0):
    for c in page.objects.get("char", []):
        c["upright"] = True  # pdfminer mis-flags these Excel-paste glyphs as rotated
    words = page.extract_words(x_tolerance=x_tol, y_tolerance=y_tol, keep_blank_chars=False)
    # split glued unit tokens such as "$/kVA/month$/kVA/month"
    fixed = []
    for w in words:
        t = w["text"]
        parts = re.split(r"(?<=[A-Za-z])(?=\$/)", t)
        if len(parts) > 1:
            width = (w["x1"] - w["x0"]) / len(t)
            pos = w["x0"]
            for p in parts:
                fixed.append(dict(w, text=p, x0=pos, x1=pos + width * len(p)))
                pos += width * len(p)
        else:
            fixed.append(w)
    for w in fixed:
        w["xc"] = (w["x0"] + w["x1"]) / 2
    fixed.sort(key=lambda w: (w["top"], w["x0"]))
    lines, cur, cur_top = [], [], None
    for w in fixed:
        if cur and abs(w["top"] - cur_top) > 2.5:
            lines.append(sorted(cur, key=lambda z: z["x0"]))
            cur = []
        if not cur:
            cur_top = w["top"]
        cur.append(w)
    if cur:
        lines.append(sorted(cur, key=lambda z: z["x0"]))
    return lines


def line_text(line) -> str:
    return " ".join(w["text"] for w in line)


def phrases(line, gap=6.0):
    """Join horizontally adjacent words into phrases (gap below `gap` points)."""
    out, cur = [], []
    for w in line:
        if cur and w["x0"] - cur[-1]["x1"] > gap:
            out.append(cur)
            cur = []
        cur.append(w)
    if cur:
        out.append(cur)
    return [{"text": " ".join(w["text"] for w in p), "x0": p[0]["x0"], "x1": p[-1]["x1"],
             "xc": (p[0]["x0"] + p[-1]["x1"]) / 2, "top": p[0]["top"], "words": p} for p in out]


def unit_type(u: str) -> str:
    ul = u.lower()
    if "kwh" in ul or "mwh" in ul:
        return "usage"
    if "kva" in ul or "kw" in ul:
        return "demand"
    return "fixed"


def group_type(g: str) -> str:
    gl = g.lower()
    return "fixed" if gl.startswith("fixed") else "usage" if gl.startswith("usage") else "demand"


def assign_groups(cols: list[dict], groups: list[dict]) -> list[int]:
    """Assign each column (sorted by x) to a group index, monotonically (merged header cells cover
    contiguous columns).  Hard constraint: the column unit type must match the group (c/day <->
    Fixed, kWh <-> Usage, kW/kVA <-> Demand/Capacity).  Soft: a merged header cell does not repeat
    a (band, unit) pair, and the mean of its columns' centres should be close to the label centre
    (the latter alone is ambiguous when a source's merged cell is off-centre)."""
    import itertools
    n, m = len(cols), len(groups)
    best, best_cost = None, None
    for cuts in itertools.combinations(range(1, n), m - 1):
        bounds = (0,) + cuts + (n,)
        cost = 0.0
        for g in range(m):
            cs = cols[bounds[g]:bounds[g + 1]]
            if any(unit_type(c["unit"]) != group_type(groups[g]["text"]) for c in cs):
                cost += 1e6
            pairs = [(c["band"], c["unit"]) for c in cs]
            cost += 1e3 * (len(pairs) - len(set(pairs)))
            cost += abs(sum(c["xc"] for c in cs) / len(cs) - groups[g]["xc"])
        if best_cost is None or cost < best_cost:
            best, best_cost = bounds, cost
    if best_cost >= 1e6:
        raise ValueError("no unit-consistent group assignment")
    assign = []
    for g in range(m):
        assign += [g] * (best[g + 1] - best[g])
    return assign


def parse_pdf_table(lines, title_idx: int, vlines=()):
    """Rebuild the price table that starts below lines[title_idx]. Returns (columns, data_rows)."""
    # unit line: first line below the title with >= 3 unit tokens
    def n_units(line):
        return sum(1 for w in line if UNIT_RE.match(w["text"]))
    u = next(i for i in range(title_idx + 1, min(title_idx + 14, len(lines))) if n_units(lines[i]) >= 3)
    header_lines = lines[title_idx + 1:u]
    # continuation lines made only of unit fragments (e.g. "month month")
    cont = []
    k = u + 1
    while k < len(lines) and all(PERIOD_FRAG_RE.match(w["text"]) or UNIT_FRAG_RE.match(w["text"]) for w in lines[k]):
        cont.append(lines[k])
        k += 1
    data_start = k
    # --- unit fragments -> columns
    frags = []
    for ln in header_lines + [lines[u]] + cont:
        ws = list(ln)
        i = 0
        while i < len(ws):
            w = ws[i]
            if UNIT_FRAG_RE.match(w["text"]) and i + 1 < len(ws) and PERIOD_FRAG_RE.match(ws[i + 1]["text"]) and ws[i + 1]["x0"] - w["x1"] < 8:
                frags.append({"text": w["text"] + ws[i + 1]["text"], "xc": (w["x0"] + ws[i + 1]["x1"]) / 2, "top": w["top"], "w": [w, ws[i + 1]]})
                i += 2
                continue
            if UNIT_RE.match(w["text"]) or UNIT_FRAG_RE.match(w["text"]) or (PERIOD_FRAG_RE.match(w["text"]) and ln in cont + header_lines):
                frags.append({"text": w["text"], "xc": w["xc"], "top": w["top"], "w": [w]})
            i += 1
    frags.sort(key=lambda f: f["xc"])
    cols = []
    for f in frags:
        if cols and abs(f["xc"] - cols[-1]["xc"]) < 12:
            cols[-1]["frags"].append(f)
        else:
            cols.append({"xc": f["xc"], "frags": [f]})
    for c in cols:
        c["frags"].sort(key=lambda f: f["top"])
        c["unit"] = "".join(f["text"] for f in c["frags"]).strip("()")
        main = [f for f in c["frags"] if UNIT_RE.match(f["text"])]
        c["xc"] = (main[0]["xc"] if main else sum(f["xc"] for f in c["frags"]) / len(c["frags"]))
        if not UNIT_RE.match(c["unit"]):
            raise ValueError(f"unit column could not be rebuilt: {c['unit']!r}")
    cols.sort(key=lambda c: c["xc"])
    spacing = min(b["xc"] - a["xc"] for a, b in zip(cols, cols[1:]))
    price_left = cols[0]["xc"] - spacing / 2
    hdr_words = {id(w) for c in cols for f in c["frags"] for w in f["w"]}
    # --- group phrases and band words
    groups, band_words = [], []
    for ln in header_lines:
        for p in phrases(ln):
            if p["xc"] < price_left:
                continue
            if all(id(w) in hdr_words for w in p["words"]):
                continue
            if GROUP_RE.match(p["text"]):
                groups.append(p)
            else:
                band_words += [w for w in p["words"] if id(w) not in hdr_words]
    groups.sort(key=lambda g: g["xc"])
    if not groups:
        raise ValueError("no group headers found")
    for c in cols:
        words = sorted((w for w in band_words if abs(w["xc"] - c["xc"]) < spacing / 2), key=lambda w: (w["top"], w["x0"]))
        c["band"] = norm_label(" ".join(w["text"] for w in words))
        c["unit"] = norm_unit(c["unit"])
    gidx = assign_groups(cols, groups)
    for c, g in zip(cols, gidx):
        c["group"] = groups[g]["text"]
        gl, bl = c["group"].rstrip("*"), c["band"].rstrip("*")
        c["component"] = f"{gl} - {bl}" if bl else gl
        c["star"] = "*" in c["group"] + c["band"]
    dedupe_components(cols)
    if len({c["component"] for c in cols}) != len(cols):
        raise ValueError("duplicate component labels after group assignment")
    # --- left header: Code column and extra columns (Status / PFIT / availability)
    left = [w for ln in header_lines + [lines[u]] for w in ln if w["xc"] < price_left]
    code_w = next(w for w in left if w["text"] == "Code")
    same_line = [w for w in left if abs(w["top"] - code_w["top"]) < 2.5]
    name_hdr_x1 = max([w["x1"] for w in same_line if w["x1"] < code_w["x0"]] or [code_w["x0"] - 40])
    name_right = (name_hdr_x1 + code_w["x0"]) / 2
    extra_hdr = [w for w in left if w["x0"] > code_w["x1"]]
    # cluster extra header words into columns by x-centre
    extra_cols = []
    for w in sorted(extra_hdr, key=lambda w: w["xc"]):
        if extra_cols and abs(w["xc"] - extra_cols[-1]["xc"]) < 25:
            extra_cols[-1]["w"].append(w)
            extra_cols[-1]["xc"] = sum(z["xc"] for z in extra_cols[-1]["w"]) / len(extra_cols[-1]["w"])
        else:
            extra_cols.append({"xc": w["xc"], "w": [w]})
    for e in extra_cols:
        e["label"] = norm_label(" ".join(w["text"] for w in sorted(e["w"], key=lambda w: (w["top"], w["x0"]))))
    # --- data rows
    def row_rulings(top):
        return sorted({round(x, 1) for x, t, b in vlines if t <= top + 3 and b >= top - 3})

    data = []
    data_end = data_start
    for ln in lines[data_start:]:
        xs = row_rulings(ln[0]["top"])
        vals = [w for w in ln if w["xc"] >= price_left]
        rest = [w for w in ln if w["xc"] < price_left]
        if len(xs) >= 3 and xs[0] <= code_w["xc"] <= xs[-1]:
            # exact cell split from the table's vertical rulings
            cell = lambda w: bisect.bisect_right(xs, w["xc"])  # noqa: E731
            code_cell = bisect.bisect_right(xs, code_w["xc"])
            name = [w for w in rest if cell(w) < code_cell]
            code = [w for w in rest if cell(w) == code_cell]
            extras = [w for w in rest if cell(w) > code_cell]
            hdr_cells = {}
            for w in extra_hdr:
                hdr_cells.setdefault(bisect.bisect_right(xs, w["xc"]), []).append(w)
            labels = {k: norm_label(" ".join(w["text"] for w in sorted(ws, key=lambda w: (w["top"], w["x0"])))) for k, ws in hdr_cells.items()}
            label_of = lambda w: labels.get(cell(w), "extra")  # noqa: E731
        else:
            # fallback: name left of the Code header, then word-gap phrases (code first, extras after)
            name = [w for w in rest if w["x0"] < name_right]
            mid = phrases([w for w in rest if w["x0"] >= name_right])
            code = mid[0]["words"] if mid else []
            extras = [w for p in mid[1:] for w in p["words"]]
            label_of = lambda w: (min(extra_cols, key=lambda e: abs(e["xc"] - w["xc"]))["label"] if extra_cols else "extra")  # noqa: E731
        if not code or not vals:
            break
        data_end += 1
        cells = {}
        for w in vals:
            c = min(range(len(cols)), key=lambda i: abs(cols[i]["xc"] - w["xc"]))
            if abs(cols[c]["xc"] - w["xc"]) > spacing * 0.75:
                raise ValueError(f"token {w['text']!r} at x={w['xc']:.1f} not near any column")
            cells.setdefault(c, []).append(w)
        values = {}
        for c, ws in cells.items():
            s = "".join(w["text"] for w in sorted(ws, key=lambda w: w["x0"])).replace(",", "")
            if s in ("-", "–", "—", ""):
                continue
            if not NUM_RE.match(s):
                raise ValueError(f"unparseable cell {s!r} for code {line_text(code)!r}")
            v = ("-" if s.startswith("(") else "") + s.strip("()")
            values[c] = v
        ex = []
        for w in extras:
            ex.append((label_of(w), w["text"]))
        data.append({"name": line_text(name), "code": re.sub(r"-$", "", line_text(code)).strip(), "values": values, "extras": ex})
    foot = [line_text(l) for l in lines[data_end:data_end + 8] if line_text(l).startswith("*")]
    for c in cols:
        c["note"] = re.sub(r"^[*\s]+", "", foot[0]) if c.get("star") and foot else ""
    return cols, data


def parse_pricing_pdf(path: str, dist: str, fin_year: str, url: str, version_note: str, log: list[str]) -> list[dict]:
    out: list[dict] = []
    pdf = pdfplumber.open(path)
    gst_pages = []
    found = 0
    for pno, page in enumerate(pdf.pages, 1):
        lines = page_lines(page)
        vlines = page_vlines(page)
        texts = [line_text(l) for l in lines]
        if any("exclusive of GST" in t.lower().replace("gst", "GST") for t in texts):
            gst_pages.append(pno)
        for i, t in enumerate(texts):
            for rx, basis, extra_note in PDF_TABLES:
                if not rx.search(t):
                    continue
                if fin_year_in(t) not in (None, fin_year):
                    log.append(f"  p{pno}: skip {t!r} (not {fin_year})")
                    continue
                found += 1
                cols, data = parse_pdf_table(lines, i, vlines)
                n_rows = 0
                for d in data:
                    extras = "; ".join(f"{k}: {v}" for k, v in d["extras"] if v not in ("-", "na"))
                    note_parts = [f"page {pno} {t}", version_note]
                    if extra_note:
                        note_parts.append(extra_note)
                        jsa = [x for x in texts if "do not attract JSA" in x]
                        if jsa:
                            note_parts.append(f"document: {jsa[0].lstrip('• ').strip()!r}")
                    if extras:
                        note_parts.append(extras)
                    note_parts.append("GST not stated for this table; assumed excl (document marks its other price tables GST exclusive)")
                    for ci, v in d["values"].items():
                        if float(v) == 0:
                            continue
                        c = cols[ci]
                        out.append(make_row(dist=dist, fin_year=fin_year, code=d["code"], name=d["name"], component=c["component"],
                                            unit=c["unit"], value=v, gst="excl", basis=basis, source_file=path, url=url,
                                            note="; ".join(note_parts + ([c["note"]] if c.get("note") else []))))
                        n_rows += 1
                log.append(f"  p{pno}: {t!r} -> basis {basis}, {len(cols)} price columns "
                           f"[{', '.join(c['component'] + ' ' + c['unit'] for c in cols)}], {len(data)} tariff rows, {n_rows} rows")
    if not found:
        log.append("  NO price tables found in this PDF")
    return out


# ---------------------------------------------------------------------------------------------

def main() -> None:
    urls = load_inventory()
    all_rows: list[dict] = []
    log: list[str] = []
    for dist, fy, path, version_note in FILES:
        url = urls.get(path)
        if url is None:
            raise SystemExit(f"{path} not in {INVENTORY}")
        if not Path(path).exists():
            raise SystemExit(f"missing file {path}")
        log.append(f"== {dist} {fy} {path}")
        if path.endswith(".xlsx"):
            rows = parse_summary_xlsx(path, dist, fy, url, version_note, log)
        else:
            rows = parse_pricing_pdf(path, dist, fy, url, version_note, log)
        log.append(f"  -> {len(rows)} rows, {len({r['tariff_code'] for r in rows})} codes")
        all_rows.extend(rows)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=schema.COLUMNS)
        w.writeheader()
        for r in all_rows:
            w.writerow({k: r.get(k, "") for k in schema.COLUMNS})
    print("\n".join(log))
    print(f"wrote {OUT} ({len(all_rows)} rows)")


if __name__ == "__main__":
    main()
