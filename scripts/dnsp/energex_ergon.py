#!/usr/bin/env python
"""Parse Energex and Ergon Energy "Network Price List" workbooks into the normalised long CSV.

Run from the repo root:  .venv/bin/python scripts/dnsp/energex_ergon.py  ->  out/dnsp/energex_ergon.csv

Workbook shape (all years, both distributors): one sheet per tariff class ('SACS Residential',
'SACS Business', 'SAC Large', 'SAC Unmetered', 'CAC', trial sheet). Each in-scope sheet stacks
several price blocks vertically, each introduced by a basis label in the tariff column
("NUOS", "DUOS", "TUOS", "JS" and, from 2025-26, "Metering"):

    <basis>      | Network Access | Network Energy | ... | Network Demand        <- group header
    Tariff | [Pricing Zone/Region] | NTC | Fixed Charge | Volume Charge | ...   <- component header
    <class label> |                 |     | $/day        | $/kWh         | ...   <- unit row
    <tariff rows ...>  (sub-group header rows such as "Controlled load (<100MWh pa)" carry no NTC)
    *Grandfathered                                                               <- footnote

Ergon rows forward-fill the tariff name and pricing zone (East/West/Mt Isa). The "Metering" block
is an alternative-control metering charge: it is not emitted, but its value is used to annotate
NUoS fixed charges that include it (from 2025-26, NUoS fixed = DUoS + TUoS + JSA + Metering).
ACS sheets, the NTC mapping table and the classes/codes sheets are skipped by exact sheet name.

Zero/blank handling: the matrices are wide, so most cells are 0 or blank placeholders for
"component not applicable to this tariff". A (tariff, component) is emitted - for every basis
block that lists it, including zero sub-basis values - only if at least one basis block prices it
non-zero. Components that are zero in every basis block are dropped.
"""
import csv
import os
import re
import sys
import warnings

sys.path.insert(0, "scripts")
import openpyxl  # noqa: E402

from published import cell_value
import schema  # noqa: E402
import units  # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

SLUG = "energex_ergon"
OUT_PATH = os.path.join("out", "dnsp", f"{SLUG}.csv")
INVENTORY = os.path.join("sources", "inventory.csv")

# (distributor, fin_year, repo path)
FILES = [
    ("Energex", "2023-24", "sources/dnsp/energex/Energex_Network_Price_List_2023-24_Sch8_wayback.xlsx"),
    ("Energex", "2023-24", "sources/aer/2023-24_price_lists/Energex_2023-24_Network_Price_List_27Apr2023.xlsx"),
    ("Energex", "2024-25", "sources/dnsp/energex/Energex_Network_Price_List_2024-25_Sch8_wayback.xlsx"),
    ("Energex", "2024-25", "sources/aer/dnsp_copies/Energex_2024-25_Network_Price_List_28Mar2024.xlsx"),
    ("Energex", "2025-26", "sources/dnsp/energex/Energex_Network_Price_List_2025-26_Sch8_wayback.xlsx"),
    ("Energex", "2026-27", "sources/dnsp/energex/Energex_Network_Price_List_2026-27_Sch8_wayback.xlsx"),
    ("Ergon Energy", "2023-24", "sources/dnsp/ergon/Ergon_Network_Price_List_2023-24_Sch8_wayback.xlsx"),
    ("Ergon Energy", "2023-24", "sources/aer/2023-24_price_lists/Ergon_2023-24_Network_Price_List_31Mar2023.xlsx"),
    ("Ergon Energy", "2024-25", "sources/dnsp/ergon/Ergon_Network_Price_List_2024-25_Sch8_wayback.xlsx"),
    ("Ergon Energy", "2024-25", "sources/aer/dnsp_copies/Ergon_2024-25_Network_Price_List_28Mar2024.xlsx"),
    ("Ergon Energy", "2025-26", "sources/dnsp/ergon/Ergon_Network_Price_List_2025-26_Sch8_wayback.xlsx"),
    ("Ergon Energy", "2026-27", "sources/dnsp/ergon/Ergon_Network_Price_List_2026-27_Sch8_wayback.xlsx"),
]

# Standard control service sheets (exact names as they appear across 2023-24 .. 2026-27).
IN_SCOPE_SHEETS = {
    "SACS Residential", "SACS Business", "SAC Large", "SAC Unmetered", "CAC",
    "CAC Trial", "Trial Tariffs", "Tariff Trials",
}
TRIAL_SHEETS = {"CAC Trial", "Trial Tariffs", "Tariff Trials"}
# Known out-of-scope sheets (ACS, mapping tables, code lists). Anything else is reported.
OUT_OF_SCOPE_SHEETS = {
    "Process", "NTC Table", "Network Tariff Classes & Codes", "ACS Fee-Based", "ACS Public Lighting",
    "ACS Metering", "ACS Security Lighting", "ACS Tariff Classes & Codes",
}

BASIS_MAP = {"NUOS": "NUoS", "DUOS": "DUoS", "TUOS": "TUoS", "JS": "JSA", "JSA": "JSA"}
METERING_BLOCK = "METERING"  # ACS metering charge block (2025-26 onwards) - captured for notes only
ZONE_HEADERS = ("pricing zone", "region", "zone")


def side_for(path: str) -> str:
    return "DNSP" if path.startswith("sources/dnsp/") else "AER_HOSTED"


def load_source_urls() -> dict:
    with open(INVENTORY, newline="", encoding="utf-8") as fh:
        return {row["local_path"]: row["source_url"] for row in csv.DictReader(fh)}


def clean(s) -> str:
    """Collapse whitespace (labels like 'Peak Demand Charge \\nkW')."""
    return re.sub(r"\s+", " ", str(s)).strip()


def is_blank(v) -> bool:
    """None, whitespace, or an explicit not-applicable marker ('N/A', '-')."""
    return v is None or (isinstance(v, str) and (not v.strip() or v.strip().upper() in ("N/A", "NA", "-")))


def as_number(v):
    """Return float for numeric or numeric-looking text, else None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(",", "").replace("$", ""))
        except ValueError:
            return None
    return None


def is_unit_text(v) -> bool:
    return isinstance(v, str) and bool(re.match(r"^\s*(\$|c/|cents?)", v, re.I))


def code_str(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return clean(v)


def comp_key(label: str) -> str:
    return re.sub(r"[^a-z0-9]", "", label.lower())


def fmt(v: float) -> str:
    return f"{v:.10g}"


def charge_type(label: str, unit: str) -> str:
    """schema helper, with an override so an *import* reward is not classed as export."""
    ll = label.lower()
    if "import" in ll and re.search(r"reward|credit|rebate", ll):
        return schema.charge_type_from_label(re.sub(r"\b(reward|credit|rebate)\b", "", ll), unit)
    return schema.charge_type_from_label(label, unit)


def time_band(label: str) -> str:
    tb = schema.time_band_from_label(label)
    if tb:
        return tb
    m = re.search(r"block\s*(\d)", label.lower())
    if m:
        return f"block{m.group(1)}"
    if comp_key(label) == "volumecharge":  # flat anytime energy rate
        return "anytime"
    return ""


class Sheet:
    """Raw extraction for one sheet, before zero-filtering."""

    def __init__(self):
        self.records = []      # dicts: basis, code, name, zone, cls, comp, unit, value, row, notes
        self.metering = {}     # (tid, comp_key) -> metering block value
        self.footnotes = []    # e.g. '*Grandfathered'
        self.site_specific = {}  # tid -> [note, ...]
        self.gst = None
        self.title_year = None
        self.subtitle = ""     # e.g. 'Standard Asset Customers (SAC) - Residential (<100MWh pa)'
        self.issues = []


def tid(r) -> tuple:
    return (r["code"], r["name"], r["zone"])


def parse_sheet(ws) -> Sheet:
    out = Sheet()
    cell_rows = list(ws.iter_rows())
    rows = [[c.value for c in r] for r in cell_rows]

    # Title rows: GST statement, year, class subtitle
    titles = [clean(v) for r in rows[:4] for v in r if isinstance(v, str) and v.strip()]
    for t in titles:
        if re.search(r"excluding\s+gst|excl\.?\s*gst|ex\s*gst", t, re.I):
            out.gst = "excl"
        elif re.search(r"including\s+gst|incl\.?\s*gst", t, re.I):
            out.gst = "incl"
        m = re.search(r"(20\d\d)\s*[/\-]\s*(\d\d)\b", t)
        if m and out.title_year is None:
            out.title_year = f"{m.group(1)}-{m.group(2)}"
    for t in titles[1:2]:
        out.subtitle = re.sub(r"\s*-\s*(Energex|Ergon Energy( Network)?)\s*$", "", t)

    state = "seek_basis"
    basis = None
    metering = False
    hdr = {}             # col index -> component label
    tcol = ncol = zcol = None
    unit_of = {}         # col index -> unit text
    cls = ""
    cur_name = cur_zone = ""

    for ri, row in enumerate(rows, start=1):
        if all(is_blank(v) for v in row):
            continue
        first_idx = next(i for i, v in enumerate(row) if not is_blank(v))
        first = row[first_idx]
        first_txt = clean(first).upper() if isinstance(first, str) else ""

        # 1) basis label row (starts a new block)
        if first_txt in BASIS_MAP or first_txt == METERING_BLOCK:
            basis = BASIS_MAP.get(first_txt)
            metering = first_txt == METERING_BLOCK
            state = "seek_header"
            continue

        # 2) component header row: contains an 'NTC' cell
        if state == "seek_header":
            ntc_idx = next((i for i, v in enumerate(row) if isinstance(v, str) and clean(v).upper() == "NTC"), None)
            if ntc_idx is None:
                continue
            ncol = ntc_idx
            tcol = next((i for i, v in enumerate(row) if isinstance(v, str) and clean(v).lower() == "tariff"), first_idx)
            zcol = next((i for i, v in enumerate(row) if isinstance(v, str) and clean(v).lower() in ZONE_HEADERS), None)
            hdr = {i: clean(v) for i, v in enumerate(row) if i > ncol and isinstance(v, str) and not is_blank(v)}
            state = "units"
            continue

        # 3) unit row (immediately follows the header)
        if state == "units":
            unit_of = {i: clean(v) for i, v in enumerate(row) if i in hdr and not is_blank(v)}
            cls = clean(row[tcol]) if not is_blank(row[tcol]) else ""
            cur_name = cur_zone = ""
            state = "data"
            continue
        if state != "data":
            continue

        code_v = row[ncol] if ncol < len(row) else None
        name_v = row[tcol] if tcol < len(row) else None
        zone_v = row[zcol] if zcol is not None and zcol < len(row) else None

        # 4) rows without an NTC: footnote, repeated unit row, or sub-group heading
        if is_blank(code_v):
            if isinstance(name_v, str) and name_v.strip():
                txt = clean(name_v)
                comp_cells = [row[i] for i in hdr if i < len(row) and not is_blank(row[i])]
                if txt.startswith("*"):
                    out.footnotes.append(txt)
                elif comp_cells and all(is_unit_text(v) for v in comp_cells):
                    unit_of = {i: clean(row[i]) for i in hdr if i < len(row) and not is_blank(row[i])}
                    cls = txt
                    cur_name = cur_zone = ""
                elif not any(as_number(v) is not None for v in comp_cells):
                    cls = txt
                    cur_name = cur_zone = ""
                else:
                    out.issues.append(f"row {ri}: priced row without NTC ({txt!r}) skipped")
            continue

        # 5) data row
        if not is_blank(name_v):
            cur_name = clean(name_v)
            cur_zone = ""  # a new tariff resets the forward-filled zone
        if not is_blank(zone_v):
            cur_zone = clean(zone_v)
        code = code_str(code_v)
        rec_id = (code, cur_name, cur_zone)
        for i, v in enumerate(row):
            if i > ncol and i not in hdr and (as_number(v) or 0) != 0:
                out.issues.append(f"row {ri} {code}: non-zero value {v!r} in unlabelled column {i + 1}")
        for i, comp in hdr.items():
            v = row[i] if i < len(row) else None
            if is_blank(v):
                continue
            num = as_number(v)
            if num is None:
                txt = clean(v)
                if "site specific" in txt.lower():
                    note = f"site-specific: {comp} published as '{txt}'"
                    if note not in out.site_specific.setdefault(rec_id, []):
                        out.site_specific[rec_id].append(note)
                else:
                    out.issues.append(f"row {ri} {code} {comp}: non-numeric value {txt!r} skipped")
                continue
            if metering:
                out.metering[(rec_id, comp_key(comp))] = num
            elif basis is not None:
                out.records.append({
                    "basis": basis, "code": code, "name": cur_name, "zone": cur_zone, "cls": cls,
                    "comp": comp, "unit": unit_of.get(i, ""), "value": num, "published": cell_value(cell_rows[ri - 1][i]), "row": ri,
                })
    return out


def rows_for_sheet(sheet: Sheet, sheet_name, distributor, fin_year, side, source_file, source_url):
    gst = sheet.gst or "excl"
    gst_note = "" if sheet.gst else "GST not stated in sheet; assumed excl"

    def ck(r):
        return (tid(r), comp_key(r["comp"]))

    # Prefer the NUoS block's component spelling, class label and unit for a tariff/component.
    nuos_first = sorted(sheet.records, key=lambda r: 0 if r["basis"] == "NUoS" else 1)
    label_for, cls_for, unit_for = {}, {}, {}
    for r in nuos_first:
        label_for.setdefault(comp_key(r["comp"]), r["comp"])
        cls_for.setdefault(tid(r), r["cls"])
        unit_for.setdefault(ck(r), r["unit"])

    # Basis arithmetic per tariff/component: NUoS vs DUoS + TUoS + JSA (+ Metering block).
    by_basis = {}
    for r in sheet.records:
        by_basis.setdefault(ck(r), {})[r["basis"]] = r["value"]

    def sum_note(r):
        if r["basis"] != "NUoS":
            return ""
        parts = by_basis[ck(r)]
        if not any(b in parts for b in ("DUoS", "TUoS", "JSA")):
            return ""
        sub = sum(parts.get(b, 0.0) for b in ("DUoS", "TUoS", "JSA"))
        tol = 2e-5 if "kwh" in r["unit"].lower() else 1.5e-3
        if abs(r["value"] - sub) <= tol:
            return ""
        m = sheet.metering.get(ck(r))
        if m is not None and abs(r["value"] - sub - m) <= tol:
            return f"NUoS includes Metering block charge {fmt(m)} {r['unit']} (NUoS = DUoS+TUoS+JSA+Metering)"
        return f"NUoS {fmt(r['value'])} != DUoS+TUoS+JSA {fmt(sub)} as published"

    starred = {tid(r) for r in sheet.records if r["name"].endswith("*")}
    footnote = "; ".join(dict.fromkeys(sheet.footnotes)) or "asterisk footnote"
    priced = {ck(r) for r in sheet.records if r["value"] != 0}

    rows = []
    for r in sheet.records:
        if ck(r) not in priced:
            continue
        label = label_for[comp_key(r["comp"])]
        unit = r["unit"]
        cls_unit = unit_for.get(ck(r)) or unit
        value_std, unit_std = units.to_std(r["published"], unit, label)
        notes = [f"sheet '{sheet_name}' {r['basis']} block row {r['row']}"]
        if r["zone"]:
            notes.append(f"zone: {r['zone']}")
        if tid(r) in starred:
            notes.append(footnote)
        if sheet_name in TRIAL_SHEETS:
            notes.append("trial tariff")
        if r["code"].upper().startswith("TBA"):
            notes.append("NTC published as TBA (not yet assigned)")
        notes.extend(sheet.site_specific.get(tid(r), []))
        if cls_unit and unit and cls_unit != unit:
            notes.append(f"unit as published differs from NUoS block unit '{cls_unit}' for this component")
        if (sn := sum_note(r)):
            notes.append(sn)
        if gst_note:
            notes.append(gst_note)
        name = r["name"].rstrip("*").strip()
        rows.append({
            "side": side,
            "distributor": distributor,
            "fin_year": fin_year,
            "tariff_code": r["code"],
            "tariff_name": name,
            "customer_class": cls_for.get(tid(r)) or r["cls"] or sheet.subtitle,
            "component": label,
            "charge_type": charge_type(label, cls_unit),
            "time_band": time_band(label),
            "season": schema.season_from_label(label),
            "unit": unit,
            "value": r["published"],
            "value_std": "" if value_std is None else str(value_std),
            "unit_std": unit_std,
            "gst": gst,
            "basis": r["basis"],
            "source_file": source_file,
            "source_url": source_url,
            "note": "; ".join(notes),
            "_sheet": sheet_name,
            "_zone": r["zone"],
        })
    return rows


def dedupe_across_sheets(rows):
    """Controlled-load tariffs are repeated on the Residential and Business sheets with identical
    prices: keep the first listing, note the other sheet. Differing values are kept and flagged."""
    seen = {}
    kept = []
    for r in rows:
        key = (r["tariff_code"], r["tariff_name"], r["_zone"], r["component"], r["basis"])
        prev = seen.get(key)
        if prev is None:
            seen[key] = r
            kept.append(r)
        elif prev["value"] == r["value"] and prev["unit"] == r["unit"]:
            prev["note"] += f"; also listed on sheet '{r['_sheet']}' (identical)"
        else:
            prev["note"] += f"; also listed on sheet '{r['_sheet']}' with a different value ({r['value']})"
            r["note"] += f"; also listed on sheet '{prev['_sheet']}' with a different value ({prev['value']})"
            kept.append(r)
    return kept


def parse_file(distributor, fin_year, path, source_url):
    """Parse one workbook; returns (rows, report)."""
    report = {"sheets": {}, "issues": [], "skipped": [], "unknown": []}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        name = ws.title.strip()
        if name in IN_SCOPE_SHEETS:
            sheet = parse_sheet(ws)
            report["issues"].extend(f"{name}: {m}" for m in sheet.issues)
            if sheet.title_year and sheet.title_year != fin_year:
                report["issues"].append(f"{name}: title year {sheet.title_year} != {fin_year}")
            r = rows_for_sheet(sheet, name, distributor, fin_year, side_for(path), path, source_url)
            report["sheets"][name] = len(r)
            rows.extend(r)
        elif name in OUT_OF_SCOPE_SHEETS:
            report["skipped"].append(name)
        else:
            report["unknown"].append(name)
    wb.close()
    return dedupe_across_sheets(rows), report


def main():
    urls = load_source_urls()
    all_rows = []
    for distributor, fin_year, path in FILES:
        if path not in urls:
            raise SystemExit(f"{path} missing from {INVENTORY}")
        rows, rep = parse_file(distributor, fin_year, path, urls[path])
        all_rows.extend(rows)
        codes = {r["tariff_code"] for r in rows}
        print(f"{path}\n  side={side_for(path)} rows={len(rows)} codes={len(codes)} "
              f"bases={sorted({r['basis'] for r in rows})} gst={sorted({r['gst'] for r in rows})}")
        print("  sheets: " + ", ".join(f"{k}={v}" for k, v in rep["sheets"].items()))
        if rep["unknown"]:
            print(f"  UNKNOWN SHEETS (not parsed): {rep['unknown']}")
        for m in rep["issues"]:
            print(f"  ! {m}")
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=schema.COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(all_rows)
    print(f"\nwrote {len(all_rows)} rows -> {OUT_PATH}")


if __name__ == "__main__":
    main()
