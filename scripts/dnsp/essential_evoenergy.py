"""Parse Essential Energy (NSW) and Evoenergy (ACT) network tariff price documents into the
normalised long CSV defined in scripts/schema.py.

Run from the repo root:  .venv/bin/python scripts/dnsp/essential_evoenergy.py
Output: out/dnsp/essential_evoenergy.csv

Files parsed (see FILES below). One function per file format:
  - Essential xlsx price lists (2023-24 AER-hosted, 2024-25 AER-hosted): stacked tables with
    multi-row headers; per-component sheets (DUOS/TUOS/CCF/QSS/Roadmap) are emitted with their basis.
  - Essential PDF price lists (2024-25..2026-27): ruled tables read with pdfplumber find_tables();
    column labels are rebuilt from header characters by x-range because header cells wrap/merge.
  - Evoenergy 2023-24 Statement of Tariff Classes and Tariffs: Table 2.6 (DUOS/TUOS/JS/NUOS) and
    Table 2.7 (NUOS excluding metering) parsed from word positions.
  - Evoenergy 2023-24 pricing proposal (AER-hosted): Table 4.2 is embedded as page images only, so
    it is OCR'd (rapidocr) and every row is validated with NUOS == DUOS + TUOS + JS before emission.
  - Evoenergy 2024-25 schedule PDF: per-tariff blocks parsed from word positions.
  - Evoenergy 2025-26 / 2026-27 schedule xlsx: per-tariff blocks (Component/Unit/Rate columns).
"""
import csv
import os
import re
import sys

import openpyxl
import pdfplumber

sys.path.insert(0, "scripts")
from published import cell_value
import schema  # noqa: E402
import units  # noqa: E402

ROOT = os.getcwd()
OUT = os.path.join("out", "dnsp", "essential_evoenergy.csv")
INVENTORY = os.path.join("sources", "inventory.csv")

# (fin_year, side, path, parser-key, extra)
FILES = [
    ("2023-24", "AER_HOSTED", "sources/aer/2023-24_price_lists/Essential_2023-24_Network_Price_List_31Mar2023.xlsx", "ess_xlsx", {}),
    ("2024-25", "AER_HOSTED", "sources/aer/dnsp_copies/Essential_2024-25_NUoS_Price_List_20May2024.xlsx", "ess_xlsx", {}),
    ("2024-25", "DNSP", "sources/dnsp/essential/Essential_Price_List_and_Explanatory_Notes_2024-25.pdf", "ess_pdf", {}),
    ("2025-26", "DNSP", "sources/dnsp/essential/Essential_Price_List_and_Explanatory_Notes_2025-26.pdf", "ess_pdf", {}),
    ("2026-27", "DNSP", "sources/dnsp/essential/Essential_Price_List_and_Explanatory_Notes_2026-27.pdf", "ess_pdf", {}),
    ("2023-24", "AER_HOSTED", "sources/aer/2023-24_price_lists/Evoenergy_2023-24_Electricity_network_pricing_proposal_5May2023.pdf", "evo_proposal", {}),
    ("2023-24", "DNSP", "sources/dnsp/evoenergy/Evoenergy_Statement_of_Tariff_Classes_and_Tariffs_2023-24.pdf", "evo_statement", {}),
    ("2024-25", "DNSP", "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2024-25_LFiT_adjusted.pdf", "evo_pdf", {"lfit": "LFiT included"}),
    ("2025-26", "DNSP", "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2025-26_incl_LFiT_May2025.xlsx", "evo_xlsx", {"lfit": "LFiT included"}),
    ("2026-27", "DNSP", "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2026-27_incl_LFiT_June2026.xlsx", "evo_xlsx", {"lfit": "LFiT included"}),
    ("2026-27", "DNSP", "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2026-27_AER_approved_April2026.xlsx", "evo_xlsx", {"lfit": "AER-approved excl LFiT"}),
]

EMIT_INCL_GST = False  # the documents also publish GST-inclusive copies of the same prices; not emitted
UNIT_TOKEN = re.compile(r"(\$|c|cents)\s*/\s*(kWh|kVAh|kVA|kW|Day|M\b|month|year)", re.I)
ESS_CODE = re.compile(r"^B[A-Z]{2,3}\d{0,2}[A-Z]{0,3}\d?$")
EVO_CODE = re.compile(r"^\d{3}\*?$")
NUM = re.compile(r"^-?\s?\d[\d,]*(\.\d+)?\^?\*?$")


# ----------------------------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------------------------
def load_inventory():
    m = {}
    with open(INVENTORY, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r.get("local_path"):
                m[r["local_path"].strip()] = r["source_url"].strip()
    return m


def clean(s):
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip()


def to_num(v):
    """Return float or None for a published cell value ('-' / 'Various' / blank -> None)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("^", "").replace("*", "").replace(",", "")
    s = re.sub(r"^-\s+(?=\d)", "-", s)  # '- 4.8793' -> '-4.8793'
    if not s or not re.match(r"^-?\d+(\.\d+)?$", s):
        return None
    return float(s)


def std_unit_for(unit, label):
    """Normalise a published unit for units.to_std(); keep the published string in `unit`."""
    u = unit.strip()
    u = re.sub(r"/\s*M$", "/month", u, flags=re.I)  # $/kVA/M -> $/kVA/month
    return u


def charge_type(label, unit):
    l = label.lower()
    if "net energy" in l:
        return "energy"
    if "export threshold" in l:
        return "export"
    return schema.charge_type_from_label(label, unit)


def time_band(label):
    l = label.lower()
    tb = schema.time_band_from_label(label)
    if tb:
        if tb == "peak" and ("peak period maximum demand" in l):
            return "peak"
        return tb
    if "max times" in l or "business times" in l:
        return "peak"
    if "mid times" in l or "evening times" in l:
        return "shoulder"
    if "economy times" in l:
        return "offpeak"
    if "controlled times" in l or "controlled load" in l:
        return "anytime"
    if "first" in l and "kwh" in l:
        return "block1"
    if "above" in l and "kwh" in l:
        return "block2"
    if "net energy" in l or re.search(r"\bany ?time\b", l) or re.search(r"\ball\b", l):
        return "anytime"
    return ""


def season(label):
    s = schema.season_from_label(label)
    if s:
        return s
    l = label.lower()
    if "(high)" in l:
        return "high"
    if "(low)" in l:
        return "low"
    return ""


def make_row(distributor, fin_year, side, path, url, code, name, cls, component, unit, value,
             gst, basis, note):
    unit_pub = clean(unit)
    value = re.sub(r"^-\s+(?=\d)", "-", str(value).replace(",", "").strip())
    if re.search(r"kvah", unit_pub, re.I):
        v_std, u_std = (float(value) * (100.0 if unit_pub.strip().startswith("$") else 1.0)), "c/kVAh"
    elif re.search(r"kka", unit_pub, re.I):  # Evoenergy 2023-24 typo 'c/KkA/day' (tariff 124)
        v_std, u_std = float(value), "c/kVA/day"
        note = (note + "; " if note else "") + "unit published as 'c/KkA/day' (document typo; tariff 123 shows c/kVA/day)"
    else:
        v_std, u_std = units.to_std(value, std_unit_for(unit_pub, component), component)
    return {
        "side": side,
        "distributor": distributor,
        "fin_year": fin_year,
        "tariff_code": clean(code),
        "tariff_name": clean(name),
        "customer_class": clean(cls),
        "component": clean(component),
        "charge_type": charge_type(component, unit_pub),
        "time_band": time_band(component),
        "season": season(component),
        "unit": unit_pub,
        "value": value,
        "value_std": "" if v_std is None else repr(v_std),
        "unit_std": u_std,
        "gst": gst,
        "basis": basis,
        "source_file": path,
        "source_url": url,
        "note": note,
    }


def split_label_unit(text):
    """'Peak $/kVA/M' -> ('Peak', '$/kVA/M'); 'Charge <=7.5kWh* c/kWh' -> ('Charge <=7.5kWh*', 'c/kWh')."""
    t = clean(text)
    m = re.search(r"((?:\$|c)/(?:kWh|kVA|kW)(?:/M)?|\$/Day|c/kWh)\s*$", t, re.I)
    if m:
        return clean(t[: m.start()]), m.group(1)
    return t, ""


# ----------------------------------------------------------------------------------------------
# Essential Energy - xlsx price lists (2023-24 and 2024-25 AER-hosted workbooks)
# ----------------------------------------------------------------------------------------------
ESS_SHEET_BASIS = [
    (re.compile(r"duos", re.I), "DUoS", "DUOS component sheet"),
    (re.compile(r"tuos", re.I), "TUoS", "TUOS component sheet"),
    (re.compile(r"ccf", re.I), "JSA", "CCF (Climate Change Fund) jurisdictional scheme component sheet"),
    (re.compile(r"qss", re.I), "JSA", "QSS component sheet (jurisdictional scheme component as labelled in workbook; scheme not defined in document)"),
    (re.compile(r"roadmap", re.I), "JSA", "NSW Electricity Infrastructure Roadmap jurisdictional scheme component sheet"),
]


def sheet_basis(title):
    for rx, basis, note in ESS_SHEET_BASIS:
        if rx.search(title):
            return basis, note
    return "NUoS", "NUOS network price (total network price incl. transmission and jurisdictional scheme components)"


def is_unit_row(texts):
    return sum(1 for t in texts.values() if UNIT_TOKEN.search(t)) >= 2


def is_header_start(texts):
    vals = [t.lower() for t in texts.values()]
    if any(re.match(r"^(obsolete\s+)?tariff\s*code$", v) or v == "network charge" for v in vals):
        return True
    return "obsolete" in vals and any(v in ("network", "energy", "description") for v in vals)


def build_header(header_rows):
    """header_rows: list of {col: text}. Returns dict col -> (label, unit)."""
    unit_row = header_rows[-1]
    label_rows = header_rows[:-1]
    cols = set()
    for r in header_rows:
        cols |= set(r)
    # forward-fill first label row across empty columns (merged group headers)
    filled_first = {}
    if label_rows:
        last = None
        for c in range(min(cols), max(cols) + 1):
            if c in label_rows[0]:
                last = label_rows[0][c]
            if last is not None:
                filled_first[c] = last
    out = {}
    for c in sorted(cols):
        parts = []
        if filled_first.get(c):
            parts.append(filled_first[c])
        for r in label_rows[1:]:
            if r.get(c):
                parts.append(r[c])
        unit = ""
        u = unit_row.get(c, "")
        if u and re.match(r"^(\$|c)/\S+$", u):
            unit = u  # pure unit cell, e.g. '$/kVA/M'
        elif u:
            parts.append(u)  # label with embedded unit, e.g. 'Peak c/kWh' (split below)
        label = " - ".join(p for p in parts if p)
        if not unit:
            label, unit = split_label_unit(label) if UNIT_TOKEN.search(label) else (label, "")
        out[c] = (label, unit)
    return out


def parse_essential_xlsx(path, fin_year, side, url):
    rows_out = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        if not re.search(r"price ?list|tariffs", ws.title, re.I) or re.search(r"details|explanatory|time periods", ws.title, re.I):
            continue
        basis, basis_note = sheet_basis(ws.title)
        gst = "incl" if re.search(r"incl", ws.title, re.I) else "excl"
        if gst == "incl" and not EMIT_INCL_GST:
            continue
        page_note = f"sheet '{ws.title}'"
        header = None
        header_rows = []
        in_header = False
        section = ""
        obsolete = False
        overrides = {}
        table_note = ""
        code_col = desc_col = same_rate_col = tp_col = app_col = None
        footnotes = {}
        for cell_row in ws.iter_rows():
            raw = [c.value for c in cell_row]
            cells = {j: v for j, v in enumerate(raw, 1) if v is not None and str(v).strip() != ""}
            if not cells:
                continue
            texts = {j: clean(v) for j, v in cells.items() if isinstance(v, str)}
            nums = {j: v for j, v in cells.items() if isinstance(v, (int, float))}
            first = texts.get(min(cells))
            # title rows
            if first and re.search(r"price ?list", first, re.I) and len(cells) == 1:
                gst = "incl" if re.search(r"including gst", first, re.I) else "excl"
                table_note = ""
                if re.search(r"interim", first, re.I):
                    table_note = "Interim Network Price List table"
                continue
            if first and re.match(r"^effective ", first, re.I) and len(cells) == 1:
                if "2015" in first:
                    table_note = (table_note + "; " if table_note else "") + "effective 1 July 2015 (legacy table republished)"
                continue
            if in_header:
                header_rows.append(texts)
                if is_unit_row(texts):
                    header = build_header(header_rows)
                    in_header = False
                    code_col = min((c for c, (l, u) in header.items() if re.search(r"tariff|network charge|^obsolete", l, re.I)), default=min(header))
                    desc_col = next((c for c, (l, u) in header.items() if re.match(r"^description", l, re.I)), None)
                    same_rate_col = next((c for c, (l, u) in header.items() if re.search(r"same rate", l, re.I)), None)
                    tp_col = next((c for c, (l, u) in header.items() if re.search(r"time period", l, re.I)), None)
                    app_col = next((c for c, (l, u) in header.items() if re.search(r"application", l, re.I)), None)
                    overrides = {}
                    section = ""
                    obsolete = bool(re.search(r"obsolete", header[code_col][0], re.I))
                continue
            if is_header_start(texts):
                in_header = True
                header_rows = [texts]
                continue
            if header is None:
                continue
            if all(t.lower() in ("excluding gst", "including gst") for t in texts.values()):
                continue
            code_text = texts.get(code_col, "")
            if first and first.startswith(("*", "^")):
                footnotes[first[0]] = first
                continue
            if code_text and not any(ESS_CODE.match(tok) for tok in code_text.split()):
                # section heading (may carry header overrides for some columns) or footnote
                if nums or re.search(r"various", " ".join(texts.values()), re.I):
                    continue  # 'Customer specific prices' row: no numeric prices published
                section = code_text
                obsolete = obsolete or bool(re.search(r"obsolete", section, re.I))
                overrides = {}
                for c, t in texts.items():
                    if c != code_col and UNIT_TOKEN.search(t):
                        grp = header[c][0].split(" - ")[0] if c in header else ""
                        lab, u = split_label_unit(t)
                        overrides[c] = (f"{grp} - {lab}" if grp else lab, u)
                continue
            if not code_text:
                continue
            codes = [tok for tok in code_text.split() if ESS_CODE.match(tok)]
            name = texts.get(desc_col, "") if desc_col else ""
            name_note = ""
            if not name:
                # obsolete tables on the component sheets have no 'Description' header; the name sits in the
                # unlabelled column right after the code column (before the first priced column)
                first_priced = min((c for c, (l, u) in header.items() if u), default=None)
                cand = texts.get(code_col + 1, "")
                if cand and (first_priced is None or code_col + 1 < first_priced) and not UNIT_TOKEN.search(cand) and not ESS_CODE.match(cand):
                    name = cand
                    name_note = "tariff name read from the unlabelled column next to the code (no 'Description' header on this table)"
            notes = [basis_note, page_note] + ([name_note] if name_note else [])
            if table_note:
                notes.append(table_note)
            if obsolete:
                notes.append("obsolete tariff - not applicable to new connections")
            if tp_col and texts.get(tp_col):
                notes.append(f"time period: {texts[tp_col]}")
            if app_col and texts.get(app_col):
                notes.append("application: " + texts[app_col])
            if gst == "incl":
                notes.append("prices including GST as published")
            same_rate = [tok for tok in texts.get(same_rate_col, "").split() if ESS_CODE.match(tok)] if same_rate_col else []
            emit_codes = [(c, "") for c in codes]
            if len(codes) > 1:
                emit_codes = [(c, f"published jointly with {', '.join(x for x in codes if x != c)} (same rate)") for c in codes]
            for sc in same_rate:
                emit_codes.append((sc, f"obsolete tariff on same rate as {codes[0]} (published in 'Obsolete tariffs on same rate' column)"))
            for c, (label, unit) in header.items():
                if c in (code_col, desc_col, same_rate_col, tp_col, app_col) or c not in cells:
                    continue
                if c in overrides:
                    label, unit = overrides[c]
                if not unit:
                    continue
                val = to_num(cells[c])
                if val is None:
                    continue
                comp_note = []
                if "*" in label and "*" in footnotes:
                    comp_note.append(footnotes["*"].lstrip("* "))
                if unit.strip() in ("$/kW", "$/kVA"):
                    comp_note.append("period not stated in unit as published (PDF price list shows $/kW/M)")
                for code, cnote in emit_codes:
                    rows_out.append(make_row("Essential Energy", fin_year, side, path, url, code, name, section,
                                             label, unit, cell_value(cell_row[c - 1]), gst, basis,
                                             "; ".join(n for n in notes + comp_note + ([cnote] if cnote else []) if n)))
    return rows_out


# ----------------------------------------------------------------------------------------------
# Essential Energy - PDF price lists (2024-25 .. 2026-27)
# ----------------------------------------------------------------------------------------------
GROUP_WORDS = {"energy", "demand", "export"}


def _chars_text(chars):
    """Reading-order text for a set of chars (lines by top, then x)."""
    lines = {}
    for c in chars:
        lines.setdefault(round(c["top"] / 4), []).append(c)
    out = []
    for k in sorted(lines):
        cs = sorted(lines[k], key=lambda c: c["x0"])
        s = ""
        prev = None
        for c in cs:
            if prev is not None and c["x0"] - prev > 1.0:
                s += " "
            s += c["text"]
            prev = c["x1"]
        out.append(s)
    return clean(" ".join(out))


def parse_essential_pdf(path, fin_year, side, url):
    rows_out = []
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            m = re.search(r"Network Price List \((Excluding|Including) GST\)", text)
            if not m or not re.search(r"\bB[LHS][A-Z]{2}\d?[A-Z0-9]{2,4}\b", text):
                continue
            gst = "excl" if m.group(1) == "Excluding" else "incl"
            if gst == "incl" and not EMIT_INCL_GST:
                continue  # Attachment 1 'Network Price List including GST' (same prices x 1.1)
            tables = page.find_tables()
            if not tables:
                continue
            table = max(tables, key=lambda t: (t.bbox[2] - t.bbox[0]) * (t.bbox[3] - t.bbox[1]))
            grid = table.extract()
            rows = table.rows
            words = page.extract_words()
            # classify rows
            def row_text(i):
                return [clean(c) for c in grid[i]]
            kinds = []
            for i in range(len(grid)):
                t = row_text(i)
                nonempty = [x for x in t if x]
                first = nonempty[0] if nonempty else ""
                if first and ESS_CODE.match(first.split()[0]) or (first == "Customer specific"):
                    kinds.append("data")
                elif any(re.search(r"Tariffs|Customer Specific", x) for x in nonempty):
                    kinds.append("section")
                elif any(re.search(r"Description", x) for x in nonempty):
                    kinds.append("header")
                elif nonempty:
                    kinds.append("text")
                else:
                    kinds.append("empty")
            if "data" not in kinds or "header" not in kinds:
                continue
            h0 = kinds.index("header")
            first_body = next(i for i, k in enumerate(kinds) if k in ("data", "section") and i > h0)
            # data columns: indices with numeric-looking cells in data rows (beyond the first 3 text columns)
            data_idx = [i for i, k in enumerate(kinds) if k == "data" and grid[i][0] and ESS_CODE.match(clean(grid[i][0]).split()[0])]
            textcols = set()
            for i in data_idx:
                t = row_text(i)
                ne = [j for j, x in enumerate(t) if x]
                textcols |= set(ne[:3])
            data_cols = {}
            for i in data_idx:
                t = row_text(i)
                for j, x in enumerate(t):
                    if j in textcols or not x:
                        continue
                    if NUM.match(x.replace(" ", "")) and rows[i].cells[j]:
                        data_cols[j] = rows[i].cells[j]
            data_cols = {j: (b[0], b[2]) for j, b in data_cols.items()}
            # header band: from the header row top to the first section/data row top
            band_top = min(b[1] for b in rows[h0].cells if b)
            first_data = next(i for i, k in enumerate(kinds) if k == "data" and i > h0)
            band_bot = min(b[1] for b in rows[first_data].cells if b)
            # the first section heading may share a ruled row with the unit line: cut the band at that word line
            sec_words = [w for w in words if re.search(r"Tariffs", w["text"]) and band_top < w["top"] < band_bot]
            if sec_words:
                band_bot = min(band_bot, min(w["top"] for w in sec_words) - 1)
            band_chars = [c for c in page.chars if band_top - 1 <= c["top"] and c["bottom"] <= band_bot + 1]
            group_words = [w for w in words if w["text"].lower() in GROUP_WORDS and band_top - 1 <= w["top"] <= band_bot]
            # group row cells (merged) for group assignment
            group_cells = []
            for i in range(h0, first_body):
                for j, cell in enumerate(rows[i].cells):
                    if cell and re.search(r"\b(Energy|Demand|Export)\b", clean(grid[i][j]) or ""):
                        gw = re.search(r"\b(Energy|Demand|Export)\b", clean(grid[i][j])).group(1)
                        group_cells.append((cell[0], cell[2], gw))

            def in_group_word(c):
                return any(w["x0"] - 0.5 <= c["x0"] and c["x1"] <= w["x1"] + 0.5 and abs(w["top"] - c["top"]) < 2 for w in group_words)

            base_labels = {}
            for j, (x0, x1) in data_cols.items():
                cx = (x0 + x1) / 2
                cs = [c for c in band_chars if x0 - 0.5 <= (c["x0"] + c["x1"]) / 2 <= x1 + 0.5 and not in_group_word(c)]
                sub = _chars_text(cs)
                grp = next((g for (gx0, gx1, g) in group_cells if gx0 - 0.5 <= cx <= gx1 + 0.5), "")
                lab, unit = split_label_unit(sub)
                base_labels[j] = ((f"{grp} - {lab}" if grp else lab), unit)
            labels = dict(base_labels)
            section = ""
            pending_override = {}
            foot = {}
            key = None
            for line in text.splitlines():
                if line[:1] in ("*", "^"):
                    key = line[0]
                    foot[key] = clean(line[1:])
                elif key and not re.match(r"^(Network Pricelist|FOR PUBLIC|Page \d|Effective )", line):
                    foot[key] = clean(foot[key] + " " + line)
                else:
                    key = None
            if "^" in foot:
                foot["^"] = "daily charge includes legacy metering cost (document does not show the split): ^" + foot["^"]
            for i in range(first_body, len(grid)):
                k = kinds[i]
                t = row_text(i)
                if k == "section":
                    section = next(x for x in t if x)
                    labels = dict(base_labels)
                    pending_override = {}
                if k in ("section", "text"):
                    # header override cells for export columns (e.g. 'Charge <=1.5KW $/kW/M' over several rows)
                    for j, x in enumerate(t):
                        if not x or (k == "section" and x == section) or not rows[i].cells[j]:
                            continue
                        cx0, _, cx1, _ = rows[i].cells[j]
                        for dj, (x0, x1) in data_cols.items():
                            if min(cx1, x1) - max(cx0, x0) > 3:
                                pending_override.setdefault(dj, []).append(x)
                    continue
                if k != "data":
                    continue
                if pending_override:
                    for dj, parts in pending_override.items():
                        grp = base_labels[dj][0].split(" - ")[0] if " - " in base_labels[dj][0] else ""
                        lab, unit = split_label_unit(" ".join(parts))
                        labels[dj] = ((f"{grp} - {lab}" if grp else lab), unit or base_labels[dj][1])
                    pending_override = {}
                ne = [j for j, x in enumerate(t) if x]
                code = t[ne[0]]
                if not ESS_CODE.match(code.split()[0]):
                    continue  # 'Customer specific' - 'Various', not priced
                name = t[ne[1]] if len(ne) > 1 else ""
                tp = t[ne[2]] if len(ne) > 2 and len(t[ne[2]]) <= 12 else ""
                notes = ["NUOS network price (total network price incl. transmission and jurisdictional scheme components)",
                         f"page {pno}"]
                if re.search(r"obsolete", section, re.I):
                    notes.append("obsolete tariff - not available for new connects or tariff changes")
                if tp:
                    notes.append(f"time period: {tp}")
                if gst == "incl":
                    notes.append("prices including GST as published (Attachment 1)")
                cls = re.split(r"\s[–-]\s", section)[0] if section else ""
                for j, (label, unit) in labels.items():
                    raw = t[j] if j < len(t) else ""
                    if not raw or not unit:
                        continue
                    val = to_num(raw)
                    if val is None:
                        continue
                    cnote = []
                    if "^" in raw and "^" in foot:
                        cnote.append(foot["^"])
                    if "*" in label and "*" in foot:
                        cnote.append(foot["*"])
                    rows_out.append(make_row("Essential Energy", fin_year, side, path, url, code, name, cls,
                                             label.replace("*", "").strip(), unit, str(raw).replace("^", "").replace("*", "").replace(",", "").strip(), gst, "NUoS",
                                             "; ".join(notes + cnote)))
    return rows_out


# ----------------------------------------------------------------------------------------------
# Evoenergy - 2023-24 Statement of Tariff Classes and Tariffs (Tables 2.6 and 2.7)
# ----------------------------------------------------------------------------------------------
EVO_UNIT = re.compile(r"^(cents/day|cents/kWh|c/kW/day|c/kVA/day|c/KkA/day|\$/day|c/day|c/kWh|c/kVAh)$", re.I)
EVO_SECTION = re.compile(r"^(Residential|LV Commercial|HV Commercial) Tariffs$", re.I)


def _lines(words, tol=3.0):
    lines = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(lines[-1][0] - w["top"]) <= tol:
            lines[-1][1].append(w)
        else:
            lines.append([w["top"], [w]])
    return [(t, sorted(ws, key=lambda w: w["x0"])) for t, ws in lines]


def _header_groups(line_words, skip=("Description", "Units")):
    """Group adjacent header words into column headers; return list of (text, x0, x1)."""
    groups = []
    for w in line_words:
        if w["text"] in skip:
            continue
        if groups and w["x0"] - groups[-1][2] < 8:
            groups[-1] = (groups[-1][0] + " " + w["text"], groups[-1][1], w["x1"])
        else:
            groups.append((w["text"], w["x0"], w["x1"]))
    seen = {}
    out = []
    for text, x0, x1 in groups:  # make duplicate header names unique (e.g. several 'NUOS' columns)
        n = seen.get(text, 0)
        seen[text] = n + 1
        out.append((text if n == 0 else f"{text} #{n}", x0, x1))
    return out


def _parse_evo_table(pages_words, header_rx, col_names, stop_rx):
    """Generic parser for Evoenergy 'Description Units <numeric cols>' tables spread over pages.
    pages_words: list of (page_no, words). Returns list of dicts with code, name, section, label, unit, values{col:val}, page.
    """
    out = []
    code = name = section = None
    cur = None
    for pno, words in pages_words:
        lines = _lines(words)
        header = None
        for top, lw in lines:
            txt = " ".join(w["text"] for w in lw)
            if stop_rx.search(txt):
                code = None
                header = None
                break
            if header_rx.search(txt):
                header = _header_groups(lw)
                cur = None
                continue
            if header is None:
                continue
            if re.match(r"^\d+ \| Evoenergy", txt) or re.match(r"^(price|capital|non-|\(units\)|\(%\))", txt):
                continue  # footer / header continuation lines
            if EVO_SECTION.match(txt):
                section = txt
                cur = None
                continue
            m = re.match(r"^(\d{3}\*?)\s+([A-Z].*)$", txt)
            if m and "kWh" not in txt and not any(EVO_UNIT.match(w["text"]) for w in lw):
                code, name = m.group(1), m.group(2)
                cur = None
                continue
            uw = next((w for w in lw if EVO_UNIT.match(w["text"])), None)
            nums = [w for w in lw if re.match(r"^-?\d+\.\d+%?$", w["text"]) and (uw is None or w["x0"] > uw["x1"])]
            if uw and nums and code:
                label = " ".join(w["text"] for w in lw if w["x1"] <= uw["x0"] + 1)
                values = {}
                for w in nums:
                    g = min(header, key=lambda g: abs(g[2] - w["x1"]))
                    values[g[0]] = w["text"]
                cur = {"code": code, "name": name, "section": section, "label": label, "unit": uw["text"],
                       "values": values, "page": pno}
                out.append(cur)
            elif cur is not None and not nums and code:
                # continuation of the label (wrapped text below the price line)
                cur["label"] = clean(cur["label"] + " " + txt)
    return out


def parse_evo_statement(path, fin_year, side, url):
    rows_out = []
    with pdfplumber.open(path) as pdf:
        pages = [(i, p.extract_words(), p.extract_text() or "") for i, p in enumerate(pdf.pages, 1)]
    t26 = [(i, w) for i, w, t in pages if re.search(r"Description Units DUOS TUOS JS price Metering", t)]
    t27 = [(i, w) for i, w, t in pages if re.search(r"Description Units NUOS 2022/23 NUOS 2023/24", t)]
    recs26 = _parse_evo_table(t26, re.compile(r"^Description Units DUOS TUOS"), None, re.compile(r"XMC tariffs exclude metering|^2\.5 Changes"))
    recs27 = _parse_evo_table(t27, re.compile(r"^Description Units NUOS 2022/23"), None, re.compile(r"^2\.6 |Table 2\.8"))
    base = "Table 2.6 Network use of system (NUOS) charges 2023/24, excluding GST"
    n_chk = n_fail = 0
    for r in recs26:
        vals = {k: to_num(v) for k, v in r["values"].items()}
        parts = [v for k, v in vals.items() if not k.startswith("NUOS") and v is not None]
        nuos = next((v for k, v in vals.items() if k.startswith("NUOS")), None)
        if nuos is not None and parts:
            n_chk += 1
            if abs(sum(parts) - nuos) > 0.0015:
                n_fail += 1
                diff = sum(parts) - nuos
                if abs(diff - 2.270) <= 0.002:
                    why = "the JS price column for this tariff appears to exclude the LFiT rebate (difference 2.270 c/kWh = the average LFiT rebate stated in the document) while the NUOS column includes it"
                elif abs(diff) <= 0.003:
                    why = "rounding"
                else:
                    why = "unexplained"
                r["sum_note"] = f"document inconsistency: published NUOS {nuos:g} != DUOS+TUOS+JS(+metering) = {sum(parts):.3f} ({why}); values kept as published"
                print(f"WARN statement row fails NUOS = sum of columns: p{r['page']} {r['code']} {r['label']} {r['values']}", file=sys.stderr)
    print(f"statement: Table 2.6 column-mapping check NUOS == DUOS+TUOS+JS+metering: {n_chk - n_fail}/{n_chk} rows pass")
    for r in recs26:
        name = r["name"]
        cls = r["section"] or ""
        notes_common = [base, f"page {r['page']}"]
        if "XMC" in name:
            notes_common.append("XMC tariff excludes metering capital charges")
        if r.get("sum_note"):
            notes_common.append(r["sum_note"])
        for col, val in r["values"].items():
            v = to_num(val)
            if v is None:
                continue
            if col.startswith("DUOS"):
                basis, note = "DUoS", "DUOS price column (distribution only; unaffected by LFiT)"
            elif col.startswith("TUOS"):
                basis, note = "TUoS", "TUOS price column (transmission; unaffected by LFiT)"
            elif col.startswith("JS"):
                basis, note = "JSA", "LFiT included: JS price column = jurisdictional schemes incl. the LFiT rebate (negative adjustment, avg -2.27 c/kWh) plus feed-in tariff schemes, Utilities Network Facilities Tax and Energy Industry Levy"
            elif col.startswith("NUOS"):
                basis, note = "NUoS", "LFiT included: NUOS price column as published = DUOS + TUOS + JS + metering capital + metering non-capital (LFiT rebate applied via the JS price; metering (ACS) included where published; the excl-metering NUOS is in the 'NUOS 2023/24 column' rows from Table 2-7)"
            else:
                continue  # metering capital / non-capital columns (alternative control) not emitted
            rows_out.append(make_row("Evoenergy", fin_year, side, path, url, r["code"], name, cls, r["label"], r["unit"], val,
                                     "excl", basis, "; ".join([note] + notes_common)))
    base27 = "Table 2.7 2022/23 and 2023/24 NUOS tariffs, excluding metering (nominal)"
    for r in recs27:
        for col, val in r["values"].items():
            if not col.startswith("NUOS 2023/24"):
                continue
            v = to_num(val)
            if v is None:
                continue
            rows_out.append(make_row("Evoenergy", fin_year, side, path, url, r["code"], r["name"], r["section"] or "", r["label"], r["unit"], val,
                                     "excl", "NUoS", "; ".join(["LFiT included: NUOS 2023/24 column (excludes metering; LFiT rebate applied via the JS price)", base27, f"page {r['page']}"])))
    return rows_out


# ----------------------------------------------------------------------------------------------
# Evoenergy - 2023-24 pricing proposal (AER-hosted): Table 4.2 is image-only -> OCR + validation
# ----------------------------------------------------------------------------------------------
def parse_evo_proposal(path, fin_year, side, url, reference_rows):
    """Table 4.2 'Proposed 2023/24 prices and revenue, excluding metering' (pages 40-45) is image-only.
    Two OCR passes (288 and 400 dpi); a row is emitted only when NUOS == DUOS + TUOS + JS (+-0.0015),
    which also repairs decimal points dropped by the OCR. DUOS/TUOS are cross-checked against the
    Statement of Tariff Classes (same code, same ordinal component) and Table 4.3 (pages 46-50) is
    used as a second cross-check of the NUOS column."""
    rows_out = []
    from rapidocr_onnxruntime import RapidOCR
    from difflib import SequenceMatcher
    from itertools import combinations
    ocr = RapidOCR()
    ref = {}  # code -> list of (label, DUOS, TUOS) from the Statement's Table 2.6, in document order
    tmp = {}
    for r in reference_rows:
        if r["basis"] in ("DUoS", "TUoS") and "Table 2.6" in r["note"]:
            tmp.setdefault((r["tariff_code"], r["basis"]), []).append((r["component"], float(r["value"])))
    for (code_, basis_), lst in tmp.items():
        if basis_ == "DUoS":
            tu = tmp.get((code_, "TUoS"), [])
            ref[code_] = [(lab, d, tu[i][1] if i < len(tu) else None) for i, (lab, d) in enumerate(lst)]
    names = {r["tariff_code"]: r["tariff_name"] for r in reference_rows}
    UNIT_RX = re.compile(r"^(cents/day|cents/kWh|cents/kVAh|c/kVAh|c/kW/day|c/kVA/day|c/KkA/day|\$/day)$", re.I)

    def ocr_pages(resolution):
        recs, recs43 = [], []
        header = header43 = None
        desc_right = None
        code = name = section = None
        cur = None
        with pdfplumber.open(path) as pdf:
            for pno, page in enumerate(pdf.pages, 1):
                if page.chars or not page.images or pno < 30 or pno > 60:
                    continue
                res, _ = ocr(page.to_image(resolution=resolution).original)
                if not res:
                    continue
                items = []
                for box, txt, conf in res:
                    xs = [pt[0] for pt in box]
                    ys = [pt[1] for pt in box]
                    items.append({"top": sum(ys) / 4, "x0": min(xs), "x1": max(xs), "text": txt.strip(), "conf": conf})
                for top, lw in _lines(items, tol=14 * resolution / 288):
                    txt = " ".join(w["text"] for w in lw)
                    if re.search(r"DUOS prices", txt) and re.search(r"NUOS prices", txt):
                        header = _header_groups(lw, skip=("Description", "Units"))
                        desc_right = next(w["x0"] for w in lw if w["text"] in ("Units", "Unit"))
                        header43, code = None, None
                        continue
                    if re.search(r"NUOS actual", txt) and re.search(r"Change", txt):
                        header43 = _header_groups(lw, skip=("Description", "Unit", "Units"))
                        desc_right = next(w["x0"] for w in lw if w["text"] in ("Units", "Unit"))
                        header, code = None, None
                        continue
                    if header is None and header43 is None:
                        continue
                    if re.search(r"Evoenergy \| Electricity Network Pricing Proposal", txt) or \
                            re.match(r"^(forecast|volumes|revenue|\(per|20\d\d/\d\d|23/24|evoenerg|\*|Table )", txt, re.I):
                        continue
                    if re.match(r"^(Residential|LV|HV)\s*(Commercial)?\s*Tariffs$", txt, re.I):
                        section = txt
                        continue
                    m = re.match(r"^(\d{3})\s*([A-Za-z].*)$", txt)
                    if m and not re.search(r"cents/|c/k|\$/|kWh", txt):
                        code, name = m.group(1), m.group(2)
                        ref_name = names.get(code, "")
                        if ref_name and re.sub(r"\W", "", ref_name).lower() == re.sub(r"\W", "", name).lower():
                            name = ref_name  # OCR drops the spaces in bold headings; identical text in the Statement
                        cur = None
                        continue
                    uw = next((w for w in lw if UNIT_RX.match(w["text"])), None)
                    if uw and code:
                        hdr_use = header if header is not None else header43
                        values = {}
                        for w in lw:
                            if w["x0"] > uw["x1"]:
                                g = min(hdr_use, key=lambda g: abs(g[2] - w["x1"]))
                                values[g[0]] = w["text"]
                        cur = {"code": code, "name": name, "section": section, "unit": uw["text"], "values": values, "page": pno,
                               "label": " ".join(w["text"] for w in lw if w["x1"] <= uw["x0"] + 2)}
                        (recs if header is not None else recs43).append(cur)
                    elif cur is not None and code and not re.search(r"\d\.\d{3}|\$", txt) and all(w["x1"] <= desc_right + 2 for w in lw):
                        cur["label"] = clean(cur["label"] + " " + txt)
        return recs, recs43

    def fix(sv):
        return re.sub(r"[^\d.\-]", "", sv.replace(",", "").replace("$", "").replace(" ", ""))

    def check(n):
        return abs(n["NUOS prices"] - (n["DUOS prices"] + n["TUOS prices"] + n["JS prices"])) <= 0.0015

    def validate(recs):
        order = {}
        for r in recs:
            r["ordinal"] = order.get(r["code"], 0)
            order[r["code"]] = r["ordinal"] + 1
            vals = {}
            for key in ("DUOS prices", "TUOS prices", "JS prices", "NUOS prices"):
                raw = next((v for k, v in r["values"].items() if k.startswith(key)), None)
                vals[key] = fix(raw) if raw is not None else None
            if any(v in (None, "") for v in vals.values()):
                r["status"] = "missing column"
                continue
            try:
                nums = {k: float(v) for k, v in vals.items()}
            except ValueError:
                r["status"] = "unparseable"
                continue
            if not check(nums):
                cands = [k for k, v in vals.items() if "." not in v and nums[k] != 0]
                fixed = False
                for n in range(1, len(cands) + 1):
                    for combo in combinations(cands, n):
                        alt = dict(nums)
                        for k in combo:
                            alt[k] = nums[k] / 1000.0  # OCR dropped the decimal point (3 dp in the document)
                        if check(alt):
                            nums, fixed = alt, True
                            break
                    if fixed:
                        break
                if not fixed:
                    r["status"] = "sum check failed"
                    continue
                r["repair"] = "OCR decimal point repaired via sum check"
            r["status"] = "ok"
            r["nums"] = nums

    recs_a, recs43 = ocr_pages(288)
    recs_b, _ = ocr_pages(400)
    validate(recs_a)
    validate(recs_b)
    by_key_b = {(r["code"], r["ordinal"]): r for r in recs_b}
    final = []
    for r in recs_a:
        if r["status"] != "ok":
            alt = by_key_b.get((r["code"], r["ordinal"]))
            if alt is not None and alt["status"] == "ok" and alt["unit"].lower() == r["unit"].lower():
                alt["repair"] = (alt.get("repair", "") + "; " if alt.get("repair") else "") + "value taken from second OCR pass (400 dpi) after first pass failed the sum check"
                final.append(alt)
                continue
            print(f"WARN proposal OCR row dropped ({r['status']}): p{r['page']} {r['code']} {r['label']} {r['values']}", file=sys.stderr)
            continue
        final.append(r)
    # Table 4.3 cross-check of the NUOS column (2nd header group = 'NUOS proposed 2023/24')
    t43 = {}
    for r in recs43:
        keys = list(r["values"].keys())
        if len(keys) >= 2:
            t43.setdefault(r["code"], []).append(fix(r["values"][keys[1]]) if keys[1] in r["values"] else "")
    n_ok = n_bad = 0
    for r in final:
        lst = t43.get(r["code"], [])
        if r["ordinal"] < len(lst):
            try:
                same = abs(float(lst[r["ordinal"]]) - r["nums"]["NUOS prices"]) <= 0.0015
            except ValueError:
                same = False
            n_ok += same
            n_bad += not same
    print(f"proposal: Table 4.2 rows ok={len(final)} dropped={len(recs_a) - len(final)}; Table 4.3 NUOS cross-check: {n_ok} match, {n_bad} differ")

    def norm(t):
        return re.sub(r"[^a-z]", "", t.lower())
    for r in final:
        label = r["label"]
        # Statement rows for this code with the same DUOS and TUOS values (both documents publish identical DUOS/TUOS);
        # among those, the one whose label is most similar to the OCR text supplies the clean wording.
        cands = [(SequenceMatcher(None, norm(label), norm(lab)).ratio(), lab) for lab, d, t in ref.get(r["code"], [])
                 if abs(d - r["nums"]["DUOS prices"]) <= 0.0005 and t is not None and abs(t - r["nums"]["TUOS prices"]) <= 0.0005]
        if cands:
            ratio, lab = max(cands)
            xnote = "DUOS/TUOS match a Statement of Tariff Classes row for this code"
            if ratio >= 0.6 and lab != label:
                label = lab
            elif ratio < 0.6:
                xnote += f" (label similarity low: '{lab}')"
        elif ref.get(r["code"]):
            xnote = "DUOS/TUOS DIFFER from every Statement of Tariff Classes row for this code"
            print(f"WARN proposal OCR row differs from Statement: p{r['page']} {r['code']} {r['label']} {r['nums']}", file=sys.stderr)
        else:
            xnote = "no Statement row to cross-check"
        notes = ["excl LFiT (proposed prices as submitted to the AER; the ACT LFiT scheme amount is not included)",
                 "Table 4.2 Proposed 2023/24 prices and revenue, excluding metering (nominal)", f"page {r['page']}",
                 "OCR of image-only page (rapidocr); validated NUOS = DUOS + TUOS + JS", xnote]
        if r.get("repair"):
            notes.append(r["repair"])
        if label != r["label"]:
            notes.append(f"component label aligned to Statement wording (OCR text: '{r['label']}')")
        for b, key, note in (("DUoS", "DUOS prices", "DUOS prices column"), ("TUoS", "TUOS prices", "TUOS prices column"),
                             ("JSA", "JS prices", "JS prices column (jurisdictional schemes excluding LFiT)"),
                             ("NUoS", "NUOS prices", "NUOS prices column = DUOS + TUOS + JS; excludes metering")):
            rows_out.append(make_row("Evoenergy", fin_year, side, path, url, r["code"], r["name"], r["section"] or "", label, r["unit"],
                                     f'{r["nums"][key]:.3f}', "excl", b, "; ".join(notes + [note])))
    if not rows_out:
        raise RuntimeError("Evoenergy proposal OCR produced no validated price rows")
    return rows_out


# ----------------------------------------------------------------------------------------------
# Evoenergy - 2024-25 schedule of charges PDF (per-tariff blocks)
# ----------------------------------------------------------------------------------------------
def parse_evo_pdf(path, fin_year, side, url, lfit_note):
    rows_out = []
    adder = None
    with pdfplumber.open(path) as pdf:
        for p in pdf.pages[:4]:
            m = re.search(r"equivalent to an additional ([\d.]+) cents per kilowatt-hour", (p.extract_text() or "").replace("\n", " "))
            if m:
                adder = m.group(1)
                break
        cls = ""
        hdr = None
        pend_meter = None  # ('Metering' x1, 'Rate +' x1) - this header line precedes the 'Component' line
        code = name = None
        closed = False
        tariff_note = ""
        block = []  # lines in current block: (top, words); tops are offset per page so blocks can span pages
        price_lines = []
        pno = 0
        for pno, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            mcls = re.search(r"Network Charges \(excluding GST\): (.+)", text)
            if mcls:
                cls = clean(mcls.group(1))
            if not re.search(r"c/kWh|c/day", text):
                continue
            if re.search(r"Schedule of (Metering|Fee-based)", text):
                break
            words = page.extract_words()
            lines = [(top + pno * 10000.0, lw) for top, lw in _lines(words, tol=2.5)]

            def flush():
                nonlocal block, price_lines
                if not price_lines or hdr is None:
                    block, price_lines = [], []
                    return
                # assign non-price lines to the nearest price line (by vertical distance)
                comp_words = {id(pl): [] for pl in price_lines}
                app_words = {id(pl): [] for pl in price_lines}
                for top, lw in block:
                    pl = min(price_lines, key=lambda pl: abs(pl[0] - top))
                    for w in lw:
                        if w["x1"] <= hdr["app"] - 2:
                            comp_words[id(pl)].append(w)
                        elif w["x1"] <= hdr["unit"] - 2:
                            app_words[id(pl)].append(w)
                for pl in price_lines:
                    cw = sorted(comp_words[id(pl)], key=lambda w: (round(w["top"]), w["x0"]))
                    aw = sorted(app_words[id(pl)], key=lambda w: (round(w["top"]), w["x0"]))
                    comp = clean(" ".join(w["text"] for w in cw))
                    app = clean(" ".join(w["text"] for w in aw))
                    uw = next(w for w in pl[1] if EVO_UNIT.match(w["text"]))
                    nums = [w for w in pl[1] if w["x0"] > uw["x1"] and re.match(r"^-?\d+(\.\d+)?$", w["text"])]
                    rate = None
                    for w in nums:
                        col = min(("rate", "meter", "total"), key=lambda k: abs(hdr[k] - w["x1"]))
                        if col == "rate":
                            rate = w["text"]
                    if rate is None:
                        continue
                    notes = [lfit_note + (f" (uniform adder of {adder} c/kWh applied to c/kWh consumption charges per document)" if adder else ""),
                             "Rate column (network charge excluding metering; 'Rate + metering' column not emitted)", f"page {pno}"]
                    if closed:
                        notes.append("closed to new customers")
                    if "XMC" in (name or ""):
                        notes.append("XMC tariff excludes metering charges")
                    if tariff_note:
                        notes.append(tariff_note)
                    if app:
                        notes.append("applicability: " + app)
                    rows_out.append(make_row("Evoenergy", fin_year, side, path, url, code, name, cls, comp, uw["text"], rate,
                                             "excl", "NUoS", "; ".join(notes)))
                block, price_lines = [], []

            for top, lw in lines:
                txt = " ".join(w["text"] for w in lw)
                if re.match(r"^\d+ \| ", txt) or re.search(r"Network Charges \(excluding GST\)|^Rates apply from", txt):
                    continue
                m = re.match(r"^(\d{3}\*?)\s+(.+)$", txt)
                if m and lw[0]["x0"] < 95 and not any(EVO_UNIT.match(w["text"]) for w in lw):
                    flush()
                    code, name = m.group(1), m.group(2)
                    closed = False
                    tariff_note = ""
                    continue
                if txt.startswith("This tariff is closed"):
                    closed = True
                    continue
                if hdr is not None and lw[0]["x0"] < hdr["comp"] - 5:
                    # tariff note lines sit left of the Component column (e.g. 'This tariff is a secondary tariff ...')
                    if txt.startswith("This tariff"):
                        tariff_note = clean(tariff_note + " " + txt)[:160]
                    continue
                if re.match(r"^Metering Rate \+$", txt):
                    pend_meter = (lw[0]["x1"], lw[-1]["x1"])
                    continue
                if txt.startswith("Component") and "Unit" in txt:
                    xs = {w["text"]: w for w in lw}
                    hdr = {"comp": xs["Component"]["x0"], "app": xs["Charge"]["x0"], "unit": xs["Unit"]["x0"],
                           "rate": xs["Rate"]["x1"]}
                    if pend_meter is None:
                        raise RuntimeError(f"page {pno}: 'Metering / Rate +' header line not found before 'Component' line")
                    hdr["meter"], hdr["total"] = pend_meter
                    continue
                if txt in ("charge metering", "charge", "metering") or txt.startswith("*XMC") or txt.startswith("* XMC"):
                    continue
                if hdr is None or code is None:
                    continue
                if any(EVO_UNIT.match(w["text"]) for w in lw):
                    price_lines.append((top, lw))
                block.append((top, lw))
        flush()
    return rows_out


# ----------------------------------------------------------------------------------------------
# Evoenergy - 2025-26 / 2026-27 schedule of charges xlsx
# ----------------------------------------------------------------------------------------------
def parse_evo_xlsx(path, fin_year, side, url, lfit_note):
    rows_out = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Network tariffs"]
    adder = None
    gst = "excl"
    section = ""
    code = name = None
    closed = False
    tariff_note = ""
    hdr = None
    for cell_row in ws.iter_rows():
        raw = [c.value for c in cell_row]
        cells = {j: v for j, v in enumerate(raw, 1) if v is not None and str(v).strip() != ""}
        if not cells:
            continue
        texts = {j: clean(v) for j, v in cells.items() if isinstance(v, str)}
        c2 = texts.get(2, "")
        code_numeric = False
        if not c2 and isinstance(cells.get(2), (int, float)) and float(cells[2]).is_integer() and 0 < cells[2] < 1000:
            c2, code_numeric = f"{int(cells[2]):03d}", True  # code cell stored as a number in some blocks
        if "Prices exclude GST" in " ".join(texts.values()):
            gst = "excl"
        m = re.search(r"equivalent to an additional ([\d.]+) cents per kilowatt-hour", c2)
        if m:
            adder = m.group(1)
            continue
        if c2 and EVO_CODE.match(c2):
            code, name = c2, texts.get(3, "")
            closed = False
            tariff_note = "tariff code cell is stored as a number in the spreadsheet (read as %s)" % c2 if code_numeric else ""
            hdr = None
            continue
        if c2 and re.match(r"^(Residential|LV commercial|HV commercial) tariffs$", c2, re.I):
            section = c2
            continue
        if c2 and code and hdr is None:
            if c2.startswith("This tariff is closed"):
                closed = True
            elif len(c2) > 30:
                tariff_note = c2[:160]
            continue
        if texts.get(3) == "Component":
            hdr = {j: t for j, t in texts.items()}
            continue
        if hdr is None or code is None or 5 not in texts:
            continue
        comp = texts.get(3, "")
        app = texts.get(4, "")
        unit = texts.get(5, "")
        rate_col = next((j for j, t in hdr.items() if t == "Rate"), 6)
        rate = to_num(cells.get(rate_col))
        if rate is None:
            continue
        notes = [lfit_note + (f" (uniform adder of {adder} c/kWh applied to c/kWh consumption charges per document)" if adder else ""),
                 "Rate column (network charge excluding metering; 'Rate + metering' column not emitted)", "sheet 'Network tariffs'"]
        if closed:
            notes.append("closed to new customers")
        if "XMC" in name:
            notes.append("XMC tariff excludes metering charges")
        if tariff_note:
            notes.append(tariff_note)
        if app:
            notes.append("applicability: " + app.replace("\n", " "))
        rows_out.append(make_row("Evoenergy", fin_year, side, path, url, code, name, section, comp, unit, cell_value(cell_row[rate_col - 1]), gst, "NUoS", "; ".join(notes)))
    names_per_code = {}
    for r in rows_out:
        names_per_code.setdefault(r["tariff_code"], set()).add(r["tariff_name"])
    for r in rows_out:
        if len(names_per_code[r["tariff_code"]]) > 1:
            r["note"] += f"; code {r['tariff_code']} is used by {len(names_per_code[r['tariff_code']])} tariff blocks in this schedule (names differ)"
    return rows_out


# ----------------------------------------------------------------------------------------------
def main():
    inv = load_inventory()
    all_rows = []
    statement_rows = []
    for fin_year, side, path, kind, extra in FILES:
        url = inv.get(path, "")
        if not url:
            print(f"WARN: no inventory url for {path}", file=sys.stderr)
        if not os.path.exists(path):
            if kind == "evo_proposal":
                raise FileNotFoundError(path)
            print(f"WARN: missing file {path}", file=sys.stderr)
            continue
        if kind == "ess_xlsx":
            rows = parse_essential_xlsx(path, fin_year, side, url)
        elif kind == "ess_pdf":
            rows = parse_essential_pdf(path, fin_year, side, url)
        elif kind == "evo_statement":
            rows = parse_evo_statement(path, fin_year, side, url)
            statement_rows = rows
        elif kind == "evo_proposal":
            # needs the Statement rows for cross-validation -> parsed after the loop
            continue
        elif kind == "evo_pdf":
            rows = parse_evo_pdf(path, fin_year, side, url, extra["lfit"])
        elif kind == "evo_xlsx":
            rows = parse_evo_xlsx(path, fin_year, side, url, extra["lfit"])
        else:
            rows = []
        print(f"{path}: {len(rows)} rows, {len({r['tariff_code'] for r in rows})} codes")
        all_rows.extend(rows)
    for fin_year, side, path, kind, extra in FILES:
        if kind == "evo_proposal" and os.path.exists(path):
            rows = parse_evo_proposal(path, fin_year, side, inv.get(path, ""), statement_rows)
            print(f"{path}: {len(rows)} rows, {len({r['tariff_code'] for r in rows})} codes")
            all_rows.extend(rows)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=schema.COLUMNS)
        w.writeheader()
        for r in all_rows:
            w.writerow({k: r.get(k, "") for k in schema.COLUMNS})
    print(f"wrote {OUT}: {len(all_rows)} rows")


if __name__ == "__main__":
    main()
