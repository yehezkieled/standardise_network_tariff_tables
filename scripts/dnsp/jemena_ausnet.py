#!/usr/bin/env python
"""Parse Jemena and AusNet Services network tariff schedules into the normalised long CSV.

Run from the repo root:  .venv/bin/python scripts/dnsp/jemena_ausnet.py
Writes out/dnsp/jemena_ausnet.csv (columns = schema.COLUMNS).

Sources handled
- Jemena "Network Tariff Schedule" PDFs (2023-24 .. 2026-27). Vertical list layout: class heading,
  "CODE Name" line, description lines, then "- component  unit  value" lines. The 2023-24 edition
  repeats the whole list four times (Network = NUoS, Distribution = DUoS, Transmission = TUoS,
  Jurisdictional = JSA); later editions publish the Network (NUoS) list only. Footnote letters are
  small-font superscripts; "+", "CR", "MS_CR" etc. are small-font parts of tariff names and are kept.
- AusNet "Schedule of tariffs" 2023-24 PDF (AER hosted) and the identical appendix (10.1-10.4) of the
  2023-24 Annual Pricing Proposal: landscape wide tables with one column per charging parameter and
  four blocks (Network/Distribution/Transmission/Jurisdictional scheme). Parsed from word positions
  anchored on the drawn horizontal rules (row rules start at the code column; class-group rules span
  the full width).
- AusNet Annual Pricing Proposal "Tariff trials proposed tariffs" tables (Table 4.7 2023-24, Table 4.8
  2024-25): the only priced tariff codes in the 2024-25 proposal PDF (its schedule of tariffs is a
  separate attachment that is not in the inventory).
- AusNet xlsx schedules (2024-25 AER hosted, 2025-26, 2026-27): sheets "Network/Distribution/
  Transmission/Jurisdictional <year>" -> NUoS/DUoS/TUoS/JSA. "Indicative prices" and
  "Tariff structures" sheets are skipped.
"""
import csv
import os
import re
import sys
from collections import OrderedDict

sys.path.insert(0, "scripts")
import schema  # noqa: E402
import units  # noqa: E402
from tariffdb import locators  # noqa: E402

import openpyxl  # noqa: E402
import pdfplumber  # noqa: E402

OUT = "out/dnsp/jemena_ausnet.csv"
INVENTORY = "sources/inventory.csv"

JEMENA_FILES = [
    # (fin_year, path). Byte-identical pairs are parsed once and emitted for both copies.
    ("2023-24", "sources/aer/2023-24_price_lists/Jemena_2023-24_Network_Tariff_Schedule_31Mar2023.pdf"),
    ("2023-24", "sources/dnsp/jemena/Jemena_Network_Tariff_Schedule_2023-24.pdf"),
    ("2024-25", "sources/aer/dnsp_copies/Jemena_2024-25_Network_Tariff_Schedule_28Mar2024.pdf"),
    ("2024-25", "sources/dnsp/jemena/Jemena_Network_Tariff_Schedule_2024-25.pdf"),
    ("2025-26", "sources/dnsp/jemena/Jemena_Network_Tariff_Schedule_2025-26.pdf"),
    ("2026-27", "sources/dnsp/jemena/Jemena_Network_Tariff_Schedule_2026-27.pdf"),
]
AUSNET_PDF_SCHEDULES = [
    ("2023-24", "sources/aer/2023-24_price_lists/AusNet_2023-24_Schedule_of_tariffs_31Mar2023.pdf"),
    ("2023-24", "sources/dnsp/ausnet/AusNet_Annual_Pricing_Proposal_2023-24.pdf"),
    ("2024-25", "sources/dnsp/ausnet/AusNet_Annual_Pricing_Proposal_2024-25.pdf"),
]
AUSNET_XLSX = [
    ("2024-25", "sources/aer/dnsp_copies/AusNet_2024-25_Schedule_of_tariffs_28Mar2024.xlsx"),
    ("2025-26", "sources/dnsp/ausnet/AusNet_Network_Tariff_Schedule_2025-26.xlsx"),
    ("2026-27", "sources/dnsp/ausnet/AusNet_Network_Tariff_Schedule_2026-27_20260515.xlsx"),
]

NUM_RE = re.compile(r"^-?\d[\d,]*(\.\d+)?$")
BASIS_WORDS = (("jurisdictional", "JSA"), ("transmission", "TUoS"), ("distribution", "DUoS"), ("network", "NUoS"))


# --------------------------------------------------------------------------------------- helpers
def load_inventory():
    inv = {}
    with open(INVENTORY, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            inv[r["local_path"].strip()] = r["source_url"].strip()
    return inv


def side_for(path):
    return "DNSP" if path.startswith("sources/dnsp/") else "AER_HOSTED"


def fmt_num(v):
    """Numeric value as published -> compact string (no thousands separators, no float noise)."""
    if isinstance(v, str):
        s = v.replace(",", "").replace("$", "").strip()
        float(s)  # validate
        return s
    return str(v)


def basis_from_title(text):
    t = (text or "").lower()
    for word, basis in BASIS_WORDS:
        if word in t:
            return basis
    return None


def make_row(*, distributor, fin_year, code, name, customer_class, component, unit, value, gst, basis,
             source_file, source_url, note, locator, charge_type=None, time_band=None, season=None):
    label = component
    ct = charge_type or schema.charge_type_from_label(label, unit)
    tb = schema.time_band_from_label(label) if time_band is None else time_band
    sn = schema.season_from_label(label) if season is None else season
    vstr = fmt_num(value)
    vstd, ustd = units.to_std(vstr, unit, label)
    return OrderedDict([
        ("side", side_for(source_file)), ("distributor", distributor), ("fin_year", fin_year),
        ("tariff_code", code.strip()), ("tariff_name", name.strip()), ("customer_class", customer_class.strip()),
        ("component", label.strip()), ("charge_type", ct), ("time_band", tb), ("season", sn),
        ("unit", unit.strip()), ("value", vstr),
        ("value_std", "" if vstd is None else fmt_num(vstd)), ("unit_std", ustd),
        ("gst", gst), ("basis", basis), ("source_file", source_file), ("source_url", source_url),
        ("note", note.strip("; ").strip()), ("locator", locator),
    ])


def cluster_lines(items, tol=3.0, key="top"):
    """Group word-like dicts into lines by vertical position (sorted by top, then x0)."""
    lines = []
    for it in sorted(items, key=lambda w: (w[key], w["x0"])):
        if lines and abs(it[key] - lines[-1][0][key]) <= tol:
            lines[-1].append(it)
        else:
            lines.append([it])
    return [sorted(l, key=lambda w: w["x0"]) for l in lines]


SYMBOL_GLYPHS = {"": "≤", "": "≥"}  # Symbol-font private-use code points used in Jemena names


def join_tokens(tokens, glue_gap=1.5):
    """Join tokens left-to-right. pdfplumber words are >= 3pt apart, so only fragments split off a word
    (super/subscripts) can be closer than glue_gap; glue them only when they continue a number ("2.2+"),
    otherwise separate with a space ("LV CR", "Subtransmission MA CR")."""
    out = ""
    prev = None
    for t in tokens:
        txt = "".join(SYMBOL_GLYPHS.get(ch, ch) for ch in t["text"])
        if prev is None:
            out = txt
        elif t["x0"] - prev["x1"] < glue_gap and (txt == "+" or prev["text"][-1:].isdigit()):
            out += txt
        else:
            out += " " + txt
        prev = t
    return out


# ------------------------------------------------------------------------------------------ Jemena
JEM_CODE_RE = re.compile(r"^[AF]\d{2}[0-9A-Z]$")
JEM_SMALL = 6.0  # font size below which a glyph is a super/subscript (body 7.0-7.6, super 4.7)


def jemena_page_lines(page):
    """Lines of tokens with footnote superscripts removed and other small-font fragments kept.

    pdfplumber merges a superscript footnote letter into its host word ("F100a", "Chargef"). We split
    each word into same-size runs, drop single lowercase-letter small runs (footnote markers) and
    re-attach the remaining small runs ("+", "CR", "MS_CR", "EN", "MS") to the nearest body line.
    """
    words = page.extract_words(extra_attrs=["size", "fontname"], return_chars=True)
    big, small = [], []
    for w in words:
        runs = []
        for ch in w["chars"]:
            is_small = ch["size"] < JEM_SMALL
            if runs and runs[-1]["small"] == is_small:
                r = runs[-1]
                r["text"] += ch["text"]
                r["x1"] = ch["x1"]
            else:
                runs.append({"text": ch["text"], "x0": ch["x0"], "x1": ch["x1"], "top": ch["top"],
                             "bottom": ch["bottom"], "size": ch["size"], "fontname": ch["fontname"],
                             "small": is_small})
        for r in runs:
            if r["text"].strip() == "":
                continue
            (small if r["small"] else big).append(r)
    lines = cluster_lines(big, tol=3.0)
    for s in small:
        if re.fullmatch(r"[a-h]", s["text"]):
            continue  # footnote marker
        best, bestd = None, 99
        for ln in lines:
            d = min(abs(s["top"] - t["top"]) for t in ln)
            if d < bestd:
                best, bestd = ln, d
        if best is not None and bestd <= 6:
            best.append(s)
            best.sort(key=lambda t: t["x0"])
    return lines


def parse_jemena_pdf(path, fin_year):
    """Return rows without source columns (added per copy by the caller).

    Font names are not reliable across editions (2024-25 embeds fonts as F1/F3/F4), so detection uses
    geometry and relative font sizes only: codes sit in the code column, class headings sit left of it
    in a larger font than the body (code) size, component lines start with "-", and a wrapped tariff
    name is a line in the same font as the code line before any description/component line.
    """
    rows = []
    gaps = []
    state = {"tariff": None}

    def flush():
        tariff = state["tariff"]
        if not tariff:
            return
        for code in tariff["codes"]:
            paired = ""
            if len(tariff["codes"]) > 1:
                paired = ("published as paired code '" + " / ".join(tariff["codes"]) + "' with one price set; "
                          "footnote a: a tariff code starting with 'F' indicates the tariff attracts the "
                          "Premium Feed-In-Tariff rebate")
            for label, unit, value, cnote, comp_page in tariff["comps"]:
                note = "; ".join(x for x in (f"p{tariff['page']}", tariff["min_demand"], paired, cnote, tariff["gst_note"]) if x)
                rows.append(dict(code=code, name=tariff["name"], customer_class=tariff["class"],
                                 component=label, unit=unit, value=value, gst=tariff["gst"], basis=tariff["basis"],
                                 note=note, fin_year=fin_year, locator=locators.pdf(comp_page)))
            if not tariff["comps"]:
                gaps.append(f"p{tariff['page']} {code}: no priced components found")
        state["tariff"] = None

    with pdfplumber.open(path) as pdf:
        basis = None
        gst, gst_note = "excl", "GST not stated; assumed excl"
        current_class = ""
        for pno, page in enumerate(pdf.pages, start=1):
            lines = jemena_page_lines(page)
            code_x = code_font = body_size = None
            for ln in lines:
                if JEM_CODE_RE.match(ln[0]["text"]):
                    code_x, code_font, body_size = ln[0]["x0"], ln[0]["fontname"], ln[0]["size"]
                    break
            if code_x is None:
                gaps.append(f"p{pno}: no tariff code lines found; page skipped")
                continue
            prev_was_class = False
            for ln in lines:
                first = ln[0]
                txt = join_tokens(ln)
                if "Jemena Electricity Networks" in txt and " - " in txt:
                    b = basis_from_title(txt.split(" - ", 1)[1])
                    if b:
                        basis = b
                    continue
                if txt.startswith("Effective"):
                    if "exclusive of gst" in txt.lower():
                        gst, gst_note = "excl", ""
                    elif "inclusive of gst" in txt.lower():
                        gst, gst_note = "incl", ""
                    continue
                if txt.startswith("Tariff Class") or re.match(r"^Page \d+ of \d+$", txt):
                    continue
                if first["x0"] < code_x - 10:
                    # class column: heading (larger than body font) or class description / footnotes
                    if body_size + 0.3 < first["size"] < 9.6:
                        flush()
                        current_class = (current_class + " " + txt).strip() if prev_was_class else txt
                        prev_was_class = True
                    continue
                prev_was_class = False
                # tariff header: CODE [/ FCODE] Name
                if abs(first["x0"] - code_x) < 6 and JEM_CODE_RE.match(first["text"]):
                    flush()
                    codes, name_toks = [], []
                    for t in ln:
                        if not name_toks and (JEM_CODE_RE.match(t["text"]) or t["text"] == "/"):
                            if t["text"] != "/":
                                codes.append(t["text"])
                        else:
                            name_toks.append(t)
                    state["tariff"] = {"codes": codes, "name": join_tokens(name_toks), "comps": [], "min_demand": "",
                                       "class": current_class, "page": pno, "header_open": True, "basis": basis,
                                       "gst": gst, "gst_note": gst_note}
                    continue
                tariff = state["tariff"]
                if tariff is None:
                    continue
                if first["text"] == "-" and first["x0"] > code_x + 40:
                    tariff["header_open"] = False
                    body = ln[1:]
                    unit_idx = next((i for i, t in enumerate(body) if t["text"][:1] in ("$", "¢") or t["text"].startswith("c/")), None)
                    if unit_idx is None:
                        gaps.append(f"p{pno} {'/'.join(tariff['codes'])}: component line without unit: '{txt}'")
                        continue
                    label_toks, rest = body[:unit_idx], body[unit_idx:]
                    value_tok = None
                    if len(rest) >= 2 and NUM_RE.match(rest[-1]["text"]) and rest[-1]["x0"] > rest[0]["x0"] + 30:
                        value_tok, rest = rest[-1], rest[:-1]
                    label, unit = join_tokens(label_toks), join_tokens(rest)
                    cnote = ""
                    if re.search(r"Chargef$", label):  # footnote marker typeset at body size (JSA p38)
                        label, cnote = label[:-1], "footnote marker 'f' typeset full-size in source label"
                    if value_tok is None:
                        gaps.append(f"p{pno} {'/'.join(tariff['codes'])}: component '{label}' unit='{unit}' has no value")
                        continue
                    tariff["comps"].append((label, unit, value_tok["text"], cnote, pno))
                    continue
                if txt.startswith("Minimum Chargeable"):
                    tariff["header_open"] = False
                    tariff["min_demand"] = txt
                    continue
                if tariff["header_open"] and first["fontname"] == code_font and first["x0"] > code_x + 40:
                    tariff["name"] = (tariff["name"] + " " + txt).strip()  # wrapped tariff name
                    continue
                tariff["header_open"] = False  # description line
        flush()
    return rows, gaps


def emit_jemena(inv):
    out, report = [], []
    cache = {}
    for fin_year, path in JEMENA_FILES:
        with open(path, "rb") as fh:
            import hashlib
            digest = hashlib.md5(fh.read()).hexdigest()
        if digest not in cache:
            cache[digest] = parse_jemena_pdf(path, fin_year)
        rows, gaps = cache[digest]
        n = 0
        for r in rows:
            out.append(make_row(distributor="Jemena", fin_year=fin_year, code=r["code"], name=r["name"],
                                customer_class=r["customer_class"], component=r["component"], unit=r["unit"],
                                value=r["value"], gst=r["gst"], basis=r["basis"], source_file=path,
                                source_url=inv[path], note=r["note"], locator=r["locator"]))
            n += 1
        codes = sorted({r["code"] for r in rows})
        report.append(f"Jemena {fin_year} {path}: rows={n} codes={len(codes)} bases={sorted({r['basis'] for r in rows})} gaps={gaps}")
    return out, report


# ------------------------------------------------------------------------------------ AusNet PDFs
AUS_CODE_RE = re.compile(r"^[A-Z]{3,4}\d{2}[A-Z]?$")
AUS_TRIAL_CODE_RE = re.compile(r"^[A-Z]{3,4}\d{2}T$")
UNIT_RE = re.compile(r"^(\$|c)/\S+$")  # "$/year", "c/kWh", "$/kVA/yr", "$/kW/mth", wrapped "$/kVA/ye"


def unit_columns(unit_words, all_words):
    """Columns from the unit header words; re-attach a wrapped tail fragment printed just below a unit
    ("$/kVA/ye" + "ar" -> "$/kVA/year"). Returns (cols, consumed_fragment_words)."""
    cols, consumed = [], []
    for u in unit_words:
        text = u["text"]
        for w in all_words:
            if (0 < w["top"] - u["top"] <= 10 and w["x0"] < u["x1"] and w["x1"] > u["x0"]
                    and re.fullmatch(r"[a-z]{1,3}", w["text"])):
                text += w["text"]
                consumed.append(w)
        cols.append({"x0": u["x0"], "x1": u["x1"], "unit": text})
    return cols, consumed


def h_rules(page):
    """Sorted unique tops of thin horizontal drawn segments -> list of (top, x0, x1)."""
    segs = []
    for c in page.curves + page.lines + page.rects:
        if (c["bottom"] - c["top"]) < 1.5 and (c["x1"] - c["x0"]) > 10:
            segs.append((c["top"], c["x0"], c["x1"]))
    segs.sort()
    out = []
    for s in segs:
        if out and abs(s[0] - out[-1][0]) < 1.0:
            out[-1] = (out[-1][0], min(out[-1][1], s[1]), max(out[-1][2], s[2]))
        else:
            out.append(s)
    return out


def assign_columns(words, cols):
    """Map words to column index by horizontal centre (containment first, else nearest centre)."""
    res = {}
    for w in words:
        cx = (w["x0"] + w["x1"]) / 2
        idx = None
        for i, c in enumerate(cols):
            if c["x0"] - 4 <= cx <= c["x1"] + 4:
                idx = i
                break
        if idx is None:
            idx = min(range(len(cols)), key=lambda i: abs(cx - (cols[i]["x0"] + cols[i]["x1"]) / 2))
        res.setdefault(idx, []).append(w)
    return res


def header_labels(hwords, cols):
    labels = {}
    for i, ws in assign_columns(hwords, cols).items():
        txt = " ".join(w["text"] for w in sorted(ws, key=lambda w: (round(w["top"]), w["x0"])))
        txt = re.sub(r"\s+", " ", txt).replace("Dedicate d", "Dedicated").strip()
        labels[i] = txt
    return labels


def parse_ausnet_schedule_page(page, pno):
    """Parse one wide-table schedule page. Returns (title_or_None, records, notes_text) or None."""
    words = page.extract_words()
    unit_lines = [l for l in cluster_lines([w for w in words if UNIT_RE.match(w["text"])], tol=2.0) if len(l) >= 10]
    if not unit_lines:
        return None
    unit_line = max(unit_lines, key=len)
    unit_top = unit_line[0]["top"]
    rules = h_rules(page)
    full = [r for r in rules if r[1] < 60]
    tops_above = [r[0] for r in full if r[0] < unit_top]
    below = [r[0] for r in full if r[0] > unit_top]
    if not tops_above or not below:
        return None
    table_top, header_bottom = tops_above[-1], below[0]
    cols, consumed = unit_columns(unit_line, words)
    words = [w for w in words if w not in consumed]
    first_price_x0 = cols[0]["x0"] - 6
    hdr_words = [w for w in words if table_top <= w["top"] < header_bottom and w["top"] < unit_top - 1]
    # non-price column anchors from header words
    def hx(name):
        c = [w["x0"] for w in hdr_words if w["text"].lower().startswith(name)]
        return min(c) if c else None
    desc_x0 = hx("description")
    closed_x0 = hx("closed")
    struct_x0 = hx("structure")
    code_x0 = hx("code")
    if None in (desc_x0, closed_x0, struct_x0, code_x0):
        return None
    labels = header_labels([w for w in hdr_words if w["x0"] >= first_price_x0], cols)
    title_words = [w for w in words if w["top"] < table_top - 2]
    title = " ".join(w["text"] for w in sorted(title_words, key=lambda w: (round(w["top"]), w["x0"])))

    row_edges = [r[0] for r in rules if r[0] >= header_bottom - 0.5]
    group_edges = [r[0] for r in full if r[0] >= header_bottom - 0.5]
    body_words = [w for w in words if w["top"] >= header_bottom - 1]

    def centre(w):
        return (w["top"] + w["bottom"]) / 2

    # class label per group band
    groups = []
    for a, b in zip(group_edges, group_edges[1:]):
        cw = [w for w in body_words if w["x1"] <= code_x0 - 2 and a <= centre(w) < b]
        label = " ".join(w["text"] for w in sorted(cw, key=lambda w: (round(w["top"]), w["x0"])))
        groups.append((a, b, re.sub(r"\s+", " ", label).strip()))

    records = []
    notes = []
    for a, b in zip(row_edges, row_edges[1:]):
        rw = [w for w in body_words if a <= centre(w) < b]
        codes = [w for w in rw if AUS_CODE_RE.match(w["text"]) and w["x0"] < struct_x0 - 3 and w["x1"] > code_x0 - 5]
        if not codes:
            txt = " ".join(w["text"] for w in sorted(rw, key=lambda w: (round(w["top"]), w["x0"])))
            if txt:
                notes.append(txt)
            continue
        code = codes[0]
        cls = ""
        for ga, gb, label in groups:
            if ga - 0.5 <= a and b <= gb + 0.5:
                cls = label
        def col_text(lo, hi):
            ws = [w for w in rw if lo <= w["x0"] < hi]
            return re.sub(r"\s+", " ", " ".join(w["text"] for w in sorted(ws, key=lambda w: (round(w["top"]), w["x0"])))).strip()
        structure = col_text(struct_x0 - 3, desc_x0 - 3)
        desc = col_text(desc_x0 - 3, closed_x0 - 3)
        closed = col_text(closed_x0 - 3, first_price_x0)
        vals = []
        price_ws = [w for w in rw if w["x1"] > cols[0]["x0"] - 2]  # x1-based: wide standing charges start left of the unit
        numeric = [w for w in price_ws if NUM_RE.match(w["text"])]
        stray = [w["text"] for w in price_ws if not NUM_RE.match(w["text"])]
        for i, ws in assign_columns(numeric, cols).items():
            for w in ws:
                vals.append((labels.get(i, f"col{i}"), cols[i]["unit"], w["text"]))
        records.append({"code": code["text"], "desc": desc, "structure": structure, "closed": closed, "class": cls,
                        "vals": vals, "stray": stray, "page": pno})
    tail = [w for w in words if w["top"] > group_edges[-1]] if group_edges else []
    notes.append(" ".join(w["text"] for w in sorted(tail, key=lambda w: (round(w["top"]), w["x0"]))))
    return title, records, " ".join(notes)


def parse_ausnet_schedule_pdf(path, fin_year, inv):
    rows, gaps = [], []
    seen_pages = []
    with pdfplumber.open(path) as pdf:
        basis = None
        gst_seen = False
        pending = []  # records waiting for gst statement (collected per file; gst note decided at end)
        for pno, page in enumerate(pdf.pages, start=1):
            res = parse_ausnet_schedule_page(page, pno)
            if res is None:
                continue
            title, records, notes = res
            b = basis_from_title(title) if "tariff schedule" in title.lower() else None
            if b:
                basis = b
            if basis is None:
                gaps.append(f"p{pno}: schedule table without a section title; skipped")
                continue
            if re.search(r"ex\.?\s*gst|excl\w* of gst|exclusive of gst", notes, re.I):
                gst_seen = True
            seen_pages.append((pno, basis, len(records)))
            for rec in records:
                if rec["stray"]:
                    gaps.append(f"p{pno} {rec['code']}: non-numeric price cells {rec['stray']}")
                if not rec["vals"]:
                    gaps.append(f"p{pno} {rec['code']}: no numeric values")
                pending.append((pno, basis, rec))
    gst_note = "" if gst_seen else "GST not stated in document; assumed excl"
    for pno, basis, rec in pending:
        note_bits = [f"p{pno}", f"tariff structure {rec['structure']}" if rec["structure"] else ""]
        if rec["closed"].lower().startswith("yes"):
            note_bits.append("closed to new entrants")
        if "*" in rec["desc"]:
            note_bits.append("description footnote '*' (see page notes)")
        if gst_note:
            note_bits.append(gst_note)
        for label, unit, val in rec["vals"]:
            rows.append(make_row(distributor="AusNet Services", fin_year=fin_year, code=rec["code"], name=rec["desc"],
                                 customer_class=rec["class"], component=label, unit=unit, value=val, gst="excl",
                                 basis=basis, source_file=path, source_url=inv[path], note="; ".join(x for x in note_bits if x),
                                 locator=locators.pdf(pno)))
    return rows, gaps, seen_pages


def parse_ausnet_trial_tables(path, fin_year, inv):
    """'Table 4.x: Tariff trials proposed tariffs for <year>' tables in the Annual Pricing Proposal."""
    rows, gaps, found = [], [], []
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            m = re.search(r"Table [\d.]+: Tariff trials proposed tariffs for (\S+)", text)
            if not m:
                continue
            table_title = m.group(0)
            words = page.extract_words()
            lines = cluster_lines(words, tol=2.0)
            # locate "Tariff trial" header starts and unit lines
            starts = [i for i, l in enumerate(lines) if join_tokens(l).startswith("Tariff trial")]
            for si, start in enumerate(starts):
                end = starts[si + 1] if si + 1 < len(starts) else len(lines)
                block = lines[start:end]
                unit_idx = next((i for i, l in enumerate(block) if sum(1 for w in l if UNIT_RE.match(w["text"])) >= 5), None)
                if unit_idx is None:
                    continue
                unit_line = [w for w in block[unit_idx] if UNIT_RE.match(w["text"])]
                cols, _ = unit_columns(unit_line, [])
                first_price_x0 = cols[0]["x0"] - 6
                hdr_words = [w for l in block[:unit_idx] for w in l if w["x0"] >= first_price_x0]
                labels = header_labels(hdr_words, cols)
                code_lines = [l for l in block[unit_idx + 1:] if any(AUS_TRIAL_CODE_RE.match(w["text"]) for w in l)]
                for l in code_lines:
                    code_w = next(w for w in l if AUS_TRIAL_CODE_RE.match(w["text"]))
                    # trial name: words left of the code within +-10pt vertically (wrapped names)
                    name_ws = [w for l2 in block[unit_idx + 1:] for w in l2
                               if w["x1"] <= code_w["x0"] - 2 and abs(w["top"] - code_w["top"]) <= 10
                               and not AUS_TRIAL_CODE_RE.match(w["text"])]
                    name = " ".join(w["text"] for w in sorted(name_ws, key=lambda w: (round(w["top"]), w["x0"])))
                    numeric = [w for w in l if w["x1"] > cols[0]["x0"] - 2 and NUM_RE.match(w["text"])]
                    if not numeric:
                        gaps.append(f"p{pno} {code_w['text']}: no numeric values")
                    for i, ws in assign_columns(numeric, cols).items():
                        for w in ws:
                            label = labels.get(i, f"col{i}")
                            ct = "other" if "event rebate" in label.lower() else None
                            note = (f"p{pno} {table_title}; tariff trial '{name}'; proposed (pre-approval) trial prices; "
                                    "basis not stated in table (values match the Network/NUoS schedule of the parent tariffs); "
                                    "GST not stated; assumed excl")
                            rows.append(make_row(distributor="AusNet Services", fin_year=fin_year, code=code_w["text"],
                                                 name=f"{name} (tariff trial)", customer_class="Tariff trial",
                                                 component=label, unit=cols[i]["unit"], value=w["text"], gst="excl",
                                                 basis="unknown", source_file=path, source_url=inv[path], note=note,
                                                 locator=locators.pdf(pno), charge_type=ct))
                    found.append((pno, code_w["text"], len(numeric)))
    return rows, gaps, found


# ------------------------------------------------------------------------------------ AusNet xlsx
def fmt_decimals(number_format):
    sec = (number_format or "General").split(";")[0]
    m = re.search(r"\.(0+)", sec)
    return len(m.group(1)) if m else None


def parse_ausnet_xlsx(path, fin_year, inv):
    rows, gaps, sheets_done = [], [], []
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        basis = None
        m = re.match(r"^(Network|Distribution|Transmission|Jurisdictional)\b", ws.title)
        if m:
            basis = basis_from_title(m.group(1))
        if basis is None:
            gaps.append(f"sheet '{ws.title}' skipped (not a current-year price schedule)")
            continue
        # class by merged ranges in column B + plain values
        class_by_row = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_col == 2 and rng.max_col == 2:
                val = ws.cell(rng.min_row, 2).value
                for r in range(rng.min_row, rng.max_row + 1):
                    class_by_row[r] = str(val or "").strip()
        colB = [str(ws.cell(r, 2).value).strip() for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value is not None]
        gst_stated = any(re.search(r"ex\.?\s*gst|excl\w* of gst", s, re.I) for s in colB)
        star_notes = " | ".join(s for s in colB if s.startswith("*"))
        header = None  # list of (col_idx, label, unit)
        current_class = ""
        n_codes = 0
        for r in range(1, ws.max_row + 1):
            b = ws.cell(r, 2).value
            c = ws.cell(r, 3).value
            if isinstance(b, str) and b.strip().lower() == "tariff class":
                header = []
                for cidx in range(7, ws.max_column + 1):
                    lab = ws.cell(r, cidx).value
                    unit = ws.cell(r + 1, cidx).value
                    if lab is not None and str(lab).strip():
                        header.append((cidx, str(lab).strip(), str(unit or "").strip()))
                continue
            if header is None:
                continue
            if r in class_by_row:
                current_class = class_by_row[r]
            elif isinstance(b, str) and b.strip() and not re.match(r"^(\*|\d+\.|notes)", b.strip(), re.I):
                current_class = b.strip()
            code = str(c).strip() if c is not None else ""
            if not AUS_CODE_RE.match(code):
                continue
            n_codes += 1
            desc = str(ws.cell(r, 5).value or "").strip()
            structure = str(ws.cell(r, 4).value or "").strip()
            closed = str(ws.cell(r, 6).value or "").strip()
            note_bits = [f"sheet '{ws.title}' row {r}"]
            if structure:
                note_bits.append(f"tariff structure {structure}")
            if closed.lower().startswith("yes"):
                note_bits.append("closed to new entrants")
            if "*" in desc and star_notes:
                note_bits.append(f"description footnote: {star_notes}")
            if not gst_stated:
                note_bits.append("GST not stated in sheet; assumed excl")
            last_hdr_col = max(h[0] for h in header)
            for cidx in range(last_hdr_col + 1, ws.max_column + 1):
                v = ws.cell(r, cidx).value
                if v is not None and str(v).strip():
                    note_bits.append(f"unlabelled cell {ws.cell(r, cidx).coordinate}='{v}' (predecessor code?)")
            text_cells = []
            comps = []
            for cidx, label, unit in header:
                v = ws.cell(r, cidx).value
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if not isinstance(v, (int, float)):
                    text_cells.append(f"{label}='{v}'")
                    continue
                dec = fmt_decimals(ws.cell(r, cidx).number_format)
                shown = round(v, dec) if dec is not None else v
                extra = ""
                if dec is not None and abs(shown - v) > 1e-9:
                    extra = f"cell stores unrounded {v!r}, displayed as {shown:.{dec}f}"
                clean_label = label.rstrip("^*").strip()
                lab_note = f"header published as '{label}'" if clean_label != label else ""
                comps.append((clean_label, unit, f"{shown:.{dec}f}" if dec is not None else str(shown), "; ".join(x for x in (lab_note, extra) if x),
                              locators.xlsx(ws, ws.cell(r, cidx))))
            if text_cells:
                gaps.append(f"sheet '{ws.title}' {code}: non-numeric price cells {text_cells}")
                note_bits.append("non-numeric cells: " + ", ".join(text_cells) + (" (site-specific)" if any("site" in t.lower() for t in text_cells) else ""))
            if not comps:
                gaps.append(f"sheet '{ws.title}' {code}: no numeric values")
            for label, unit, val, cnote, loc in comps:
                rows.append(make_row(distributor="AusNet Services", fin_year=fin_year, code=code, name=desc,
                                     customer_class=current_class, component=label, unit=unit, value=val, gst="excl",
                                     basis=basis, source_file=path, source_url=inv[path],
                                     note="; ".join(x for x in note_bits + [cnote] if x), locator=loc))
        sheets_done.append((ws.title, basis, n_codes, gst_stated))
    return rows, gaps, sheets_done


# -------------------------------------------------------------------------------------------- main
def main():
    inv = load_inventory()
    all_rows, report = [], []
    jem_rows, jem_report = emit_jemena(inv)
    all_rows += jem_rows
    report += jem_report
    for fin_year, path in AUSNET_PDF_SCHEDULES:
        rows, gaps, pages = parse_ausnet_schedule_pdf(path, fin_year, inv)
        trows, tgaps, found = parse_ausnet_trial_tables(path, fin_year, inv)
        all_rows += rows + trows
        codes = sorted({r["tariff_code"] for r in rows})
        report.append(f"AusNet {fin_year} {path}: schedule rows={len(rows)} codes={len(codes)} pages={pages} gaps={gaps}; "
                      f"trial rows={len(trows)} trial codes={sorted({r['tariff_code'] for r in trows})} trial gaps={tgaps}")
    for fin_year, path in AUSNET_XLSX:
        rows, gaps, sheets = parse_ausnet_xlsx(path, fin_year, inv)
        all_rows += rows
        codes = sorted({r["tariff_code"] for r in rows})
        report.append(f"AusNet {fin_year} {path}: rows={len(rows)} codes={len(codes)} sheets={sheets} gaps={gaps}")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=schema.COLUMNS)
        w.writeheader()
        for r in all_rows:
            assert list(r.keys()) == schema.COLUMNS
            w.writerow(r)
    for line in report:
        print(line)
    print(f"wrote {OUT}: {len(all_rows)} rows")


if __name__ == "__main__":
    main()
