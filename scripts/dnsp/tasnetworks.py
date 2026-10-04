#!/usr/bin/env python
"""TasNetworks network tariff price documents -> out/dnsp/tasnetworks.csv

Files parsed (see FILES):
  2023-24 AER_HOSTED  Network tariff application and price guide (per-tariff tables with DUoS/TUoS/NUoS columns)
  2023-24 DNSP        Indicative network tariff pricing schedule (Table 1 NUoS, Table 2 DUoS, Table 3 TUoS)
  2024-25 DNSP        Pricing schedule PDF (one page each NUoS / DUoS / TUoS + locational TUoS page)
  2025-26 DNSP        Pricing schedule xlsx (three stacked tables NUoS / DUoS / TUoS + 'Locational TUoS' sheet)
  2026-27 DNSP        Pricing schedule xlsx (same layout)

Run from the repo root:  .venv/bin/python scripts/dnsp/tasnetworks.py
"""
import csv
import os
import re
import sys

sys.path.insert(0, "scripts")
import schema  # noqa: E402
import units  # noqa: E402

import openpyxl  # noqa: E402
import pdfplumber  # noqa: E402

DIST = "TasNetworks"
OUT = "out/dnsp/tasnetworks.csv"
INVENTORY = "sources/inventory.csv"

CODE_RE = re.compile(r"^TAS[A-Z0-9]{2,6}$")
NUM_RE = re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$|^-?\d+(\.\d+)?$")
BASIS_WORDS = {"network": "NUoS", "distribution": "DUoS", "transmission": "TUoS"}

# Measure of demand per tariff, from the 2023-24 price guide Table 36 ("Measure of demand (kVA or kW)").
# Used only where a document's demand header lists several units (e.g. "c/kVA, kW, lamp watt/day").
DEMAND_MEASURE_2023_24 = {"TAS87": "kW", "TAS97": "kW", "TAS88": "kW", "TAS98": "kW",
                          "TAS82": "kVA", "TAS89": "kVA", "TASSDM": "kVA", "TAS15": "kVA"}


# ----------------------------------------------------------------------------- helpers
def source_urls():
    with open(INVENTORY, newline="", encoding="utf-8") as f:
        return {r["local_path"]: r["source_url"] for r in csv.DictReader(f)}


def num_text(s):
    """'3,170.000' -> '3170.000' (string as published, commas removed); None if not numeric."""
    s = (s or "").strip()
    if not NUM_RE.match(s):
        return None
    return s.replace(",", "")


def fmt_float(v):
    """Excel float -> compact string without binary noise (round to 10 dp)."""
    s = f"{round(float(v), 10):.10f}".rstrip("0").rstrip(".")
    return s if s not in ("", "-0") else "0"


def squash(s):
    return re.sub(r"\s+", " ", (s or "")).strip()


def strip_markers(name):
    """Remove trailing '(1)'-style footnote markers from a tariff name; return (name, [markers])."""
    marks = re.findall(r"\((\d)\)", name)
    clean = squash(re.sub(r"\s*\(\d\)", "", name))
    return clean, marks


def group_lines(words, tol=2.5):
    """Group pdfplumber words into lines by their 'top' coordinate."""
    lines = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1]["top"] - w["top"]) <= tol:
            lines[-1]["words"].append(w)
        else:
            lines.append({"top": w["top"], "words": [w]})
    for ln in lines:
        ln["words"].sort(key=lambda w: w["x0"])
        ln["text"] = " ".join(w["text"] for w in ln["words"])
    return lines


def is_small(w, thresh=7.5):
    return w.get("size", 9) < thresh


def page_footnotes(words):
    """Footnotes of the form '<small digit(s)> text ...' -> {num: text}."""
    out = {}
    for ln in group_lines(words):
        ws = ln["words"]
        if len(ws) >= 2 and ws[0]["text"].isdigit() and is_small(ws[0]) and not is_small(ws[1]):
            out[ws[0]["text"]] = " ".join(w["text"] for w in ws[1:])
    return out


def paren_footnotes(text):
    """Footnotes of the form '(1) text' -> {'1': text}."""
    out = {}
    for m in re.finditer(r"^\((\d)\)\s+(.+)$", text, re.M):
        out[m.group(1)] = m.group(2).strip()
    return out


class Emitter:
    def __init__(self, fin_year, side, path, url, gst, file_note):
        self.rows = []
        self.ctx = dict(fin_year=fin_year, side=side, path=path, url=url, gst=gst, file_note=file_note)

    def add(self, code, name, cls, component, unit, value, basis, note="", charge_type=None, time_band=None):
        if value is None or value == "":
            return
        v_std, u_std = units.to_std(value, unit, component)
        notes = [n for n in (self.ctx["file_note"], note) if n]
        self.rows.append({
            "side": self.ctx["side"], "distributor": DIST, "fin_year": self.ctx["fin_year"],
            "tariff_code": code.strip(), "tariff_name": squash(name), "customer_class": squash(cls),
            "component": squash(component),
            "charge_type": charge_type or schema.charge_type_from_label(component, unit),
            "time_band": time_band if time_band is not None else schema.time_band_from_label(component),
            "season": schema.season_from_label(component),
            "unit": unit, "value": value, "value_std": "" if v_std is None else v_std, "unit_std": u_std,
            "gst": self.ctx["gst"], "basis": basis,
            "source_file": self.ctx["path"], "source_url": self.ctx["url"], "note": "; ".join(notes),
        })


def add_locational(em, pairs, basis_note, name="kVA specified demand (>2MVA)", cls="Large business HV"):
    """Locational TUoS service charges (per transmission node) -> TUoS rows attached to TAS15."""
    for desc, node, val in pairs:
        em.add("TAS15", name, cls, f"Locational TUoS service charge - {desc} ({node})", "c/kVA/day", val, "TUoS",
               note=f"site-specific: locational TUoS demand charge for transmission node {node}; {basis_note}",
               charge_type="demand", time_band="")


# ----------------------------------------------------------------------------- xlsx 2025-26 / 2026-27
def parse_xlsx(path, fin_year, side, url):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = f"Network tariffs {fin_year}"
    ws = wb[sheet]
    grid = [[c for c in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=30, values_only=True)]
    title = next((squash(str(c)) for row in grid[:3] for c in row if c and "prices for" in str(c).lower()), "")
    em = Emitter(fin_year, side, path, url, "excl",
                 f"sheet '{sheet}' ({title}); GST not stated in workbook, assumed excl")

    r = 0
    while r < len(grid):
        row = grid[r]
        hdr = next((str(c) for c in row[:4] if c and re.search(r"prices for .*use of service", str(c), re.I)), None)
        if not hdr:
            r += 1
            continue
        basis = BASIS_WORDS[re.search(r"for (network|distribution|transmission) use", hdr, re.I).group(1).lower()]
        hdr_label = squash(hdr)
        # column map: group header on this row (merged -> forward-fill), sub-labels on the next row
        sub = grid[r + 1]
        cols = {}
        group = None
        for ci in range(len(row)):
            if row[ci] is not None and str(row[ci]).strip():
                group = squash(str(row[ci]))
            if ci >= 4 and sub[ci] is not None and str(sub[ci]).strip() and group:
                units_in_hdr = re.findall(r"\(([^)]*)\)", group)
                cols[ci] = dict(sub=squash(str(sub[ci])), group=group, unit=units_in_hdr[0] if units_in_hdr else "",
                                multi=len(units_in_hdr) > 1)
        code_col = next(ci for ci in range(len(row)) if row[ci] and str(row[ci]).strip().lower() == "tariff code")
        name_col = next(ci for ci in range(len(row)) if row[ci] and "prices for" in str(row[ci]).lower())
        class_col = next(ci for ci in range(len(row)) if row[ci] and str(row[ci]).strip().lower() == "tariff class")
        # data rows then footnotes
        data, foot = [], {}
        rr = r + 2
        while rr < len(grid):
            drow = grid[rr]
            code = str(drow[code_col]).strip() if drow[code_col] is not None else ""
            b = str(drow[name_col]).strip() if drow[name_col] is not None else ""
            if CODE_RE.match(code):
                data.append(drow)
            elif re.match(r"^\(\d\)", b):
                m = re.match(r"^\((\d)\)\s*(.+)$", b)
                foot[m.group(1)] = m.group(2).strip()
            elif b.startswith("DISCLAIMER") or b.startswith("http") or re.search(r"prices for .*use of service", b, re.I):
                break
            rr += 1
        for drow in data:
            code = str(drow[code_col]).strip()
            name, marks = strip_markers(str(drow[name_col]))
            cls = squash(str(drow[class_col] or ""))
            notes = [f"{hdr_label} table"]
            for mk in marks:
                if mk in foot:
                    notes.append(f"footnote ({mk}): {foot[mk]}" + (" -> obsolete tariff" if "obsolete" in foot[mk].lower() else ""))
            for ci, c in cols.items():
                v = drow[ci] if ci < len(drow) else None
                if v is None or str(v).strip() == "":
                    continue
                try:
                    val = fmt_float(v)
                except ValueError:
                    continue
                n = list(notes)
                if c["multi"]:
                    n.append(f"header unit published as '{c['group']}' (first unit used); measure of demand is kW for "
                             f"TAS87/TAS88/TAS98 and kVA for TAS82/TAS89/TASSDM/TAS15 per 2023-24 price guide Table 36")
                em.add(code, name, cls, c["sub"], c["unit"], val, basis, note="; ".join(n))
        r = rr
    # locational TUoS sheet
    loc = f"Locational TUoS {fin_year}"
    if loc in wb.sheetnames:
        pairs = []
        for row in wb[loc].iter_rows(min_row=1, max_col=6, values_only=True):
            vals = [c for c in row if c is not None and str(c).strip()]
            if len(vals) >= 3 and re.match(r"^T[A-Z]{2}\d$", str(vals[1]).strip()) and isinstance(vals[2], (int, float)):
                pairs.append((squash(str(vals[0])), str(vals[1]).strip(), fmt_float(vals[2])))
        add_locational(em, pairs, f"sheet '{loc}' (applies to TAS15, footnote (2) 'Additional locational TUOS demand charges apply')")
    return em.rows


# ----------------------------------------------------------------------------- 2024-25 schedule PDF
def parse_pdf_2024_25(path, fin_year, side, url):
    em = Emitter(fin_year, side, path, url, "excl", "GST not stated in document, assumed excl")
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            m = re.search(r"prices for (network|distribution|transmission) use of service", text, re.I)
            if not m:
                if "Locational TUoS charges" in text:
                    pairs = [(a, b, num_text(c)) for a, b, c in
                             re.findall(r"^(.+?)\s+(T[A-Z]{2}\d)\s+([\d,.]+)$", text, re.M)]
                    add_locational(em, pairs, "'Locational TUoS charges for 2024-25' page (applies to TAS15, footnote (2))")
                continue
            basis = BASIS_WORDS[m.group(1).lower()]
            title = squash(text.split("\n")[0])
            table = page.find_tables()[0]
            edges = sorted({round(c[0]) for c in table.cells} | {round(c[2]) for c in table.cells})
            trows = table.extract()
            grp_row = next(rw for rw in trows if any(c and "Tariff code" in c for c in rw))
            sub_row = next(rw for rw in trows if any(c and "Service charge" in c for c in rw))
            code_ci = next(i for i, c in enumerate(grp_row) if c and "Tariff code" in c)
            class_ci = next(i for i, c in enumerate(grp_row) if c and "Tariff class" in c)
            name_ci = next(i for i, c in enumerate(grp_row) if c and "prices for" in c)
            cols, group = {}, None
            for ci in range(len(grp_row)):
                if grp_row[ci] and grp_row[ci].strip():
                    group = squash(grp_row[ci])
                if ci > code_ci and sub_row[ci] and sub_row[ci].strip():
                    u = re.findall(r"\(([^)]*)\)", group)
                    cols[ci] = dict(sub=squash(sub_row[ci]), group=group, unit=u[0] if u else "")
            foot = paren_footnotes(text)

            def col_of(w):
                xc = (w["x0"] + w["x1"]) / 2
                for i in range(len(edges) - 1):
                    if edges[i] <= xc < edges[i + 1]:
                        return i
                return None

            for ln in group_lines(page.extract_words()):
                by_col = {}
                for w in ln["words"]:
                    by_col.setdefault(col_of(w), []).append(w["text"])
                code = " ".join(by_col.get(code_ci, []))
                if not CODE_RE.match(code):
                    continue
                name, marks = strip_markers(" ".join(by_col.get(name_ci, [])))
                cls = " ".join(by_col.get(class_ci, []))
                notes = [f"{squash(grp_row[name_ci])} table ({title})"]
                for mk in marks:
                    if mk in foot:
                        notes.append(f"footnote ({mk}): {foot[mk]}" + (" -> obsolete tariff" if "obsolete" in foot[mk].lower() else ""))
                for ci, c in cols.items():
                    val = num_text(" ".join(by_col.get(ci, [])))
                    if val is None:
                        continue
                    em.add(code, name, cls, c["sub"], c["unit"], val, basis, note="; ".join(notes))
    return em.rows


# ----------------------------------------------------------------------------- 2023-24 schedule PDF
def parse_pdf_2023_24_schedule(path, fin_year, side, url):
    em = Emitter(fin_year, side, path, url, "excl",
                 "indicative schedule ('Indicative network tariff pricing schedule for standard control services 2023-24'; "
                 "'This document provides the proposed tariffs for the remainder of the regulatory control period')")
    basis = None
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            m = re.search(r"Table (\d) – Standard control services – 2023-24 network tariff schedule – (NUoS|DUoS|TUoS)", text)
            if m:
                basis, table_no = m.group(2), m.group(1)
            if "Network tariff description" not in text or basis is None:
                continue
            tables = page.find_tables()
            if not tables:
                continue
            table = tables[0]
            edges = sorted({round(c[0]) for c in table.cells} | {round(c[2]) for c in table.cells})
            words = page.extract_words(extra_attrs=["size"])
            foot = page_footnotes(words)

            def col_of(w):
                xc = (w["x0"] + w["x1"]) / 2
                for i in range(len(edges) - 1):
                    if edges[i] <= xc < edges[i + 1]:
                        return i
                return None

            # logical row bands = cell rectangles inside the table, dropping rectangles nested in another
            x0, top, x1, bottom = table.bbox
            rects = [(round(r["top"]), round(r["bottom"])) for r in page.rects
                     if r["x0"] >= x0 - 2 and r["x1"] <= x1 + 2 and r["top"] >= top - 2 and r["bottom"] <= bottom + 2
                     and (r["bottom"] - r["top"]) > 5]
            rects = sorted(set(rects))
            bands = [b for b in rects if not any(o != b and o[0] <= b[0] and b[1] <= o[1] for o in rects)]
            bands = sorted(set(bands))
            # merge bands that still overlap by more than a hairline
            merged = []
            for b in bands:
                if merged and b[0] < merged[-1][1] - 3:
                    merged[-1] = (merged[-1][0], max(merged[-1][1], b[1]))
                else:
                    merged.append(b)
            bands = merged

            def band_of(w):
                yc = (w["top"] + w["bottom"]) / 2
                for i, (t, b) in enumerate(bands):
                    if t <= yc < b:
                        return i
                return None

            rows = {}
            for w in words:
                bi = band_of(w)
                if bi is not None:
                    rows.setdefault(bi, []).append(w)

            # header band: the first band that contains the sub-labels (Anytime/Peak/...)
            header_bi = next(bi for bi in sorted(rows) if any(w["text"] == "Anytime" for w in rows[bi]))
            hw = [w for bi in sorted(rows) if bi <= header_bi for w in rows[bi] if not is_small(w)]
            code_ci = next(col_of(w) for w in hw if w["text"] == "tariff" and col_of(w) is not None and col_of(w) > 0)
            name_ci = 0
            # fixed-charge column: header words in that column (label words + unit)
            fixed_ci = code_ci + 1
            fx = [w["text"] for w in sorted(hw, key=lambda w: (w["top"], w["x0"])) if col_of(w) == fixed_ci]
            fixed_label = " ".join(t for t in fx if "/" not in t)
            fixed_unit = next((t for t in fx if "/" in t), "c/day")
            # sub-labels (bottom line of the header band) in column order -> groups by order rule
            sub_top = max(w["top"] for w in hw)
            subs = sorted([w for w in hw if abs(w["top"] - sub_top) < 3 and col_of(w) is not None and col_of(w) > fixed_ci],
                          key=lambda w: w["x0"])
            htext = squash(" ".join(w["text"] for w in sorted(hw, key=lambda w: (w["top"], w["x0"]))))
            cols, group, seen_anytime = {}, "Energy charges", 0
            for w in subs:
                if w["text"] == "Anytime":
                    seen_anytime += 1
                    if seen_anytime == 2:
                        group = "Demand rates"
                if w["text"] == "Specified":
                    group = "Capacity/connection charges"
                cols[col_of(w)] = dict(sub=w["text"], group=group)
            foot_hdr = [w["text"] for bi in sorted(rows) if bi <= header_bi for w in rows[bi]
                        if is_small(w) and w["text"].isdigit()]
            hdr_note = "; ".join(f"header footnote {n}: {foot[n]}" for n in foot_hdr if n in foot)

            cls = ""
            for bi in sorted(rows):
                if bi <= header_bi:
                    continue
                ws = rows[bi]
                by_col = {}
                for w in ws:
                    by_col.setdefault(col_of(w), []).append(w)
                codes = [w for w in by_col.get(code_ci, []) if CODE_RE.match(w["text"])]
                if not codes:
                    # class header row: words only in the description column
                    if set(by_col) == {name_ci}:
                        cls = " ".join(w["text"] for w in sorted(ws, key=lambda w: w["x0"]))
                    continue
                code = codes[0]["text"]
                name = " ".join(w["text"] for w in sorted(by_col.get(name_ci, []), key=lambda w: (round(w["top"]), w["x0"]))
                                if not is_small(w))
                marks = [w["text"] for w in ws if is_small(w) and w["text"].isdigit()]
                notes = [f"Table {table_no} ({basis}) p{pno}"]
                for mk in marks:
                    if mk in foot:
                        notes.append(f"footnote {mk}: {foot[mk]}" + (" -> obsolete tariff" if "obsolete" in foot[mk].lower() else ""))
                if hdr_note:
                    notes.append(hdr_note)
                # fixed charge
                fv = [w for w in by_col.get(fixed_ci, []) if num_text(w["text"])]
                if fv:
                    em.add(code, name, cls, fixed_label, fixed_unit, num_text(fv[0]["text"]), basis, note="; ".join(notes))
                for ci, c in cols.items():
                    nums = [w for w in sorted(by_col.get(ci, []), key=lambda w: w["top"]) if num_text(w["text"])]
                    if not nums:
                        continue
                    if c["group"] == "Energy charges":
                        em.add(code, name, cls, f"{c['group']} - {c['sub']}", "c/kWh", num_text(nums[0]["text"]), basis,
                               note="; ".join(notes))
                    elif c["group"] == "Demand rates":
                        if code == "TASUMSSL":
                            unit, un = "c/lamp watt/day", "unit per footnote 'Public lighting is charged on the basis of c/lamp watt/day'"
                        else:
                            meas = DEMAND_MEASURE_2023_24.get(code, "kVA")
                            unit = f"c/{meas}/day"
                            un = (f"header unit published as 'c/kVA, kW, lamp watt/day'; {meas} per 2023-24 price guide Table 36"
                                  if code in DEMAND_MEASURE_2023_24 else
                                  "header unit published as 'c/kVA, kW, lamp watt/day'; kVA assumed [UNSURE]")
                        em.add(code, name, cls, f"{c['group']} - {c['sub']}", unit, num_text(nums[0]["text"]), basis,
                               note="; ".join(notes + [un]))
                    else:  # Capacity/connection charges: 'a / b' = demand charge / connection charge
                        labels = [f"{c['group']} - {c['sub']}", f"{c['group']} - {c['sub']} connection"]
                        for k, w in enumerate(nums[:2]):
                            n2 = list(notes)
                            if len(nums) > 1:
                                n2.append("cell published as 'demand / connection' pair (labels per 2023-24 price guide Table 27: "
                                          f"{'specified/excess daily demand charge' if k == 0 else 'specified/excess daily demand connection charge'})")
                            em.add(code, name, cls, labels[k], "c/kVA/day", num_text(w["text"]), basis,
                                   note="; ".join(n2), charge_type="demand", time_band="")
    return em.rows


# ----------------------------------------------------------------------------- 2023-24 price guide (AER hosted)
def parse_guide_2023_24(path, fin_year, side, url):
    em = Emitter(fin_year, side, path, url, "excl", "")
    with pdfplumber.open(path) as pdf:
        # Table 1: Standard control services network tariffs -> class / name / type per code
        meta = {}
        for page in pdf.pages:
            text = page.extract_text() or ""
            if "Table 1: Standard control services network tariffs" not in text:
                continue
            lines = group_lines(page.extract_words())
            hdr = next(ln for ln in lines if any(w["text"] == "Type" for w in ln["words"]))
            name_x = sorted(w["x0"] for w in hdr["words"] if w["text"] == "Network")[-1]
            cls = ""
            for ln in lines:
                if ln["top"] <= hdr["top"]:
                    continue
                m = re.match(r"^(.*?)\s*(TAS\w+|ITC)\s+(Published tariff|Published obsolete tariff|Negotiated tariff)$", ln["text"])
                if not m:
                    continue
                cw = [w["text"] for w in ln["words"] if w["x0"] < name_x - 5]
                if cw:
                    cls = " ".join(cw)
                nm = " ".join(w["text"] for w in ln["words"] if w["x0"] >= name_x - 5 and w["x0"] < ln["words"][-3]["x0"]
                              and w["text"] != m.group(2))
                meta[m.group(2)] = dict(cls=cls, name=nm, type=m.group(3))
            break

        # per-tariff price tables
        for pno, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            flat = text.replace("\n", " ")
            titles = list(re.finditer(r"Table (\d+): Tariff prices for (.+?) \((TAS\w+)\) for 2023-24", flat))
            if not titles:
                continue
            words = page.extract_words(extra_attrs=["size"])
            lines = group_lines(words)
            foot = page_footnotes(words)
            for tm in titles:
                tno, tname, code = tm.group(1), tm.group(2), tm.group(3)
                title_ln = next(ln for ln in lines if f"({code})" in ln["text"] and "Tariff prices" in ln["text"]
                                or (f"({code})" in ln["text"] and "for 2023-24" in ln["text"]))
                after = [ln for ln in lines if ln["top"] > title_ln["top"] + 2]
                # header: basis column centres and the Unit column centre
                BASIS_CANON = {"duos": "DUoS", "tuos": "TUoS", "nuos": "NUoS"}  # p52 header is typeset 'DUoS TuoS NuoS'
                hdr_lines = [ln for ln in after[:4]
                             if any(w["text"].lower() in ("duos", "tuos", "nuos", "unit", "charge") for w in ln["words"])
                             and not any(num_text(w["text"]) for w in ln["words"])]
                centres = {}
                unit_x = None
                for ln in hdr_lines:
                    for w in ln["words"]:
                        if w["text"].lower() in BASIS_CANON:
                            centres[BASIS_CANON[w["text"].lower()]] = (w["x0"] + w["x1"]) / 2
                        if w["text"] == "Unit":
                            unit_x = (w["x0"] + w["x1"]) / 2
                hdr_bottom = max(ln["top"] for ln in hdr_lines)
                first_basis_x = min(centres.values())

                def nearest_basis(xc):
                    return min(centres, key=lambda b: abs(centres[b] - xc))

                def classify(w):
                    """label | unit | <basis> (numeric or '-') | text (free text in a basis column)."""
                    xc = (w["x0"] + w["x1"]) / 2
                    if w["x1"] <= unit_x - 12:
                        return "label"
                    if xc < (unit_x + first_basis_x) / 2:
                        return "unit"
                    # every price in these tables carries decimals; bare integers belong to free text
                    # ('As per nodal charge in section 19', '5 times nodal charge')
                    if (num_text(w["text"]) is not None and "." in w["text"]) or w["text"] == "-":
                        return nearest_basis(xc)
                    return "text"

                data, frags = [], []
                for ln in after:
                    if ln["top"] <= hdr_bottom + 2:
                        continue
                    t = ln["text"]
                    if (t.startswith("Note") or t.startswith("Network tariff application") or re.match(r"^\d+(\.\d+)+\s", t)
                            or (ln["words"][0]["text"].isdigit() and is_small(ln["words"][0]))):
                        break
                    parts = {}
                    for w in ln["words"]:
                        parts.setdefault(classify(w), []).append(w)
                    if any(k in centres for k in parts) or "text" in parts:
                        data.append(dict(top=ln["top"], parts=parts))
                    else:
                        frags.append(dict(top=ln["top"], parts=parts))
                for fr in frags:  # attach label/unit fragments (wrapped cells) to the nearest data line
                    if not data:
                        break
                    tgt = min(data, key=lambda d: abs(d["top"] - fr["top"]))
                    for k, ws in fr["parts"].items():
                        tgt["parts"].setdefault(k, []).extend(ws)
                mt = meta.get(code, {})
                base_notes = [f"price guide Table {tno} p{pno}"]
                if mt.get("type") and mt["type"] != "Published tariff":
                    base_notes.append(f"Table 1 type: {mt['type']}" + (" -> obsolete tariff" if "obsolete" in mt["type"].lower() else ""))
                for d in data:
                    p = d["parts"]
                    lab_words = sorted(p.get("label", []), key=lambda w: (round(w["top"]), w["x0"]))
                    label = " ".join(w["text"] for w in lab_words if not is_small(w))
                    marks = [w["text"] for w in lab_words if is_small(w) and w["text"].isdigit()]
                    unit = " ".join(w["text"] for w in sorted(p.get("unit", []), key=lambda w: (round(w["top"]), w["x0"])))
                    notes = list(base_notes)
                    for mk in marks:
                        if mk in foot:
                            notes.append(f"footnote {mk}: {foot[mk]}")
                    text_cells = {}
                    tw = sorted(p.get("text", []), key=lambda w: (round(w["top"]), w["x0"]))
                    if tw:
                        span_mid = (min(w["x0"] for w in tw) + max(w["x1"] for w in tw)) / 2
                        text_cells[nearest_basis(span_mid)] = " ".join(w["text"] for w in tw)
                    for b in centres:
                        cell = " ".join(w["text"] for w in sorted(p.get(b, []), key=lambda w: w["x0"]))
                        val = num_text(cell)
                        if val is None:
                            continue
                        n = list(notes)
                        for tb, tc in text_cells.items():
                            n.append(f"{tb} published as text: '{tc}'")
                        if "NUoS" not in centres:
                            n.append("NUoS column not published for this tariff (TUoS is locational, see Table 33 rows)")
                        em.add(code, mt.get("name") or tname, mt.get("cls", ""), label, unit, val, b, note="; ".join(n))

        # Table 33: locational TUoS (transmission connection sites)
        pairs, grab = [], False
        for page in pdf.pages:
            text = page.extract_text() or ""
            if "Table 33: Transmission connection sites" in text:
                grab = True
            if not grab:
                continue
            for m in re.finditer(r"^(.+?)\s+(T[A-Z]{2}\d)\s+([\d,.]+|-)$", text, re.M):
                if num_text(m.group(3)) is not None:
                    pairs.append((m.group(1), m.group(2), num_text(m.group(3))))
            if "Virtual nodes" in text:
                break
        mt = meta.get("TAS15", {})
        add_locational(em, pairs, "price guide s19 Table 33 (applies to TAS15 and ITC customers; TAS15 excess daily demand TUoS = 5 x nodal charge)",
                       name=mt.get("name", "Large business high voltage specified demand > 2MVA"), cls=mt.get("cls", ""))
    return em.rows


# ----------------------------------------------------------------------------- main
FILES = [
    ("2023-24", "AER_HOSTED", "sources/aer/2023-24_price_lists/TasNetworks_2023-24_Network_tariff_application_and_price_guide_4Apr2023.pdf", parse_guide_2023_24),
    ("2023-24", "DNSP", "sources/dnsp/tasnetworks/TasNetworks_Network_Tariff_Pricing_Schedule_SCS_2023-24.pdf", parse_pdf_2023_24_schedule),
    ("2024-25", "DNSP", "sources/dnsp/tasnetworks/TasNetworks_Network_Tariff_Pricing_Schedule_SCS_2024-25.pdf", parse_pdf_2024_25),
    ("2025-26", "DNSP", "sources/dnsp/tasnetworks/TasNetworks_Network_Tariff_Pricing_Schedule_SCS_2025-26.xlsx", parse_xlsx),
    ("2026-27", "DNSP", "sources/dnsp/tasnetworks/TasNetworks_Network_Pricing_Schedule_2026-27.xlsx", parse_xlsx),
]


def main():
    urls = source_urls()
    rows = []
    for fin_year, side, path, fn in FILES:
        url = urls.get(path, "")
        if not url:
            print(f"WARNING: no inventory URL for {path}", file=sys.stderr)
        got = fn(path, fin_year, side, url)
        codes = sorted({r["tariff_code"] for r in got})
        bases = sorted({r["basis"] for r in got})
        print(f"{fin_year} {side:10s} {os.path.basename(path)}: {len(got)} rows, {len(codes)} codes, bases={bases}")
        rows.extend(got)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=schema.COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in schema.COLUMNS})
    print(f"wrote {OUT}: {len(rows)} rows")


if __name__ == "__main__":
    main()
