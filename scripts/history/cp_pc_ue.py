#!/usr/bin/env python
"""Historical network tariffs of CitiPower, Powercor and United Energy (Victoria) for the pricing years before 2023-24.

Run from the repository root:  .venv/bin/python scripts/history/cp_pc_ue.py
Writes out/history/cp_pc_ue.csv (scripts/history/CONTRACT.md).

One function per document layout. Every value is read from the page's words (pdfplumber, glyphs forced upright, words
in content-stream order so that a tariff name running into the code column stays apart from the code) and placed by
the printed labels around it: the tariff code or name of its line and the column headings above it. Column positions
are taken from the numbers themselves (a column is the horizontal span its numbers overlap) and each column is labelled
with the heading phrases printed over that span.
"""
import os
import functools
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "tariffdb"))
import common  # noqa: E402
import locators  # noqa: E402

import pdfplumber  # noqa: E402

SLUG = "cp_pc_ue"
BEFORE_GST = "before GST"
PRE_GST_YEARS = ("1996-97", "1997-98", "1998-99", "1999-00")
NUM_RE = re.compile(r"^\(?-?\$?\d[\d,]*(?:\.\d+)?\)?$|^\(?-?\$?\.\d+\)?$")
DASH = ("-", "–", "—")
ARCH = "sources/archive/"
LOG = []


# ------------------------------------------------------------------------------------------------ helpers
def clean_num(s):
    """The number as printed without '$', thousands separators or parentheses-as-negative kept as '-'."""
    s = s.strip().replace("$", "").replace(",", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if not re.match(r"^-?(\d+(\.\d+)?|\.\d+)$", s):
        return None
    if s.startswith("."):
        s = "0" + s
    return ("-" + s) if neg else s


def page_lines(page, flow=True, xtol=1):
    """Lines of a page: lists of word dicts (text, x0, x1, top), left to right. Number fragments that touch ('1 .247',
    '3 ,020.690': one printed number split into glyph runs) are joined."""
    for c in page.chars:
        c["upright"] = True
    ws = page.extract_words(x_tolerance=xtol, y_tolerance=2, keep_blank_chars=False, use_text_flow=flow)
    lines = []
    for w in sorted(ws, key=lambda w: w["top"]):
        for ln in lines:
            if abs(ln[0]["top"] - w["top"]) < 2.5:
                ln.append(w)
                break
        else:
            lines.append([w])
    out = []
    for ln in lines:
        ln.sort(key=lambda w: w["x0"])
        merged = []
        for w in ln:
            if merged and -1 < w["x0"] - merged[-1]["x1"] < 0.6 and re.search(r"[\d,(\-]$", merged[-1]["text"]) \
                    and re.match(r"^[.,]?\d", w["text"]):
                merged[-1] = dict(merged[-1], text=merged[-1]["text"] + w["text"], x1=w["x1"])
            else:
                merged.append(dict(w))
        out.append(merged)
    out.sort(key=lambda l: l[0]["top"])
    return out


PAGE_CODES = {}  # (path, page) -> the tariff codes a price table prints, priced or not
OCR_DPI = 288  # the resolution locators.ocr_text re-reads an image page at, so check.py sees the same text


@functools.lru_cache(maxsize=8)
def ocr_lines(path, page_no, dpi=OCR_DPI):
    """Lines of an image-only page read by OCR, as page_lines gives them (positions in PDF points). Each OCR box is
    one word; a box read with confidence below 0.9 stops the run."""
    from rapidocr_onnxruntime import RapidOCR
    with pdfplumber.open(path) as pdf:
        result, _ = RapidOCR()(pdf.pages[page_no - 1].to_image(resolution=dpi).original)
    k = 72 / dpi
    ws = [{"text": t.strip(), "x0": b[0][0] * k, "x1": b[2][0] * k, "top": b[0][1] * k, "conf": c}
          for b, t, c in result or [] if t.strip()]
    lines = []
    for w in sorted(ws, key=lambda w: w["top"]):
        for ln in lines:
            if abs(ln[0]["top"] - w["top"]) < 4:
                ln.append(w)
                break
        else:
            lines.append([w])
    for ln in lines:
        ln.sort(key=lambda w: w["x0"])
    return sorted(lines, key=lambda l: l[0]["top"])


def text(ws):
    return " ".join(w["text"] for w in ws)


def phrases(line, gap=3.2):
    """Words of a line joined into phrases where they are one space apart."""
    out = []
    for w in line:
        if out and w["x0"] - out[-1]["x1"] < gap:
            out[-1] = dict(out[-1], text=out[-1]["text"] + " " + w["text"], x1=w["x1"])
        else:
            out.append(dict(w))
    return out


class Doc:
    """One archived document: emits parser rows for it."""

    def __init__(self, path, note="", gst="excl"):
        self.path = path
        self.gst = gst
        self.doc = common.document(path)
        self.rows = []
        self.note = note
        self.before_gst = self.doc["pricing_year"] in PRE_GST_YEARS

    def add(self, code, component, value, unit, page_no, *, name="", customer_class="", note="", basis="NUoS",
            time_band=None, season=None, charge_type=None, ocr=False):
        notes = [n for n in (note, self.note, BEFORE_GST if self.before_gst else "") if n]
        if ocr:  # OCR drops the space and the capital of a printed unit ('$/kVApa', 'c/kwh')
            unit = re.sub(r"(?<=\S)pa$", " pa", re.sub(r"kw(h?)\b", r"kW\1", unit))
        self.rows.append(common.row(self.doc, code, component, value, unit, locators.pdf(page_no, ocr=ocr),
                                    name=name, customer_class=customer_class, basis=basis, gst=self.gst,
                                    note="; ".join(notes), time_band=time_band, season=season,
                                    charge_type=charge_type))


def energy_band(label, blocks=False):
    """time_band of an energy price from its printed label."""
    l = label.lower()
    if re.search(r"off[- ]?peak|\bopk\b", l):
        return "offpeak"
    if re.search(r"\bshoulder\b|\bsh\b", l):
        return "shoulder"
    if re.search(r"\bpeak\b|\bpk\b", l):
        return "peak"
    return "anytime"


# ---------------------------------------------------------------- 1996-97, 1997-98: all five distributors
def vic_combined(path, column):
    """ORG 'Network tariffs for use of each distributor's distribution system' (nt9697, nt9798): one table with one
    column per distributor, read by the distributor's column heading. Tariff names, no codes."""
    d = Doc(path, note="no code printed")
    with pdfplumber.open(path) as pdf:
        heads, cls, name, pending, last_name = None, "", "", None, False
        for page_no, page in enumerate(pdf.pages, 1):
            for ws in page_lines(page, flow=False):
                t = text(ws)
                if re.search(r"\bEastern\b.*\bSolaris\b", t):
                    heads = [(w["text"].lower(), w["x1"]) for w in ws if w["text"].lower() != "units"]
                    continue
                if re.match(r"^(LOW|HIGH) VOLTAGE TARIFFS|^SUBTRANSMISSION TARIFFS", t):
                    cls, last_name = t, False
                    continue
                if heads is None or re.match(r"^(NETWORK|FOR THE|These prices|http|\*|#|Minimum|demand\b|\d[\d,]* ?kW)",
                                             t):
                    continue
                vals = []
                for w in reversed(ws):
                    if clean_num(w["text"]) is not None or w["text"] in DASH:
                        vals.insert(0, w)
                    else:
                        break
                body = ws[:len(ws) - len(vals)]
                if body and not t.startswith("-"):
                    vals, body = [], ws  # a tariff name ending in '-'
                if not body and vals and pending:  # values wrapped onto the next line
                    pending[2].extend(vals)
                    continue
                if pending:
                    emit_combined(d, heads, column, *pending)
                    pending = None
                if t.startswith("-"):
                    label = text(body).lstrip("-").strip()
                    m = re.match(r"^(.*?)\s*(\$pa|c/kWh\]?|\$/kW pa|\$/kVA pa)$", label)
                    if not m:
                        raise SystemExit(f"{path} p{page_no}: unreadable component line {t!r}")
                    pending = [page_no, (cls, name, m.group(1).rstrip("]").strip(), m.group(2).rstrip("]")), vals]
                    last_name = False
                    continue
                if vals:
                    raise SystemExit(f"{path} p{page_no}: numbers on a line that is not a component: {t!r}")
                name = f"{name} {t}" if last_name and name.endswith("-") else t  # a name wrapped after '-'
                name, last_name = re.sub(r"\s+", " ", name), True
            if pending:
                emit_combined(d, heads, column, *pending)
                pending = None
    return d.rows


def emit_combined(d, heads, column, page_no, what, vals):
    cls, name, label, unit = what
    names = [h for h, _ in heads]
    if len(vals) == len(heads):
        picked = vals[names.index(column)]
    else:
        x1 = dict(heads)[column]
        near = [v for v in vals if abs(v["x1"] - x1) < 8]
        if len(near) != 1:
            raise SystemExit(f"{d.path} p{page_no}: cannot place {[v['text'] for v in vals]} under {column}")
        picked = near[0]
    if picked["text"] in DASH:
        return
    code = re.sub(r"\s*-$", "", name.replace("#", "").replace("*", "")).strip()
    note = "new tariff 1997/98" if "#" in name else ""
    if "*" in name:
        note = "restricted to customers previously on retail tariff E1 or N1 consuming more than 400 MWh pa"
    d.add(code, label, clean_num(picked["text"]), unit, page_no, name=code, customer_class=cls.title(), note=note,
          time_band=energy_band(label) if "kWh" in unit else None)


# ---------------------------------------- 1998-99 to 2003: per tariff, one 'component unit price' line per charge
LIST_UNIT = r"\$ ?/ ?cust pa|\$/KVA/p\.a|\$/kW/p\.a|\$/kW ?pa|\$/p\.a|\$ ?p\.a|\$ ?pa|c/kWh"
LIST_LINE = re.compile(rf"^(?P<label>.*?)\s*(?P<unit>{LIST_UNIT})\s+(?P<value>[\d ,.]*\d)$")
LIST_SKIP = re.compile(r"^(CitiPower|POWERCOR|UNITED ENERGY|NETWORK|DISTRIBUTION|TRANSMISSION|To Apply|Excluding GST|"
                       r"EXCLUDING GST|APPENDIX|\d+/\d+$|\d{4}/\d{2,4}$|\d+$|Low Voltage( Categories| Components)?$|"
                       r"Categories|Nominal|volts|< ?1,?000|1000 volts|22,000 volts|and [£<]|"
                       r"Component|Units|(LOW|HIGH) VOLTAGE TARIFFS|SUBTRANSMISSION TARIFFS|& < 22|Minimum|demand\.?$|"
                       r"\d[\d,]* ?k[WV]|kW$|\*|Page \d|Tariffs$|Supplies$)", re.I)
NOTE_LINE = re.compile(r"^(Closed Tariff.*|Opening.*)$")


def list_layout(path, pages, *, name_x=None, title=None, flow=False):
    """Regulator-hosted lists (CitiPower 1999-2000 and 2002, Powercor 1998-99 and 1999-2000, United Energy 1999-2000)
    and CitiPower's 2003 list: a tariff name, then one line per charge: label, unit, price. A category printed left of
    `name_x` (Residential, General Purpose, Large Low Voltage, High Voltage) prefixes the tariff name. A label printed
    alone ('Standing Charge:') takes the next price printed without its own label. `title`: a heading the page must
    print (the network table of a report that also prints distribution and transmission tables)."""
    d = Doc(path, note="no code printed")
    seen = set()
    with pdfplumber.open(path) as pdf:
        category, name, fresh, note, sub, queue = "", "", False, "", "", []
        for page_no in pages:
            page = pdf.pages[page_no - 1]
            if title and not re.search(title, page.extract_text() or ""):
                raise SystemExit(f"{path} p{page_no}: no {title!r} heading")
            for ws in page_lines(page, flow=flow):
                cat_ws = [w for w in ws if name_x is not None and w["x0"] < name_x - 4]
                rest = [w for w in ws if w not in cat_ws]
                cat_t, t = text(cat_ws), text(rest)
                if cat_t:
                    m = re.match(r"^(.*?)\s+Tariffs?$", cat_t)
                    new = m.group(1) if m else None if LIST_SKIP.match(cat_t) or cat_t.startswith("(") else cat_t
                    if new is not None:
                        if new != category and not fresh:
                            name = ""
                        category = new
                if re.search(r"\bComponents?\b.*\bUnits\b|\bUnits\s+Price\b", text(ws)) or \
                        re.match(r"^(To Apply|EXCLUDING GST|Excluding GST|APPENDIX|NETWORK|CitiPower$|\d+$)", text(ws)):
                    continue
                if not t:
                    continue
                m = LIST_LINE.match(t)
                if not m:
                    if NOTE_LINE.match(t):
                        note = t
                    elif re.match(r"^Unit charge:?$", t):
                        sub = "Unit charge"
                    elif t.endswith(":") or re.match(r"^(Standing Charge|Peak unit rate|Off Peak unit rate):?$", t):
                        queue.append(t.rstrip(":").strip())
                    elif LIST_SKIP.match(t):
                        pass
                    else:
                        name = f"{name} {t}" if fresh else t
                        fresh, note, sub, queue = True, "", "", []
                    continue
                label, unit = m.group("label").strip(" :"), m.group("unit")
                value = clean_num(m.group("value").replace(" ", ""))
                if value is None:
                    raise SystemExit(f"{path} p{page_no}: unreadable value in {t!r}")
                if queue:
                    if label:
                        queue.append(label)
                    label = queue.pop(0)
                elif sub and re.match(r"^(First|Balance)", label):
                    label = f"{sub} {label}"
                if not label:
                    raise SystemExit(f"{path} p{page_no}: price without a label: {t!r}")
                tname = name
                if category and not (name and name.split()[0].lower() == category.split()[0].lower()):
                    tname = f"{category} {name}".strip()
                tname = re.sub(r"\s+", " ", tname).strip()
                if not tname:
                    raise SystemExit(f"{path} p{page_no}: price without a tariff: {t!r}")
                if (tname, label) in seen:
                    raise SystemExit(f"{path} p{page_no}: {tname} / {label} printed twice")
                seen.add((tname, label))
                band = None
                if "kWh" in unit:
                    band = energy_band(label)
                    if re.search(r"first|block 1", label, re.I):
                        band = "block1"
                    elif re.search(r"balance|block 2", label, re.I):
                        band = "block2"
                code = tname.replace("*", "").strip()
                d.add(code, label.replace("*", "").strip(), value, unit.replace(" ", " "), page_no, name=code,
                      note=note, time_band=band)
                fresh = False
    return d.rows


# ------------------------------------------------------------- tables: one line per tariff, one column per charge
class Grid:
    """A price table of one page: rows are the lines that print a tariff code (`code_re`, a whole word, optionally
    left of `code_x1`), their prices are the numbers right of the code. A column is the horizontal span its numbers
    overlap; its heading is every heading phrase (a run of words one space apart) printed over that span between the
    line matching `header_re` and the first row. A '-' printed in a column is no price."""

    def __init__(self, page, *, code_re, header_re, code_x1=None, code_x0=0, stop_re=None, flow=True, extra_re=None,
                 heads=None, top=0, bottom=None, xtol=1, after_header=False, lines=None, strict=True, skip_re=None):
        lines = [l for l in (lines or page_lines(page, flow=flow, xtol=xtol))
                 if l[0]["top"] >= top and (bottom is None or l[0]["top"] < bottom)]
        start = next((i for i, l in enumerate(lines) if re.search(header_re, text(l))), None)
        if start is None and heads is None:
            raise SystemExit(f"no heading {header_re!r} on page {page.page_number}")
        start = (start or 0) + (1 if after_header else 0)
        end = next((i for i, l in enumerate(lines) if i > start and stop_re and re.search(stop_re, text(l))),
                   len(lines))
        if code_x1 == "Code":  # codes start under the 'Code' heading
            cw = next((w for l in lines[start:start + 8] for w in l if w["text"] == "Code"), None)
            if cw is None:
                raise SystemExit(f"page {page.page_number}: no 'Code' heading")
            code_x0, code_x1 = cw["x0"] - 30, cw["x1"] + 10
        skip = None  # skip_re: the heading of a column that holds no price ('Available to new customers': Yes/No,
        # which OCR can misread as a number); the words under it are left out of every row
        if skip_re:
            hw = [w for l in lines[max(0, start - 1):start + 8] for w in l if re.fullmatch(skip_re, w["text"])]
            if hw:
                skip = (min(w["x0"] for w in hw) - 3, max(w["x1"] for w in hw) + 3)
        self.rows, self.other = [], []  # rows: (line index, code, name words, value words)
        for i in range(start, end):
            l = lines[i]
            if skip:
                l = [w for w in l if not (w["x0"] >= skip[0] and w["x1"] <= skip[1])]
            ks = [k for k, w in enumerate(l) if re.fullmatch(code_re, w["text"])
                  and (code_x1 is None or w["x0"] < code_x1) and w["x0"] >= code_x0]
            ks = ks[-1:] if code_x1 is not None else ks[:1]
            k = ks[0] if ks else None
            if k is None:
                self.other.append((i, l))
                continue
            vals = [w for w in l[k + 1:] if clean_num(w["text"]) is not None or w["text"] in DASH]
            extra = [w["text"] for w in l[k + 1:] if w not in vals]
            if extra and not (extra_re and all(re.fullmatch(extra_re, e) for e in extra)):
                if not strict:  # a check page: leave out a line OCR garbled
                    continue
                raise SystemExit(f"page {page.page_number}: words {extra} among the prices of {text(l)!r}")
            low = [w["text"] for w in vals if w.get("conf", 1) < 0.9]
            if low and not strict:
                continue
            if low:
                raise SystemExit(f"page {page.page_number}: OCR is unsure of {low} in {text(l)!r}")
            self.rows.append((i, l[k]["text"], l[:k], vals))
        if not self.rows:
            raise SystemExit(f"page {page.page_number}: no tariff rows")
        for i, l in self.other:  # a priced line whose code was not recognised
            if strict and i > self.rows[0][0] and sum(clean_num(w["text"]) is not None for w in l
                                                     if code_x1 is None or w["x0"] >= code_x1) >= 2:
                raise SystemExit(f"page {page.page_number}: prices on a line without a code: {text(l)!r}")
        nums = sorted((w for r in self.rows for w in r[3] if w["text"] not in DASH), key=lambda w: w["x0"])
        spans = []
        for w in nums:
            if spans and w["x0"] < spans[-1][1] - 0.5:
                spans[-1][1] = max(spans[-1][1], w["x1"])
            else:
                spans.append([w["x0"], w["x1"]])
        self.spans = spans
        first = self.rows[0][0]
        if heads is not None:
            self.heads = [self.match_heads(heads, sp) for sp in spans]
        else:
            self.heads = [[] for _ in spans]
            self.head_phrases = []
            for l in lines[start:first]:
                for ph in phrases(l):
                    self.head_phrases.append(ph)
                    for j, (a, b) in enumerate(spans):
                        if ph["x0"] < b + 2 and ph["x1"] > a - 2:
                            self.heads[j].append(ph["text"])
        self.lines = lines

    @staticmethod
    def match_heads(prev, span):
        """Headings of a continuation page: those of the previous page's column at the same place."""
        for (a, b), h in prev:
            if span[0] < b + 2 and span[1] > a - 2:
                return h
        raise SystemExit(f"no column heading over x={span[0]:.0f}-{span[1]:.0f}")

    def column(self, w):
        for j, (a, b) in enumerate(self.spans):
            if w["x0"] < b and w["x1"] > a:
                return j
        raise AssertionError(w)

    def prices(self, row):
        """(column index, heading text, value word) of one row's printed prices."""
        return [(self.column(w), " | ".join(self.heads[self.column(w)]), w) for w in row[3] if w["text"] not in DASH]

    def saved_heads(self):
        return list(zip([tuple(s) for s in self.spans], self.heads))


UNIT_RE = re.compile(r"\$ ?/ ?cust ?/? ?p\.?a\.?|\$/kVA/p\.?a|\$/kW/p\.?a|\$/kW/month|\$/kVA/month|c/kWh?r?|c/Wh|"
                     r"\$ ?pa|\$/year|c/day|c/kVA/day|c/kW/day|\$/kvA/pa|\$/kVA ?pa|\$/kW ?pa|\$/kVAr/pa", re.I)


VALID_BLOCK_BANDS = {"block1", "block2", "block3", "peak_block1", "peak_block2"}


def resolve_band(band, off, blocks, extra=""):
    """(time_band, note) of a price column's band. 'PEAK': the peak column, which holds a single-rate tariff's only
    energy price when the tariff prints no off-peak price (`off`); 'BLOCKn': the n-th energy block (`blocks`: the
    tariff's block bands), a peak block when the tariff also prints an off-peak price."""
    if band == "PEAK":
        if not off:
            return "anytime", "single rate: its energy price is printed in the peak column"
        return "peak", extra
    if band == "BLOCK1" and blocks == ["BLOCK1"]:  # one energy price, printed in the first block's column
        return ("peak" if off else "anytime"), \
            f"the only {'peak ' if off else ''}energy price, printed in the first block's column"
    if band.startswith("BLOCK"):
        n = int(band[5:])
        band = f"peak_block{n}" if off else f"block{n}"
        if band not in VALID_BLOCK_BANDS:
            return "", f"{'peak ' if off else ''}energy block {n} (the database has no band for it)"
    return band, extra


def grid_rows(d, page_no, g, spec, *, name_of=None, cls="", note="", ocr=False, code_of=None):
    """Rows of a Grid's prices. spec: (heading regex, component or None for the heading as printed, unit or None for
    the unit the heading prints, band, season) in order of precedence, or (regex, None) for a column that is not a
    price (minimum demand, availability). band 'BLOCKn' is the n-th energy block: a peak block when the tariff also
    prints an off-peak price, else a plain block."""
    def match(head, col):
        if callable(spec):
            return spec(head, col)
        for item in spec:
            if re.search(item[0], head, re.I):
                return item
        raise SystemExit(f"{d.path} p{page_no}: no rule for the column headed {head!r}")
    for row in g.rows:
        prices = [(c, h, w, match(h, c)) for c, h, w in g.prices(row)]
        off = any(len(it) > 2 and it[3] == "offpeak" for _, _, _, it in prices)
        blocks = [it[3] for _, _, _, it in prices if len(it) > 2 and it[3].startswith("BLOCK")]
        seen = set()
        for c, head, w, it in prices:
            if len(it) == 2:
                continue
            _, comp, unit, band, season = it
            comp = comp or re.sub(r"\s*\|\s*", " ", head)
            if unit is None:
                m = UNIT_RE.search(head)
                if not m:
                    raise SystemExit(f"{d.path} p{page_no}: no unit in the heading {head!r}")
                unit = m.group(0)
            extra = "the column heading prints 'c/Wh', a misprint for c/kWh" if "c/Wh" in head else ""
            band, extra = resolve_band(band, off, blocks, extra)
            if (comp, band, season) in seen:
                raise SystemExit(f"{d.path} p{page_no}: {row[1]} prints two prices under {comp!r}")
            seen.add((comp, band, season))
            code = code_of(row) if code_of else row[1]
            name = name_of(row) if name_of else text(row[2])
            d.add(code, comp, clean_num(w["text"]), unit, page_no, name=name, customer_class=cls,
                  note="; ".join(x for x in (note, extra) if x), time_band=band, season=season, ocr=ocr)


# ---------------------------------------------------------------- Powercor 2002: tariff report appendix A
PC2002 = [
    (r"^Standing", "Standing charges", "$/cust/pa", "", ""),
    (r"^Demand \| charges", "Demand charges", "$/kW/pa", "", ""),
    (r"Minimum", None),
    (r"First 4", "Peak charges First 4 MWh/yr", "c/kWh", "BLOCK1", ""),
    (r"Next 16", "Peak charges Next 16 MWh/yr", "c/kWh", "BLOCK2", ""),
    (r"Next 50", "Peak charges Next 50 MWh/yr", "c/kWh", "BLOCK3", ""),
    (r"Balance", "Peak charges Balance", "c/kWh", "BLOCK4", ""),
    (r"^Off peak", "Off peak energy", "c/kWh", "offpeak", ""),
]


def powercor_2002(path, page_no):
    """Powercor 2002 and 2003 tariff reports, Appendix A 'Network tariff schedule (GST exclusive)': NUoS tariff, code,
    standing, demand, minimum demand, four peak energy blocks (First 4 / Next 16 / Next 50 MWh/yr / Balance), off
    peak."""
    d = Doc(path)
    with pdfplumber.open(path) as pdf:
        g = Grid(pdf.pages[page_no - 1], code_re=r"[A-Z]{1,3}\d?(\.[A-Z0-9]+)?", code_x1=250,
                 header_re=r"^NUoS Tariff Code",
                 stop_re=r"^Some Tariffs", extra_re=r"kW")
        grid_rows(d, page_no, g, PC2002)
    return d.rows


# ------------------------------------------- CitiPower and Powercor 2004-2018: 'Network tariff schedule' (one page)
def block_spec(g, *, peak="Peak charges", before=(), after=()):
    """Rules for a Grid whose energy blocks are columns headed First / Next / Balance / Block n, numbered left to right,
    and whose time-of-use prices are two runs of Pk / Sh / Opk columns, the summer run left of the non-summer run (the
    season headings are centred over each run, so some columns carry none)."""
    heads = [" | ".join(h) for h in g.heads]
    blocks = [j for j, h in enumerate(heads) if re.search(r"First|Next|Balance|\b340\b|Block ?\d", h)
              and not re.search(r"Off[- ]?Peak|Time ?of ?Use|\b(Pk|Sh|Opk)\b", h, re.I)]
    by_col = {}
    for n, j in enumerate(blocks, 1):
        label = re.sub(r"\s+", " ", re.sub(r"\bPeak\b|\(c/kWh\)|charges|c/kWh|\|", " ", heads[j], flags=re.I)).strip()
        m = re.search(r"Block ?(\d)", heads[j])
        by_col[j] = (heads[j], f"{peak} {label}", "c/kWh", f"BLOCK{m.group(1) if m else n}", "")
    first = next((j for j, h in enumerate(heads) if re.search(r"Time ?of ?Use|\b(Pk|Sh|Opk)\b", h)), None)
    if first is not None:  # the time-of-use runs: every column from the first one on
        tou = list(range(first, len(heads)))
        if len(tou) != 6:
            raise SystemExit(f"{len(tou)} time-of-use columns {heads[first:]}: expected summer and non-summer peak, "
                             f"shoulder, off-peak")
        for n, j in enumerate(tou):
            kind, season = ("Pk", "Sh", "Opk")[n % 3], "summer" if n < 3 else "non_summer"
            h = heads[j]
            said = ("Opk" if re.search(r"\bOpk\b|Off[- ]?Peak", h, re.I) else "Sh" if re.search(r"\bSh\b|Shoulder", h)
                    else "Pk" if re.search(r"\bPk\b|\bPeak\b", h) else kind)
            said_season = "non_summer" if "Non-Summer" in h else "summer" if "Summer" in h else season
            if (said, said_season) != (kind, season):
                raise SystemExit(f"column {h!r} sits where the {season} {kind} column belongs")
            band = {"Pk": "peak", "Sh": "shoulder", "Opk": "offpeak"}[kind]
            by_col[j] = (h, f"{'Summer' if n < 3 else 'Non-Summer'} Time of Use {kind}", "c/kWh", band, season)
    def spec(head, col):
        if col in by_col:
            return by_col[col]
        for item in tuple(before) + SCHEDULE + tuple(after):
            if re.search(item[0], head, re.I):
                return item
        raise SystemExit(f"no rule for the column headed {head!r}")
    return spec


SCHEDULE = (
    (r"Minimum|Available|\|\s*kW$", None),
    (r"^Fixed", "Fixed", None, "", ""),
    (r"Jan-Dec", "Demand Jan-Dec", None, "", ""),
    (r"Dec-Mar", "Demand Dec-Mar", None, "", "summer"),
    (r"Apr-Nov", "Demand Apr-Nov", None, "", "non_summer"),
    (r"Anytime", "Anytime", None, "anytime", ""),
    (r"^Usage \| c/kWh$", "Usage Peak", "c/kWh", "PEAK", ""),  # 'Peak' printed off the column, under 'Usage'

    (r"\$/kvA ?/?pa|\$/kVA ?/?pa", "Demand charges kVA", None, "", ""),
    (r"Standing|\$/cust", "Standing charges", None, "", ""),
    (r"Demand|\$/kW/pa", "Demand charges", None, "", ""),
    (r"Off[- ]?peak", "Off peak charges", "c/kWh", "offpeak", ""),
    (r"Peak", "Peak charges", "c/kWh", "PEAK", ""),
)
CODE_CP = r"C[0-9A-Z]+(-B)?"
LATE = {"flow": True, "after_header": True,  # 2012-2018: headings start below the date line
        "header_re": r"(DECEMBER|December) \d{4}$|EXCLUSIVE OF GST$"}
CODE_PC = r"[A-Z][A-Z0-9]{0,7}(\.[A-Z0-9]+)*"
PP20 = {"ocr": True, "dpi": 360, "title_lines": 6, "after_header": True, "skip_re": r"customers",
        "header_re": r"\((NUoS|DUoS|TUoS|JUoS)\) ?Tariff ?2020$",  # (OCR drops some spaces of the titles)
        "title": r"Table A\. ?1 Network ?\(NUoS\) ?Tariff ?2020",
        "parts_title": r"Table A\. ?[234] ?(Distribution|Transmission|Jurisdictional)"}
NTS19 = {"ocr": True, "dpi": 360, "after_header": True, "header_re": r"GST Exclusive$",
         "title": r"(CitiPower|Powercor) 2019 Network Tariff Schedule \(NUoS\)"}


def tariff_schedule(path, page_no, code_re, note="", flow=False, header_re=r"\bCode\b", after_header=False, ocr=False,
                    parts=(), title=r"NETWORK TARIFF SCHEDULE", strict=True, dpi=OCR_DPI, title_lines=3,
                    parts_title=r"(Distribution|Transmission|Jurisdictional) tariff schedule", skip_re=None):
    """CitiPower and Powercor 'Network tariff schedule (GST exclusive)', 2004-2018: NUoS tariff, code, [available to
    new customers], fixed or standing $/cust pa, demand ($/kW pa, $/kVA pa, $/kW/month by season), minimum demand,
    peak energy blocks (CitiPower: First 1020 kWh/qtr, First 340 kWh/month / Balance, Block 1-2; Powercor: First 4 /
    Next 16 / Next 50 MWh/yr, First 333 / Next 1334 / Next 4166 kWh/month / Balance, Block 1-4), anytime, peak, off
    peak, and summer and non-summer time-of-use runs.

    ocr: the page is an image, read by OCR. `parts` then names the pages of the same document that print the
    distribution, transmission and jurisdictional parts of every price in the same layout: each NUoS price must equal
    the sum of its parts (to the rounding of the printed figures), which catches a misread digit."""
    d = Doc(path, note=note)
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_no - 1]
        first = text([w for l in (ocr_lines(path, page_no, dpi) if ocr else page_lines(page))[:title_lines] for w in l])
        if not re.search(title, first, re.I):
            raise SystemExit(f"p{page_no} is not headed {title!r}: {first[:80]!r}")
        g = Grid(page, code_re=code_re, code_x1="Code", header_re=header_re, flow=flow, xtol=2.5,
                 after_header=after_header, lines=ocr_lines(path, page_no, dpi) if ocr else None, strict=strict,
                 skip_re=skip_re,
                 extra_re=r"kW|(Yes|No)['¹1]?|[·.,'\-]",
                 stop_re=r"^(Some Tariffs|Issued by|The following|Notes?:?$|Approved by|\(\d\) )")
        if os.environ.get("CPU_DEBUG"):
            for (a, b), h in zip(g.spans, g.heads):
                print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
        grid_rows(d, page_no, g, block_spec(g), ocr=ocr)
        PAGE_CODES[(path, page_no, dpi)] = {r[1] for r in g.rows}
    if parts:
        def key(r):
            return r["tariff_code"], r["component"], r["time_band"], r["season"]
        total, seen = {}, []  # the parts' sums; the tariffs each part page prints (OCR misses a code now and then)
        for p in parts:
            codes, part = set(), {}
            # a row OCR garbles at one resolution it often reads at another (the main page's resolution first)
            for pdpi in dict.fromkeys((dpi, OCR_DPI, 360, 216, 432, 504)):
                try:
                    got = tariff_schedule(path, p, code_re, flow=flow, header_re=header_re, after_header=after_header,
                                          ocr=ocr, title=parts_title, strict=False, dpi=pdpi, title_lines=title_lines,
                                          skip_re=skip_re)
                except SystemExit:
                    continue
                new = PAGE_CODES[(path, p, pdpi)] - codes
                codes |= new
                for r in got:
                    if r["tariff_code"] in new:
                        part[key(r)] = float(r["value"])
            seen.append(codes)
            for k, v in part.items():
                total[k] = total.get(k, 0) + v
        bad, unchecked = [], set()
        for r in d.rows:
            if not all(r["tariff_code"] in codes for codes in seen):
                unchecked.add(r["tariff_code"])
                continue
            dp = len(r["value"].partition(".")[2])
            if abs(total.pop(key(r), 0) - float(r["value"])) > 1.5 * len(parts) * 10 ** -dp / 2 + 1e-9:
                bad.append(f"{r['tariff_code']} {r['component']} {r['value']}")
        codes = {r["tariff_code"] for r in d.rows} - unchecked  # (a misread code on a part page is no NUoS code)
        bad += [f"{k[0]} {k[1]}: parts only" for k, v in total.items() if abs(v) > 1e-9 and k[0] in codes]
        if bad:
            raise SystemExit(f"p{page_no}: NUoS is not the sum of the parts on pages {parts}: {bad}")
        LOG.append(f"{path} p{page_no}: NUoS = sum of the parts on pages {parts} for every price, except tariffs "
                   f"{sorted(unchecked)} whose row OCR could not place on a part page")
    return d.rows


# ------------------------------------------------------------------ United Energy 2000-H2: ORG-hosted NUoS schedule
def join_parens(lines):
    """Lines with every '(...)' run of words joined into one word (codes printed as '(LVS1R-block 1)')."""
    out = []
    for l in lines:
        new = []
        for w in l:
            if new and new[-1]["text"].startswith("(") and not new[-1]["text"].endswith(")"):
                new[-1] = dict(new[-1], text=new[-1]["text"] + " " + w["text"], x1=w["x1"])
            else:
                new.append(dict(w))
        out.append(new)
    return out


UE2000 = [
    (r"Minimum|\(kVA\)$", None),
    (r"Standing", "Standing Charge", "$/year", "", ""),
    (r"Off Peak", "Off Peak Energy", "c/kWh", "offpeak", ""),
    (r"Peak", "Peak Energy", "c/kWh", "PEAK", ""),
    (r"^Demand$", "Demand Charge", "$/kVA/year", "", ""),
]


def ue_2000(path, page_no=1):
    """United Energy 'Schedule of Network Use of System Tariffs: 1 July 2000 (Without 0.25% tax savings, and without
    10% GST additional tax)': tariff name, then its code in parentheses on the line of its prices: standing $/year,
    peak and off peak c/kWh, demand $/kVA/year, minimum chargeable demand kVA. 'LV Small 1 rate' prints its two
    energy blocks as two lines, codes 'LVS1R-block 1' and 'LVS1R-block 2'."""
    d = Doc(path, note="without the 0.25% tax savings")
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_no - 1]
        g = Grid(page, code_re=r"\([A-Za-z0-9.\- ]+\)", header_re=r"^Minimum$", stop_re=r"^\* LV Small",
                 lines=join_parens(page_lines(page, flow=False)))
        if os.environ.get("CPU_DEBUG"):
            for (a, b), h in zip(g.spans, g.heads):
                print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
        rows, g.rows = g.rows, []
        for row in rows:
            code = row[1].strip("()")
            name = re.sub(r"\s*-$", "", text(g.lines[row[0] - 1])).rstrip("*")
            m = re.match(r"^(LVS1R)-block (\d)$", code)
            g.rows = [row]
            spec = UE2000
            if m:  # one tariff's two energy blocks, printed as two lines
                block = f"block{m.group(2)}"
                spec = [(r"Standing", "Standing Charge", "$/year", "", ""),
                        (r"Peak", f"Peak Energy block {m.group(2)}", "c/kWh", block, "")]
                code, name = m.group(1), re.sub(r"\s*block \d$", "", name)
            grid_rows(d, page_no, g, spec, code_of=lambda r, c=code: c, name_of=lambda r, n=name: n,
                      note=f"printed as '{row[1].strip('()')}'" if m else "")
    return d.rows


# ------------------------------------------------------- United Energy 2001: tariff report, NUOS schedule (p26)
UE2001 = [  # headings as Grid reads them: one word per phrase (the page spaces every word wide)
    (r"Minimum|Eligibility", None),
    (r"Standing \| Charge \| \$/cust \| pa$", "Standing Charge", "$/cust pa", "", ""),
    (r"^Summer \| Demand \| Incentive \| Interruptibl \| Contract \| \$/kVA \| pm$", "Summer Interruptible Contract",
     "$/kVA/month", "", "summer"),
    (r"^Rolling \| Peak \| Demand \| Interruptibl \| Contract \| \$/kVA \| pa$",
     "Rolling Peak Demand Interruptible Contract", "$/kVA pa", "", ""),
    (r"Summer \| Demand \| Incentive \| Charge \| \*\* \| \$/kW \| pm \| \$/kVA \| pm$",
     "Summer Demand Incentive Charge", "UNIT", "", "summer"),
    (r"Rolling \| Peak \| Demand \| Charge \| \$/kVA \| pa$", "Rolling Peak Demand Charge", "$/kVA pa", "", ""),
    (r"^Summer \| Peak \| Energy \| Charge \| c/kWh$", "Summer Peak Energy Charge", "c/kWh", "PEAK", "summer"),
    (r"^Non- \| Summer \| Peak \| Energy \| Charge \| c/kWh$", "Non-Summer Peak Energy Charge", "c/kWh", "PEAK",
     "non_summer"),
    # 'Off' is printed left of the column's numbers, so the heading Grid reads is 'Peak Energy Charge'
    (r"^Peak \| Energy \| Charge \| c/kWh$", "Off Peak Energy Charge", "c/kWh", "offpeak", ""),
]


def ue_2001(path, page_no=26):
    """United Energy Tariff Report 2001, 'Schedule of Network Use of System (NUOS) Tariffs: 1 January 2001' (rates
    exclude GST): tariff name (a code such as LVS1R; '*' closed to new connections), standing $/cust pa, summer demand
    incentive charge ($/kW pm for LVkWTOU, $/kVA pm otherwise), rolling peak demand $/kVA pa, the two interruptible
    contract prices (printed in parentheses: credits), summer and non-summer peak energy, off peak energy, eligibility
    and minimum chargeable demand. Below it, the 2001 transition adjustment to the summer demand incentive charge."""
    d = Doc(path)
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_no - 1]
        lines, split = [], set()  # split: numbers the text layer breaks between two digits
        for l in page_lines(page, flow=False, xtol=3):
            name = [w for w in l if w["x1"] < 158]
            rest = []
            for w in (w for w in l if w["x1"] >= 158):  # the page prints a space inside some numbers: '5 3.402'
                if rest and w["x0"] - rest[-1]["x1"] < 2.5:
                    if rest[-1]["text"][-1:].isdigit() and w["text"][:1].isdigit():
                        split.add((rest[-1]["text"] + w["text"]).strip("()"))
                    rest[-1] = dict(rest[-1], text=rest[-1]["text"] + w["text"], x1=w["x1"])
                else:
                    rest.append(w)
            if name:
                rest.insert(0, dict(name[0], text=text(name), x1=name[-1]["x1"]))
            lines.append(rest)
        g = Grid(page, code_re=r"(LV|HV|SubT)[A-Za-z0-9 ]*\*?", code_x1=158, header_re=r"^Network Tariff Component",
                 stop_re=r"^S = Small", extra_re=r"\d+-\d+|>\d+", lines=lines)
        if os.environ.get("CPU_DEBUG"):
            for (a, b), h in zip(g.spans, g.heads):
                print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
        sdic = {}  # the transition adjustment table under the schedule: tariff group -> its monthly adjustments
        t = [text(l) for l in g.lines]
        k = next(i for i, x in enumerate(t) if x.startswith("Tariff Group Jan-01 Feb-01 Mar-01 Nov-01 Dec-01"))
        for x in t[k + 1:]:
            m = re.match(r"^(\S+) ((\( ?[\d.]+\)|-) ?){5}$", x)
            if not m:
                break
            sdic[re.sub(r"\W", "", m.group(1)).upper()] = re.findall(r"\( ?[\d.]+\)|-", x[len(m.group(1)):])
        group = {"LVKWTOU": "LVKWTOU", "LVKVATOU": "LVKVATOU", "LVINTERRUPTIBLE": "LVKVATOU", "HVKVATOU": "HVKVA",
                 "HVINTERRUPTIBLE": "HVKVA", "SUBTKVATOU": "ST22KVA"}
        for row in g.rows:
            code = row[1].rstrip("*")
            note = "closed to new connections" if row[1].endswith("*") else ""
            spec = []
            for item in UE2001:
                if item[1] == "Summer Demand Incentive Charge":
                    unit = "$/kW/month" if code == "LVkWTOU" else "$/kVA/month"
                    item = (item[0], item[1], unit, "", "summer")
                spec.append(item)
            g1 = Grid.__new__(Grid)
            g1.__dict__.update(g.__dict__, rows=[row])
            rows_before = len(d.rows)
            grid_rows(d, page_no, g1, spec, code_of=lambda r, c=code: c, name_of=lambda r, c=code: c, note=note)
            adj = sdic.get(group.get(re.sub(r"\W", "", code).upper(), ""))
            for r in d.rows[rows_before:]:
                if r["component"] == "Summer Demand Incentive Charge":
                    if adj is None:
                        raise SystemExit(f"{path} p{page_no}: no transition adjustment row for {code}")
                    r["note"] = "; ".join(x for x in (r["note"], "the transition adjustment table on this page "
                                         "changes it by " + ", ".join(f"{m} {a.replace(' ', '')}" for m, a in zip(
                                             ("Jan-01", "Feb-01", "Mar-01", "Nov-01", "Dec-01"), adj))) if x)
    for r in d.rows:  # the validator reads a number split between two digits ('2 4.640') only from the rendered page
        if r["value"].lstrip("-") in split:
            r["locator"] = locators.pdf(page_no, ocr=True)
    return d.rows


# --------------------------------------- United Energy 2002-2005: 'Schedule of Network Use of System (NUOS) Tariffs'
def join_left(lines, x):
    """Lines with the words left of x joined into one (a tariff name such as 'WET-Step 1<=4080kWh') and number
    fragments a printed space splits ('1 ,150') joined."""
    out = []
    for l in lines:
        name = [w for w in l if w["x1"] < x]
        rest = []
        for w in (w for w in l if w["x1"] >= x):
            if rest and w["x0"] - rest[-1]["x1"] < 2.5 and re.fullmatch(r"[\d,.]+", rest[-1]["text"] + w["text"]):
                rest[-1] = dict(rest[-1], text=rest[-1]["text"] + w["text"], x1=w["x1"])
            else:
                rest.append(w)
        if name:
            rest.insert(0, dict(name[0], text=text(name), x1=name[-1]["x1"]))
        out.append(rest)
    return out


def ue_nuos(path, page_no=1):
    """United Energy 'Schedule of Network Use of System (NUOS) Tariffs: 1 January <year> (GST Exclusive)', 2002-2005:
    tariff ('*' closed to new connections), standing ($/cust pa, from 2003 c/day), summer and non-summer peak energy,
    off-peak energy, rolling peak demand ($/kVA pa, from 2003 c/kVA/day), summer demand incentive charge ($/kW or
    $/kVA pm, from 2003 c/kW or c/kVA per day), eligibility and minimum chargeable demand. 'WET-Step 1' and 'WET-Step
    2' are the two energy blocks of one tariff; only the first prints a summer price."""
    d = Doc(path)
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_no - 1]
        raw = page_lines(page, flow=False)
        if not re.match(r"Schedule of Network Use of System \(NUOS\) Tariffs: 1 January \d{4} \(GST Exclusive\)",
                        text(raw[0])):
            raise SystemExit(f"{path} p{page_no} is not the NUOS schedule: {text(raw[0])!r}")
        sx = next(w["x0"] for l in raw for w in l if w["text"] == "Standing") - 2
        g = Grid(page, code_re=r"(HOT )?(LV|HV|ST22|RCAC)\S*|Ded\*?|UnMet|WET-Step [12]\S*", code_x1=sx,
                 header_re=r"^Network Tariff Component", after_header=True, stop_re=r"^\*Tariff closed",
                 extra_re=r"[<>]\d+|\d+-\d+|Residential|Business|Customer", lines=join_left(raw, sx))
        if os.environ.get("CPU_DEBUG"):
            for (a, b), h in zip(g.spans, g.heads):
                print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
        for row in g.rows:
            code, note, step = row[1].rstrip("*"), "closed to new connections" if row[1].endswith("*") else "", None
            m = re.match(r"WET-Step ([12])(\S*?)\**$", code)
            if m:
                step, note, code = m.group(1), f"printed as '{row[1]}'", "WET"
            kw = "kW" in re.sub(r"kVA", "", code)

            def spec(head, col):
                day = bool(re.search(r"/ ?da", head))
                if re.search(r"Minimum|Eligibility|^(kVA|kWh|MWh|mWh pa)$|Customer", head):
                    return (head, None)
                if "Standing" in head:
                    return (head, "Standing Charge", "c/day" if day else "$/cust pa", "", "")
                if "Incentive" in head:
                    unit = ("c/kW/day" if kw else "c/kVA/day") if day else ("$/kW/month" if kw else "$/kVA/month")
                    return (head, "Summer Demand Incentive Charge", unit, "", "summer")
                if "Rolling" in head:
                    return (head, "Rolling Peak Demand", "c/kVA/day" if day else "$/kVA pa", "", "")
                if "Off-Peak" in head:
                    return (head, "Off-Peak Energy", "c/kWh", "offpeak", "")
                if re.match(r"Non(-sum| ?-) \| (Summer \| )?Peak \| Energy \| c/kWh$", head):
                    season = "non_summer"
                elif re.match(r"Summer \| Peak \| Energy \| c/kWh$", head):
                    season = "summer"
                else:
                    raise SystemExit(f"{path} p{page_no}: no rule for the column headed {head!r}")
                comp = f"{'Non-Summer' if season == 'non_summer' else 'Summer'} Peak Energy"
                if step:
                    return (head, f"{comp} step {step}", "c/kWh", f"block{step}", season)
                return (head, comp, "c/kWh", "PEAK", season)
            g1 = Grid.__new__(Grid)
            g1.__dict__.update(g.__dict__, rows=[row])
            grid_rows(d, page_no, g1, spec, code_of=lambda r, c=code: c, name_of=lambda r, c=code: c, note=note)
    return d.rows


# ------------------------------------------- United Energy 2006: tariff report, Table 6.1 NUoS schedule (p14)
UE2006 = [
    (r"Eligibility|Minimum|^(kVA|MWh \| pa)$", None),
    # the heading prints '(C/kWh)' under 'Standing Charge'; the report's abbreviations table (p15) gives the fixed
    # charge in c/day, and the price is the DUoS fixed charge of the GST-exclusive schedule plus 10% GST
    (r"^Standing \| Charge", "Standing Charge", "c/day", "", ""),
    (r"^Summer \| Peak \| Energy", "Summer Peak Energy", "c/kWh", "PEAK", "summer"),
    (r"^Non \| Summer \| Peak \| Energy", "Non Summer Peak Energy", "c/kWh", "PEAK", "non_summer"),
    (r"^Off peak \| Energy", "Off peak Energy", "c/kWh", "offpeak", ""),
    (r"^Rolling \| Peak \| Demand", "Rolling Peak Demand", "c/kVA/day", "", ""),
    (r"^Summer \| Demand \| Incentive", "Summer Demand Incentive Charge", "UNIT", "", "summer"),
]


def ue_2006(path, page_no=14):
    """United Energy Distribution Tariff Report 2006, 'Table 6.1 Schedule of Network Use of System (NUoS) Tariffs: 1
    January 2006 (GST inclusive)': tariff names (no codes) wrapped over two or three lines, the prices printed on the
    first; columns as in 2003-2005. 'WET-Step 1' and 'WET-Step 2' are the two energy blocks of one tariff."""
    d = Doc(path, note="no code printed", gst="incl")
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_no - 1]
        raw = page_lines(page, flow=False)
        if not any(re.match(r"Table 6.1 Schedule of Network Use of System \(NUoS\) Tariffs: 1 January 2006$", text(l))
                   for l in raw[:4]) or not any(text(l) == "(GST inclusive)" for l in raw[:4]):
            raise SystemExit(f"{path} p{page_no} is not Table 6.1 (GST inclusive)")
        nx = 130  # tariff names end left of this, prices start right of it
        start = next(i for i, l in enumerate(raw) if text(l).startswith("Network Tariff Component"))
        first = next(i for i, l in enumerate(raw) if i > start and l[0]["x1"] < nx and text(l) != "Tariffs" and
                     not text(l).startswith("Tariffs "))
        lines, entries = raw[:first], []  # entries: [name words, line]

        def priced(ws):
            return any(clean_num(w["text"]) is not None for w in ws)
        for i in range(first, len(raw)):
            l = raw[i]
            left, right = [w for w in l if w["x1"] < nx], [w for w in l if w["x0"] >= nx]
            nxt = raw[i + 1] if i + 1 < len(raw) else []
            if re.fullmatch(r"\d+ of \d+", text(l)):
                lines.append(l)
                break
            if left and priced(right):  # a name and its prices on one line
                ln = [dict(left[0], text="")] + right
                entries.append([left, ln])
                lines.append(ln)
            elif left and not priced(right) and nxt and not [w for w in nxt if w["x1"] < nx] and priced(nxt) \
                    and abs(nxt[0]["top"] - l[0]["top"]) < 5:  # a name whose prices print a little lower
                ln = [dict(left[0], text="")] + right
                entries.append([left, ln])
                lines.append(ln)
            elif not left and priced(right) and entries and len(entries[-1][1]) - 1 == len(
                    [w for w in raw[i - 1] if w["x0"] >= nx]):
                entries[-1][1].extend(right)
                entries[-1][1].sort(key=lambda w: w["x0"])
            elif left and entries:  # the name continued on the next line (and an eligibility word beside it)
                entries[-1][0].extend(left)
                entries[-1][1].extend(w for w in right)
                entries[-1][1].sort(key=lambda w: w["x0"])
            else:
                raise SystemExit(f"{path} p{page_no}: cannot place the line {text(l)!r}")
        for words, ln in entries:
            ln[0]["text"] = text(words)
            ln.sort(key=lambda w: w["x0"])
        g = Grid(page, code_re=r"(Low|High|Sub|Dedicated|Reverse|Unmetered|WET-Step)\b.+", code_x1=nx,
                 header_re=r"^Network Tariff Component", after_header=True, stop_re=r"^\d+ of \d+$",
                 extra_re=r"[<>]\d+|Business|Residential|Customer", lines=lines)
        if os.environ.get("CPU_DEBUG"):
            for (a, b), h in zip(g.spans, g.heads):
                print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
        for row in g.rows:
            name = row[1]
            code, note, step = name.rstrip("*"), "closed to new connections" if "*" in name else "", None
            m = re.match(r"WET-Step ([12])", name)
            if m:
                step, note, code = m.group(1), f"printed as '{name}'", "WET-Step"
            spec = []
            for rx, *rest in UE2006:
                if len(rest) > 1 and rest[1] == "UNIT":  # kW for the kW tariffs (RCACkWTOU in the 2002-05 lists)
                    rest[1] = "c/kW/day" if " KW " in f" {code} " or code.startswith("Reverse cycle") else "c/kVA/day"
                if len(rest) > 1 and step and rest[2] == "PEAK":
                    rest = [f"{rest[0]} step {step}", rest[1], f"block{step}", rest[3]]
                spec.append((rx, *rest))
            g1 = Grid.__new__(Grid)
            g1.__dict__.update(g.__dict__, rows=[row])
            n0 = len(d.rows)
            grid_rows(d, page_no, g1, spec, code_of=lambda r, c=code: c, name_of=lambda r, c=code: c, note=note)
            for r in d.rows[n0:]:
                if r["component"] == "Standing Charge":
                    r["note"] = "the column heading prints '(C/kWh)'; the abbreviations table on p15 gives the fixed " \
                                "charge in c/day; " + r["note"]
    return d.rows


# ------------------------- United Energy 2007-2012: 'Schedule of Network Use of System (NUOS) Tariffs' with codes
def split_heads(g, start):
    """Recompute a Grid's column headings where one heading phrase runs over several columns (2017-18: 'Demand Max
    Demand Max Overrun Max Overrun Max'): each of its words goes to the column whose centre is nearest."""
    heads = [[] for _ in g.spans]
    for l in g.lines[start:g.rows[0][0]]:
        for ph in phrases(l):
            cols = [j for j, (a, b) in enumerate(g.spans) if ph["x0"] < b + 2 and ph["x1"] > a - 2]
            if len(cols) <= 1:
                for j in cols:
                    heads[j].append(ph["text"])
                continue
            parts = {}
            for w in l:
                near = [j for j in cols if w["x0"] < g.spans[j][1] + 2 and w["x1"] > g.spans[j][0] - 2]
                if near and w["x0"] >= ph["x0"] - 0.1 and w["x1"] <= ph["x1"] + 0.1:
                    c = (w["x0"] + w["x1"]) / 2
                    j = min(near, key=lambda j: abs(c - (g.spans[j][0] + g.spans[j][1]) / 2))
                    parts.setdefault(j, []).append(w["text"])
            for j in cols:
                if j in parts:
                    heads[j].append(" ".join(parts[j]))
    g.heads = heads


def ue_coded(path, page_no, cont=(), title=r"Network Use of System \(NUOS\)"):
    """United Energy 'Schedule of Network Use of System (NUOS) Tariffs: 1 January <year> (GST Exclusive)', 2007-2012:
    description (wrapping onto the code's line), tariff code ('*' closed to new connections), [2012: 'F,T' where the
    tariff is also offered with the premium (F) and transitional (T) feed-in tariff prefix], standing c/day, summer
    peak, non-summer peak blocks 1 and 2, [from 2010: summer and non-summer shoulder], off peak (c/kWh), rolling peak
    demand c/kVA/day, summer demand incentive c/kW/day or c/kVA/day, eligibility, minimum chargeable demand. `cont`:
    pages the table continues on, without headings."""
    d = Doc(path)
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_no - 1]
        raw = []
        for l in page_lines(page, flow=False):  # footnote lists: 'RESKWTOU3,' '4' -> 'RESKWTOU3,4'
            out = []
            for w in l:
                if out and re.fullmatch(r"\S*\d,", out[-1]["text"]) and re.fullmatch(r"\d", w["text"]):
                    out[-1] = dict(out[-1], text=out[-1]["text"] + w["text"], x1=w["x1"])
                else:
                    out.append(w)
            raw.append(out)
        if not re.search(r"Schedule of " + title + r" Tariffs: (\d\.\d\. )?1 January \d{4} \(GST Exclusive\)( |$)",
                         " ".join(text(l) for l in raw[:4])):  # (from 2013 the title wraps after 'Tariffs:')
            raise SystemExit(f"{path} p{page_no} is not the GST-exclusive schedule {title!r}")
        hl = next(l for l in raw if re.match(r"Description Tariff Code\b", text(l)))
        tw = next(w for w in hl if w["text"] == "Tariff")
        cw = next(w for w in hl if w["text"] == "Code")
        kw_ = dict(code_re=r"(?!(Code|Tariff|Charge|PFIT|TFIT|Description)\b)[A-Z][A-Za-z0-9]+(\d,\d)?[*#]{0,2}",
                   code_x0=tw["x0"] - 20, code_x1=cw["x1"] + 10, stop_re=r"^\*", flow=False,
                   extra_re=r"F,T|F,|T|F|NA|[<>]\d+|>\d+ & <\d+|&|Residential|Business|Customer|LVS1R\)?|\(on")
        g = Grid(page, header_re=r"\(GST Exclusive\)$", after_header=True, lines=raw, **kw_)
        split_heads(g, next(i for i, l in enumerate(raw) if re.search(r"\(GST Exclusive\)$", text(l))) + 1)
        if os.environ.get("CPU_DEBUG"):
            for (a, b), h in zip(g.spans, g.heads):
                print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
        grids = [(page_no, g)] + [(p, Grid(pdf.pages[p - 1], header_re=r"$^", heads=g.saved_heads(), **kw_))
                                  for p in cont]
        for pno, g in grids:
            ue_coded_rows(d, path, pno, g)
    return d.rows


def ue_coded_rows(d, path, page_no, g):
    """The rows of one page of ue_coded's table."""
    rowlines = {r[0] for r in g.rows}
    cx = min(g.lines[r[0]][len(r[2])]["x0"] for r in g.rows)  # where the codes start
    for row in g.rows:
        i0 = row[0]
        name = row[2]
        prev = g.lines[i0 - 1]
        if i0 > 0 and i0 - 1 not in rowlines and all(w["x1"] < cx for w in prev):
            name = prev + name  # the description begun on the line above
        code = row[1]
        m = re.search(r"(\d(,\d)?)$", code)
        if m and re.search(re.escape(m.group(1).replace(",", ", ")) + r"$|" + re.escape(m.group(1)) + "$", text(name)):
            # a footnote mark printed after both the description and the code: 'rate1 RESKW1R1'
            name = [dict(w) for w in name]
            name[-1]["text"] = re.sub(r",?\s*\d$", "", name[-1]["text"]) if "," not in m.group(1) else \
                name[-1]["text"]
            nm = re.sub(r"\s*" + re.escape(m.group(1)).replace(",", ", ?") + "$", "", text(name))
            name = [dict(name[0], text=nm)]
            code = code[:m.start()]
        note = ["closed to new connections"] if re.search(r"(?<!\*)\*$", code) else []  # ("**": a footnote)
        if code.endswith("#"):
            note = ["open to new connections from 1 January 2007 (see eligibility conditions)"]
        ft = "".join(w["text"] for w in g.lines[i0][len(row[2]) + 1:] if re.fullmatch(r"F,T|F,|T|F", w["text"]))
        if ft:
            note.append(f"'{ft}': also offered as the " + " and ".join(
                {"F": "PFIT (F)", "T": "TFIT (T)"}[c] for c in "FT" if c in ft) + " variant of this tariff")
        kw = "kW" in code or code.rstrip("*#") in ("TOU", "FTOU")
        if code.rstrip("*#") in ("TOU", "FTOU"):  # (its code names no kW or kVA)
            note.append("demand incentive in c/kW/day: the 2010 tariff report (Table 4.3) measures the TOU "
                        "tariff's SDIC at maximum kW")

        def spec(head, col):
            if re.search(r"Minimum|Eligibility|^(kVA|MWh pa|category\)|Customer)", head) or head in ("", "kVA"):
                return (head, None)
            if re.search(r"Standi( \| )?ng", head):  # (2011 wraps it: 'Standi | ng | Charg | e')
                return (head, "Standing Charge", "c/day", "", "")
            # (from 2015) seasonal maximum demand, and from 2016 its overrun
            if re.search(r"Max( \| | )?kW", head, re.I):
                season = "non_summer" if re.match(r"Non", head) else "summer"
                comp = ("Non-summer" if season == "non_summer" else "Summer") + " Demand" + (
                    " Overrun" if "Overrun" in head else "") + " Max kW"
                return (head, comp, "c/kW/day", "", season)
            if "Incentive" in head:
                return (head, "Summer Demand Incentive Charge", "c/kW/day" if kw else "c/kVA/day", "", "summer")
            if "Rolling" in head and "Peak" in head and "Demand" in head:
                return (head, "Rolling Peak Demand", "c/kVA/day", "", "")
            if re.search(r"Off Peak", head):
                return (head, "Off Peak Energy", "c/kWh", "offpeak", "")
            non = bool(re.search(r"^Non\b|\bNon \| Summer", head))
            season = "non_summer" if non else "summer"
            label = "Non Summer" if non else "Summer"
            if "Shoulde" in head:  # (2011: 'Shoulde | r')
                return (head, f"{label} Shoulder Energy", "c/kWh", "shoulder", season)
            m = re.search(r"Block (\d)", head)
            if m and non:
                return (head, f"Non Summer Peak Energy Block {m.group(1)}", "c/kWh", f"BLOCK{m.group(1)}",
                        "non_summer")
            if "Peak" in head and "Energy" in head and not non:
                return (head, "Summer Peak Energy", "c/kWh", "PEAK", "summer")
            raise SystemExit(f"{path} p{page_no}: no rule for the column headed {head!r}")
        g1 = Grid.__new__(Grid)
        g1.__dict__.update(g.__dict__, rows=[row])
        grid_rows(d, page_no, g1, spec, code_of=lambda r, c=code.rstrip("*#"): c,
                  name_of=lambda r, n=text(name).rstrip("*#"): n, note="; ".join(note))


# ---------------------------------- United Energy 2011: tariff report, NUOS schedule with wrapped codes (pp69-72)
def ue_2011(path, pages, parts):
    """United Energy 2011 tariff report, 'Schedule of Network Use of System (NUOS) Tariffs: 1 January 2011 (GST
    Exclusive)' over `pages`: the 2010-2012 columns in narrow cells, so descriptions and tariff codes wrap over two
    or three lines ('SubTk' / 'VATO' / 'U*') and the prices print on the last. A tariff's lines run from the line after
    the previous tariff's prices to its own prices; its code is the words of the code column joined.

    The report is UE's 2011 pricing proposal (its pages are headed 'UED Pricing Proposal 2011'); `parts` is UE's
    published 2011 tariff schedule (path, DUOS page, TUOS page), and every NUOS price must be the sum of its DUOS and
    TUOS prices there, so the stored prices are the ones UE charged."""
    d = Doc(path)
    with pdfplumber.open(path) as pdf:
        for k, page_no in enumerate(pages):
            page = pdf.pages[page_no - 1]
            raw = page_lines(page, flow=False)
            if k == 0 and not any(text(l) == "Schedule of Network Use of System (NUOS) Tariffs: 1 January 2011 (GST "
                                  "Exclusive)" for l in raw[:3]):
                raise SystemExit(f"{path} p{page_no} is not the GST-exclusive NUOS schedule")
            hl = next(l for l in raw if text(l) == "Code")  # the second line of the 'Tariff Code' heading
            c0, c1 = hl[0]["x0"] - 9, hl[0]["x1"] + 12  # the code column
            head_end = next(i for i, l in enumerate(raw) if text(l).startswith("(c/day)"))
            elig = max(w["x1"] for w in raw[head_end] if w["text"] == "c/kVA/day") + 2  # eligibility from here
            lines, part = raw[:head_end + 1], []
            for l in raw[head_end + 1:]:
                if re.match(r"\*", text(l)):
                    break
                part.append(l)
                prices = [w for l2 in [l] for w in l2 if w["x0"] >= c1 and (w["x0"] < elig or clean_num(w["text"]))]
                if not any(clean_num(w["text"]) for w in prices):
                    continue
                words = [w for l2 in part for w in l2]
                code = "".join(w["text"] for w in words if w["x0"] >= c0 and w["x1"] <= c1)
                name = text([w for w in words if w["x1"] < c0])
                if not code or not name:
                    raise SystemExit(f"{path} p{page_no}: no code or name for the prices of {text(l)!r}")
                lines.append([dict(prices[0], text=name, x0=c0 - 30, x1=c0 - 20),
                              dict(prices[0], text=code, x0=c0 + 1, x1=c0 + 2)] + prices)
                part = []
            if any(clean_num(w["text"]) for l in part for w in l if w["x0"] >= c1):
                raise SystemExit(f"{path} p{page_no}: prices left over")
            g = Grid(page, code_re=r"(?!Tariff$|Code$)\S+", code_x0=c0, code_x1=c1, header_re=r"^Minimum$",
                     lines=lines)
            if os.environ.get("CPU_DEBUG"):
                for (a, b), h in zip(g.spans, g.heads):
                    print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
            ue_coded_rows(d, path, page_no, g)
    key = lambda r: (r["tariff_code"], r["component"], r["season"])  # (a part without off peak reads "anytime")
    total = {}
    for p, title in zip(parts[1:], (r"Distribution Use of System \(DUOS\)", r"Transmission Use of System \(TUOS\)")):
        for r in ue_coded(parts[0], p, title=title):
            total[key(r)] = round(total.get(key(r), 0) + float(r["value"]), 6)
    bad = [f"{key(r)} {r['value']} != {total.get(key(r))}" for r in d.rows
           if abs(total.pop(key(r), 0) - float(r["value"])) > 0.0011]
    bad += [f"{k} {v}: DUOS + TUOS only" for k, v in total.items() if v]
    if bad:
        raise SystemExit(f"{path}: NUOS is not DUOS + TUOS of {parts[0]}: {bad}")
    return d.rows


# ------------------------------------------------------------------------------------------------ documents
CP, PC, UE = ARCH + "citipower/", ARCH + "powercor/", ARCH + "unitedenergy/"
# ------------------------------------- United Energy 2020: AER-approved pricing proposal, Table A.1 (text layer)
def ue_2020(path, page_no=19):
    """United Energy 2020 pricing proposal (AER approved), 'Table A. 1 Network (NUoS) Tariffs 2020': tariff, code,
    PFIT variant code, available to new customers (Yes/No), fixed c/day, jurisdictional fixed (PFIT recovery) c/day,
    demand (rolling peak c/kVA/day, summer incentive 'c/kW/day or c/kVA/day', summer and non-summer c/kW/day), usage
    (anytime, peak, off-peak c/kWh), summer time of use (Pk, Sh, Opk) and non-summer time of use (Pk - Block1,
    Pk - Block2, Sh, Opk) c/kWh. The NUoS fixed charge is printed in two parts (network and jurisdictional); both are
    emitted as printed. Codes 'LVDed *': see the '*' footnote under the table."""
    d = Doc(path)
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_no - 1]
        g = Grid(page, code_re=r"(?!(Code|PFIT)\b)[A-Z][A-Za-z0-9]+", code_x1="Code",
                 header_re=r"^Table A\. ?1 Network \(NUoS\) Tariffs 2020$", after_header=True,
                 extra_re=r"F[A-Za-z0-9-]+|Yes|No|\*", stop_re=r"^\* |^United Energy \|")
        star = next((text(l)[2:] for l in g.lines if text(l).startswith("* ")), "")
        if os.environ.get("CPU_DEBUG"):
            for (a, b), h in zip(g.spans, g.heads):
                print(f"{a:6.1f}-{b:6.1f}: {' | '.join(h)}")
        for row in g.rows:
            i0, code = row[0], row[1]
            words = [w["text"] for w in g.lines[i0] if w not in row[3] and w not in row[2] and w["text"] != code]
            note = []
            if "*" in words:
                note.append(star or "see the '*' footnote under the table")
            note += [f"PFIT variant {w}" for w in words if re.fullmatch(r"F[A-Za-z0-9-]+", w)]
            if "No" in words:
                note.append("not available to new customers")
            kw = "kw" in code.lower() or code in ("TOU", "FTOU")

            def spec(head, col, kw=kw):
                if "Juris" in head:
                    return (head, "Jurisdictional Fixed (PFIT recovery)", "c/day", "", "")
                if re.match(r"Fixed \| c/day$", head):
                    return (head, "Fixed", "c/day", "", "")
                if "Rolling peak" in head:
                    return (head, "Rolling Peak Demand", "c/kVA/day", "", "")
                if "incentive" in head:  # unit printed 'c/kW/day or c/kVA/day'
                    return (head, "Summer Demand Incentive Charge", "c/kW/day" if kw else "c/kVA/day", "", "summer")
                if re.match(r"Non- \| Summer \| c/kW/day$", head):
                    return (head, "Non-Summer Demand", "c/kW/day", "", "non_summer")
                if re.match(r"Demand Charges \| Summer \| c/kW/day$", head):
                    return (head, "Summer Demand", "c/kW/day", "", "summer")
                if re.match(r"(Usage \| )?Off-peak \| c/kWh$", head):
                    return (head, "Off-peak", "c/kWh", "offpeak", "")
                if re.match(r"(Usage \| )?Anytime \| c/kWh$", head):
                    return (head, "Anytime", "c/kWh", "anytime", "")
                m = re.match(r"(Non-)?Summer Time of Use Tariffs \| (Pk|Sh|Opk|Pk - \| Block(\d)) \| c/kWh$", head)
                if m:
                    season = "non_summer" if m.group(1) else "summer"
                    label = ("Non-Summer" if m.group(1) else "Summer") + " " + m.group(2).replace(" | ", " ")
                    band = f"BLOCK{m.group(3)}" if m.group(3) else {"Pk": "PEAK", "Sh": "shoulder", "Opk": "offpeak"}[
                        m.group(2)]
                    return (head, label, "c/kWh", band, season)
                raise SystemExit(f"{path} p{page_no}: no rule for the column headed {head!r}")
            g1 = Grid.__new__(Grid)
            g1.__dict__.update(g.__dict__, rows=[row])
            grid_rows(d, page_no, g1, spec, code_of=lambda r, c=code: c,
                      name_of=lambda r, n=text(row[2]).rstrip(" *"): n, note="; ".join(note))
    return d.rows


# ------------------------------- CitiPower, Powercor, United Energy 2019-2022-23: AER-hosted 'Attachment B' summaries
SUMMARY_COLS = (  # (group heading, column heading) -> component, time_band, season ('' = all year)
    (r"^Fixed", r"", "", ""),
    (r"^Demand", r"^(Jan-Dec|Rolling peak)$", "", ""),
    (r"^Demand", r"^(Dec-Mar|Summer incentive|Summer)$", "", "summer"),
    (r"^Demand", r"^(Apr-Nov|Non-Summer)$", "", "non_summer"),
    (r"^Usage", r"^Anytime$", "anytime", ""),
    (r"^Usage", r"^Peak$", "peak", ""),
    (r"^Usage", r"^Off-peak$", "offpeak", ""),
    (r"^Summer Time of Use", r"^Pk$", "PEAK", "summer"),
    (r"^Summer Time of Use", r"^Sh$", "shoulder", "summer"),
    (r"^Summer Time of Use", r"^Opk$", "offpeak", "summer"),
    (r"^Non-Summer Time of Use", r"^Pk$", "PEAK", "non_summer"),
    (r"^Non-Summer Time of Use", r"^Pk - Block(\d)$", "BLOCK", "non_summer"),
    (r"^Non-Summer Time of Use", r"^Sh$", "shoulder", "non_summer"),
    (r"^Non-Summer Time of Use", r"^Opk$", "offpeak", "non_summer"),
)


def tariff_summary(path, sheet):
    """The NUOS sheet of a distributor's tariff summary ('Attachment B') to its annual pricing proposal, as the AER
    hosts it: title row ('2021/22 POWERCOR NETWORK TARIFF SCHEDULE'), 'ALL PRICES SHOWN ARE EXCLUSIVE OF GST', then
    three heading rows (group: Fixed / Demand Charges / Usage / Summer and Non-Summer Time of Use Tariffs; column:
    Jan-Dec, Dec-Mar, Rolling peak, Pk, Pk - Block1 ...; unit) over one row per tariff: name, code, [PFIT code],
    [available to new customers], prices. A price the sheet shows as '-' (a zero) is not charged and is skipped.
    United Energy's summer incentive demand column is headed 'c/kW/day or c/kVA/day' to 2021-H1: c/kW/day for the
    tariffs measured in kW (as in its schedules to 2018), c/kVA/day otherwise."""
    import openpyxl
    from published import cell_value
    d = Doc(path, note="proposed prices" if common.document(path)["price_status"] == "proposed" else "")
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    title, gst = (str(ws["B1"].value or ""), str(ws["B2"].value or ""))
    if not re.search(r"NETWORK (\(NUOS = DUOS \+ TUOS\) )?TARIFF SCHEDULE$", title) or \
            gst != "ALL PRICES SHOWN ARE EXCLUSIVE OF GST":
        raise SystemExit(f"{sheet}: not a GST-exclusive network tariff schedule: {title!r} {gst!r}")
    h = next(r for r in range(1, 10) if ws.cell(r, 3).value == "Code")
    cols, group = [], ""
    for c in range(4, ws.max_column + 1):
        group = str(ws.cell(h, c).value or "").strip() or group
        sub, unit = (str(ws.cell(h + k, c).value or "").strip() for k in (1, 2))
        if not unit:
            continue
        hit = [it for it in SUMMARY_COLS if re.search(it[0], group) and re.search(it[1], sub)]
        if len(hit) != 1:
            raise SystemExit(f"{sheet}: no rule for the column {group!r} / {sub!r} ({unit})")
        _, sub_re, band, season = hit[0]
        if band == "BLOCK":
            band = "BLOCK" + re.search(sub_re, sub).group(1)
        cols.append((c, f"{group.rstrip('*')} {sub}".strip(), unit, band, season))
    heads = {str(ws.cell(h, c).value or "").strip(): c for c in range(2, 7)}
    footnotes = {m.group(1): m.group(2).strip() for r in range(h + 3, ws.max_row + 1)
                 for m in [re.match(r"^(\*) (.+)$", str(ws.cell(r, 2).value or ""))] if m}
    for r in range(h + 3, ws.max_row + 1):
        code = str(ws.cell(r, 3).value or "").strip()
        if not code:
            continue
        name = str(ws.cell(r, 2).value or "").strip()
        note = []
        if code.endswith("*"):
            code, name = code.rstrip(" *"), name.rstrip(" *")
            note.append(footnotes.get("*", "see the sheet's '*' footnote"))
        name = re.sub(r"\(\d\)$", "", name).strip()
        pfit = ws.cell(r, heads["PFIT"]).value if "PFIT" in heads else None
        if pfit not in (None, "", "-"):
            note.append(f"PFIT variant {str(pfit).strip()}")
        avail = next((ws.cell(r, c).value for k, c in heads.items() if k.startswith("Available")), None)
        if avail == "No":
            note.append("not available to new customers")
        prices = [(c, comp, unit, band, season) for c, comp, unit, band, season in cols
                  if isinstance(ws.cell(r, c).value, (int, float)) and ws.cell(r, c).value != 0]
        blocks = [band for _, _, _, band, _ in prices if band.startswith("BLOCK")]
        off = any(band == "offpeak" for _, _, _, band, _ in prices)
        for c, comp, unit, band, season in prices:
            band, why = resolve_band(band, off, blocks)
            extra = note + [why] if why else list(note)
            if unit.startswith("c/kW/day or"):
                kw = "kw" in code.lower() or code in ("TOU", "FTOU")
                extra.append(f"unit printed '{unit.strip()}'")
                unit = "c/kW/day" if kw else "c/kVA/day"
            cell = ws.cell(r, c)
            charge = "fixed" if comp.startswith("Fixed") else "demand" if comp.startswith("Demand") else "energy"
            d.rows.append(common.row(d.doc, code, comp, cell_value(cell), unit, locators.xlsx(ws, cell), name=name,
                                     note="; ".join(x for x in extra + [d.note] if x), charge_type=charge,
                                     time_band=band, season=season))
    return d.rows


DOCS = [
    # (layout, path, arguments)
    (vic_combined, CP + "1996-97/nt9697.pdf", {"column": "citipower"}),
    (vic_combined, PC + "1996-97/nt9697.pdf", {"column": "powercor"}),
    (vic_combined, UE + "1996-97/nt9697.pdf", {"column": "united"}),
    (vic_combined, CP + "1997-98/nt9798.pdf", {"column": "citipower"}),
    (vic_combined, PC + "1997-98/nt9798.pdf", {"column": "powercor"}),
    (vic_combined, UE + "1997-98/nt9798.pdf", {"column": "united"}),
    (list_layout, PC + "1998-99/pco9899.pdf", {"pages": (1, 2, 3), "name_x": 190}),
    (list_layout, CP + "1999-00/cpnt9920.pdf", {"pages": (1, 2, 3), "name_x": 275}),
    (list_layout, PC + "1999-00/pc9900net.pdf", {"pages": (1, 2)}),
    (list_layout, UE + "1999-00/ue9900net.pdf", {"pages": (1, 2)}),
    (list_layout, CP + "2002/CitiPowerNetworkTariffReport02.pdf",
     {"pages": (20, 21, 22), "name_x": 270, "title": r"NETWORK\s+TARIFFS"}),
    (powercor_2002, PC + "2002/TariffReport2002.pdf", {"page_no": 15}),
    # (2003: the report's Appendix A NUoS schedule; network_tariffs_table2003.pdf prints its tables as images)
    (powercor_2002, PC + "2003/TariffReport2003.pdf", {"page_no": 15}),
    (list_layout, CP + "2003/CP_Network_Tariffs_2003.pdf",
     {"pages": (1, 2, 3), "name_x": 235, "title": r"NETWORK\s+TARIFFS"}),
    (ue_2000, UE + "2000-H2/uenetwork00.pdf", {}),
    (ue_2001, UE + "2001/2001_tariff_report.pdf", {}),
    (ue_nuos, UE + "2002/UE_Nuos2002.pdf", {}),
    (ue_nuos, UE + "2003/Nuos_2003.pdf", {}),
    (ue_nuos, UE + "2004/internetsitetariffs_2004.pdf", {}),
    (ue_nuos, UE + "2005/2005GSTExclusiveTariffs.pdf", {}),
    (ue_2006, UE + "2006/2006TariffReport.pdf", {}),
    (ue_coded, UE + "2007/070205NetworkTarrifs.pdf", {"page_no": 3}),  # (070124's prices, two codes renamed)
    (ue_coded, UE + "2008/UEDTariffSchedule2008.pdf", {"page_no": 5, "cont": (6,)}),
    (ue_coded, UE + "2009/2009_UED_Approved_Tariffs_-_DUOS_TUOS_NUOS.pdf", {"page_no": 3}),
    (ue_coded, UE + "2010/2010_UED_Rates_DUOS_TUOS_NUOS_and_PFIT.pdf", {"page_no": 3}),
    (ue_2011, UE + "2011/2011Tariff_Report.pdf",
     {"pages": (69, 70, 71, 72), "parts": (UE + "2011/2011_Tariff_Schedule.pdf", 1, 2)}),
    (ue_coded, UE + "2012/2012_Tariff_Schedule.pdf", {"page_no": 3}),
    (ue_coded, UE + "2013/2013_-_tariff_schedule.pdf", {"page_no": 5}),
    (ue_coded, UE + "2014/2014_-_tariff_schedule_v2.pdf", {"page_no": 5}),  # (v1's prices)
    # (2015_-_tariff_schedule.pdf's prices but for the seasonal demand tariff, renamed RESKWTOU)
    (ue_coded, UE + "2015/UE-tariff-schedule-2015.pdf", {"page_no": 5}),
    (ue_coded, UE + "2016/2016-Tariff-Schedule-1.pdf", {"page_no": 5}),
    (ue_coded, UE + "2017/2017-Tariff-Schedule.pdf", {"page_no": 5}),  # (v1's prices)
    (ue_coded, UE + "2018/2018-Tariff-Schedule.pdf", {"page_no": 7}),
    (tariff_schedule, CP + "2004/cp_network_tariffs_2004.pdf", {"page_no": 3, "code_re": CODE_CP}),
    (tariff_schedule, CP + "2005/CitiPower_2005_Network_Tariff_23-11-04.pdf", {"page_no": 3, "code_re": CODE_CP}),
    (tariff_schedule, CP + "2006/2006_CitiPower_Tariff_Schedule.pdf", {"page_no": 3, "code_re": CODE_CP}),
    (tariff_schedule, CP + "2007/2007_CitiPowerTariff_Schedule.pdf", {"page_no": 3, "code_re": CODE_CP}),
    (tariff_schedule, CP + "2008/2008_CitiPower_Tariff_Schedule.pdf", {"page_no": 3, "code_re": CODE_CP}),
    (tariff_schedule, CP + "2009/2009_CitiPower_Tariff_Schedule.pdf", {"page_no": 3, "code_re": CODE_CP}),
    (tariff_schedule, CP + "2010/2010_CitiPower_Tariff_Schedule_reformatted.pdf", {"page_no": 4, "code_re": CODE_CP,
                                                                     "note": "NUoS including the PFIT fee"}),
    (tariff_schedule, CP + "2011/2011_CP_Tariff_Schedule.pdf", {"page_no": 4, "code_re": CODE_CP}),
    (tariff_schedule, PC + "2004/network_tariffs_table2004.pdf", {"page_no": 3, "code_re": CODE_PC}),
    (tariff_schedule, PC + "2005/Powercor_2005_Network_Tariffs_23-11-04.pdf", {"page_no": 3, "code_re": CODE_PC}),
    (tariff_schedule, PC + "2006/Powercor_2006_Tariff_Schedule_201205.pdf", {"page_no": 3, "code_re": CODE_PC}),
    (tariff_schedule, PC + "2007/2007_Powercor_Tariff_Schedule.pdf", {"page_no": 3, "code_re": CODE_PC}),
    (tariff_schedule, PC + "2008/2008_Powercor_Tariff_Schedule.pdf", {"page_no": 3, "code_re": CODE_PC}),
    (tariff_schedule, PC + "2009/2009_Powercor_Tariff_Schedule.pdf", {"page_no": 3, "code_re": CODE_PC}),
    (tariff_schedule, PC + "2010/2010_Powercor_Tariff_Schedule.pdf", {"page_no": 4, "code_re": CODE_PC,
                                                              "note": "NUoS including the PFiT fee"}),
    (tariff_schedule, PC + "2011/2011_PAL_Tariff_Schedule.pdf", {"page_no": 4, "code_re": CODE_PC}),
    (tariff_schedule, CP + "2012/CitiPower_2012_Network_Tariffs_Approved.pdf",
     {"page_no": 1, "code_re": CODE_CP, **LATE}),
    (tariff_schedule, PC + "2012/Powercor_Australia_2012_Network_Tariff_Schedule_Approved.pdf",
     {"page_no": 1, "code_re": CODE_PC, **LATE}),
    (tariff_schedule, CP + "2013/CitiPower_2013_Tariff_Schedule.pdf", {"page_no": 1, "code_re": CODE_CP, **LATE}),
    (tariff_schedule, PC + "2013/Powercor_2013_Tariff_Schedule.pdf", {"page_no": 1, "code_re": CODE_PC, **LATE}),
    (tariff_schedule, CP + "2014/CitiPower_2014_Network_Pricing_Schedule.pdf",
     {"page_no": 1, "code_re": CODE_CP, **LATE}),
    (tariff_schedule, PC + "2014/Powercor_Australia_2014_Network_Tariff_Schedule.pdf",
     {"page_no": 1, "code_re": CODE_PC, **LATE}),
    (tariff_schedule, CP + "2015/citipower-2015-network-tariff-schedule.pdf",
     {"page_no": 1, "code_re": CODE_CP, **LATE}),
    (tariff_schedule, PC + "2015/powercor-2015-network-tariff-schedule.pdf",
     {"page_no": 1, "code_re": CODE_PC, **LATE}),
    (tariff_schedule, CP + "2016/citipower-2016-network-tariff-schedule.pdf",
     {"page_no": 1, "code_re": CODE_CP, "ocr": True, "parts": (2, 3, 4), "after_header": True,
      "header_re": r"December \d{4}$"}),
    (tariff_schedule, PC + "2016/powercor-2016-network-tariff-schedule.pdf",
     {"page_no": 1, "code_re": CODE_PC, **LATE}),
    (tariff_schedule, CP + "2017/2017-cp-network-tariff-schedule-website-version.pdf",
     {"page_no": 1, "code_re": CODE_CP, **LATE}),
    (tariff_schedule, PC + "2017/pal-2017-network-tariff-schedule-website-version.pdf",
     {"page_no": 1, "code_re": CODE_PC, **LATE}),
    (tariff_schedule, CP + "2018/2018-citipower-network-tariffs.pdf", {"page_no": 1, "code_re": CODE_CP, **LATE}),
    (tariff_schedule, PC + "2018/2018-powercor-network-tariffs.pdf", {"page_no": 1, "code_re": CODE_PC, **LATE}),
    # 2019: the joint 'Network Tariff Schedules 2016-2019' (image pages; OCR at 288 dpi misreads a 'Yes' column cell,
    # 360 reads every row). Every price equals Table A.1 of the AER-approved 2019 pricing proposal.
    (tariff_schedule, CP + "2019/Network-Tariff-Schedule-2016-2019.pdf", {"page_no": 2, "code_re": CODE_CP, **NTS19}),
    (tariff_schedule, PC + "2019/Network-Tariff-Schedule-2016-2019.pdf", {"page_no": 6, "code_re": CODE_PC, **NTS19}),
    # Provisional years (no distributor or state price list held): the AER-hosted approved document, else the latest
    # proposed one. 2020: Table A.1 of the AER-approved pricing proposal (CitiPower and Powercor print it as an image:
    # OCR, every price checked against the sum of the DUoS, TUoS and JUoS tables A.2-A.4 that follow it).
    (tariff_summary,
     UE + "2019/AER_approved_-_United_Energy_-_Attachment_B_Tariff_summary_-_last_updated_22_November_2018.xlsx",
     {"sheet": "UE_2019_NUOS"}),
    (tariff_schedule, CP + "2020/AER_-_Approved_CitiPower_Pricing_Proposal_2020_-_November_2019.pdf",
     {"page_no": 19, "code_re": CODE_CP, "parts": (20, 21, 22), **PP20}),
    (tariff_schedule, PC + "2020/AER_-_Approved_Powercor_PricingProposal_2020_-_November_2019.pdf",
     {"page_no": 20, "code_re": CODE_PC, "parts": (21, 22, 23), **dict(PP20, dpi=OCR_DPI)}),  # (360: 'D3' -> '80')
    (ue_2020, UE + "2020/AER_-_Approved_United_Energy_-_2020_Pricing_Proposal_-_November_2019.pdf", {"page_no": 19}),
    (tariff_summary, CP + "2021-H1/Attachment_B_-_HY_2021_Tariff_Summary_CP.xlsx", {"sheet": "CP_HY 2021_NUOS"}),
    (tariff_summary, PC + "2021-H1/Attachment_B_-_HY_2021_Tariff_Summary_PAL_0.xlsx", {"sheet": "PAL_HY 2021_NUOS"}),
    (tariff_summary, UE + "2021-H1/Attachment_B_HY_2021_Tariff_Summary_UE.xlsx", {"sheet": "UE_2021_NUOS"}),
    (tariff_summary, CP + "2021-22/CitiPower_-_Attachment_B_-_2021-22_Tariff_Summary_-_June_2021.xlsx",
     {"sheet": "CP_2021_22_NUOS"}),
    (tariff_summary, PC + "2021-22/Powercor_-_Attachment_B_-_2021-22_Tariff_Summary_-_June_2021_0.xlsx",
     {"sheet": "PAL_2021_22_NUOS"}),
    (tariff_summary, UE + "2021-22/United_Energy_-_Attachment_B_-2021-22_Tariff_Tables_-_June_2021.xlsx",
     {"sheet": "UE_2021-22_NUOS"}),
    (tariff_summary, CP + "2022-23/Attachment_B_-_CitiPower_-_2022-23_Tariff_Summary_-_6_April_2022.xlsx",
     {"sheet": "CP_2022_23_NUOS"}),
    (tariff_summary, PC + "2022-23/Attachment_B_-_Powercor_-_2022-23_Tariff_Summary_-_6_April_2022.xlsx",
     {"sheet": "PAL_2022_23_NUOS"}),
    (tariff_summary, UE + "2022-23/Attachment_B_United_Energy_2022-23_Tariff_Summary_-_6_April_2022.xlsx",
     {"sheet": "UE_2022-23_NUOS"}),
]


def main(argv=None):
    only = set((argv if argv is not None else sys.argv[1:]))
    rows = []
    for layout, path, kw in DOCS:
        if only and not any(o in path for o in only):
            continue
        try:
            got = layout(path, **kw)
        except SystemExit as e:
            raise SystemExit(f"{path}: {e}" if not str(e).startswith(path) else str(e))
        if not got:
            raise SystemExit(f"{path}: no prices read")
        LOG.append(f"{path}: {len(got)} rows, {len({r['tariff_code'] for r in got})} tariffs")
        rows += got
    for line in LOG:
        print(line)
    if not only:
        common.write(SLUG, rows)
    return rows


if __name__ == "__main__":
    main()
