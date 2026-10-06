"""Parse the AER-authored price files into the normalised long table (out/aer_long.csv).

Sources handled:
  * AER consolidated stakeholder report 2025-26 (v5) and 2026-27 (26 Aug 2026) - sheet 'Tariff schedule'
    (total network prices = distribution + transmission + jurisdictional scheme; excl GST).
  * AER per-DNSP stakeholder reports 2024-25 - sheet 'Tariff schedule' (Tariff schedule 3 = all prices,
    Tariff schedule 4 = site specific; each with Total / Distribution / DPPC / JSA sections).
There is no AER-authored price file for 2023-24 (the AER only hosts the DNSP-submitted documents); those are
parsed by the DNSP parsers with side=AER_HOSTED.
"""
import csv, re, sys, os
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
from schema import COLUMNS, charge_type_from_label, time_band_from_label, season_from_label
from units import to_std
from published import cell_value
from tariffdb import locators

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

CONSOLIDATED = [
    ("2025-26", "sources/aer/AER_Consolidated_stakeholder_report_2025-26_v5.xlsx",
     "https://www.aer.gov.au/system/files/2025-05/AER%20-%20Consolidated%20stakeholder%20report%202025%E2%80%9326%20v5%C2%A0.xlsx"),
    ("2026-27", "sources/aer/AER_Consolidated_stakeholder_report_2026-27_26Aug2026.xlsx",
     "https://www.aer.gov.au/system/files/2026-08/AER%20%E2%80%93%202026%E2%80%9327%20%E2%80%93%20Consolidated%20stakeholder%20report%20%E2%80%93%2026%20August%202026.xlsx"),
]

# Superseded versions: not used by the reconciliation (which compares the final AER version), written to
# out/aer_versions_long.csv for the tariff database (scripts/tariffdb/build.py) so proposed and approved prices coexist.
SUPERSEDED = [
    ("2025-26", "sources/aer/AER_Consolidated_stakeholder_report_2025-26_v1_wayback.xlsx",
     "https://web.archive.org/web/20250409020106id_/https://www.aer.gov.au/system/files/2025-04/Consolidated%C2%A0stakeholder%20report%202025%E2%80%9326.xlsx",
     "AER consolidated stakeholder report v1 (8 Apr 2025): proposed prices"),
]

STAKEHOLDER_2024_25 = {
    "Ausgrid": ("AER_Stakeholder_report_Ausgrid_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Ausgrid%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "AusNet Services": ("AER_Stakeholder_report_AusNet_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20AusNet%20Services%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "CitiPower": ("AER_Stakeholder_report_CitiPower_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20CitiPower-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Endeavour Energy": ("AER_Stakeholder_report_Endeavour_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Endeavour%20Energy%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Energex": ("AER_Stakeholder_report_Energex_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Energex%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Ergon Energy": ("AER_Stakeholder_report_Ergon_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Ergon%20Energy%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Essential Energy": ("AER_Stakeholder_report_Essential_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Essential%20Energy%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Evoenergy": ("AER_Stakeholder_report_Evoenergy_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Evoenergy%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Jemena": ("AER_Stakeholder_report_Jemena_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Jemena%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Power and Water Corporation": ("AER_Stakeholder_report_PWC_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Power%20and%20Water%20Corporation%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "Powercor": ("AER_Stakeholder_report_Powercor_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20Powercor%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "SA Power Networks": ("AER_Stakeholder_report_SAPN_2024-25_updated17Jul2024.xlsx", "https://www.aer.gov.au/system/files/2024-07/AER%20-%20Stakeholder%20report%20-%20SA%20Power%20Networks%20-%20%202024%E2%80%9325%20Annual%20Pricing%20Proposal%20%28updated%2017%20July%202024%29.xlsx"),
    "TasNetworks": ("AER_Stakeholder_report_TasNetworks_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20TasNetworks%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
    "United Energy": ("AER_Stakeholder_report_UnitedEnergy_2024-25.xlsx", "https://www.aer.gov.au/system/files/2024-05/AER%20-%20Stakeholder%20report%20-%20United%20Energy%20-%202024%E2%80%9325%20Annual%20Pricing%20Proposal.xlsx"),
}

SECTION_BASIS = {
    "Total network prices": "NUoS",
    "Distribution": "DUoS",
    "Designated Pricing Proposal Costs": "DPPC",
    "Jurisdictional Scheme Amounts": "JSA",
}


def nz(v):
    return v is not None and v != "" and not (isinstance(v, (int, float)) and v == 0)


def row(dnsp, fy, code, name, cls, comp, unit, val, basis, src, url, note="", side="AER", locator=""):
    vs, us = to_std(val, unit, comp)
    ct = charge_type_from_label(comp, unit)
    if dnsp == "Ausgrid" and str(code).strip().rstrip("*") == "EA029" and ct == "energy":
        ct = "export"
    return {
        "side": side, "distributor": dnsp, "fin_year": fy, "tariff_code": str(code).strip(), "tariff_name": (name or "").strip() if isinstance(name, str) else str(name or ""),
        "customer_class": cls or "", "component": comp.strip(), "charge_type": ct,
        "time_band": time_band_from_label(comp), "season": season_from_label(comp) or season_from_label(unit), "unit": unit, "value": val,
        "value_std": vs, "unit_std": us, "gst": "excl", "basis": basis, "source_file": src, "source_url": url, "note": note,
        "locator": locator,
    }


def consolidated_columns(ws, hr):
    """Columns of one distributor block, found by header label: the layout moves between versions (2025-26 v1 has
    'Tariff class' and 'Code' in D/E and prices from G; v5 and 2026-27 have 'Tariff code' (SA Power Networks: 'Code SA')
    in D and prices from H)."""
    labels = {c: ws.cell(hr, c).value.strip() for c in range(3, ws.max_column + 1) if isinstance(ws.cell(hr, c).value, str)}
    code_col = next(c for c, v in labels.items() if v in ("Tariff code", "Code", "Code SA"))
    top_col = next(c for c, v in labels.items() if v == "Top")
    return code_col, top_col + 1


def parse_consolidated(fy, path, url, note_prefix=""):
    wb = openpyxl.load_workbook(os.path.join(ROOT, path), data_only=True)
    ws = wb["Tariff schedule"]
    out = []
    r = 1
    maxr = ws.max_row
    blocks = []
    for rr in range(1, maxr + 1):
        b = ws.cell(rr, 2).value
        if isinstance(b, str) and re.search(r"\d{4}.\d{2} network prices$", b.strip()):
            blocks.append(rr)
    for bi, hr in enumerate(blocks):
        dnsp = re.sub(r"\s+\d{4}.\d{2} network prices$", "", ws.cell(hr, 2).value.strip())
        code_col, first_price_col = consolidated_columns(ws, hr)
        comps = {}
        for c in range(first_price_col, ws.max_column + 1):
            lab = ws.cell(hr, c).value
            if isinstance(lab, str) and lab.strip():
                comps[c] = (lab.strip(), (ws.cell(hr + 1, c).value or "").strip() if isinstance(ws.cell(hr + 1, c).value, str) else str(ws.cell(hr + 1, c).value or ""))
        end = blocks[bi + 1] if bi + 1 < len(blocks) else maxr + 1
        for rr in range(hr + 2, end):
            code = ws.cell(rr, code_col).value
            name = ws.cell(rr, 3).value
            if ws.cell(rr, 2).value == "End":
                break
            if code == "#REF!":
                code = ""  # 2025-26 v1: the code column is a broken formula; the row is identified by its AER tariff ID
            if not nz(code) and not (isinstance(name, str) and name.strip()):
                continue
            if code in (None, "", 0) and isinstance(name, str):
                code = ""
            aer_id = ws.cell(rr, 2).value if isinstance(ws.cell(rr, 2).value, str) else ""
            any_val = False
            for c, (lab, unit) in comps.items():
                v = ws.cell(rr, c).value
                if isinstance(v, (int, float)) and v != 0:
                    any_val = True
                    note = "; ".join(x for x in (note_prefix, f"aer_id={aer_id}" if aer_id else "") if x)
                    out.append(row(dnsp, fy, code, name, "", lab, unit, cell_value(ws.cell(rr, c)), "NUoS", path, url, note,
                                   locator=locators.xlsx(ws, ws.cell(rr, c))))
            # a superseded version without codes (v1) keeps its zero-priced rows by AER tariff ID
            if not any_val and (nz(code) or (note_prefix and aer_id)):
                note = "; ".join(x for x in (note_prefix, "zero-priced row", f"aer_id={aer_id}" if aer_id else "") if x)
                out.append(row(dnsp, fy, code, name, "", "(no non-zero components)", "", 0, "NUoS", path, url, note,
                               locator=locators.xlsx(ws, ws.cell(rr, code_col if nz(code) else 2))))
    return out


def parse_stakeholder_2024_25(dnsp, fname, url):
    path = f"sources/aer/2024-25_stakeholder_reports/{fname}"
    wb = openpyxl.load_workbook(os.path.join(ROOT, path), data_only=True)
    ws = wb["Tariff schedule"]
    out = []
    schedules = []
    for rr in range(1, ws.max_row + 1):
        b = ws.cell(rr, 2).value
        if isinstance(b, str) and re.match(r"Tariff schedule \d", b):
            schedules.append((rr, b))
    for si, (hr, title) in enumerate(schedules):
        if "Metering" in title or "DMO" in title or "VDO" in title:
            continue  # metering handled elsewhere; DMO/VDO schedules repeat a subset of schedule 3
        note_sched = "site-specific" if "Site specific" in title else ""
        comps = {}
        for c in range(9, ws.max_column + 1):
            lab = ws.cell(hr, c).value
            if isinstance(lab, str) and lab.strip():
                u = ws.cell(hr + 1, c).value
                comps[c] = (lab.strip(), u.strip() if isinstance(u, str) else "")
        end = schedules[si + 1][0] if si + 1 < len(schedules) else ws.max_row + 1
        basis = "NUoS"
        for rr in range(hr + 2, end):
            cval = ws.cell(rr, 3).value
            code = ws.cell(rr, 5).value
            if isinstance(cval, str) and cval.strip() in SECTION_BASIS and not nz(code):
                basis = SECTION_BASIS[cval.strip()]
                continue
            if not nz(code):
                continue
            name = cval if isinstance(cval, str) else ""
            cls = ws.cell(rr, 4).value or ""
            other = " / ".join(str(ws.cell(rr, c).value) for c in (6, 7) if nz(ws.cell(rr, c).value))
            note = "; ".join(x for x in (note_sched, f"other_id={other}" if other else "") if x)
            any_val = False
            for c, (lab, unit) in comps.items():
                v = ws.cell(rr, c).value
                if isinstance(v, (int, float)) and v != 0:
                    any_val = True
                    out.append(row(dnsp, "2024-25", code, name, cls, lab, unit, cell_value(ws.cell(rr, c)), basis, path, url, note,
                                   locator=locators.xlsx(ws, ws.cell(rr, c))))
            if not any_val and basis == "NUoS":
                out.append(row(dnsp, "2024-25", code, name, cls, "(no non-zero components)", "", 0, basis, path, url, ("zero-priced row; " + note).strip("; "),
                               locator=locators.xlsx(ws, ws.cell(rr, 5))))
    return out


def main():
    rows = []
    for fy, path, url in CONSOLIDATED:
        rows += parse_consolidated(fy, path, url)
    for dnsp, (fname, url) in STAKEHOLDER_2024_25.items():
        rows += parse_stakeholder_2024_25(dnsp, fname, url)
    os.makedirs(os.path.join(ROOT, "out"), exist_ok=True)
    with open(os.path.join(ROOT, "out/aer_long.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    versions = []
    for fy, path, url, note in SUPERSEDED:
        versions += parse_consolidated(fy, path, url, note_prefix=note)
    with open(os.path.join(ROOT, "out/aer_versions_long.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(versions)
    print("superseded-version rows", len(versions))
    # summary
    from collections import Counter
    c = Counter((r["fin_year"], r["distributor"], r["basis"]) for r in rows)
    uniq = {}
    for r in rows:
        uniq.setdefault((r["fin_year"], r["distributor"], r["basis"]), set()).add(r["tariff_code"])
    for k in sorted(uniq):
        print(f"{k[0]} {k[1]:30s} {k[2]:5s} tariffs={len(uniq[k]):4d} rows={c[k]}")
    print("total rows", len(rows))


if __name__ == "__main__":
    main()
