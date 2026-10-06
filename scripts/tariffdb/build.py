"""Build the tariff database under data/tariffdb/: one CSV per table (tables/*.csv), the DDL for SQLite and PostgreSQL
(schema.sqlite.sql, schema.postgres.sql), the PostgreSQL import script (load.postgres.sql) and the machine-readable
table spec (schema.json).

Inputs:
  - the parser outputs out/aer_long.csv, out/aer_versions_long.csv and out/dnsp/*.csv (every row carries a locator);
  - the AER Metering worksheet / 2024-25 'Tariff schedule 1' and the Energex/Ergon per-tariff Metering blocks (read here);
  - the curated facts in data/tariffdb/curated/*.yaml (TOU windows, demand rules, eligibility, relations, flags);
  - the document registry (scripts/tariffdb/build_support.py over sources/inventory.csv).

Output is deterministic (rows sorted by primary key; keys derived from content), so a rebuild from the same inputs is
byte-identical and a changed source shows up as added rows, never as rewritten ones.

  .venv/bin/python scripts/tariffdb/build.py                           rebuild data/tariffdb/
  .venv/bin/python scripts/tariffdb/build.py --out DIR                 write the same files under DIR instead
  .venv/bin/python scripts/tariffdb/build.py --check-append-only REF   fail when a row of a source-fact table committed
                                                                       at git REF was changed or removed (history is
                                                                       append-only; derived tables/columns, which are
                                                                       recomputed from the facts, are not compared)
"""
import argparse
import csv
import functools
import glob
import hashlib
import io
import json
import math
import os
import re
import statistics
import subprocess
import sys
from collections import Counter, defaultdict
from decimal import Decimal

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, SCRIPTS)
import adjustments  # noqa: E402
import build_support as bs  # noqa: E402
import curated  # noqa: E402
import locators  # noqa: E402
import rates  # noqa: E402
import spec  # noqa: E402
from published import cell_value  # noqa: E402
from reconcile import code_alts, code_key, digits, sim  # noqa: E402
from schema import REPEATED_PRINTING  # noqa: E402
from units import to_std  # noqa: E402

ROOT = bs.ROOT
OUT_DIR = os.path.join(ROOT, "data", "tariffdb")
PARSER_OUTPUTS = ["out/aer_long.csv", "out/aer_versions_long.csv"]  # + out/dnsp/*.csv
DNSP_SIDES = ("DNSP", "AER_HOSTED")
PLACEHOLDER = "(no non-zero components)"

# Distributor-years whose distributor daily charge = AER daily charge + metering, and the Evoenergy LFiT adders: the
# documented adjustments of scripts/adjustments.py (shared with the reconciliation). The price_adjustment_tariff rows
# below are computed, not listed, and the tests re-derive every one.
METERING_ADDERS = [(bs.ID_BY_NAME[name], fy) for name, (years, *_rest) in adjustments.METERING.items() for fy in years]
ESSENTIAL_METERING_NOTE = adjustments.ESSENTIAL_METERING_NOTE
ESSENTIAL_NOTE_DOCS = adjustments.ESSENTIAL_NOTE_DOCS
# Evoenergy: what each document says about the ACT Large-scale Feed-in Tariff (LFiT); (path, locator, quote, status)
EVO_LFIT_DOCS = {
    "sources/aer/2023-24_price_lists/Evoenergy_2023-24_Electricity_network_pricing_proposal_5May2023.pdf": (
        "sources/dnsp/evoenergy/Evoenergy_Statement_of_Tariff_Classes_and_Tariffs_2023-24.pdf", "pdf:p4",
        "Evoenergy's 2023/24 regulated electricity network prices approved by the Australian Energy Regulator (AER) do "
        "not include any amounts for the Australian Capital Territory (ACT) Government's Large-scale Feed-in Tariff "
        "(LFiT) scheme.", "excluded"),
    "sources/dnsp/evoenergy/Evoenergy_Statement_of_Tariff_Classes_and_Tariffs_2023-24.pdf": (
        "sources/dnsp/evoenergy/Evoenergy_Statement_of_Tariff_Classes_and_Tariffs_2023-24.pdf", "pdf:p4",
        "The prices presented in this document include the LFiT rebate and are therefore different from the 2023/24 "
        "network charges approved by the AER.", "rebate"),
    **{doc: (doc, loc, quote, "included") for _amount, doc, loc, quote in adjustments.LFIT_ADDERS.values()},
    "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2026-27_AER_approved_April2026.xlsx": (
        "sources/dnsp/evoenergy/Evoenergy_Schedule_of_Charges_2026-27_AER_approved_April2026.xlsx",
        "xlsx:Network tariffs!B6",
        "It presents Evoenergy's 2026-27 network charges approved by the Australian Energy Regulator (AER), which do not "
        "include costs for the ACT Government's Large-scale Feed-in Tariff (LFiT) Scheme.", "excluded"),
}
AER_LFIT_QUOTE = "Do not include Evoenergy's application of costs related to the ACT large-scale feed-in tariff scheme."
LFIT_ADDERS = [(fy, amount, doc) for fy, (amount, doc, _loc, _quote) in sorted(adjustments.LFIT_ADDERS.items())]
LFIT_REBATE_QUOTE = ("The LFiT rebate has been applied as a negative adjustment to the AER's approved charges for 2023/24 "
                     "and is equivalent to a reduction of 2.27 cents per kilowatt-hour (kWh) excluding Goods and "
                     "Services Tax (GST), on average, across Evoenergy's tariffs.")
# Pricing PDFs whose price tables are in a separate Tariff Summary workbook that is not retrievable: (path, locator, quote)
PRICE_ATTACHMENTS_NOT_HELD = [
    ("sources/dnsp/citipower/CitiPower_Pricing_Proposal_2025-26_31Mar2025_wayback.pdf", "pdf:p5",
     "Attachment ‘CitiPower - 2025-26 Tariff Summary’ sets out the 2025/26 network tariff pricing schedule"),
    ("sources/aer/dnsp_copies/UnitedEnergy_2025-26_Final_Pricing_31Mar2025.pdf", "pdf:p5",
     "Attachment United Energy - 2025-26 Tariff Summary - 31 March 2025 sets out the 2025/26 network tariff pricing "
     "schedule"),
]
STATUS_FLAGS = [  # (flag, pattern) matched in published tariff names and parser notes
    ("withdrawn", r"\bwithdrawn\b"), ("closed_to_new", r"closed to new|\bclosed\b|not available (?:on application|for new)"),
    ("obsolete", r"\bobsolete\b"), ("grandfathered", r"grandfather"), ("trial", r"(?<![-\w])trial\b(?![- ]rebate)"),
    ("transitional", r"\btransitional\b"), ("site_specific", r"site[- ]specific"), ("indicative", r"\bindicative\b"),
    # AER 2024-25 'Tariff schedule 2 | 2024–25 DMO tariffs' / '... VDO tariffs': tariffs the default market offer
    # (DMO) or Victorian default offer (VDO) is set on
    ("dmo_vdo_tariff", r"^Tariff schedule 2 \| \S+ (?:DMO|VDO) tariffs"),
]


def rel(p):
    return os.path.relpath(p, ROOT)


def short_hash(*parts):
    return hashlib.sha1("\x1f".join(str(p) for p in parts).encode()).hexdigest()[:10]


def decimals(s):
    try:
        return max(0, -Decimal(str(s).strip()).as_tuple().exponent)
    except Exception:
        return 0


def unique_ids(items):
    """Ids from content, never from row order. items: [(base_id, [disambiguators in order of preference])]. An item
    keeps its base id when no other item shares it; items sharing a base get the first disambiguator that tells them
    all apart (as a slug), else a hash of all of them; items whose content is identical are numbered, which is
    order-free because they are interchangeable."""
    groups = defaultdict(list)
    for i, (base, _) in enumerate(items):
        groups[base].append(i)
    out = [None] * len(items)
    for base, idx in groups.items():
        if len(idx) == 1:
            out[idx[0]] = base
            continue
        for j in range(len(items[idx[0]][1])):
            vals = [bs.slug(str(items[i][1][j] or ""))[:60] for i in idx]
            if all(vals) and len(set(vals)) == len(vals):
                for i, v in zip(idx, vals):
                    out[i] = f"{base}/{v}"
                break
        else:
            seen = Counter()
            for h, i in sorted((short_hash(*items[i][1]), i) for i in idx):
                seen[h] += 1
                out[i] = f"{base}/{h}" + (f"-{seen[h]}" if seen[h] > 1 else "")
    return out


def num_text(x):
    """Canonical text for a computed number: rounded to 15 significant digits, which drops binary-float noise
    (205.79000000000002 -> 205.79) while keeping every digit a double carries reliably."""
    if x is None:
        return None
    d = Decimal(repr(float(x)))
    if not d:
        return "0"
    return format(d.quantize(Decimal(1).scaleb(d.adjusted() - 14)).normalize(), "f")


@functools.lru_cache(maxsize=None)
def document_prints(path, code):
    """Does the document print `code` as a word anywhere (any PDF page as text, any spreadsheet cell)?"""
    if path.endswith(".xlsx"):
        wb = locators._workbook(path)
        return any(locators.verify_quote(path, locators.xlsx(ws, c), code)[0] for ws in wb.worksheets
                   for row in ws.iter_rows() for c in row if isinstance(c.value, str) and code in c.value)
    page = 1
    while True:
        ok, why = locators.verify_quote(path, locators.pdf(page), code)
        if ok:
            return True
        if "beyond end" in why:
            return False
        page += 1


class Tables:
    def __init__(self):
        self.rows = {t: {} for t in spec.TABLE_ORDER}
        self.pk = {t["name"]: [c["name"] for c in t["columns"] if c["primary_key"]] for t in spec.TABLES}
        self.cols = {t["name"]: [c["name"] for c in t["columns"]] for t in spec.TABLES}

    def add(self, table, row, replace=False, same_ok=False):
        extra = set(row) - set(self.cols[table])
        if extra:
            raise KeyError(f"{table}: unknown columns {sorted(extra)}")
        full = {c: (None if row.get(c) == "" else row.get(c)) for c in self.cols[table]}
        key = tuple(full[c] for c in self.pk[table])
        if key in self.rows[table] and not replace:
            if same_ok and self.rows[table][key] == full:
                return full
            raise ValueError(f"duplicate {table} key {key}:\n  {self.rows[table][key]}\n  {full}")
        self.rows[table][key] = full
        return full

    def get(self, table, *key):
        return self.rows[table].get(tuple(key))

    def all(self, table):
        return list(self.rows[table].values())


def load_parser_rows():
    rows = []
    for p in PARSER_OUTPUTS + sorted(rel(x) for x in glob.glob(os.path.join(ROOT, "out", "dnsp", "*.csv"))):
        with open(os.path.join(ROOT, p), newline="", encoding="utf-8") as f:
            for i, r in enumerate(csv.DictReader(f), 2):
                r["_file"], r["_line"] = p, i
                rows.append(r)
    return rows


def region_of(r):
    m = re.search(r"zone: ([^;]+)", r["note"] or "") or re.search(r"/\s*(T[1-4])\b", r["note"] or "")
    return m.group(1).strip() if m else None


def aer_id_of(r):
    m = re.search(r"aer_id=([^;\s]+)", r["note"] or "")
    return m.group(1) if m else None


def quantity_period(unit_std):
    u = unit_std or ""
    if u == "c/kWh":
        return "kWh", "none", 0
    if u == "c/kVAh":
        return "kVAh", "none", 0
    if u == "c/day":
        return "none", "day", 0
    m = re.match(r"^c/(kW|kVA|k\?|lamp)/(day|month|year|season|\?)(\??)$", u)
    if not m:
        raise ValueError(f"unknown standard unit {unit_std!r}")
    q = {"kW": "kW", "kVA": "kVA", "k?": "kW_or_kVA", "lamp": "lamp"}[m.group(1)]
    if m.group(2) == "?":
        return q, "unstated", 0
    return q, m.group(2), 1 if m.group(3) else 0


class Builder:
    def __init__(self):
        self.t = Tables()
        self.docs = bs.documents()
        self.doc_by_id = {d["document_id"]: d for d in self.docs}
        self.doc_by_path = {d["local_path"]: d for d in self.docs if d["local_path"]}
        self.dist = {d["distributor_id"]: d for d in bs.DISTRIBUTORS}
        self.rows = load_parser_rows()
        self.listing_rows = defaultdict(list)    # listing_id -> parser rows
        self.listings_by_doc_tariff = defaultdict(list)
        self.alias_to_tariff = defaultdict(set)  # (did, alias label key) -> tariff ids
        self.instances = []

    # ------------------------------------------------------------------ reference data
    def reference(self):
        for fy, (a, b) in bs.FIN_YEAR_DATES.items():
            self.t.add("financial_year", {"fin_year": fy, "start_date": a, "end_date": b})
        for d in bs.DISTRIBUTORS:
            self.t.add("distributor", dict(d))
        for s in bs.series_rows(self.docs):
            self.t.add("document_series", s)
        cols = set(self.t.cols["source_document"])
        for d in self.docs:
            self.t.add("source_document", {k: v for k, v in d.items() if k in cols})
        for c in bs.version_coverage(self.docs):
            self.t.add("document_coverage", c)

    # ------------------------------------------------------------------ tariffs and listings
    def tariffs_and_listings(self):
        did_of = {d["name"]: d["distributor_id"] for d in bs.DISTRIBUTORS}
        for r in self.rows:
            r["_did"] = did_of[r["distributor"]]
            if r["source_file"] not in self.doc_by_path:
                raise SystemExit(f"{r['_file']}:{r['_line']}: source {r['source_file']} not in the document registry")
            r["_doc"] = self.doc_by_path[r["source_file"]]
        # canonical codes: what distributors themselves publish (own site or AER-hosted copies), every year
        published = defaultdict(Counter)
        for r in self.rows:
            if r["side"] in DNSP_SIDES and r["tariff_code"].strip():
                published[(r["_did"], code_key(r["tariff_code"]))][r["tariff_code"].strip()] += 1
        self.universe = defaultdict(set)
        for (did, k), forms in published.items():
            self.universe[did].add(k)
            form = sorted(forms.items(), key=lambda kv: (-kv[1], kv[0]))[0][0]
            self.t.add("tariff", {"tariff_id": f"{did}:{k}", "distributor_id": did, "tariff_code": form,
                                  "identity_basis": "distributor_code"})
        groups = defaultdict(list)
        for r in self.rows:
            key = (r["source_file"], r["tariff_code"].strip(), r["tariff_name"].strip(), r["customer_class"].strip(),
                   region_of(r), aer_id_of(r) if r["side"] == "AER" else None)
            groups[key].append(r)
        # a GST-inclusive table restates the exclusive one: its rows join the exclusive listing of the same code even
        # where the two printings word the tariff's name or class differently (each such charge notes the wording)
        excl_groups = defaultdict(list)
        for key, rs in groups.items():
            if key[1] and any(r["gst"] == "excl" for r in rs):
                excl_groups[(key[0], key[1], key[4], key[5])].append(key)
        for key in sorted(groups, key=lambda k: tuple("" if x is None else x for x in k)):
            rs = groups[key]
            target = [k for k in excl_groups.get((key[0], key[1], key[4], key[5]), []) if k != key]
            if not key[1] or len(target) != 1 or any(r["gst"] != "incl" for r in rs):
                continue
            t = target[0]
            wording = [f"{what} as {mine!r}" for what, mine, theirs in (("name", key[2], t[2]), ("class", key[3], t[3]))
                       if mine != theirs]
            for r in rs:
                r["note"] = "; ".join(filter(None, [r["note"], "the GST-inclusive table prints the tariff "
                                                    + " and ".join(wording)]))
            groups[t].extend(rs)
            del groups[key]
        # AER aer_id -> code in the latest version of the same year, to resolve the codeless 2025-26 v1 rows
        latest_code = {}
        latest_name = defaultdict(set)
        for (src, code, name, cls, region, aid), rs in groups.items():
            d = rs[0]["_doc"]
            if d["author"] == "AER" and d["document_type"] == "aer_consolidated_stakeholder_report" and code:
                if d["version_seq"] == max(x["version_seq"] for x in self.docs if x["series_id"] == d["series_id"]
                                           and x["retrieval_status"] == "retrieved"):
                    if aid:
                        latest_code[(d["fin_year"], rs[0]["_did"], aid)] = code
                    latest_name[(d["fin_year"], rs[0]["_did"], name)].add(code)
        # document -> AER labels resolved directly, for the name-matching pass
        resolved = {}
        for key in sorted(groups, key=lambda k: tuple("" if x is None else x for x in k)):
            rs = groups[key]
            src, code, name, cls, region, aid = key
            did, d = rs[0]["_did"], rs[0]["_doc"]
            if d["author"] != "AER":
                resolved[key] = [(f"{did}:{code_key(code)}", None, None, code)]
                continue
            label, via = code, None
            if not code:  # AER 2025-26 v1 prints no codes: resolve its AER tariff ID through the approved version
                code = latest_code.get((d["fin_year"], did, aid))
                via = "aer_tariff_id"
                if code is None and len(latest_name.get((d["fin_year"], did, name), ())) == 1:
                    code = next(iter(latest_name[(d["fin_year"], did, name)]))
                    via = "aer_tariff_id_by_name"
                if code is None:
                    resolved[key] = [(f"{did}:{aid}", "aer_tariff_id", "AER tariff ID with no code in any version", aid)]
                    self.ensure_tariff(did, aid, "aer_tariff_id")
                    continue
            members = self.resolve_aer_label(did, rs[0]["distributor"], code, rs[0]["note"])
            out = []
            for tid, kind, note in members:
                if via:
                    note = (f"v1 row with AER ID {aid}: code {code} from the same ID in the approved version"
                            if via == "aer_tariff_id" else
                            f"v1 row with AER ID {aid}: that ID is absent from the approved version; matched by the "
                            f"unique tariff name {name!r} (code {code})") + (f"; {note}" if note else "")
                    kind = kind or "aer_tariff_id"
                out.append((tid, kind, note, label))
            resolved[key] = out
        # name-matching pass for AER labels that matched nothing (same rule as the reconciliation)
        for key, members in resolved.items():
            src, code, name, cls, region, aid = key
            tid, kind, note, label = members[0]
            if kind != "aer_label_identity":
                continue
            d = groups[key][0]["_doc"]
            did = groups[key][0]["_did"]
            v1_note = None
            if not code and latest_code.get((d["fin_year"], did, aid)):  # codeless v1 row, code from its AER ID
                code = latest_code[(d["fin_year"], did, aid)]
                v1_note = f"v1 row with AER ID {aid}: code {code} from the same ID in the approved version"
            direct = {m[0] for k2, ms in resolved.items() if k2[0] == src for m in ms if m[1] != "aer_label_identity"}
            cands = set()
            for k2, ms in resolved.items():
                d2 = groups[k2][0]["_doc"]
                if d2["author"] == "AER" or d2["distributor_id"] != did or d2["fin_year"] != d["fin_year"]:
                    continue
                n2 = k2[2]
                if not name or not n2:
                    continue
                sc = sim(re.sub(r"\(.*?\)|\*", "", name), re.sub(r"\(.*?\)|\*", "", n2))
                if sc >= 0.9 and digits(name) == digits(n2) and digits(code) == digits(k2[1]) and ms[0][0] not in direct:
                    cands.add((ms[0][0], round(sc, 2)))
            if len({c[0] for c in cands}) == 1:
                t2, sc = sorted(cands)[0]
                resolved[key] = [(t2, "name_matched", f"AER label {code} matched by tariff name (similarity {sc})", label)]
            else:
                resolved[key] = [(self.ensure_tariff(did, code or aid, "aer_label" if code else "aer_tariff_id"), None,
                                  "; ".join(filter(None, [v1_note, "AER label that no distributor document uses"])),
                                  label)]
        # listings; a document that prints one label on several rows gets one listing per row, told apart by content
        entries = []
        for key in sorted(groups, key=lambda k: tuple("" if x is None else x for x in k)):
            src, code, name, cls, region, aid = key
            for tid, kind, note, label in resolved[key]:
                base = f"{groups[key][0]['_doc']['document_id']}/{label or aid}"
                if len(resolved[key]) > 1:
                    base += f"/{tid.split(':', 1)[1]}"
                entries.append((key, (tid, kind, note, label), base, [aid, region, cls, name]))
        lids = unique_ids([(base, opts) for _, _, base, opts in entries])
        for (key, member, _, _), lid in zip(entries, lids):
            rs = groups[key]
            src, code, name, cls, region, aid = key
            d = rs[0]["_doc"]
            fy = rs[0]["fin_year"]
            joint = len(resolved[key]) > 1
            tid, kind, note, label = member
            start, end = bs.FIN_YEAR_DATES[fy]
            priced = [r for r in rs if not r["component"].startswith(PLACEHOLDER)]
            self.t.add("tariff_listing", {
                "listing_id": lid, "document_id": d["document_id"], "tariff_id": tid, "code_published": code or None,
                "name_published": name or None, "class_published": cls or None, "region": region,
                "effective_from": start, "effective_to": end,
                "price_availability": "priced" if priced else "placeholder",
                "locator": rs[0]["locator"], "note": listing_note(rs),
            })
            self.listing_rows[lid] = priced
            self.listings_by_doc_tariff[(d["document_id"], tid)].append(lid)
            if kind and kind != "aer_label_identity":
                alias_label = label or aid
                akind = "joint_label_member" if joint else kind
                if akind == "aer_tariff_id_by_name":
                    akind = "aer_tariff_id"
                self.t.add("tariff_alias", {
                    "alias_id": f"{d['document_id']}/{akind}/{alias_label}/{tid}", "tariff_id": tid,
                    "alias_label": alias_label, "alias_kind": akind, "document_id": d["document_id"], "note": note,
                }, same_ok=True)
                self.alias_to_tariff[(rs[0]["_did"], code_key(alias_label))].add(tid)
            if aid and d["author"] == "AER" and aid != (label or aid):
                self.t.add("tariff_alias", {
                    "alias_id": f"{d['document_id']}/aer_tariff_id/{aid}/{tid}", "tariff_id": tid, "alias_label": aid,
                    "alias_kind": "aer_tariff_id", "document_id": d["document_id"],
                    "note": "AER tariff ID (column B of the AER Tariff schedule)"}, same_ok=True)
            if joint:
                self.flag(lid, "joint_label_member", "document",
                          f"the AER row is labelled {label!r}, which names several distributor tariffs")
                self.instance("joint_code_label", did=rs[0]["_did"], fy=fy, tariff=tid, listing=lid,
                              doc=d["document_id"], detail=f"AER label {label!r} -> {tid}; one listing per member, "
                              f"charges repeated from the same cells")
            elif kind in ("code_variant", "regional_suffix", "name_matched"):
                self.instance("code_label_quirk", did=rs[0]["_did"], fy=fy, tariff=tid, listing=lid,
                              doc=d["document_id"], detail=f"{kind}: {note or label}")
        for (src, code, name, cls, region, aid), members in resolved.items():
            if any(m[2] and "absent from the approved version" in m[2] for m in members):
                d = groups[(src, code, name, cls, region, aid)][0]["_doc"]
                self.instance("aer_id_changed_between_versions", did=groups[(src, code, name, cls, region, aid)][0]["_did"],
                              fy=d["fin_year"], tariff=members[0][0], doc=d["document_id"],
                              detail=members[0][2])

    def ensure_tariff(self, did, code, basis):
        k = code_key(code)
        tid = f"{did}:{k}"
        if not self.t.get("tariff", tid):
            self.t.add("tariff", {"tariff_id": tid, "distributor_id": did, "tariff_code": code.strip(),
                                  "identity_basis": basis})
        return tid

    def resolve_aer_label(self, did, dname, label, note):
        """[(tariff_id, alias_kind | None | 'aer_label_identity', note)] for an AER code label."""
        U = self.universe[did]
        k = code_key(label)
        if k in U:
            return [(f"{did}:{k}", "code_variant" if label.strip() != self.t.get("tariff", f"{did}:{k}")["tariff_code"]
                     else None, "AER prints the code differently" if label.strip() != self.t.get(
                         "tariff", f"{did}:{k}")["tariff_code"] else None)]
        if re.search(r"[/,]", k):
            parts = [p.rstrip("*") for p in re.split(r"[/,]", k) if p]
            members = [p for p in parts if p in U]
            if members:
                return [(f"{did}:{p}", "joint_label_member", f"member {p} of joint AER label {label!r}") for p in members]
        for alt in code_alts(label, note, dname)[1:]:
            if alt in U:
                kind = "regional_suffix" if re.sub(r"T[1-4]$", "", alt) == re.sub(r"T[1-4]$", "", k) and alt != k \
                    else "code_variant"
                return [(f"{did}:{alt}", kind, f"AER label {label!r} -> distributor code {alt}")]
        return [(None, "aer_label_identity", None)]

    # ------------------------------------------------------------------ charges
    def charges(self):
        for lid, rs in sorted(self.listing_rows.items()):
            listing = self.t.get("tariff_listing", lid)
            d = self.doc_by_id[listing["document_id"]]
            did = d["distributor_id"] or self.t.get("tariff", listing["tariff_id"])["distributor_id"]
            cids = unique_ids([(f"{lid}/{r['basis']}/{r['gst']}/{r['locator']}",
                                [r["component"], f"{r['component']} {r['unit']}",
                                 f"{r['component']} {r['unit']} {r['time_band']} {r['season']}",
                                 f"{r['component']} {r['unit']} {r['value']}"]) for r in rs])
            for r, cid in zip(rs, cids):
                loc = locators.parse(r["locator"])
                q, period, inferred = quantity_period(r["unit_std"])
                input_unit = r["unit"]
                _, direct_unit = to_std(r["value"], input_unit, r["component"])
                normalisation_note = None
                if direct_unit != r["unit_std"]:
                    input_unit = interpreted_unit(r)
                    normalisation_note = f"Source parser interprets {r['unit']!r} as {input_unit!r}; see parser note and original unit."
                    try:  # the period is inferred only where the printed unit does not give it
                        printed_period = quantity_period(direct_unit)[1]
                    except ValueError:
                        printed_period = None
                    if printed_period != period:
                        inferred = 1
                raw, value, value_std = None, r["value"], num_text(r["value_std"]) if r["value_std"] else None
                if loc["kind"] == "xlsx":
                    raw_v, excel = locators.read_cell_excel(os.path.join(ROOT, r["source_file"]), loc["sheet"],
                                                            loc["cell"])
                    raw = repr(raw_v) if isinstance(raw_v, float) else str(raw_v)
                    if str(excel) != r["value"]:
                        raise SystemExit(f"{cid}: parser value {r['value']!r} is not the cell as Excel displays it "
                                         f"({excel!r}); format spreadsheet numbers with scripts/published.py")
                    naive = locators.binary_float_display(raw_v, str(excel))
                    if naive is not None and naive != str(excel):
                        self.instance("display_rounds_half_way", did=did, fy=d["fin_year"],
                                      tariff=listing["tariff_id"], listing=lid, charge=cid, doc=d["document_id"],
                                      detail=f"cell {raw} shows {excel} in Excel; formatting the binary float gives "
                                             f"{naive}")
                self.t.add("charge", {
                    "charge_id": cid, "listing_id": lid, "price_basis": r["basis"], "component_label": r["component"],
                    "charge_type": r["charge_type"], "time_band": r["time_band"] or None, "season": r["season"] or None,
                    "value_published": value, "unit_published": r["unit"] or None, "value_raw": raw,
                    "unit_interpreted": input_unit or None, "normalisation_note": normalisation_note,
                    "value_num": value, "value_std": value_std, "unit_std": r["unit_std"] or None,
                    "quantity": q, "period": period, "period_inferred": inferred, "gst": r["gst"],
                    "includes_metering": self.metering_inclusion(r, d), "includes_lfit": self.lfit_inclusion(r, d, did),
                    "effective_from": listing["effective_from"], "effective_to": listing["effective_to"],
                    "locator": r["locator"], "locator_kind": loc["kind"], "sheet": loc["sheet"], "cell": loc["cell"],
                    "page": loc["page"], "verification": {"xlsx": "cell_display", "pdf": "page_text",
                                                          "pdf-ocr": "ocr_sum_check"}[loc["kind"]],
                    "note": r["note"] or None,
                })

    @staticmethod
    def metering_inclusion(r, d):
        if d["author"] == "AER":
            return "no"  # the AER prices metering on its own sheet / schedule 1 (metering_price)
        note = r["note"] or ""
        if re.search(r"exclud\w*\s+metering|excl\.?\s+metering|without\s+metering", note, re.I):
            return "no"
        if re.search(r"includ\w*\s+(legacy\s+)?metering|metering\s+(charge\s+)?included|NUoS includes Metering block",
                     note, re.I):
            return "yes"
        if r["charge_type"] != "fixed":
            return "no"  # metering is priced per meter per year/day in every source, never per kWh or kW
        if re.search(r"column as published = [^;]*\+ metering", note, re.I):
            return "yes"  # the column is defined as including metering (Evoenergy statement 2023-24 p25-26)
        return "unknown"  # set to 'yes' below where the metering adjustment reproduces the value

    @staticmethod
    def lfit_inclusion(r, d, did):
        if did != "evoenergy":
            return "not_applicable"
        if d["author"] == "AER":
            return "no"
        status = EVO_LFIT_DOCS.get(d["local_path"], (None, None, None, None))[3]
        if status == "excluded":
            return "no"
        if status == "included":
            return "yes" if r["charge_type"] == "energy" else "no"
        if status == "rebate":
            return "yes" if r["charge_type"] == "energy" else "unknown"
        return "unknown"

    # ------------------------------------------------------------------ flags
    def flag(self, lid, flag, kind, evidence, replace=False):
        if not replace and self.t.get("listing_flag", lid, flag):
            return
        self.t.add("listing_flag", {"listing_id": lid, "flag": flag, "evidence_kind": kind, "evidence": evidence},
                   replace=True)

    def parser_flags(self):
        cov = {(c["document_id"], c["distributor_id"]): c["price_status"] for c in self.t.all("document_coverage")}
        for l in self.t.all("tariff_listing"):
            lid = l["listing_id"]
            d = self.doc_by_id[l["document_id"]]
            did = self.t.get("tariff", l["tariff_id"])["distributor_id"]
            if l["price_availability"] == "placeholder":
                self.flag(lid, "zero_priced_placeholder", "document", "every price cell of the row is zero or blank")
            for flag, pat in STATUS_FLAGS:
                if l["name_published"] and re.search(pat, l["name_published"], re.I):
                    self.flag(lid, flag, "published_text", f"tariff name: {l['name_published']}")
            notes = sorted({r["note"] for r in self.listing_rows[lid] if r["note"]} | ({l["note"]} if l["note"] else set()))
            for flag, pat in STATUS_FLAGS:
                hit = next((m for m in (re.search(pat, n, re.I) for n in notes) if m), None)
                if hit:
                    self.flag(lid, flag, "parser_note", evidence_excerpt(hit))
            status = cov.get((d["document_id"], did), d["price_status"])
            if status == "proposed":
                self.flag(lid, "proposed_price", "document", f"{d['document_id']} carries {did} prices as proposed")
            if l["region"]:
                self.flag(lid, "regional_variant", "parser_note", f"pricing region {l['region']}")
            if did == "evoenergy" and d["local_path"] in EVO_LFIT_DOCS:
                path, loc, quote, st = EVO_LFIT_DOCS[d["local_path"]]
                self.flag(lid, "excludes_lfit" if st == "excluded" else "includes_lfit", "published_text",
                          f"{rel(os.path.join(ROOT, path))} {loc}: {quote}")
            if did == "evoenergy" and d["author"] == "AER" and d["fin_year"] in bs.LANDING:
                self.flag(lid, "excludes_lfit", "published_text", f"AER landing page: {AER_LFIT_QUOTE}")
            ms = {c["includes_metering"] for c in self.charges_of(lid) if c["charge_type"] == "fixed"}
            if "yes" in ms:
                self.flag(lid, "includes_metering", "parser_note", "fixed charge includes metering (parser note)")

    def charges_of(self, lid):
        if not hasattr(self, "_charges_by_listing"):
            self._charges_by_listing = defaultdict(list)
            for c in self.t.all("charge"):
                self._charges_by_listing[c["listing_id"]].append(c)
        return self._charges_by_listing[lid]

    # ------------------------------------------------------------------ metering prices
    def metering_prices(self):
        for d in self.docs:
            if d["retrieval_status"] != "retrieved":
                continue
            path = os.path.join(ROOT, d["local_path"])
            if d["document_type"] == "aer_consolidated_stakeholder_report" and d["local_path"].endswith(".xlsx"):
                self.aer_metering_sheet(d, path)
            elif d["document_type"] == "aer_stakeholder_report":
                self.aer_schedule_1(d, path)
        import dnsp.energex_ergon as ee
        import openpyxl
        for name, fy, p in ee.FILES:
            d = self.doc_by_path[p]
            wb = openpyxl.load_workbook(os.path.join(ROOT, p), read_only=True, data_only=True)
            for ws in wb.worksheets:
                if ws.title.strip() not in ee.IN_SCOPE_SHEETS:
                    continue
                for m in ee.parse_sheet(ws).metering_cells:
                    tid = f"{d['distributor_id']}:{code_key(m['code'])}"
                    loc = locators.parse(m["locator"])
                    raw, shown = locators.read_cell(os.path.join(ROOT, p), loc["sheet"], loc["cell"])
                    per_day = (m["unit"] or "").strip().lower() in ("$/day", "$ /day", "$/ day")
                    self.t.add("metering_price", {
                        "metering_price_id": f"{d['document_id']}/{m['locator']}", "document_id": d["document_id"],
                        "distributor_id": d["distributor_id"], "fin_year": fy, "source_block": "distributor_metering_block",
                        "meter_class": m["cls"] or ws.title.strip(), "tariff_codes_published": m["code"],
                        "tariff_id": tid if self.t.get("tariff", tid) else None, "component_label": m["comp"],
                        "charge_basis": "per_day" if per_day else "unstated", "value_published": str(shown),
                        "unit_published": m["unit"] or None, "gst": "excl",
                        "value_raw": repr(raw) if isinstance(raw, float) else str(raw),
                        "value_num": num_text(raw), "value_c_per_day": num_text(float(raw) * 100) if per_day else None,
                        "locator": m["locator"], "locator_kind": "xlsx", "sheet": loc["sheet"], "cell": loc["cell"],
                    })
            wb.close()
        self.distributor_price_table_metering()

    def distributor_price_table_metering(self):
        """Metering columns of distributor network price tables, from the parsers' side output (out/dnsp_metering)."""
        import schema
        did_of = {d["name"]: d["distributor_id"] for d in bs.DISTRIBUTORS}
        rows = []
        for path in sorted(glob.glob(os.path.join(ROOT, schema.METERING_OUT_DIR, "*.csv"))):
            with open(path, newline="", encoding="utf-8") as f:
                rows += [dict(r, _file=rel(path), _line=i) for i, r in enumerate(csv.DictReader(f), 2)]
        ids = unique_ids([(f"{self.doc_by_path[r['source_file']]['document_id']}/{r['locator']}/{code_key(r['tariff_code'])}",
                           [r["component"], f"{r['component']} {r['meter_class']}", f"{r['component']} {r['value']}"])
                          for r in rows])
        for r, mid in zip(rows, ids):
            where = f"{r['_file']}:{r['_line']}"
            if r["gst"] not in ("excl", "incl"):
                raise SystemExit(f"{where}: gst must be excl or incl, got {r['gst']!r}")
            d = self.doc_by_path[r["source_file"]]
            did = did_of[r["distributor"]]
            loc = locators.parse(r["locator"])
            raw = None
            if loc["kind"] == "xlsx":
                raw_v, excel = locators.read_cell_excel(os.path.join(ROOT, r["source_file"]), loc["sheet"], loc["cell"])
                if str(excel) != r["value"]:
                    raise SystemExit(f"{where}: value {r['value']!r} is not the cell as Excel displays it ({excel!r})")
                raw = repr(raw_v) if isinstance(raw_v, float) else str(raw_v)
            unit = r["unit"].strip()
            per_day = re.fullmatch(r"(?:c|cents|\$)\s*/\s*day", unit, re.I)
            c_day = None
            if per_day:
                c_day = num_text(Decimal(r["value"]) * (100 if unit.startswith("$") else 1))
            tid = None
            k = code_key(r["tariff_code"])
            if self.t.get("tariff", f"{did}:{k}"):
                tid = f"{did}:{k}"
            elif len(self.alias_to_tariff.get((did, k), ())) == 1:
                tid = next(iter(self.alias_to_tariff[(did, k)]))
            self.t.add("metering_price", {
                "metering_price_id": mid, "document_id": d["document_id"], "distributor_id": did,
                "fin_year": r["fin_year"], "source_block": "distributor_price_table", "meter_class": r["meter_class"],
                "tariff_codes_published": r["tariff_code"], "tariff_id": tid, "component_label": r["component"],
                "charge_basis": "per_day" if per_day else "unstated", "value_published": r["value"],
                "unit_published": unit or None, "value_raw": raw, "value_num": num_text(r["value"]), "gst": r["gst"],
                "value_c_per_day": c_day, "locator": r["locator"], "locator_kind": loc["kind"], "sheet": loc["sheet"],
                "cell": loc["cell"], "page": loc["page"], "note": r["note"] or None,
            })

    def _metering_row(self, d, did, ws, label_cell, value_cell, codes, unit, basis, block):
        raw = value_cell.value
        shown = cell_value(value_cell)
        loc = locators.xlsx(ws, value_cell)
        per_day = num_text(float(raw) * 100 / 365) if basis == "per_year" else None
        self.t.add("metering_price", {
            "metering_price_id": f"{d['document_id']}/{loc}", "document_id": d["document_id"], "distributor_id": did,
            "fin_year": d["fin_year"], "source_block": block, "meter_class": str(label_cell.value),
            "tariff_codes_published": codes, "tariff_id": None, "component_label": None, "charge_basis": basis,
            "value_published": str(shown), "unit_published": unit, "gst": "excl", "value_raw": repr(raw) if isinstance(raw, float)
            else str(raw), "value_num": num_text(raw), "value_c_per_day": per_day, "locator": loc, "locator_kind": "xlsx",
            "sheet": ws.title, "cell": value_cell.coordinate,
        })

    def aer_metering_sheet(self, d, path):
        wb = locators._workbook(path)
        if "Metering" not in wb.sheetnames:
            return
        ws = wb["Metering"]
        by_label = {x["aer_label"]: x["distributor_id"] for x in bs.DISTRIBUTORS}
        cur = None
        for r in range(13, ws.max_row + 1):
            b = ws.cell(r, 2).value
            if isinstance(b, str) and "Metering prices" in b:
                cur = next(did for lab, did in by_label.items() if b.startswith(lab + " "))
                continue
            if isinstance(b, str) and b.strip() == "End":
                break
            if cur is None or not isinstance(b, (int, float)) or isinstance(b, bool):
                continue
            label, h, unit, per, val = (ws.cell(r, c) for c in (3, 8, 9, 10, 12))
            if val.value in (None, "") or not isinstance(val.value, (int, float)):
                if isinstance(label.value, str) and label.value.strip():
                    self.instance("metering_sheet_quirk", did=cur, fy=d["fin_year"], doc=d["document_id"],
                                  detail=f"Metering!{val.coordinate} is blank for {label.value.strip()!r}: no price row")
                continue
            if not (isinstance(label.value, str) and label.value.strip()):
                continue
            codes = h.value.strip() if isinstance(h.value, str) and h.value.strip() and h.value.strip() != "Exit fee" \
                else None
            perv = per.value.strip() if isinstance(per.value, str) else None
            basis = "per_meter" if (h.value == "Exit fee" or perv == "per meter") else \
                "per_year" if perv == "per year" else "unstated"
            self._metering_row(d, cur, ws, label, val, codes, unit.value if isinstance(unit.value, str) else None, basis,
                               "aer_metering_sheet")
            if h.value == "Exit fee" or basis == "unstated":
                self.instance("metering_sheet_quirk", did=cur, fy=d["fin_year"], doc=d["document_id"],
                              detail=(f"Metering!{h.coordinate} holds 'Exit fee' in the 'Tariff code' column (charge kind, "
                                      f"not a code)" if h.value == "Exit fee" else
                                      f"Metering!{per.coordinate} 'Charge' column is {per.value!r}: the price period is "
                                      f"not stated") + f"; {label.value.strip()!r}")

    def aer_schedule_1(self, d, path):
        ws = locators._workbook(path)["Tariff schedule"]
        start = next(r for r in range(1, 40) if str(ws.cell(r, 2).value or "").startswith("Tariff schedule 1"))
        ycol = next(c for c in range(3, 20) if str(ws.cell(start, c).value or "").strip() == d["fin_year"].replace("-", "–"))
        for r in range(start + 1, start + 40):
            if str(ws.cell(r, 2).value or "").startswith("Tariff schedule 2"):
                break
            label, val = ws.cell(r, 3), ws.cell(r, ycol)
            if isinstance(label.value, str) and label.value.strip() and isinstance(val.value, (int, float)):
                self._metering_row(d, d["distributor_id"], ws, label, val, None, None, "unstated", "aer_tariff_schedule_1")

    # ------------------------------------------------------------------ adjustments
    def aer_doc(self, did, fy):
        if fy == "2024-25":
            return next(d for d in self.docs if d["document_type"] == "aer_stakeholder_report" and d["distributor_id"] == did
                        and d["retrieval_status"] == "retrieved")
        return max((d for d in self.docs if d["document_type"] == "aer_consolidated_stakeholder_report"
                    and d["fin_year"] == fy and d["retrieval_status"] == "retrieved" and d["local_path"].endswith(".xlsx")),
                   key=lambda d: d["version_seq"])

    def nuos_charges(self, doc_id, tid, charge_type, unit_std):
        out = []
        for lid in self.listings_by_doc_tariff.get((doc_id, tid), []):
            for c in self.charges_of(lid):
                if c["price_basis"] == "NUoS" and c["gst"] == "excl" and c["charge_type"] == charge_type \
                        and c["unit_std"] == unit_std:
                    out.append(c)
        return out

    @staticmethod
    def half_unit(c):
        v, s = float(c["value_num"]), float(c["value_std"])
        factor = s / v if v else 1.0
        return 0.5 * 10 ** (-decimals(c["value_published"])) * abs(factor)

    def tariff_pairs(self, aer_doc, dnsp_doc, charge_type, unit_std):
        tids = sorted({t for (doc, t) in self.listings_by_doc_tariff if doc == dnsp_doc["document_id"]}
                      & {t for (doc, t) in self.listings_by_doc_tariff if doc == aer_doc["document_id"]})
        for tid in tids:
            a = self.nuos_charges(aer_doc["document_id"], tid, charge_type, unit_std)
            dd = self.nuos_charges(dnsp_doc["document_id"], tid, charge_type, unit_std)
            if a and dd:
                yield tid, a, dd

    def adjustments(self):
        mp = self.t.all("metering_price")
        for did, fy in METERING_ADDERS:
            ad = self.aer_doc(did, fy)
            rows = [m for m in mp if m["document_id"] == ad["document_id"] and m["distributor_id"] == did
                    and m["charge_basis"] != "per_meter" and float(m["value_num"]) != 0]
            amounts = {m["value_num"] for m in rows}
            if len(amounts) != 1:
                raise SystemExit(f"metering adder {did} {fy}: expected one metering amount in {ad['document_id']}, "
                                 f"found {sorted(amounts)}")
            M = float(amounts.pop())
            amount = M * 100 / 365
            block = {m["tariff_id"]: m for m in mp if m["source_block"] == "distributor_metering_block"
                     and m["distributor_id"] == did and m["fin_year"] == fy and m["tariff_id"]
                     and m["charge_basis"] == "per_day"}
            def metering_hits(dd):
                hits = []
                for tid, a, dch in self.tariff_pairs(ad, dd, "fixed", "c/day"):
                    # one price on each side (a repeated printing of the same price is still one price)
                    if len({c["value_std"] for c in a}) != 1 or len({c["value_std"] for c in dch}) != 1:
                        continue
                    if not adjustments.metering_scope(self.dist[did]["name"], fy, tid.split(":", 1)[1]):
                        continue  # outside the verified metering scope
                    exp = float(block[tid]["value_c_per_day"]) if tid in block else amount
                    delta = float(dch[0]["value_std"]) - float(a[0]["value_std"])
                    if abs(delta - exp) <= self.half_unit(a[0]) + self.half_unit(dch[0]) + 1e-9:
                        hits.append((tid, exp, dch))
                return hits

            best = None
            for dd in self.dnsp_docs(did, fy):
                hits = metering_hits(dd)
                if hits and (best is None or len(hits) > len(best[1])):
                    best = (dd, hits)
            if best is None:
                raise SystemExit(f"metering adder {did} {fy}: no tariff reproduces it")
            dd, _ = best
            # every distributor document of the year whose daily charge reproduces AER + metering includes metering
            # (the AER-hosted copy and the distributor's own list alike), not only the adjustment's reference document
            hits_by_doc = [(d2, metering_hits(d2)) for d2 in self.distributor_docs(did, fy)]
            hits = sorted({(tid, exp): dch for _, hs in hits_by_doc for tid, exp, dch in hs}.items())
            src = rows[0]
            if did == "essential" and fy in ESSENTIAL_NOTE_DOCS:
                ev_doc, ev_loc, ev_quote = self.doc_by_path[ESSENTIAL_NOTE_DOCS[fy]]["document_id"], "pdf:p1", \
                    ESSENTIAL_METERING_NOTE
            else:
                ev_doc, ev_loc, ev_quote = ad["document_id"], src["locator"], src["value_raw"]
            adj = f"metering_adder/{did}/{fy}"
            self.t.add("price_adjustment", {
                "adjustment_id": adj, "kind": "metering_adder", "distributor_id": did, "fin_year": fy,
                "amount": num_text(amount), "amount_unit": "c/day",
                "formula": f"distributor fixed c/day = AER fixed c/day + 100 x {num_text(M)} $/yr / 365"
                           + ("; per tariff, 100 x the $/day of the distributor's Metering block for that tariff "
                              "(price_adjustment_tariff.expected_delta_std)" if block else ""),
                "aer_document_id": ad["document_id"], "distributor_document_id": dd["document_id"],
                "evidence_document_id": ev_doc, "locator": ev_loc, "quote": ev_quote,
            })
            for (tid, exp), _ in hits:
                m = block.get(tid) or next((x for x in rows if x["tariff_codes_published"] and
                                            tid.split(":", 1)[1] in [c.strip() for c in x["tariff_codes_published"].split(",")]),
                                           src)
                self.t.add("price_adjustment_tariff", {
                    "adjustment_id": adj, "tariff_id": tid, "metering_price_id": m["metering_price_id"],
                    "expected_delta_std": num_text(exp), "delta_unit": "c/day"})
            for d2, hs in hits_by_doc:
                for tid, _, dch in hs:
                    for c in dch:
                        if c["includes_metering"] == "unknown":
                            c["includes_metering"] = "yes"
                    for lid in self.listings_by_doc_tariff[(d2["document_id"], tid)]:
                        self.flag(lid, "includes_metering", "document", f"daily charge = AER daily charge + metering "
                                  f"({adj}), reproduced to the published digit")
            for lid in sorted({l for (t, _), _ in hits for l in self.listings_by_doc_tariff[(ad["document_id"], t)]}):
                self.flag(lid, "excludes_metering", "document",
                          f"AER prices metering separately ({src['metering_price_id']})")
            self.instance("metering_excluded_by_aer", did=did, fy=fy, doc=ad["document_id"],
                          related=dd["document_id"], quantity=num_text(amount), unit="c/day",
                          detail=f"{len({t for (t, _), _ in hits})} tariffs: distributor daily charge = AER + "
                                 f"{num_text(amount)} c/day")
        for fy, amount, dpath in LFIT_ADDERS:
            ad, dd = self.aer_doc("evoenergy", fy), self.doc_by_path[dpath]
            path, loc, quote, _ = EVO_LFIT_DOCS[dpath]
            self.lfit(f"lfit_adder/evoenergy/{fy}", "lfit_adder", fy, ad, dd, float(amount), loc, quote,
                      f"distributor c/kWh = AER c/kWh + {amount} c/kWh on every consumption charge; fixed and demand "
                      f"charges unchanged")
        ad = self.doc_by_path["sources/aer/2023-24_price_lists/Evoenergy_2023-24_Electricity_network_pricing_proposal_5May2023.pdf"]
        dd = self.doc_by_path["sources/dnsp/evoenergy/Evoenergy_Statement_of_Tariff_Classes_and_Tariffs_2023-24.pdf"]
        self.lfit("lfit_rebate/evoenergy/2023-24", "lfit_rebate", "2023-24", ad, dd, None, "pdf:p4", LFIT_REBATE_QUOTE,
                  "distributor c/kWh = c/kWh of the AER-hosted pricing proposal (status unverified) - tariff-specific "
                  "LFiT rebate (2.27 c/kWh on average; not uniform; applied to consumption charges where possible); "
                  "the quoted statement applies the rebate to the AER's approved charges, which no held document lists")

    def metering_by_aer_equality(self):
        """A distributor fixed charge equal (to the published digits) to the AER's fixed charge for the same tariff, year,
        basis and GST excludes metering: the AER prices metering separately (metering_price)."""
        aer = defaultdict(list)
        for (doc, tid), lids in self.listings_by_doc_tariff.items():
            d = self.doc_by_id[doc]
            if d["author"] == "AER":
                for lid in lids:
                    for c in self.charges_of(lid):
                        if c["charge_type"] == "fixed" and c["unit_std"] == "c/day":
                            aer[(tid, d["fin_year"], c["price_basis"], c["gst"])].append(c)
        for (doc, tid), lids in self.listings_by_doc_tariff.items():
            d = self.doc_by_id[doc]
            if d["author"] == "AER":
                continue
            for lid in lids:
                for c in self.charges_of(lid):
                    if c["includes_metering"] != "unknown" or c["unit_std"] != "c/day":
                        continue
                    if any(abs(float(c["value_std"]) - float(a["value_std"])) <= self.half_unit(a) + self.half_unit(c)
                           + 1e-9 for a in aer[(tid, d["fin_year"], c["price_basis"], c["gst"])]):
                        c["includes_metering"] = "no"

    def lfit(self, adj, kind, fy, ad, dd, amount, loc, quote, formula):
        hits = []
        for tid, a, dch in self.tariff_pairs(ad, dd, "energy", "c/kWh"):
            pairs = []
            for x in a:
                ys = [y for y in dch if (y["time_band"], y["season"]) == (x["time_band"], x["season"])]
                pairs += [(x, y) for y in ys]
            if not pairs:
                continue
            deltas = [float(y["value_std"]) - float(x["value_std"]) for x, y in pairs]
            tol = max(self.half_unit(x) + self.half_unit(y) for x, y in pairs) + 1e-9
            if amount is not None and all(abs(dl - amount) <= tol for dl in deltas):
                hits.append(tid)
            elif amount is None and all(dl < 0 for dl in deltas):
                hits.append(tid)
        self.t.add("price_adjustment", {
            "adjustment_id": adj, "kind": kind, "distributor_id": "evoenergy", "fin_year": fy,
            "amount": None if amount is None else num_text(amount), "amount_unit": None if amount is None else "c/kWh",
            "formula": formula, "aer_document_id": ad["document_id"], "distributor_document_id": dd["document_id"],
            "evidence_document_id": dd["document_id"], "locator": loc, "quote": quote,
        })
        for tid in hits:
            self.t.add("price_adjustment_tariff", {
                "adjustment_id": adj, "tariff_id": tid, "metering_price_id": None,
                "expected_delta_std": None if amount is None else num_text(amount),
                "delta_unit": None if amount is None else "c/kWh"})
        self.instance("act_lfit", did="evoenergy", fy=fy, doc=ad["document_id"], related=dd["document_id"],
                      quantity=None if amount is None else num_text(amount), unit=None if amount is None else "c/kWh",
                      detail=f"{kind}: {len(hits)} tariffs; {formula}")

    def distributor_docs(self, did, fy):
        return sorted((d for d in self.docs if d["distributor_id"] == did and d["fin_year"] == fy
                       and d["author"] == "distributor" and d["retrieval_status"] == "retrieved"),
                      key=lambda d: d["document_id"])

    def dnsp_docs(self, did, fy):
        own = self.distributor_docs(did, fy)
        return [d for d in own if d["recon_side"] == "DNSP"] or own

    # ------------------------------------------------------------------ curated facts
    def tariff_for_code(self, did, code, where):
        k = code_key(str(code))
        if self.t.get("tariff", f"{did}:{k}"):
            return f"{did}:{k}"
        hits = self.alias_to_tariff.get((did, k), set())
        if len(hits) == 1:
            return next(iter(hits))
        raise SystemExit(f"{where}: tariff code {code!r} does not resolve to one {did} tariff")

    def tariffs_for_code(self, did, code, doc, locator, where):
        """Tariffs a curated fact about `code` in document `doc` applies to: the distributor's own code, else the
        tariffs listed under that label in that document (an AER label such as Ergon 'ERIB' covers one row per region;
        a fact quoted from one row applies to that row's tariff)."""
        k = code_key(str(code))
        if self.t.get("tariff", f"{did}:{k}") and self.t.get("tariff", f"{did}:{k}")["identity_basis"] == "distributor_code":
            return [f"{did}:{k}"]
        cands = [self.t.get("tariff_listing", lid) for (d2, _), lids in self.listings_by_doc_tariff.items() if d2 == doc
                 for lid in lids]
        cands = [l for l in cands if l["code_published"] and code_key(l["code_published"]) == k]
        if not cands:
            if not self.t.get("tariff", f"{did}:{k}") and self.doc_by_id[doc]["author"] == "distributor":
                path = os.path.join(ROOT, self.doc_by_id[doc]["local_path"])
                ok, _ = locators.verify_quote(path, locator, str(code))
                if ok:
                    return [self.ensure_tariff(did, str(code), "distributor_code")]
            return [self.tariff_for_code(did, code, where)]
        loc = locators.parse(locator)
        if loc["kind"] == "xlsx":
            row = re.sub(r"^[A-Z]+", "", loc["cell"])
            same = [l for l in cands if l["locator"].startswith(f"xlsx:{loc['sheet']}!")
                    and re.sub(r"^xlsx:.*![A-Z]+", "", l["locator"]) == row]
            if same:
                cands = same
        return sorted({l["tariff_id"] for l in cands})

    def curated(self):
        for name, data in sorted(curated.load_all().items()):
            did = data.get("distributor")
            where = f"data/tariffdb/curated/{name}.yaml"
            if did != name:
                raise SystemExit(f"{where}: distributor {did!r} does not match the file name")
            errs = curated.validate(data, check_quotes=True)
            if errs:
                raise SystemExit(f"{where}: " + "; ".join(errs[:10]))
            self.curated_file(did, data, where)

    def doc_id(self, path, where):
        if path not in self.doc_by_path:
            raise SystemExit(f"{where}: document {path!r} not in the registry")
        return self.doc_by_path[path]["document_id"]

    def curated_file(self, did, data, where):
        for r in data.get("steps") or []:
            doc = self.doc_id(r["doc"], where)
            start, end = bs.FIN_YEAR_DATES[r["fin_year"]]
            for code in r["codes"]:
                tid = self.tariff_for_code(did, code, where)
                self.t.add("charge_step", {
                    "step_id": f"{doc}/{tid}/{short_hash(r['locator'], r['component_label'], r['step_index'])}",
                    "tariff_id": tid, "document_id": doc, "fin_year": r["fin_year"],
                    "effective_from": r.get("effective_from", start), "effective_to": r.get("effective_to", end),
                    "lower_inclusive": int(bool(r["lower_inclusive"])), "upper_inclusive": int(bool(r["upper_inclusive"])),
                    **{k: r.get(k) for k in ("step_group", "component_label", "step_index", "lower_bound", "upper_bound",
                                           "quantity_unit", "reset_period", "locator", "quote")}})
        for s in data.get("tou_schedules") or []:
            doc = self.doc_id(s["doc"], where)
            sid = s["id"]
            self.t.add("tou_schedule", {
                "tou_schedule_id": sid, "distributor_id": did, "document_id": doc, "fin_year": s["fin_year"],
                "name": s["name"], "time_basis": s["time_basis"], "public_holidays": s["public_holidays"],
                "covers_full_day": 1 if s.get("covers_full_day") else 0, "locator": s["locator"],
                "quote": str(s["quote"]), "note": s.get("note")})
            for w in s["windows"]:
                ms = curated.months_of(w.get("months"))
                months = None if ms is None else ",".join(str(m) for m in ms)
                self.t.add("tou_window", {
                    "window_id": f"{sid}/{w['period']}/{w['days']}/{w['start']}-{w['end']}/"
                                 f"{'months-not-stated' if months is None else months.replace(',', '.')}",
                    "tou_schedule_id": sid, "period": w["period"], "period_label": str(w["label"]),
                    "day_type": w["days"], "start_time": str(w["start"]), "end_time": str(w["end"]), "months": months,
                    "season": w.get("season"), "locator": w.get("locator") or s["locator"],
                    "quote": str(w.get("quote") or s["quote"])})
            start, end = bs.FIN_YEAR_DATES[s["fin_year"]]
            for t in s.get("tariffs") or []:
                for tid in [x for code in t["codes"] for x in self.tariffs_for_code(
                        did, code, doc, t.get("locator") or s["locator"], f"{where} {sid}")]:
                    self.t.add("tariff_tou", {
                        "tariff_tou_id": f"{tid}|{sid}|{t['applies_to']}", "tariff_id": tid, "tou_schedule_id": sid,
                        "applies_to": t["applies_to"], "effective_from": start, "effective_to": end, "document_id": doc,
                        "locator": t.get("locator") or s["locator"], "quote": str(t.get("quote") or s["quote"])},
                        same_ok=True)
        for r in data.get("demand_rules") or []:
            doc = self.doc_id(r["doc"], where)
            self.t.add("demand_rule", {
                "demand_rule_id": r["id"], "distributor_id": did, "document_id": doc, "fin_year": r["fin_year"],
                "measure": r["measure"], "interval_minutes": r.get("interval_minutes"), "aggregation": r["aggregation"],
                "aggregation_count": r.get("aggregation_count"),
                "window_tou_schedule_id": r.get("window_schedule"), "window_period": r.get("window_period"),
                "months": ",".join(str(m) for m in curated.months_of(r["months"])) if r.get("months") else None,
                "minimum_chargeable": r.get("minimum_chargeable"), "minimum_unit": r.get("minimum_unit"),
                "locator": r["locator"], "quote": str(r["quote"]), "note": r.get("note")})
            start, end = bs.FIN_YEAR_DATES[r["fin_year"]]
            for t in r["tariffs"]:
                for tid in [x for code in t["codes"] for x in self.tariffs_for_code(
                        did, code, doc, r["locator"], f"{where} {r['id']}")]:
                    self.t.add("tariff_demand_rule", {
                        "tariff_demand_rule_id": f"{tid}|{r['id']}|{t.get('time_band') or ''}|{t.get('season') or ''}",
                        "tariff_id": tid, "demand_rule_id": r["id"], "time_band": t.get("time_band"),
                        "season": t.get("season"), "effective_from": start, "effective_to": end, "document_id": doc},
                        same_ok=True)
        for r in data.get("eligibility") or []:
            doc = self.doc_id(r["doc"], where)
            start, end = bs.FIN_YEAR_DATES[r["fin_year"]]
            target = self.tariff_for_code(did, r["target_code"], where) if r.get("target_code") else None
            for tid in [x for code in r["codes"] for x in self.tariffs_for_code(did, code, doc, r["locator"],
                                                                                f"{where} eligibility")]:
                h = short_hash(doc, r["rule_type"], r.get("operator"), r.get("value_num"), r.get("value_unit"),
                               r.get("value_text"), target, r["locator"], r["quote"])
                self.t.add("eligibility_rule", {
                    "rule_id": f"{tid}|{r['fin_year']}|{r['rule_type']}|{h}", "tariff_id": tid, "document_id": doc,
                    "fin_year": r["fin_year"], "effective_from": start, "effective_to": end, "rule_type": r["rule_type"],
                    "operator": r.get("operator"), "value_num": r.get("value_num"), "value_unit": r.get("value_unit"),
                    "value_text": r.get("value_text"), "target_tariff_id": target, "locator": r["locator"],
                    "quote": str(r["quote"]), "note": r.get("note")}, same_ok=True)
        for r in data.get("relations") or []:
            doc = self.doc_id(r["doc"], where)
            a = self.tariffs_for_code(did, r["from"], doc, r["locator"], where)
            b = self.tariffs_for_code(did, r["to"], doc, r["locator"], where)
            if len(a) != 1 or len(b) != 1:
                raise SystemExit(f"{where}: relation {r['from']} -> {r['to']} does not resolve to one tariff each")
            a, b = a[0], b[0]
            self.t.add("tariff_relation", {
                "relation_id": f"{a}|{r['type']}|{b}|{doc}", "from_tariff_id": a, "relation_type": r["type"],
                "to_tariff_id": b, "fin_year": r["fin_year"], "document_id": doc, "locator": r["locator"],
                "quote": str(r["quote"]), "note": r.get("note")}, same_ok=True)
        for r in data.get("flags") or []:
            doc = self.doc_id(r["doc"], where)
            for tid in [x for code in r["codes"] for x in self.tariffs_for_code(did, code, doc, r["locator"], where)]:
                lids = self.listings_by_doc_tariff.get((doc, tid))
                if not lids:
                    start, end = bs.FIN_YEAR_DATES[r["fin_year"]]
                    code = self.t.get("tariff", tid)["tariff_code"]
                    lid = f"{doc}/{code}/rules"
                    # code_published only where the document prints the code (Jemena 2025-26 names its F-codes as
                    # 'Tariffs starting with "F"', never one by one)
                    printed = any(re.search(rf"(?<![\w]){re.escape(code)}(?![\w])", str(q or ""))
                                  for q in (r["quote"], r.get("column_quote"))) or document_prints(
                        os.path.join(ROOT, self.doc_by_id[doc]["local_path"]), code)
                    self.t.add("tariff_listing", {
                        "listing_id": lid, "document_id": doc, "tariff_id": tid,
                        "code_published": code if printed else None,
                        "effective_from": start, "effective_to": end,
                        "price_availability": "rules_only", "locator": r["locator"], "note": str(r["quote"])})
                    self.listings_by_doc_tariff[(doc, tid)].append(lid)
                    lids = [lid]
                evidence = f"{r['doc']} {r['locator']}: {r['quote']}"
                if r.get("column_locator"):
                    evidence += f" (column {r['column_locator']}: {r['column_quote']})"
                for lid in lids:
                    self.flag(lid, r["flag"], "curated", evidence, replace=True)

    # ------------------------------------------------------------------ AER sibling code columns (SAPN)
    def aer_sibling_codes(self):
        """SA Power Networks rows of the AER workbooks print sibling codes next to 'Code SA': 'Code CBD' and '>30kW'
        (consolidated reports), 'Code CBD' and 'Other identifier' (2024-25 stakeholder report, e.g. 'LBAD-CBD')."""
        for d in self.docs:
            if d["retrieval_status"] != "retrieved" or not d["local_path"].endswith(".xlsx") or d["document_type"] not in (
                    "aer_consolidated_stakeholder_report", "aer_stakeholder_report"):
                continue
            ws = locators._workbook(os.path.join(ROOT, d["local_path"]))["Tariff schedule"]
            headers = [c for row in ws.iter_rows() for c in row if c.value == "Code SA"]
            seen = set()
            for k, hdr in enumerate(headers):
                labels = {hdr.column + i: str(ws.cell(hdr.row, hdr.column + i).value) for i in (1, 2)}
                last = headers[k + 1].row if k + 1 < len(headers) else ws.max_row + 1
                for r in range(hdr.row + 1, last):
                    main = ws.cell(r, hdr.column).value
                    if isinstance(main, str) and main.strip() == "Tariff code":
                        break
                    if not (isinstance(main, str) and main.strip()):
                        continue
                    for col, lab in labels.items():
                        sib = ws.cell(r, col)
                        if not (isinstance(sib.value, str) and sib.value.strip()) or code_key(sib.value) == code_key(main):
                            continue
                        a = self.tariff_for_code("sapn", main, d["document_id"]) if self.t.get(
                            "tariff", f"sapn:{code_key(main)}") or self.alias_to_tariff.get(("sapn", code_key(main))) \
                            else None
                        if a is None:
                            continue
                        sk = code_key(sib.value)
                        printed = ""
                        if sk not in self.universe["sapn"] and sk.replace("-", "") in self.universe["sapn"]:
                            sk = sk.replace("-", "")
                            printed = f" (printed {sib.value.strip()!r})"
                        known = sk in self.universe["sapn"]
                        b = self.ensure_tariff("sapn", sk if known else sib.value, "aer_label")
                        if (a, b) in seen:
                            continue
                        seen.add((a, b))
                        self.t.add("tariff_relation", {
                            "relation_id": f"{a}|aer_sibling_code|{b}|{d['document_id']}", "from_tariff_id": a,
                            "relation_type": "aer_sibling_code", "to_tariff_id": b, "fin_year": d["fin_year"],
                            "document_id": d["document_id"], "locator": locators.xlsx(ws, sib), "quote": sib.value.strip(),
                            "note": f"printed in the AER '{lab}' column of the {main.strip()} row{printed}"}, same_ok=True)
                        if not known:
                            self.instance("code_label_quirk", did="sapn", fy=d["fin_year"], tariff=b, doc=d["document_id"],
                                          detail=f"AER '{lab}' column prints {sib.value.strip()!r} (row {r}); SA Power "
                                                 f"Networks publishes no such code")

    def printed_identities(self):
        """A tariff first met as an AER label is a distributor code once a distributor document prints that code, even
        in rules text only (e.g. CitiPower 'CFTUOS', priced individually)."""
        for l in self.t.all("tariff_listing"):
            t = self.t.get("tariff", l["tariff_id"])
            if t["identity_basis"] == "aer_label" and self.doc_by_id[l["document_id"]]["author"] == "distributor" \
                    and l["code_published"] and code_key(l["code_published"]) == l["tariff_id"].split(":", 1)[1]:
                self.t.add("tariff", dict(t, identity_basis="distributor_code"), replace=True)

    # ------------------------------------------------------------------ exceptions
    def instance(self, code, did=None, fy=None, tariff=None, listing=None, charge=None, doc=None, related=None,
                 quantity=None, unit=None, detail=""):
        self.instances.append({"exception_code": code, "distributor_id": did, "fin_year": fy, "tariff_id": tariff,
                               "listing_id": listing, "charge_id": charge, "document_id": doc,
                               "related_document_id": related, "quantity": quantity, "quantity_unit": unit,
                               "detail": detail})

    def derived_exceptions(self):
        from ambiguities import AMBIGUITIES
        for a in AMBIGUITIES:
            d = self.doc_by_path[a["doc"]]
            ok, why = locators.verify_quote(os.path.join(ROOT, a["doc"]), a["locator"], a["quote"])
            if not ok:
                raise SystemExit(f"ambiguities.py {a['finding']}: {why} ({a['doc']} {a['locator']}: {a['quote']!r})")
            tids = sorted({t for c in a.get("codes") or [] for t in self.tariffs_for_code(
                a["distributor"], c, d["document_id"], a["locator"], f"ambiguities.py {a['finding']}")}) or [None]
            for tid in tids:
                self.instance("source_ambiguous", did=a["distributor"], fy=a["fin_year"], tariff=tid,
                              doc=d["document_id"], detail=f"{a['finding']} {a['locator']}: '{a['quote']}'; stored: "
                              f"{a['stored']}; also readable as: {a['alternative']}")
        for path, locator, quote in PRICE_ATTACHMENTS_NOT_HELD:
            d = self.doc_by_path[path]
            self.instance("price_attachment_not_held", did=d["distributor_id"], fy=d["fin_year"],
                          doc=d["document_id"], detail=f"{locator}: '{quote}'; that Tariff Summary attachment is neither "
                                                       f"archived nor hosted by the AER, so this document carries rules only")
        for r in self.t.all("charge_step"):
            self.instance("quantity_blocks", did=r["tariff_id"].split(":")[0], fy=r["fin_year"],
                          tariff=r["tariff_id"], doc=r["document_id"], quantity=r["upper_bound"] or r["lower_bound"],
                          unit=r["quantity_unit"], detail=f"{r['locator']}: {r['component_label']}, reset {r['reset_period']}")
        listings = self.t.all("tariff_listing")
        by_doc = defaultdict(list)
        for l in listings:
            by_doc[l["document_id"]].append(l)
        flags = defaultdict(set)
        for f in self.t.all("listing_flag"):
            flags[f["listing_id"]].add(f["flag"])
        # v1 (proposed) vs approved version
        v1 = self.doc_by_path["sources/aer/AER_Consolidated_stakeholder_report_2025-26_v1_wayback.xlsx"]
        v5 = self.doc_by_path["sources/aer/AER_Consolidated_stakeholder_report_2025-26_v5.xlsx"]
        key = lambda c, l: (l["tariff_id"], c["price_basis"], c["gst"], c["component_label"])  # noqa: E731

        def index(doc):
            out = defaultdict(list)
            for l in by_doc[doc["document_id"]]:
                for c in self.charges_of(l["listing_id"]):
                    out[(self.t.get("tariff", l["tariff_id"])["distributor_id"],) + key(c, l)].append(c)
            return out
        i1, i5 = index(v1), index(v5)
        per = defaultdict(lambda: [0, 0, []])
        for k, cs in i1.items():
            if k in i5 and len(cs) == 1 and len(i5[k]) == 1:
                a, b = float(cs[0]["value_num"]), float(i5[k][0]["value_num"])
                per[k[0]][0] += 1
                if a != b:
                    per[k[0]][1] += 1
                    if a:
                        per[k[0]][2].append((b - a) / abs(a) * 100)
        for did in sorted({k[0] for k in i5} | {k[0] for k in i1}):
            n, changed, pcts = per.get(did, (0, 0, []))
            if did not in {k[0] for k in i1}:
                detail = "no prices in v1 (the changelog adds this jurisdiction's prices in a later version)"
            else:
                detail = f"{changed} of {n} matched components differ between v1 (proposed) and v5 (approved)"
                if pcts:
                    detail += (f"; median {statistics.median(pcts):+.1f}%, largest {max(pcts, key=abs):+.1f}%")
            self.instance("aer_version_differs", did=did, fy="2025-26", doc=v1["document_id"], related=v5["document_id"],
                          quantity=num_text(max(pcts, key=abs)) if pcts else None, unit="%" if pcts else None,
                          detail=detail)
        for l in listings:
            if l["price_availability"] == "placeholder":
                self.instance("zero_priced_placeholder", did=self.t.get("tariff", l["tariff_id"])["distributor_id"],
                              fy=self.doc_by_id[l["document_id"]]["fin_year"], tariff=l["tariff_id"],
                              listing=l["listing_id"], doc=l["document_id"],
                              detail=f"{l['code_published'] or l['listing_id']} {l['name_published'] or ''}: no non-zero "
                                     f"price; price_availability = placeholder, no charge rows")
            fl = flags[l["listing_id"]] & {"withdrawn", "closed_to_new", "obsolete", "grandfathered"}
            if fl and l["price_availability"] != "rules_only":  # a tariff named only in rules text is not listed
                self.instance("withdrawn_tariff_listed", did=self.t.get("tariff", l["tariff_id"])["distributor_id"],
                              fy=self.doc_by_id[l["document_id"]]["fin_year"], tariff=l["tariff_id"],
                              listing=l["listing_id"], doc=l["document_id"], detail="flags: " + ", ".join(sorted(fl)))
        # tariffs missing from the AER side / only on the AER side, per distributor-year with an AER-authored file
        # (listings a document names only in its rules text do not count as published tariffs on either side)
        # (a code the distributor prints only in rules text still means the tariff is not AER-only)
        aer_t, dn_t, dn_named = defaultdict(set), defaultdict(set), defaultdict(set)
        first = {}
        for l in listings:
            d = self.doc_by_id[l["document_id"]]
            did = self.t.get("tariff", l["tariff_id"])["distributor_id"]
            k = (did, d["fin_year"])
            if l["price_availability"] == "rules_only":
                if d["author"] != "AER":
                    dn_named[k].add(l["tariff_id"])
                continue
            if d["author"] == "AER":
                if d["price_status"] != "proposed":
                    aer_t[k].add(l["tariff_id"])
            else:
                dn_t[k].add(l["tariff_id"])
                dn_named[k].add(l["tariff_id"])
                first.setdefault((k, l["tariff_id"]), l)
        # the AER also covers the sibling codes it prints on a row ('Code CBD', '>30kW') and the members of a label it
        # prints once for several distributor tariffs (aer_combined_label)
        aer_cover = defaultdict(set)
        members = defaultdict(set)
        for k, ts in aer_t.items():
            aer_cover[k] |= ts
        for rel in self.t.all("tariff_relation"):
            did = self.t.get("tariff", rel["from_tariff_id"])["distributor_id"]
            k = (did, rel["fin_year"])
            if rel["relation_type"] == "aer_sibling_code" and rel["from_tariff_id"] in aer_t[k] \
                    and self.doc_by_id[rel["document_id"]]["price_status"] != "proposed":
                aer_cover[k].add(rel["to_tariff_id"])
            elif rel["relation_type"] == "aer_combined_label":
                members[(k, rel["from_tariff_id"])].add(rel["to_tariff_id"])
        for (k, label), ms in members.items():
            if label in aer_t[k]:
                aer_cover[k] |= ms
        for k in sorted(aer_t):
            for tid in sorted(dn_t[k] - aer_cover[k]):
                l = first[(k, tid)]
                fl = sorted(flags[l["listing_id"]] & {"site_specific", "withdrawn", "closed_to_new", "obsolete", "trial"})
                self.instance("aer_missing_tariff", did=k[0], fy=k[1], tariff=tid, listing=l["listing_id"],
                              doc=l["document_id"], detail=" ".join(filter(None, [
                                  l["name_published"] or l["code_published"],
                                  f"[{l['class_published']}]" if l["class_published"] else None,
                                  f"flags: {', '.join(fl)}" if fl else None])))
            for tid in sorted(aer_t[k] - dn_named[k]):
                if members.get((k, tid), set()) & dn_named[k]:
                    continue  # the distributor prints the label's member tariffs
                if dn_t[k]:
                    self.instance("aer_only_tariff", did=k[0], fy=k[1], tariff=tid,
                                  detail=f"{self.t.get('tariff', tid)['identity_basis']}: listed by the AER, not in the "
                                         f"distributor's own {k[1]} documents")
        # layouts of the AER-authored workbooks
        import parse_aer
        for d in self.docs:
            if d["author"] != "AER" or d["retrieval_status"] != "retrieved" or not d["local_path"].endswith(".xlsx"):
                continue
            wb = locators._workbook(os.path.join(ROOT, d["local_path"]))
            ws = wb["Tariff schedule"]
            if d["document_type"] == "aer_consolidated_stakeholder_report":
                hr = next(c.row for row in ws.iter_rows() for c in row if c.value == "Top")
                cc, pc = parse_aer.consolidated_columns(ws, hr)
                from openpyxl.utils import get_column_letter as L
                detail = (f"consolidated layout: code column {L(cc)} labelled {ws.cell(hr, cc).value!r}, prices from "
                          f"column {L(pc)}; sheets {', '.join(wb.sheetnames)}")
                if any(ws.cell(r, cc).value == "#REF!" for r in range(hr, min(ws.max_row, hr + 40))):
                    detail += "; code cells hold #REF! (no codes printed), rows identified by AER tariff ID in column B"
            else:
                detail = ("per-distributor layout: 'Tariff schedule 1' metering, 'Tariff schedule 3' prices in four basis "
                          f"sections, code column E; sheets {', '.join(wb.sheetnames)}")
            self.instance("aer_layout_change", did=d["distributor_id"], fy=d["fin_year"], doc=d["document_id"],
                          detail=detail)
        for did in sorted(self.dist):
            hosted = [d for d in self.docs if d["distributor_id"] == did and d["fin_year"] == "2023-24"
                      and d["recon_side"] == "AER_HOSTED"]
            self.instance("no_aer_file_2023_24", did=did, fy="2023-24", doc=hosted[0]["document_id"] if hosted else None,
                          detail="no AER-authored price file for 2023-24; the AER side is the distributor's own proposal "
                                 "hosted on aer.gov.au" if hosted else "no AER-authored price file for 2023-24 and no "
                                 "AER-hosted distributor document")
        for d in self.docs:
            if d["retrieval_status"] == "not_retrievable":
                self.instance("document_not_retrievable", did=d["distributor_id"], fy=d["fin_year"],
                              doc=d["document_id"], detail=f"{d['title']} ({d['access_note'] or 'no copy found'})")
            if d["author"] == "distributor" and d["price_status"] == "proposed" and d["recon_side"] == "DNSP":
                self.instance("proposed_distributor_document", did=d["distributor_id"], fy=d["fin_year"],
                              doc=d["document_id"], detail=f"{d['title']}: the distributor's own published document is a "
                              f"proposal; price_status = proposed")
            if d["price_status"] == "unverified":
                self.instance("price_status_unverified", did=d["distributor_id"], fy=d["fin_year"],
                              doc=d["document_id"], detail=f"{d['title']}: no held source states whether these prices "
                              f"were proposed, approved or final")
        # Evoenergy 2024-25 schedule: parser notes name the next page for rows printed at the foot of a page
        for c in self.t.all("charge"):
            m = re.search(r"\bpage (\d+)\b", c["note"] or "")
            if c["locator_kind"] == "pdf" and m and int(m.group(1)) != c["page"] and "Evoenergy_Schedule_of_Charges_2024-25" \
                    in self.doc_by_id[self.t.get("tariff_listing", c["listing_id"])["document_id"]]["local_path"]:
                self.instance("parser_note_page_offset", did="evoenergy", fy="2024-25", charge=c["charge_id"],
                              listing=c["listing_id"], doc=self.t.get("tariff_listing", c["listing_id"])["document_id"],
                              detail=f"note says page {m.group(1)}, value is on page {c['page']} (locator is correct)")
        # GST-inclusive values, unstated demand periods
        per_doc = defaultdict(Counter)
        for c in self.t.all("charge"):
            doc = self.t.get("tariff_listing", c["listing_id"])["document_id"]
            if c["gst"] == "incl":
                per_doc[doc]["gst_incl"] += 1
            if c["period"] == "unstated" or c["period_inferred"]:
                per_doc[doc]["period"] += 1
        for doc, cnt in sorted(per_doc.items()):
            d = self.doc_by_id[doc]
            if cnt["gst_incl"]:
                self.instance("gst_inclusive_prices", did=d["distributor_id"], fy=d["fin_year"], doc=doc,
                              quantity=cnt["gst_incl"], unit="charges", detail="document prints GST-inclusive prices; "
                              "gst = incl on those charges, the reconciliation compares excl only")
            if cnt["period"]:
                self.instance("demand_period_unstated", did=d["distributor_id"], fy=d["fin_year"], doc=doc,
                              quantity=cnt["period"], unit="charges", detail="demand/capacity unit does not state the "
                              "billing period: period = unstated, or taken from the component label (period_inferred = 1)")
        tou = defaultdict(set)
        windows = defaultdict(set)
        for w in self.t.all("tou_window"):
            windows[w["tou_schedule_id"]].add(w["period"].replace("_", ""))
        # window periods that cover every band of a charge kind (export windows only bands of their own sign)
        any_band = {("demand", "demandwindow"): "*", ("controlled_load", "controlledloadsupply"): "*",
                    ("export", "exportchargewindow"): "*charge", ("export", "exportrewardwindow"): "*credit"}
        for t in self.t.all("tariff_tou"):
            schedule = self.t.get("tou_schedule", t["tou_schedule_id"])
            periods = windows[t["tou_schedule_id"]]
            kinds = {"energy", "demand", "export"} if t["applies_to"] == "all" else {t["applies_to"]}
            for kind in kinds:
                covered = tou[(t["tariff_id"], schedule["fin_year"], "energy" if kind == "controlled_load" else kind)]
                covered.update(periods)
                covered.update(any_band[(kind, p)] for p in periods if (kind, p) in any_band)
        # an AER label printed once for several distributor tariffs is covered by its members' windows
        for rel in self.t.all("tariff_relation"):
            if rel["relation_type"] == "aer_combined_label":
                for (tid, fy, kind), periods in list(tou.items()):
                    if tid == rel["to_tariff_id"] and fy == rel["fin_year"]:
                        tou[(rel["from_tariff_id"], fy, kind)] |= periods
        gaps = defaultdict(set)
        gap_listing = {}
        # tariff-years the distributor prices itself are checked on its documents; the others (AER-only tariff-years)
        # on the AER's
        distributor_priced = {(l["tariff_id"], self.doc_by_id[l["document_id"]]["fin_year"]) for l in listings
                              if self.doc_by_id[l["document_id"]]["author"] != "AER"
                              and l["price_availability"] == "priced"}
        for l in listings:
            d = self.doc_by_id[l["document_id"]]
            k = (l["tariff_id"], d["fin_year"])
            if d["author"] == "AER" and k in distributor_priced:
                continue
            for c in self.charges_of(l["listing_id"]):
                band = (c["time_band"] or "").replace("_", "")
                if band not in {"peak", "offpeak", "shoulder", "solarsoak", "criticalpeak", "superoffpeak"}:
                    continue
                kind = "demand" if c["charge_type"] == "capacity" else c["charge_type"]
                covered = tou[(k[0], k[1], kind)]
                wildcard = "*" if kind != "export" else "*credit" if float(c["value_num"]) < 0 else "*charge"
                if band not in covered and wildcard not in covered:
                    gaps[k].add(f"{kind}:{c['time_band']}")
                    gap_listing.setdefault(k, l)
        for k, missing in sorted(gaps.items()):
            l = gap_listing[k]
            self.instance("tou_definition_missing", did=self.t.get("tariff", k[0])["distributor_id"],
                          fy=k[1], tariff=k[0], listing=l["listing_id"], doc=l["document_id"],
                          detail=f"time-of-use components ({', '.join(sorted(missing))}) with no structured window "
                                 "covering that charge kind and band in this tariff-year; check source before use")
        for w in self.t.all("tou_window"):
            if w["months"] is None:
                s = self.t.get("tou_schedule", w["tou_schedule_id"])
                self.instance("season_months_not_stated", did=s["distributor_id"], fy=s["fin_year"], doc=s["document_id"],
                              detail=f"{w['window_id']} ({w['locator']}): window {w['start_time']}-{w['end_time']} in season "
                                     f"'{w['season']}'; the document does not list that season's months")
        for s in self.t.all("tou_schedule"):
            if s["time_basis"] == "daylight_time":
                self.instance("time_stated_in_daylight_time", did=s["distributor_id"], fy=s["fin_year"],
                              doc=s["document_id"], detail=f"{s['tou_schedule_id']} ({s['locator']}): {s['quote']!r}; "
                              f"the times are kept as stated, and the document does not say which times apply while "
                              f"daylight saving is off")
        boundaries = defaultdict(list)
        for r in self.t.all("eligibility_rule"):
            if r["rule_type"] in curated.BOUNDARY_RULES:
                boundaries[(r["document_id"], r["fin_year"], r["rule_type"].rsplit("_", 1)[0],
                            r["value_unit"], r["value_num"])].append(r)
        for group in boundaries.values():
            for r in group:
                if r["operator"] not in ("ge_unstated", "le_unstated"):
                    continue
                overlaps = sorted({other["rule_id"] for other in group
                                   if other["tariff_id"] != r["tariff_id"]
                                   and other["rule_type"] != r["rule_type"]
                                   and other["operator"] not in ("gt", "lt")})
                self.instance("boundary_inclusivity_unstated", did=r["tariff_id"].split(":")[0],
                              fy=r["fin_year"], tariff=r["tariff_id"], doc=r["document_id"],
                              quantity=r["value_num"], unit=r["value_unit"],
                              detail=f"{r['rule_id']}: {r['rule_type']} {r['operator']} {r['value_num']} "
                                     f"{r['value_unit']}; endpoint inclusion is unstated; potential shared-boundary "
                                     f"overlaps: {', '.join(overlaps) or 'none'}")
        # CitiPower CMG assignment rules
        for r in self.t.all("eligibility_rule"):
            if r["tariff_id"] in ("citipower:CMG", "citipower:CMGO21"):
                self.instance("medium_business_demand_assignment", did="citipower", fy=r["fin_year"],
                              tariff=r["tariff_id"], doc=r["document_id"],
                              quantity=r["value_num"], unit=r["value_unit"],
                              detail=f"{r['rule_type']} {r['operator'] or ''} {r['value_num'] or ''} "
                                     f"{r['value_unit'] or ''} {r['value_text'] or ''}".strip())
        # incl/excl metering tables of one document (same component listed twice)
        for lid in sorted(self.listing_rows):
            groups = defaultdict(list)
            for c in self.charges_of(lid):
                groups[(c["price_basis"], c["gst"], c["component_label"], c["unit_published"], c["includes_metering"],
                        c["includes_lfit"], c["season"], c["time_band"])].append(c)
            for k, cs in groups.items():
                if len(cs) > 1:
                    l = self.t.get("tariff_listing", lid)
                    self.instance("component_repeated_in_document",
                                  did=self.t.get("tariff", l["tariff_id"])["distributor_id"],
                                  fy=self.doc_by_id[l["document_id"]]["fin_year"], listing=lid, doc=l["document_id"],
                                  charge=sorted(cs, key=lambda c: c["charge_id"])[1]["charge_id"], quantity=len(cs), unit="charges",
                                  detail=f"{k[2]!r} ({k[0]}) printed {len(cs)} times: "
                                         + ", ".join(f"{c['value_published']} at {c['locator']}" for c in cs))

    def exceptions(self):
        from exceptions import EXCEPTIONS
        for e in EXCEPTIONS:
            self.t.add("exception_type", e)
        codes = {e["exception_code"] for e in EXCEPTIONS}
        n = Counter()
        for i in sorted(self.instances, key=lambda x: (x["exception_code"], str(x["distributor_id"]), str(x["fin_year"]),
                                                       str(x["tariff_id"]), str(x["listing_id"]), str(x["charge_id"]),
                                                       str(x["document_id"]), x["detail"])):
            if i["exception_code"] not in codes:
                raise SystemExit(f"exception {i['exception_code']} not catalogued")
            n[i["exception_code"]] += 1
            self.t.add("exception_instance", dict(i, instance_id=f"{i['exception_code']}/{short_hash(*i.values())}"),
                       same_ok=True)
        missing = codes - set(n)
        if missing:
            raise SystemExit(f"catalogued exceptions with no instance: {sorted(missing)}")

    # ------------------------------------------------------------------ run
    def build(self):
        self.reference()
        self.tariffs_and_listings()
        self.charges()
        self.metering_prices()
        self.aer_sibling_codes()
        self.curated()
        self.printed_identities()
        for w in self.t.all("tou_window"):
            for m in (w["months"] or "").split(",") if w["months"] else []:
                self.t.add("tou_window_month", {"window_id": w["window_id"], "month": int(m)})
        self.adjustments()
        self.metering_by_aer_equality()
        self.parser_flags()
        self.derived_exceptions()
        self.exceptions()
        counts = {table: Counter(r["document_id"] for r in self.t.all(table))
                  for table in ("tariff_listing", "eligibility_rule", "tou_schedule", "metering_price")}
        prices = Counter(self.t.get("tariff_listing", c["listing_id"])["document_id"] for c in self.t.all("charge"))
        for d in self.docs:
            doc = d["document_id"]
            status = "unavailable" if d["retrieval_status"] != "retrieved" else "prices" if prices[doc] else (
                "rules_only" if counts["eligibility_rule"][doc] or counts["tou_schedule"][doc] else "metadata_only")
            self.t.add("document_ingestion", {"document_id": doc, "charge_count": prices[doc],
                        "listing_count": counts["tariff_listing"][doc], "eligibility_count": counts["eligibility_rule"][doc],
                        "tou_schedule_count": counts["tou_schedule"][doc], "metering_count": counts["metering_price"][doc],
                        "status": status})
        for table, rows in rates.compute(self.t.all).items():
            for r in rows:
                self.t.add(table, r)
        return self


def write_csv(path, cols, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(cols)
    for r in rows:
        w.writerow(["" if r[c] is None else r[c] for c in cols])
    with open(path, "w", newline="", encoding="utf-8") as f:
        f.write(buf.getvalue())


def sort_key(row, pk):
    return tuple(str(row[c]) for c in pk)


def write_all(b, out_dir=OUT_DIR):
    os.makedirs(os.path.join(out_dir, "tables"), exist_ok=True)
    for t in spec.TABLES:
        name = t["name"]
        cols = b.t.cols[name]
        rows = sorted(b.t.all(name), key=lambda r: sort_key(r, b.t.pk[name]))
        write_csv(os.path.join(out_dir, "tables", f"{name}.csv"), cols, rows)
    for dialect in ("sqlite", "postgres"):
        with open(os.path.join(out_dir, f"schema.{dialect}.sql"), "w", encoding="utf-8") as f:
            f.write(spec.ddl(dialect))
    with open(os.path.join(out_dir, "schema.json"), "w", encoding="utf-8") as f:
        json.dump(spec.json_spec(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    with open(os.path.join(out_dir, "load.postgres.sql"), "w", encoding="utf-8") as f:
        f.write(spec.postgres_load())


def interpreted_unit(r):
    """The unit the parser read a charge in: its standard unit, in dollars where the printed unit is. A column that
    prints no unit (Power and Water) was read in the unit that reproduces the parser's standard value; a zero
    reproduces either, and the parser's assumed units for such columns are in dollars."""
    cents, dollars = r["unit_std"], r["unit_std"].replace("c/", "$/", 1)
    if r["unit"]:
        return dollars if "$" in r["unit"] else cents
    for unit in (dollars, cents):
        v = to_std(r["value"], unit, r["component"])[0]
        if v is not None and r["value_std"] and math.isclose(v, float(r["value_std"]), rel_tol=1e-9):
            return unit
    return dollars


def evidence_excerpt(match, width=300):
    """The text a pattern matched in, cut to `width` characters around the match so the evidence shows it."""
    text = match.string
    if len(text) <= width:
        return text
    start = max(0, min(match.start() - (width - (match.end() - match.start())) // 2, len(text) - width))
    return ("..." if start else "") + text[start:start + width] + ("..." if start + width < len(text) else "")


def listing_note(rows):
    """The note of a listing: what its rows' notes say in common. Repeated printings and GST-inclusive copies (rows that
    restate a price printed elsewhere) are left out when the listing has other rows; of the rest, the '; '-separated
    parts every row carries, in the order of the first row."""
    core = [r for r in rows if REPEATED_PRINTING not in (r["note"] or "")] or rows
    core = [r for r in core if r.get("gst", "excl") == "excl"] or core
    parts = [[p for p in (r["note"] or "").split("; ") if p] for r in core]
    common = [p for p in parts[0] if all(p in other for other in parts[1:])]
    return "; ".join(common) or None


def read_table_text(text):
    rows = list(csv.reader(io.StringIO(text)))
    return rows[0], rows[1:]


def check_append_only(ref):
    """Every row of every source-fact table committed at `ref` must still exist, unchanged, in the working tree.
    Derived tables and columns (spec: derived) are recomputed from the facts on every build and are not compared."""
    def git(*args):
        return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True)

    if git("rev-parse", "--verify", "--quiet", f"{ref}^{{commit}}").returncode:
        raise SystemExit(f"--check-append-only: {ref!r} is not a commit in this repository")
    listed = git("ls-tree", "--name-only", f"{ref}:data/tariffdb/tables")
    if listed.returncode:
        raise SystemExit(f"--check-append-only: {ref} has no data/tariffdb/tables")
    bad = []
    for name in sorted(n[:-4] for n in listed.stdout.split() if n.endswith(".csv")):
        if name not in spec.BY_NAME:
            bad.append(f"{name}: table removed")
            continue
        table = spec.BY_NAME[name]
        if table.get("derived"):
            continue
        rp = f"data/tariffdb/tables/{name}.csv"
        old = git("show", f"{ref}:{rp}")
        if old.returncode:
            raise SystemExit(f"--check-append-only: cannot read {rp} at {ref}: {old.stderr.strip()}")
        with open(os.path.join(ROOT, rp), newline="", encoding="utf-8") as f:
            bad += append_only_violations(table, old.stdout, f.read())
    return bad


# In-place changes to committed source-fact rows that correct a transcription error of this repository's tooling (the
# document did not change, so a new document version would be wrong). data/tariffdb/transcription_fixes.csv lists each
# one with the exact old and new value and the finding it fixes; a row removed by a fix (an id that changed because
# its content was corrected) is listed with column '' and its replacement's key. The append-only check accepts these
# and nothing else. The file is written by scripts/tariffdb/fixes.py, which refuses a change no finding explains.
FIXES_PATH = os.path.join(OUT_DIR, "transcription_fixes.csv")
FIXES_COLUMNS = ["table", "key", "column", "old", "new", "finding", "why"]


def load_transcription_fixes(path=FIXES_PATH):
    """{(table, primary key tuple, column or None for a removed row): (old, new, why)}"""
    out = {}
    if not os.path.exists(path):
        return out
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            out[(r["table"], tuple(json.loads(r["key"])), r["column"] or None)] = (r["old"], r["new"],
                                                                                 f"{r['finding']}: {r['why']}")
    return out


TRANSCRIPTION_FIXES = load_transcription_fixes()


def append_only_violations(table, old_text, new_text, fixes=TRANSCRIPTION_FIXES):
    pk = [c["name"] for c in table["columns"] if c["primary_key"]]
    derived = {c["name"] for c in table["columns"] if c.get("derived")}
    oh, orows = read_table_text(old_text)
    nh, nrows = read_table_text(new_text)
    out, newmap = [], {}
    for r in nrows:
        n = dict(zip(nh, r))
        k = tuple(n[c] for c in pk)
        if k in newmap:
            out.append(f"{table['name']} {k}: duplicate key")
        newmap[k] = n
    for r in orows:
        o = dict(zip(oh, r))
        k = tuple(o[c] for c in pk)
        n = newmap.get(k)
        if n is None:
            if (table["name"], k, None) not in fixes:
                out.append(f"{table['name']} {k}: removed")
            continue
        diff = [f"{c}: {v!r} -> {n.get(c)!r}" for c, v in o.items() if c not in derived and n.get(c, "") != v
                and fixes.get((table["name"], k, c), (None, None))[:2] != (v, n.get(c, ""))]
        if diff:
            out.append(f"{table['name']} {k}: changed " + ", ".join(diff))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check-append-only", metavar="REF")
    ap.add_argument("--out", metavar="DIR", default=OUT_DIR, help="write the database here instead of data/tariffdb")
    a = ap.parse_args()
    if a.check_append_only:
        bad = check_append_only(a.check_append_only)
        for x in bad:
            print(x)
        print(f"{len(bad)} source-fact rows changed or removed since {a.check_append_only}")
        sys.exit(1 if bad else 0)
    b = Builder().build()
    write_all(b, a.out)
    for t in spec.TABLE_ORDER:
        print(f"{t:24s} {len(b.t.rows[t]):7d}")


if __name__ == "__main__":
    main()
