"""Tests of the tariff database in data/tariffdb/ (built by scripts/tariffdb/build.py).

  .venv/bin/python -m unittest tests/test_tariffdb.py

Schema, load, key, history, TOU and exception tests read the committed CSVs. The source re-read tests open every
retrieved source document and FAIL when one is missing (./run.sh or scripts/fetch_sources.py fetches them). Set
TARIFFDB_SOURCES=committed to re-read only the documents committed to the repository (AER files and Wayback copies), as CI
does:
tests then say which documents they leave out. The parser-count and rebuild tests also need the parser outputs (out/,
built by ./run.sh).
"""
import csv
import glob
import json
import os
import re
import sqlite3
import sys
import unittest
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "scripts" / "tariffdb"))
import spec  # noqa: E402
import exceptions as catalogue  # noqa: E402

DB_DIR = ROOT / "data" / "tariffdb"
TABLES = DB_DIR / "tables"
COMMITTED_ONLY = os.environ.get("TARIFFDB_SOURCES") == "committed"


def read(table):
    with open(TABLES / f"{table}.csv", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


_cache = {}


def rows(table):
    if table not in _cache:
        _cache[table] = read(table)
    return _cache[table]


def by(table, key):
    return {r[key]: r for r in rows(table)}


def source(doc_id):
    """Path of a document the source tests must re-read, or None when TARIFFDB_SOURCES=committed leaves it out.
    A required document that is missing fails the test (fetch it with scripts/fetch_sources.py)."""
    d = by("source_document", "document_id")[doc_id]
    if not d["local_path"] or (COMMITTED_ONLY and d["committed_in_repo"] != "1"):
        return None
    p = ROOT / d["local_path"]
    if not p.exists():
        raise AssertionError(f"{d['local_path']} is missing: run .venv/bin/python scripts/fetch_sources.py "
                             f"(or set TARIFFDB_SOURCES=committed)")
    return p


def required_source(test, doc_id):
    p = source(doc_id)
    if p is None:
        test.skipTest(f"TARIFFDB_SOURCES=committed: needs {by('source_document', 'document_id')[doc_id]['local_path']}")
    return p


def half_unit(c):
    v, s = float(c["value_num"]), float(c["value_std"])
    factor = s / v if v else 1.0
    return 0.5 * 10 ** (-max(0, -Decimal(c["value_published"]).as_tuple().exponent)) * abs(factor)


def src(doc):
    return ROOT / doc["local_path"] if doc["local_path"] else None


class TestSchemaFiles(unittest.TestCase):
    def test_generated_files_are_current(self):
        self.assertEqual((DB_DIR / "schema.sqlite.sql").read_text(), spec.ddl("sqlite"))
        self.assertEqual((DB_DIR / "schema.postgres.sql").read_text(), spec.ddl("postgres"))
        self.assertEqual((DB_DIR / "load.postgres.sql").read_text(), spec.postgres_load())
        self.assertEqual(json.loads((DB_DIR / "schema.json").read_text()), json.loads(json.dumps(spec.json_spec())))

    def test_docs_are_current(self):
        import docs
        self.assertEqual((ROOT / "docs" / "tariffdb.md").read_text(encoding="utf-8"), docs.markdown(docs.Data()),
                         "run .venv/bin/python scripts/tariffdb/docs.py")

    def test_csv_headers_match_spec(self):
        for t in spec.TABLES:
            with open(TABLES / f"{t['name']}.csv", newline="", encoding="utf-8") as f:
                self.assertEqual(next(csv.reader(f)), [c["name"] for c in t["columns"]], t["name"])

    def test_data_dictionary_describes_every_column(self):
        """schema.json, the published data dictionary, gives every table its reasons and every CSV column a type,
        nullability and description."""
        dictionary = {t["name"]: t for t in json.loads((DB_DIR / "schema.json").read_text())["tables"]}
        self.assertEqual(set(dictionary), {Path(p).stem for p in glob.glob(str(TABLES / "*.csv"))})
        for name, t in dictionary.items():
            self.assertTrue(t["description"] and t["why"], name)
            with open(DB_DIR / t["file"], newline="", encoding="utf-8") as f:
                header = next(csv.reader(f))
            cols = {c["name"]: c for c in t["columns"]}
            self.assertEqual(header, list(cols), name)
            for c in cols.values():
                self.assertTrue(c["description"] and c["type"] in spec.SQL_TYPES["sqlite"], f"{name}.{c['name']}")


    def test_rebuild_reproduces_every_table(self):
        """build.py run afresh writes every table byte for byte: the committed data, derived tables included, follows
        from the parser outputs, the curated files and the sources alone."""
        if COMMITTED_ONLY:
            self.skipTest("TARIFFDB_SOURCES=committed: the rebuild needs the parser outputs in out/ and every source")
        import subprocess
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run([sys.executable, str(ROOT / "scripts" / "tariffdb" / "build.py"), "--out", tmp], cwd=ROOT,
                           check=True, capture_output=True)
            for p in sorted(DB_DIR.rglob("*")):
                if p.is_file():
                    rel = p.relative_to(DB_DIR)
                    if rel.parts[0] == "curated":
                        continue
                    self.assertEqual((Path(tmp) / rel).read_bytes(), p.read_bytes(), str(rel))


class TestLoad(unittest.TestCase):
    """The CSVs import into SQLite with every PRIMARY KEY, FOREIGN KEY, UNIQUE, NOT NULL and CHECK enforced."""

    @classmethod
    def setUpClass(cls):
        import load
        cls.con = load.load(DB_DIR)

    def test_row_counts_survive_the_load(self):
        for t in spec.TABLES:
            n = self.con.execute(f"SELECT count(*) FROM {t['name']}").fetchone()[0]
            self.assertEqual(n, len(rows(t["name"])), t["name"])
            self.assertGreater(n, 0, t["name"])

    def test_no_orphan_keys(self):
        self.assertEqual(self.con.execute("PRAGMA foreign_key_check").fetchall(), [])

    def test_constraints_reject_bad_rows(self):
        with self.assertRaises(sqlite3.IntegrityError):
            self.con.execute("INSERT INTO tariff VALUES ('x:1', 'nowhere', '1', 'distributor_code')")
        with self.assertRaises(sqlite3.IntegrityError):
            self.con.execute("INSERT INTO financial_year VALUES ('2030-31', '2031-06-30', '2030-07-01')")
        with self.assertRaises(sqlite3.IntegrityError):
            self.con.execute("INSERT INTO tariff VALUES ('ausgrid:Z', 'ausgrid', 'Z', 'guess')")
        with self.assertRaises(sqlite3.IntegrityError):  # boolean stored as text
            self.con.execute("UPDATE tou_schedule SET covers_full_day = 'True' WHERE rowid = 1")
        with self.assertRaises(sqlite3.IntegrityError):  # number that is not a number
            self.con.execute("UPDATE charge SET value_num = '1,234' WHERE rowid = 1")

    def test_empty_field_means_null(self):
        """CSV cannot tell '' from NULL, so no column stores an empty string: every '' is NULL and allowed."""
        for t in spec.TABLES:
            nullable = {c["name"] for c in t["columns"] if c["nullable"]}
            for r in rows(t["name"]):
                for c, v in r.items():
                    if v == "":
                        self.assertIn(c, nullable, f"{t['name']} {c}")

    @unittest.skipUnless(os.environ.get("TARIFFDB_PG_BIN"), "set TARIFFDB_PG_BIN to a PostgreSQL bin directory "
                         "(initdb, pg_ctl, psql), e.g. from `pip install pgserver`")
    def test_postgres_load(self):
        """The CSVs load into a throwaway PostgreSQL through load.postgres.sql with every constraint enforced."""
        import subprocess
        import tempfile
        pg = Path(os.environ["TARIFFDB_PG_BIN"])
        with tempfile.TemporaryDirectory() as tmp:
            data, sock = Path(tmp) / "data", Path(tmp) / "sock"
            sock.mkdir()
            run = lambda *a, **kw: subprocess.run([str(x) for x in a], check=True, capture_output=True, text=True, **kw)
            run(pg / "initdb", "-D", data, "-U", "postgres", "-A", "trust")
            run(pg / "pg_ctl", "-D", data, "-o", f"-k {sock} -c listen_addresses=''", "-l", Path(tmp) / "log", "start")
            try:
                psql = [pg / "psql", "-h", sock, "-U", "postgres", "-d", "postgres", "-v", "ON_ERROR_STOP=1", "-q"]
                run(*psql, "-f", "load.postgres.sql", cwd=DB_DIR)
                for t in spec.TABLES:
                    n = run(*psql, "-At", "-c", f"SELECT count(*) FROM {t['name']}").stdout.strip()
                    self.assertEqual(int(n), len(rows(t["name"])), t["name"])
            finally:
                run(pg / "pg_ctl", "-D", data, "stop", "-m", "fast")

    def test_load_script_saves_a_database(self):
        """scripts/tariffdb/load.py --out writes a SQLite file holding every row, with no orphan keys."""
        import subprocess
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "tariffs.sqlite"
            run = subprocess.run([sys.executable, str(ROOT / "scripts" / "tariffdb" / "load.py"), "--out", str(out)],
                                 cwd=ROOT, capture_output=True, text=True)
            self.assertEqual(run.returncode, 0, run.stderr)
            con = sqlite3.connect(out)
            try:
                for t in spec.TABLES:
                    self.assertEqual(con.execute(f"SELECT count(*) FROM {t['name']}").fetchone()[0], len(rows(t["name"])))
                con.execute("PRAGMA foreign_keys = ON")
                self.assertEqual(con.execute("PRAGMA foreign_key_check").fetchall(), [])
            finally:
                con.close()

    def test_load_rejects_bad_rows(self):
        """The loader refuses a CSV row that breaks a key or a CHECK, naming the file and line."""
        import shutil
        import tempfile
        import load
        with tempfile.TemporaryDirectory() as tmp:
            shutil.copytree(DB_DIR, Path(tmp) / "db", ignore=shutil.ignore_patterns("curated"))
            with open(Path(tmp) / "db" / "tables" / "tariff.csv", "a", encoding="utf-8") as f:
                f.write("nowhere:X,nowhere,X,distributor_code\n")
            with self.assertRaisesRegex(ValueError, r"tables/tariff\.csv:\d+"):
                load.load(Path(tmp) / "db")


class TestHistory(unittest.TestCase):
    def test_every_value_is_tied_to_a_year_dates_and_a_document_version(self):
        docs = by("source_document", "document_id")
        fy = by("financial_year", "fin_year")
        for l in rows("tariff_listing"):
            d = docs[l["document_id"]]
            self.assertEqual(d["retrieval_status"], "retrieved")
            self.assertTrue(d["sha256"] and d["version_label"])
            self.assertGreaterEqual(l["effective_from"], fy[d["fin_year"]]["start_date"])
            self.assertLessEqual(l["effective_to"], fy[d["fin_year"]]["end_date"])
        listings = by("tariff_listing", "listing_id")
        for c in rows("charge"):
            l = listings[c["listing_id"]]
            self.assertEqual((c["effective_from"], c["effective_to"]), (l["effective_from"], l["effective_to"]))

    def test_no_duplicate_effective_ranges(self):
        """Within one document version a component has one value per (basis, GST, metering, LFiT) for a date range;
        every exception is a catalogued 'component_repeated_in_document' instance."""
        listings = by("tariff_listing", "listing_id")
        groups = defaultdict(list)
        for c in rows("charge"):
            groups[(c["listing_id"], c["price_basis"], c["gst"], c["component_label"], c["unit_published"],
                    c["includes_metering"], c["includes_lfit"], c["season"], c["time_band"])].append(c)
        catalogued = {(i["listing_id"], i["charge_id"]) for i in rows("exception_instance")
                      if i["exception_code"] == "component_repeated_in_document"}
        for k, cs in groups.items():
            if len(cs) < 2:
                continue
            spans = sorted((c["effective_from"], c["effective_to"]) for c in cs)
            if any(b[0] <= a[1] for a, b in zip(spans, spans[1:])):
                self.assertIn((k[0], cs[1]["charge_id"]), catalogued, f"{k} in {listings[k[0]]['document_id']}")
        # one listing per (document, tariff, name, class, region)
        seen = Counter((l["document_id"], l["tariff_id"], l["name_published"], l["class_published"], l["region"],
                        l["code_published"]) for l in rows("tariff_listing"))
        aer_ids = Counter()
        for a in rows("tariff_alias"):
            if a["alias_kind"] == "aer_tariff_id":
                aer_ids[(a["document_id"], a["tariff_id"])] += 1
        for k, n in seen.items():
            if n > 1:  # only where the AER prints one tariff on several rows with different AER tariff IDs
                self.assertGreaterEqual(aer_ids[(k[0], k[1])], n, k)
        # one rule per tariff, schedule and charge kind; one full-day partition per tariff, year and charge kind
        full = Counter()
        sched = by("tou_schedule", "tou_schedule_id")
        for t in rows("tariff_tou"):
            s = sched[t["tou_schedule_id"]]
            if s["covers_full_day"] == "1":
                full[(t["tariff_id"], s["document_id"], t["applies_to"], s["name"])] += 1
        self.assertFalse([k for k, n in full.items() if n > 1])

    def test_versions_are_kept_side_by_side(self):
        series = defaultdict(list)
        for d in rows("source_document"):
            series[d["series_id"]].append(int(d["version_seq"]))
        for s, seqs in series.items():
            self.assertEqual(len(seqs), len(set(seqs)), s)
        v = sorted(series["aer-all-2025-26-consolidated"])
        self.assertEqual(v, [1, 2, 3, 4, 5])

    def test_append_only_check(self):
        import build
        t = spec.BY_NAME["tariff"]
        old = "tariff_id,distributor_id,tariff_code,identity_basis\na:1,a,1,distributor_code\na:2,a,2,distributor_code\n"
        self.assertEqual(build.append_only_violations(t, old, old + "a:3,a,3,distributor_code\n"), [])
        changed = old.replace("a:2,a,2,distributor_code", "a:2,a,2,aer_label")
        self.assertEqual(len(build.append_only_violations(t, old, changed)), 1)
        removed = "tariff_id,distributor_id,tariff_code,identity_basis\na:1,a,1,distributor_code\n"
        self.assertIn("removed", build.append_only_violations(t, old, removed)[0])
        dup = old + "a:2,a,2,distributor_code\n"
        self.assertIn("duplicate key", build.append_only_violations(t, old, dup)[0])
        # a derived column (recomputed from the facts) may change as data grows; a fact column may not
        c = spec.BY_NAME["charge"]
        cols = [x["name"] for x in c["columns"]]
        row = {k: "x" for k in cols}

        def text(**kw):
            return ",".join(cols) + "\n" + ",".join({**row, **kw}[k] for k in cols) + "\n"
        self.assertEqual(build.append_only_violations(c, text(includes_metering="unknown"),
                                                      text(includes_metering="yes")), [])
        self.assertEqual(len(build.append_only_violations(c, text(value_published="1"), text(value_published="2"))), 1)
        # a listed transcription fix is accepted for exactly its old and new value
        key = ("x",) * len([x for x in c["columns"] if x["primary_key"]])
        fixes = {("charge", key, "value_published"): ("1", "2", "why")}
        self.assertEqual(build.append_only_violations(c, text(value_published="1"), text(value_published="2"), fixes), [])
        self.assertEqual(len(build.append_only_violations(c, text(value_published="1"), text(value_published="3"),
                                                          fixes)), 1)
        for (table, k, col), (before, after, why) in build.TRANSCRIPTION_FIXES.items():
            pk = [x["name"] for x in spec.BY_NAME[table]["columns"] if x["primary_key"]]
            (r,) = [r for r in rows(table) if tuple(r[n] for n in pk) == k]
            self.assertEqual(r[col], after, why)

    def test_append_only_check_against_git(self):
        """Run against this commit: an unknown ref fails loudly; derived tables are skipped, fact tables compared."""
        import build
        with self.assertRaises(SystemExit):
            build.check_append_only("no-such-ref-for-tariffdb")
        self.assertTrue(spec.DERIVED_TABLES)
        self.assertNotIn("charge", spec.DERIVED_TABLES)

    def test_every_aer_version_has_a_versioned_url(self):
        """Every AER consolidated version, held or not, has its own publisher URL with a dated server answer; the
        held ones are committed, the others answer 307 (exists, login-gated)."""
        docs = by("source_document", "document_id")
        checks = defaultdict(list)
        for c in rows("document_url_check"):
            checks[c["document_id"]].append(c)
            self.assertEqual(c["outcome"], {"200": "served", "307": "login_gated", "404": "not_found"}[c["http_status"]])
            self.assertTrue(c["url"].startswith("https://www.aer.gov.au/system/files/"), c)
        versions = [d for d in docs.values() if d["series_id"].startswith("aer-all-") and d["series_id"].endswith("-consolidated")]
        self.assertEqual(Counter(d["fin_year"] for d in versions), Counter({"2025-26": 5, "2026-27": 5}))
        for d in versions + [docs["aer-stakeholder-sapn-2024-25-original"]]:
            self.assertTrue(checks[d["document_id"]], d["document_id"])
            if d["retrieval_status"] == "retrieved":
                self.assertEqual(d["committed_in_repo"], "1", d["document_id"])
            else:
                self.assertEqual({c["outcome"] for c in checks[d["document_id"]]}, {"login_gated"}, d["document_id"])
        latest = {d["document_id"] for d in versions if d["version_seq"] == "5"}
        self.assertEqual({c["document_id"] for c in rows("document_url_check") if c["outcome"] == "served"}, latest)
        self.assertTrue(all(d["committed_in_repo"] == "1" for d in docs.values()
                            if d["author"] == "AER" and d["retrieval_status"] == "retrieved"))

    def test_ids_come_from_content_not_row_order(self):
        """Shuffling the rows of a document gives every row the same id."""
        import random
        import build
        items = [("doc/A", ["peak", "c/kWh"]), ("doc/A", ["off peak", "c/kWh"]), ("doc/B", ["x", "y"]),
                 ("doc/A", ["shoulder", "c/kWh"]), ("doc/C", ["same", "1"]), ("doc/C", ["same", "1"])]
        first = build.unique_ids(items)
        ids = dict(zip(map(repr, items), first))
        self.assertEqual(ids[repr(items[0])], "doc/A/peak")
        self.assertEqual(ids[repr(items[2])], "doc/B")
        for seed in range(5):
            shuffled = items[:]
            random.Random(seed).shuffle(shuffled)
            self.assertEqual(sorted(build.unique_ids(shuffled)), sorted(first))
            for item, i in zip(shuffled, build.unique_ids(shuffled)):
                if item[0] != "doc/C":  # identical rows are interchangeable
                    self.assertEqual(i, ids[repr(item)])
        self.assertEqual(len({c["charge_id"] for c in rows("charge")}), len(rows("charge")))


class TestTOU(unittest.TestCase):
    @staticmethod
    def minutes(t):
        h, m = t.split(":")
        return int(h) * 60 + int(m)

    def test_full_day_schedules_tile_24_hours_without_overlap(self):
        import curated
        windows = defaultdict(list)
        for w in rows("tou_window"):
            windows[w["tou_schedule_id"]].append({"days": w["day_type"], "start": w["start_time"],
                                                  "end": w["end_time"], "months": w["months"], "period": w["period"]})
        full = [s for s in rows("tou_schedule") if s["covers_full_day"] == "1"]
        self.assertTrue(full)
        for s in full:
            self.assertEqual(curated.coverage_errors(windows[s["tou_schedule_id"]]), [], s["tou_schedule_id"])

    def test_windows_of_one_period_never_overlap(self):
        import curated
        spans = defaultdict(list)
        for w in rows("tou_window"):
            for d in curated.DAY_TYPE_DAYS[w["day_type"]]:
                for m in w["months"].split(","):
                    spans[(w["tou_schedule_id"], w["period"], d, m)].append(
                        (self.minutes(w["start_time"]), self.minutes(w["end_time"])))
        for k, ss in spans.items():
            ss.sort()
            for a, b in zip(ss, ss[1:]):
                self.assertLessEqual(a[1], b[0], k)

    def test_demand_rule_windows_exist(self):
        periods = {(w["tou_schedule_id"], w["period"]) for w in rows("tou_window")}
        for r in rows("demand_rule"):
            if r["window_tou_schedule_id"] and r["window_period"]:
                self.assertIn((r["window_tou_schedule_id"], r["window_period"]), periods, r["demand_rule_id"])


class TestRowCounts(unittest.TestCase):
    def test_every_distributor_and_year_has_both_sides(self):
        docs = by("source_document", "document_id")
        listings = by("tariff_listing", "listing_id")
        attachment_only = {(i["distributor_id"], i["fin_year"]) for i in instances("price_attachment_not_held")}
        n = Counter()
        for c in rows("charge"):
            d = docs[listings[c["listing_id"]]["document_id"]]
            did = d["distributor_id"] or listings[c["listing_id"]]["tariff_id"].split(":")[0]
            n[(did, d["fin_year"], d["author"])] += 1
        for did in by("distributor", "distributor_id"):
            for fy in ("2024-25", "2025-26", "2026-27"):
                self.assertGreater(n[(did, fy, "AER")], 0, (did, fy))
            for fy in ("2023-24", "2024-25", "2025-26", "2026-27"):
                if (did, fy) in attachment_only:
                    continue
                self.assertGreater(n[(did, fy, "distributor")], 0, (did, fy))
        # distributor-years whose only distributor document puts its prices in an attachment that is not held carry
        # no distributor prices, but their rules (eligibility, TOU windows) are still recorded
        self.assertEqual(attachment_only, {("citipower", "2025-26"), ("unitedenergy", "2025-26")})
        for did, fy in attachment_only:
            self.assertEqual(n[(did, fy, "distributor")], 0, (did, fy))

    @unittest.skipUnless((ROOT / "out" / "aer_long.csv").exists(), "parser outputs not built (./run.sh)")
    def test_charge_counts_equal_parser_rows(self):
        """Per distributor, year and document: charges = priced parser rows (+ one copy per extra joint-label
        member), and listings cover every parser row."""
        parser = Counter()
        for p in ["out/aer_long.csv", "out/aer_versions_long.csv"] + sorted(
                os.path.relpath(x, ROOT) for x in glob.glob(str(ROOT / "out" / "dnsp" / "*.csv"))):
            with open(ROOT / p, newline="", encoding="utf-8") as f:
                for r in csv.DictReader(f):
                    if not r["component"].startswith("(no non-zero"):
                        parser[(r["distributor"], r["fin_year"], r["source_file"])] += 1
        docs = by("source_document", "document_id")
        names = {d["distributor_id"]: d["name"] for d in rows("distributor")}
        listings = by("tariff_listing", "listing_id")
        joint = {i["listing_id"] for i in rows("exception_instance") if i["exception_code"] == "joint_code_label"}
        firsts = set()
        db = Counter()
        for c in rows("charge"):
            l = listings[c["listing_id"]]
            if c["listing_id"] in joint:
                base = c["listing_id"].rsplit("/", 1)[0]
                k = (base, c["price_basis"], c["gst"], c["locator"], c["component_label"])
                if k in firsts:
                    continue
                firsts.add(k)
            d = docs[l["document_id"]]
            db[(names[l["tariff_id"].split(":")[0]], d["fin_year"], d["local_path"])] += 1
        self.assertEqual(db, parser)


class TestSourceValues(unittest.TestCase):
    """Every number and every quote is re-read from its source document at its locator. A missing document fails;
    TARIFFDB_SOURCES=committed narrows the set to the committed documents, and each test still checks a minimum."""

    @classmethod
    def setUpClass(cls):
        import locators
        cls.loc = locators
        cls.docs = by("source_document", "document_id")
        cls.listings = by("tariff_listing", "listing_id")

    def checked(self, n, what):
        """At least the expected number of rows were re-read (all of them unless TARIFFDB_SOURCES=committed)."""
        print(f"\n  re-read {n} {what}", file=sys.stderr)
        self.assertGreater(n, 0, what)

    def test_document_checksums(self):
        import hashlib
        n = 0
        for d in rows("source_document"):
            p = source(d["document_id"])
            if p:
                self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(), d["sha256"], d["document_id"])
                n += 1
        self.checked(n, "document checksums")
        if not COMMITTED_ONLY:
            self.assertEqual(n, sum(d["retrieval_status"] == "retrieved" for d in rows("source_document")))

    def test_every_charge_value_is_in_its_source(self):
        bad, n = [], 0
        # rows of the OCR'd Evoenergy tables: NUoS = DUoS + TUoS + JS is the parser's check, so it is re-checked here
        ocr_rows = defaultdict(dict)
        for c in rows("charge"):
            if c["locator_kind"] == "pdf-ocr":
                ocr_rows[(c["listing_id"], c["component_label"], c["gst"], c["locator"])][c["price_basis"]] = c
        for c in rows("charge"):
            p = source(self.listings[c["listing_id"]]["document_id"])
            if p is None:
                continue
            path = str(p)
            n += 1
            if c["locator_kind"] == "xlsx":
                raw, shown = self.loc.read_cell_excel(path, c["sheet"], c["cell"])
                if str(shown) != c["value_published"] or (repr(raw) if isinstance(raw, float) else str(raw)) != c["value_raw"]:
                    bad.append((c["charge_id"], c["value_published"], shown))
                continue
            ok, why = self.loc.verify(path, c["locator"], c["value_published"])
            if not ok and c["verification"] == "ocr_sum_check":
                ok = self.ocr_repair_holds(path, c, ocr_rows[(c["listing_id"], c["component_label"], c["gst"],
                                                                c["locator"])])
            if not ok:
                bad.append((c["charge_id"], why))
        self.assertEqual(bad, [])
        self.checked(n, "charge values")
        if COMMITTED_ONLY:
            self.assertGreater(n, 5000)
        else:
            self.assertEqual(n, len(rows("charge")))

    def ocr_repair_holds(self, path, c, row):
        """The scan prints some decimal points too faintly to OCR ('2 311' for 2.311; Evoenergy 2023-24 proposal
        p40-45). Accept the value when its digits appear with the point dropped, or a 400 dpi pass reads it, AND the
        row still satisfies NUoS = DUoS + TUoS + JS."""
        if set(row) != {"NUoS", "DUoS", "TUoS", "JSA"}:
            return False
        parts = sum(Decimal(row[b]["value_published"]) for b in ("DUoS", "TUoS", "JSA"))
        if parts != Decimal(row["NUoS"]["value_published"]):
            return False
        ip, _, fp = c["value_published"].lstrip("-").partition(".")
        text = self.loc.ocr_text(path, int(c["page"]))
        if any(re.search(r"(?<![\d.])" + re.escape(f) + r"(?!\d)", text) for f in (f"{ip} {fp}", f"{ip}{fp}")):
            return True
        tokens = self.loc.ocr_text(path, int(c["page"]), 400).split()
        return any(f in tokens for f in self.loc.renderings(c["value_published"]))

    def test_value_num_and_std_follow_the_published_value(self):
        from units import to_std
        import build
        for c in rows("charge"):
            self.assertEqual(Decimal(c["value_num"]), Decimal(c["value_published"]), c["charge_id"])
            if c["value_std"]:
                # unit_interpreted is the published unit as the parser reads it (e.g. $/kVA/M -> $/kVA/month)
                vs, us = to_std(c["value_published"], c["unit_interpreted"] or c["unit_published"] or "",
                                c["component_label"])
                self.assertAlmostEqual(float(c["value_std"]), vs, places=9, msg=c["charge_id"])
                self.assertEqual(us, c["unit_std"], c["charge_id"])
                # canonical text: no binary-float noise such as 205.79000000000002
                self.assertEqual(c["value_std"], build.num_text(c["value_std"]), c["charge_id"])

    def test_metering_prices_are_in_their_source(self):
        n = 0
        for m in rows("metering_price"):
            p = source(m["document_id"])
            if p is None:
                continue
            raw, shown = self.loc.read_cell_excel(str(p), m["sheet"], m["cell"])
            self.assertEqual(repr(raw) if isinstance(raw, float) else str(raw), m["value_raw"], m["metering_price_id"])
            n += 1
        self.checked(n, "metering prices")

    def test_every_quote_is_in_its_source(self):
        bad, n = [], 0
        for table, doc_col in [("tou_schedule", "document_id"), ("tariff_tou", "document_id"),
                               ("demand_rule", "document_id"), ("eligibility_rule", "document_id"),
                               ("tariff_relation", "document_id"), ("price_adjustment", "evidence_document_id"),
                               ("document_coverage", "evidence_document_id"), ("charge_step", "document_id")]:
            for r in rows(table):
                p = source(r[doc_col]) if r["quote"] else None
                if p:
                    n += 1
                    ok, why = self.loc.verify_quote(str(p), r["locator"], r["quote"])
                    if not ok:
                        bad.append((table, r["locator"], why, r["quote"][:60]))
        sched = by("tou_schedule", "tou_schedule_id")
        for w in rows("tou_window"):
            p = source(sched[w["tou_schedule_id"]]["document_id"])
            if p:
                n += 1
                ok, why = self.loc.verify_quote(str(p), w["locator"], w["quote"])
                if not ok:
                    bad.append(("tou_window", w["window_id"], why))
        self.assertEqual(bad, [])
        self.checked(n, "quotes")

    def test_quote_matching_respects_word_boundaries(self):
        """A quote must start and end on word boundaries and split numbers where the source does."""
        import tempfile
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            wb = Workbook()
            wb.active.title = "S"
            wb.active["A1"], wb.active["A2"], wb.active["A3"] = 14, "Closed to New Entrants", "NEE24 4"
            path = str(Path(tmp) / "t.xlsx")
            wb.save(path)
            self.loc._workbook.cache_clear()
            self.loc._sheet_text.cache_clear()
            ok = lambda cell, q: self.loc.verify_quote(path, f"xlsx:S!{cell}", q)[0]  # noqa: E731
            self.assertTrue(ok("A1", "14"))
            self.assertFalse(ok("A1", "1"))
            self.assertFalse(ok("A1", "4"))
            self.assertTrue(ok("A2", "closed to new  entrants"))
            self.assertFalse(ok("A2", "Closed to New Entr"))
            self.assertFalse(ok("A3", "NEE2 44"))
            self.assertTrue(ok("A3", "NEE24 4"))

    def test_curated_files_validate(self):
        """Structure, enums, times and day coverage (quotes are re-read by test_every_quote_is_in_its_source)."""
        import curated
        for name, data in curated.load_all().items():
            self.assertEqual(curated.validate(data, check_quotes=False), [], name)


def instances(code):
    return [i for i in rows("exception_instance") if i["exception_code"] == code]


def charges_of(doc_id, tariff_id):
    lids = {l["listing_id"] for l in rows("tariff_listing") if l["document_id"] == doc_id and l["tariff_id"] == tariff_id}
    return [c for c in rows("charge") if c["listing_id"] in lids]


class TestExceptions(unittest.TestCase):
    def test_catalogue_is_complete(self):
        """Every catalogued exception occurs in the data, and the test the catalogue names for it runs and passes."""
        codes = {e["exception_code"] for e in catalogue.EXCEPTIONS}
        self.assertEqual({r["exception_code"] for r in rows("exception_type")}, codes)
        self.assertEqual({i["exception_code"] for i in rows("exception_instance")}, codes)
        loader = unittest.TestLoader()
        for e in catalogue.EXCEPTIONS:
            path, cls, name = e["test"].split("::")
            self.assertEqual((path, cls), ("tests/test_tariffdb.py", type(self).__name__), e["exception_code"])
            self.assertNotEqual(name, self._testMethodName)
            result = unittest.TestResult()
            loader.loadTestsFromName(name, type(self)).run(result)
            self.assertEqual((result.testsRun, result.failures, result.errors), (1, [], []), e["exception_code"])
            if not COMMITTED_ONLY:
                self.assertEqual(result.skipped, [], e["exception_code"])

    def test_aer_version_differs(self):
        docs = by("source_document", "document_id")
        self.assertEqual(docs["aer-consolidated-2025-26-v1"]["price_status"], "proposed")
        self.assertEqual(docs["aer-consolidated-2025-26-v5"]["price_status"], "approved")
        jem = next(i for i in instances("aer_version_differs") if i["distributor_id"] == "jemena")
        self.assertIn("90 of 95", jem["detail"])
        self.assertAlmostEqual(float(jem["quantity"]), -12.36, places=2)
        cov = {(c["document_id"], c["distributor_id"]): c["price_status"] for c in rows("document_coverage")}
        self.assertEqual(cov[("aer-consolidated-2025-26-v1", "jemena")], "proposed")
        self.assertEqual(cov[("aer-consolidated-2025-26-v3", "jemena")], "approved")
        self.assertEqual(cov[("aer-consolidated-2025-26-v3", "energex")], "proposed")
        self.assertNotIn(("aer-consolidated-2025-26-v1", "energex"), cov)
        # Jemena A100 fixed charge: v1 and v5 values are both kept, keyed by their own document
        v1 = [c for c in charges_of("aer-consolidated-2025-26-v1", "jemena:A100") if c["charge_type"] == "fixed"]
        v5 = [c for c in charges_of("aer-consolidated-2025-26-v5", "jemena:A100") if c["charge_type"] == "fixed"]
        self.assertTrue(v1 and v5)
        self.assertNotEqual(v1[0]["value_published"], v5[0]["value_published"])
        flags = {(f["listing_id"], f["flag"]) for f in rows("listing_flag")}
        self.assertIn((v1[0]["listing_id"], "proposed_price"), flags)

    def test_metering_excluded_by_aer(self):
        adj = {a["adjustment_id"]: a for a in rows("price_adjustment") if a["kind"] == "metering_adder"}
        self.assertEqual({(a["distributor_id"], a["fin_year"]) for a in adj.values()},
                         {("endeavour", y) for y in ("2024-25", "2025-26", "2026-27")}
                         | {("essential", y) for y in ("2024-25", "2025-26", "2026-27")}
                         | {(d, y) for d in ("energex", "ergon") for y in ("2025-26", "2026-27")})
        mp = by("metering_price", "metering_price_id")
        per = Counter()
        for t in rows("price_adjustment_tariff"):
            if t["adjustment_id"] not in adj:
                continue
            a = adj[t["adjustment_id"]]
            m = mp[t["metering_price_id"]]
            expected = float(m["value_c_per_day"]) if m["value_c_per_day"] else float(m["value_num"]) * 100 / 365
            self.assertAlmostEqual(float(t["expected_delta_std"]), expected, places=9)
            aer = [c for c in charges_of(a["aer_document_id"], t["tariff_id"]) if c["charge_type"] == "fixed"
                   and c["price_basis"] == "NUoS" and c["unit_std"] == "c/day"]
            dn = [c for c in charges_of(a["distributor_document_id"], t["tariff_id"]) if c["charge_type"] == "fixed"
                  and c["price_basis"] == "NUoS" and c["unit_std"] == "c/day"]
            delta = float(dn[0]["value_std"]) - float(aer[0]["value_std"])
            self.assertLessEqual(abs(delta - expected), half_unit(aer[0]) + half_unit(dn[0]) + 1e-9, t)
            self.assertEqual(aer[0]["includes_metering"], "no")
            self.assertEqual(dn[0]["includes_metering"], "yes")
            per[a["distributor_id"]] += 1
        self.assertEqual(per, Counter(endeavour=33, essential=49, energex=26, ergon=75))
        # Endeavour 2025-26: 12.65455 $/yr on the AER Metering sheet = 3.4670 c/day
        self.assertEqual(adj["metering_adder/endeavour/2025-26"]["amount"], "3.467")

    def test_act_lfit(self):
        adj = {a["adjustment_id"]: a for a in rows("price_adjustment") if a["distributor_id"] == "evoenergy"}
        self.assertEqual({k: v["amount"] for k, v in adj.items()},
                         {"lfit_adder/evoenergy/2024-25": "0.258", "lfit_adder/evoenergy/2025-26": "1.593",
                          "lfit_adder/evoenergy/2026-27": "3.035", "lfit_rebate/evoenergy/2023-24": ""})
        for a in adj.values():
            if a["amount"]:
                self.assertIn(a["amount"], a["quote"])
            self.assertIn("LFiT", a["quote"])
        n = Counter()
        for t in rows("price_adjustment_tariff"):
            a = adj.get(t["adjustment_id"])
            if not a:
                continue
            aer = [c for c in charges_of(a["aer_document_id"], t["tariff_id"]) if c["charge_type"] == "energy"
                   and c["price_basis"] == "NUoS"]
            dn = [c for c in charges_of(a["distributor_document_id"], t["tariff_id"]) if c["charge_type"] == "energy"
                  and c["price_basis"] == "NUoS"]
            for x in aer:
                for y in dn:
                    if (x["time_band"], x["season"]) != (y["time_band"], y["season"]):
                        continue
                    delta = float(y["value_std"]) - float(x["value_std"])
                    if a["amount"]:
                        self.assertLessEqual(abs(delta - float(a["amount"])), half_unit(x) + half_unit(y) + 1e-9)
                        self.assertEqual(y["includes_lfit"], "yes")
                    else:
                        self.assertLess(delta, 0)
                    self.assertEqual(x["includes_lfit"], "no")
            n[a["adjustment_id"]] += 1
        self.assertTrue(all(n[k] > 0 for k in adj), n)

    def test_zero_priced_placeholder(self):
        flags = defaultdict(set)
        for f in rows("listing_flag"):
            flags[f["listing_id"]].add(f["flag"])
        with_charges = {c["listing_id"] for c in rows("charge")}
        for i in instances("zero_priced_placeholder"):
            self.assertNotIn(i["listing_id"], with_charges)
            self.assertIn("zero_priced_placeholder", flags[i["listing_id"]])
        placeholders = {l["listing_id"] for l in rows("tariff_listing") if l["price_availability"] == "placeholder"}
        self.assertEqual(placeholders, {i["listing_id"] for i in instances("zero_priced_placeholder")})
        for l in rows("tariff_listing"):
            # a listing has charges exactly when it is priced
            self.assertEqual(l["listing_id"] in with_charges, l["price_availability"] == "priced", l["listing_id"])

    def test_withdrawn_tariff_listed(self):
        ev = {(f["listing_id"], f["flag"]): f["evidence"] for f in rows("listing_flag")}
        words = {"withdrawn": "withdrawn|discontinued|ceased|ended on",
                 "closed_to_new": "clos|not available|not applicable to new|no new connections|"
                                  "available to new customers|no longer",
                 "obsolete": "obsolete",
                 "grandfathered": "grandfather"}
        listings = by("tariff_listing", "listing_id")
        for i in instances("withdrawn_tariff_listed"):
            self.assertNotEqual(listings[i["listing_id"]]["price_availability"], "rules_only")
            self.assertNotIn("[]", i["detail"])
            for flag in i["detail"].removeprefix("flags: ").split(", "):
                self.assertRegex(ev[(i["listing_id"], flag)].lower(), words[flag])

    def test_aer_missing_tariff(self):
        docs = by("source_document", "document_id")
        aer = defaultdict(set)
        for l in rows("tariff_listing"):
            d = docs[l["document_id"]]
            # a tariff named only in rules text is not a published listing on either side
            if d["author"] == "AER" and d["price_status"] != "proposed" and l["price_availability"] != "rules_only":
                aer[d["fin_year"]].add(l["tariff_id"])
        site = 0
        listings = by("tariff_listing", "listing_id")
        for i in instances("aer_missing_tariff"):
            self.assertNotIn(i["tariff_id"], aer[i["fin_year"]])
            self.assertNotEqual(listings[i["listing_id"]]["price_availability"], "rules_only")
            self.assertNotIn("[]", i["detail"])
            self.assertEqual(docs[i["document_id"]]["author"], "distributor")
            site += "site_specific" in i["detail"]
        self.assertGreater(site, 0)

    def test_aer_only_tariff(self):
        docs = by("source_document", "document_id")
        dn = defaultdict(set)
        for l in rows("tariff_listing"):
            if docs[l["document_id"]]["author"] == "distributor" and l["price_availability"] != "rules_only":
                dn[docs[l["document_id"]]["fin_year"]].add(l["tariff_id"])
        for i in instances("aer_only_tariff"):
            self.assertNotIn(i["tariff_id"], dn[i["fin_year"]])
        basis = {t["tariff_id"]: t["identity_basis"] for t in rows("tariff")}
        self.assertIn("aer_label", {basis[i["tariff_id"]] for i in instances("aer_only_tariff")})

    def test_joint_code_label(self):
        doc = "aer-consolidated-2025-26-v5"
        a = charges_of(doc, "evoenergy:010")
        b = charges_of(doc, "evoenergy:011")
        self.assertTrue(a)
        self.assertEqual(sorted((c["locator"], c["value_published"]) for c in a),
                         sorted((c["locator"], c["value_published"]) for c in b))
        aliases = {(x["alias_label"], x["tariff_id"]) for x in rows("tariff_alias") if x["document_id"] == doc
                   and x["alias_kind"] == "joint_label_member"}
        self.assertIn(("010, 011*", "evoenergy:010"), aliases)
        self.assertIn(("010, 011*", "evoenergy:011"), aliases)

    def test_code_label_quirk(self):
        alias = {(x["document_id"], x["alias_label"], x["tariff_id"]): x["alias_kind"] for x in rows("tariff_alias")}
        self.assertEqual(alias[("aer-stakeholder-report-ergon-2024-25", "EBDEM", "ergon:EBDEMT1")], "regional_suffix")
        rel = {(r["from_tariff_id"], r["to_tariff_id"], r["document_id"]): r for r in rows("tariff_relation")
               if r["relation_type"] == "aer_sibling_code"}
        self.assertEqual(rel[("sapn:LBGF", "sapn:LBGFSA", "aer-consolidated-2025-26-v5")]["quote"], "LBGFSA")
        self.assertIn(("sapn:LBGF", "sapn:LBGFCBD", "aer-consolidated-2026-27-v5"), rel)
        self.assertIn(("sapn:RSR", "sapn:RSRNE", "aer-consolidated-2026-27-v5"), rel)
        basis = {t["tariff_id"]: t["identity_basis"] for t in rows("tariff")}
        self.assertEqual(basis["sapn:LBGFSA"], "aer_label")

    def test_aer_id_changed_between_versions(self):
        got = {i["tariff_id"] for i in instances("aer_id_changed_between_versions")}
        self.assertEqual(got, {"endeavour:N61", "endeavour:N95"})
        v1 = [l for l in rows("tariff_listing") if l["document_id"] == "aer-consolidated-2025-26-v1"]
        self.assertTrue(v1 and all(l["code_published"] == "" for l in v1))
        al = {(x["alias_label"], x["tariff_id"]) for x in rows("tariff_alias")
              if x["document_id"] == "aer-consolidated-2025-26-v1" and x["alias_kind"] == "aer_tariff_id"}
        self.assertIn(("TD-END26oth-Strg", "endeavour:N95"), al)
        self.assertIn(("TD-AGD26res-Flat", "ausgrid:EA010"), al)

    def test_aer_layout_change(self):
        """The column finder follows the header labels of both consolidated layouts, and every listing of a
        consolidated workbook sits on a row whose code cell (at the found column) holds the published code."""
        import parse_aer
        from openpyxl import Workbook
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        for labels, expected in [(["Tariff class", "Code", "Top", "Fixed"], (5, 7)),          # 2025-26 v1
                                 (["Tariff code", "Name", "Class", "Top", "Fixed"], (4, 8))]:  # v5 and 2026-27
            ws = Workbook().active
            for i, v in enumerate(labels):
                ws.cell(3, 4 + i, v)
            self.assertEqual(parse_aer.consolidated_columns(ws, 3), expected)
        docs = by("source_document", "document_id")
        layout = {i["document_id"] for i in instances("aer_layout_change")}
        self.assertEqual(layout, {d["document_id"] for d in docs.values() if d["author"] == "AER"
                                  and d["retrieval_status"] == "retrieved" and d["local_path"].endswith(".xlsx")})
        charges = defaultdict(list)
        for c in rows("charge"):
            charges[c["listing_id"]].append(c)
        n, ref = 0, 0
        for doc_id in sorted(layout):
            if docs[doc_id]["document_type"] != "aer_consolidated_stakeholder_report":
                continue
            p = source(doc_id)
            if p is None:
                continue
            import locators
            ws = locators._workbook(str(p))["Tariff schedule"]
            heads = [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 2).value, str)
                     and re.search(r"\d{4}.\d{2} network prices$", ws.cell(r, 2).value.strip())]
            for l in rows("tariff_listing"):
                if l["document_id"] != doc_id or not l["locator"]:
                    continue
                row = coordinate_from_string(l["locator"].split("!")[1])[1]
                hr = max(h for h in heads if h < row)
                code_col, first_price = parse_aer.consolidated_columns(ws, hr)
                code = ws.cell(row, code_col).value
                if code == "#REF!":
                    self.assertEqual(l["code_published"], "", l["listing_id"])
                    ref += 1
                else:
                    self.assertEqual(str(code).strip(), l["code_published"], l["listing_id"])
                for c in charges[l["listing_id"]]:
                    col = column_index_from_string(coordinate_from_string(c["cell"])[0])
                    self.assertGreaterEqual(col, first_price, c["charge_id"])
                n += 1
        self.checked_at_least(n, 1000, "consolidated listings")
        if not COMMITTED_ONLY:
            self.assertGreater(ref, 0)

    def checked_at_least(self, n, minimum, what):
        if COMMITTED_ONLY:
            self.assertGreater(n, 0, what)
        else:
            self.assertGreaterEqual(n, minimum, what)

    def test_no_aer_file_2023_24(self):
        docs = rows("source_document")
        self.assertFalse([d for d in docs if d["fin_year"] == "2023-24" and d["author"] == "AER"])
        self.assertEqual(len(instances("no_aer_file_2023_24")), 14)
        hosted = {d["distributor_id"] for d in docs if d["fin_year"] == "2023-24" and d["recon_side"] == "AER_HOSTED"}
        self.assertEqual(len(hosted), 14)

    def test_document_not_retrievable(self):
        docs = by("source_document", "document_id")
        for i in instances("document_not_retrievable"):
            d = docs[i["document_id"]]
            self.assertEqual((d["local_path"], d["sha256"], d["retrieval_status"]), ("", "", "not_retrievable"))
        self.assertEqual(docs["aer-consolidated-2025-26-v3"]["publication_date"], "2025-05-14")
        self.assertEqual(docs["aer-consolidated-2026-27-v4"]["publication_date"], "2026-05-20")
        self.assertFalse([l for l in rows("tariff_listing") if docs[l["document_id"]]["retrieval_status"] != "retrieved"])

    def test_proposed_distributor_document(self):
        (i,) = instances("proposed_distributor_document")
        self.assertEqual((i["distributor_id"], i["fin_year"]), ("sapn", "2025-26"))
        lids = {l["listing_id"] for l in rows("tariff_listing") if l["document_id"] == i["document_id"]}
        flagged = {f["listing_id"] for f in rows("listing_flag") if f["flag"] == "proposed_price"}
        self.assertTrue(lids and lids <= flagged)

    def test_metering_sheet_quirk(self):
        mp = rows("metering_price")
        exit_fees = [m for m in mp if m["charge_basis"] == "per_meter"]
        self.assertTrue(exit_fees)
        self.assertTrue(all(m["tariff_codes_published"] == "" and m["value_c_per_day"] == "" for m in exit_fees))
        self.assertTrue([i for i in instances("metering_sheet_quirk") if "blank" in i["detail"]])
        end = {m["tariff_codes_published"]: m for m in mp if m["document_id"] == "aer-consolidated-2026-27-v5"
               and m["distributor_id"] == "endeavour"}
        self.assertEqual(end["N70,N71,N72,N73"]["value_raw"], "13.293665")

    def test_display_rounds_half_way(self):
        c = by("charge", "charge_id")
        ins = {i["charge_id"]: i for i in instances("display_rounds_half_way")}
        sapn = [h for h in ins if "/STR/" in h and "2025-26-v5" in h]
        self.assertTrue(sapn)
        self.assertEqual((c[sapn[0]]["value_published"], c[sapn[0]]["value_raw"]), ("0.0222", "0.02215"))
        self.assertIn("formatting the binary float gives 0.0221", ins[sapn[0]]["detail"])
        self.assertEqual(len(ins), 84)
        for h, i in ins.items():  # Excel's display is one unit above the binary float's at the last digit
            shown = re.search(r"gives (\S+)$", i["detail"]).group(1)
            step = Decimal(1).scaleb(Decimal(shown).as_tuple().exponent)
            self.assertEqual(abs(Decimal(c[h]["value_published"]) - Decimal(shown)), step, i)

    def test_parser_note_page_offset(self):
        c = by("charge", "charge_id")
        ins = instances("parser_note_page_offset")
        self.assertEqual(len(ins), 47)
        for i in ins:
            note_page = int(re.search(r"page (\d+)", c[i["charge_id"]]["note"]).group(1))
            self.assertEqual(note_page, int(c[i["charge_id"]]["page"]) + 1)

    def test_gst_inclusive_prices(self):
        lst = by("tariff_listing", "listing_id")
        incl = {lst[c["listing_id"]]["document_id"] for c in rows("charge") if c["gst"] == "incl"}
        self.assertEqual(incl, {i["document_id"] for i in instances("gst_inclusive_prices")})

    def test_demand_period_unstated(self):
        lst = by("tariff_listing", "listing_id")
        docs = {lst[c["listing_id"]]["document_id"] for c in rows("charge")
                if c["period"] == "unstated" or c["period_inferred"] == "1"}
        self.assertEqual(docs, {i["document_id"] for i in instances("demand_period_unstated")})
        for c in rows("charge"):
            if c["period"] == "unstated":
                self.assertIn(c["quantity"], ("kW", "kVA", "kW_or_kVA"))

    def test_season_months_not_stated(self):
        import locators
        months = defaultdict(set)
        for m in rows("tou_window_month"):
            months[m["window_id"]].add(m["month"])
        unstated = [w for w in rows("tou_window") if not w["months"]]
        self.assertTrue(unstated)
        self.assertEqual({w["window_id"] for w in unstated},
                         {i["detail"].split(" ", 1)[0] for i in instances("season_months_not_stated")})
        for w in unstated:
            # a named season, no assumed months, and the quote really names the season without listing months
            self.assertTrue(w["season"], w["window_id"])
            self.assertFalse(months[w["window_id"]], w["window_id"])
            s = by("tou_schedule", "tou_schedule_id")[w["tou_schedule_id"]]
            self.assertEqual(s["covers_full_day"], "0", w["window_id"])
            p = source(s["document_id"])
            if p:
                ok, why = locators.verify_quote(str(p), w["locator"], w["quote"])
                self.assertTrue(ok, (w["window_id"], why))
        for w in rows("tou_window"):
            if w["months"]:
                self.assertEqual(months[w["window_id"]], set(w["months"].split(",")), w["window_id"])

    def test_boundary_inclusivity_unstated(self):
        import curated
        uncertain = {r["rule_id"]: r for r in rows("eligibility_rule")
                     if r["operator"] in ("ge_unstated", "le_unstated")}
        ins = instances("boundary_inclusivity_unstated")
        self.assertTrue(uncertain)
        self.assertEqual(Counter(i["detail"].split(": ", 1)[0] for i in ins),
                         Counter({rid: 1 for rid in uncertain}))
        for r in rows("eligibility_rule"):
            if r["rule_type"] in curated.BOUNDARY_RULES and r["operator"] not in ("ge_unstated", "le_unstated"):
                self.assertIn(r["operator"], curated.boundary_operators(r["quote"], r["value_num"], r["value_unit"]),
                              r["rule_id"])
        for quote, value, unit, op in (
                ("Low voltage businesses consuming less LVUU | LVUU24 | BSR | than 160MWh per annum", 160, "MWh/yr", "lt"),
                ("Low voltage businesses consuming more BSRT | B2RT | LBAD | than 160MWh per annum.", 160, "MWh/yr", "gt"),
                ("consume over network \uf0b7 Energy for the first 60 5,000 kWh per annum.", 5000, "kWh/yr", "gt"),
                ("over Critical peak demand tariff open to customers industrial & 4000 MWh consuming greater 4 GWh "
                 "per year", 4, "GWh/yr", "gt"),
                ("Small NASN19 Business > 40 MWh single rate Demand tariff open to small business customers "
                 "consuming between 40 MWh and 160 MWh", 160, "MWh/yr", None)):
            ops = curated.boundary_operators(quote, value, unit)
            self.assertEqual(ops & {"gt", "lt", "ge", "le"}, {op} if op else set(), quote)
        stated = {r["tariff_id"]: r["operator"] for r in rows("eligibility_rule") if r["fin_year"] == "2023-24"
                  and r["rule_type"] in curated.BOUNDARY_RULES and r["tariff_id"] in
                  ("sapn:LVUU24", "sapn:LBAD", "evoenergy:020", "ausnet:NSP78")
                  and r["value_unit"] in ("MWh/yr", "GWh/yr", "kWh/yr")}
        self.assertEqual(stated, {"sapn:LVUU24": "lt", "sapn:LBAD": "gt", "evoenergy:020": "gt", "ausnet:NSP78": "gt"})
        self.assertFalse([i for i in ins if i["tariff_id"] in ("sapn:LBAD", "sapn:SBTOU") and i["fin_year"] == "2023-24"])
        nasn19 = [r for r in rows("eligibility_rule") if r["tariff_id"] == "ausnet:NASN19"
                  and r["rule_type"] == "consumption_min" and "Business > 40 MWh" in r["quote"]]
        self.assertTrue(nasn19)
        self.assertTrue(all(r["operator"] == "gt" for r in nasn19))
        self.assertTrue(any(i["tariff_id"] == "ausgrid:EA302" and "ausgrid:EA305|" in i["detail"] for i in ins))
        base = {"codes": ["EA302"], "doc": "sources/dnsp/ausgrid/Ausgrid_Network_Price_List_2024-25_wayback20250806.pdf",
                "fin_year": "2024-25", "locator": "pdf:p1", "value_num": 160, "value_unit": "MWh/yr"}
        for rule_type in curated.BOUNDARY_RULES:
            unit = "kVA" if rule_type.startswith("demand") else "MWh/yr"
            text_unit = unit.split("/")[0]
            unknown = "ge_unstated" if rule_type.endswith("_min") else "le_unstated"
            for quote in (f"between 60 and 160 {text_unit}", f"60-160 {text_unit}",
                          f"> 40 {text_unit}, 60-160 {text_unit}"):
                for op in ("ge", "le", "gt", "lt"):
                    r = dict(base, rule_type=rule_type, value_unit=unit, quote=quote, operator=op)
                    self.assertTrue(curated.validate({"distributor": "ausgrid", "eligibility": [r]}, check_quotes=False))
                r["operator"] = unknown
                self.assertEqual(curated.validate({"distributor": "ausgrid", "eligibility": [r]}, check_quotes=False), [])
            for op, wording in (("ge", "at least"), ("gt", "more than"), ("le", "up to"), ("lt", "less than")):
                if (op in ("ge", "gt")) != rule_type.endswith("_min"):
                    continue
                r = dict(base, rule_type=rule_type, value_unit=unit, quote=f"{wording} 160 {text_unit}", operator=op)
                self.assertEqual(curated.validate({"distributor": "ausgrid", "eligibility": [r]}, check_quotes=False), [])

    def test_tou_definition_missing(self):
        ins = instances("tou_definition_missing")
        gaps = {(i["tariff_id"], i["fin_year"]): i["detail"] for i in ins}
        self.assertEqual(len(gaps), len(ins))
        ea111 = gaps[("ausgrid:EA111", "2023-24")]
        for band in ("energy:peak", "energy:offpeak", "energy:shoulder"):
            self.assertIn(band, ea111)
        self.assertNotIn("demand:", ea111)
        for code in ("027", "028"):
            self.assertIn("demand:", gaps[(f"evoenergy:{code}", "2023-24")])
        for fy in ("2023-24", "2024-25", "2025-26", "2026-27"):
            self.assertNotIn(("jemena:A180", fy), gaps)
        for key in (("unitedenergy:URCER", "2026-27"), ("sapn:RELE2W", "2023-24"), ("sapn:RELE2W", "2024-25")):
            self.assertNotIn("export:", gaps.get(key, ""), key)
        # an export_charge_window alone does not cover export credits
        for key in (("citipower:CFS", "2026-27"), ("citipower:CRCER", "2026-27")):
            self.assertIn("export:peak", gaps[key])
        for code in ("027", "028"):
            self.assertIn("export:critical_peak", gaps[(f"evoenergy:{code}", "2023-24")])
        self.assertIn("energex", {i["distributor_id"] for i in ins})

    def test_medium_business_demand_assignment(self):
        rules = defaultdict(list)
        for r in rows("eligibility_rule"):
            if r["tariff_id"] == "citipower:CMG":
                rules[r["fin_year"]].append(r)

        def has(fy, rule_type, op, num, unit):
            return any(r["rule_type"] == rule_type and r["operator"] == op and float(r["value_num"] or "nan") == num
                       and r["value_unit"] == unit for r in rules[fy])
        self.assertTrue(has("2025-26", "consumption_min", "gt", 40, "MWh/yr"))
        self.assertTrue(has("2025-26", "demand_max", "lt", 120, "kVA"))
        self.assertTrue(has("2026-27", "consumption_min", "ge", 40, "MWh/yr"))
        self.assertTrue(has("2026-27", "consumption_max", "le", 160, "MWh/yr"))
        self.assertFalse([r for r in rules["2026-27"] if r["rule_type"] == "demand_max"])
        for fy in ("2025-26", "2026-27"):
            self.assertTrue([r for r in rules[fy] if r["rule_type"] == "meter_type" and r["value_text"] == "interval"])
            self.assertTrue([r for r in rules[fy] if r["rule_type"] == "opt_out_to"
                             and r["target_tariff_id"] == "citipower:CMGO21"])
        dr = by("demand_rule", "demand_rule_id")
        win = defaultdict(list)
        for w in rows("tou_window"):
            win[(w["tou_schedule_id"], w["period"])].append(w)
        cmg = [dr[t["demand_rule_id"]] for t in rows("tariff_demand_rule") if t["tariff_id"] == "citipower:CMG"]
        windowed = set()
        self.assertTrue(cmg)
        for r in cmg:
            if "30-minute" in r["quote"]:
                self.assertEqual((r["measure"], r["interval_minutes"], r["aggregation"]), ("kW", "30", "monthly_max"))
            else:
                # the 2023-24 proposal prints only the $/kW/month unit for CG/CMG; it never states the interval
                self.assertEqual((r["measure"], r["interval_minutes"], r["aggregation"]), ("kW", "", "not_stated"))
            ws = win[(r["window_tou_schedule_id"], r["window_period"])]
            self.assertTrue(all((w["start_time"], w["end_time"]) == ("10:00", "18:00") for w in ws))
            if ws:
                windowed.add(r["fin_year"])
        # each year whose document states the basis links the 10am-6pm maximum demand period through some rule
        self.assertLessEqual({"2024-25", "2025-26", "2026-27"}, windowed)
        self.assertEqual({r["fin_year"] for r in cmg if r["aggregation"] == "monthly_max"},
                         {"2024-25", "2025-26", "2026-27"})

    def test_price_attachment_not_held(self):
        import build
        import locators
        ins = instances("price_attachment_not_held")
        docs = by("source_document", "document_id")
        self.assertEqual({docs[i["document_id"]]["local_path"] for i in ins},
                         {p for p, _, _ in build.PRICE_ATTACHMENTS_NOT_HELD})
        ingestion = by("document_ingestion", "document_id")
        by_path = {d["local_path"]: d["document_id"] for d in docs.values()}
        for path, locator, quote in build.PRICE_ATTACHMENTS_NOT_HELD:
            p = source(by_path[path])
            if p:
                self.assertTrue(locators.verify_quote(str(p), locator, quote)[0], path)
        for i in ins:
            self.assertEqual(ingestion[i["document_id"]]["charge_count"], "0")
            self.assertEqual(ingestion[i["document_id"]]["status"], "rules_only")

    def test_quantity_blocks(self):
        steps = defaultdict(list)
        for r in rows("charge_step"):
            steps[(r["tariff_id"], r["document_id"], r["step_group"])].append(r)
        self.assertTrue(steps)
        self.assertEqual({(i["tariff_id"], i["document_id"]) for i in instances("quantity_blocks")},
                         {(r["tariff_id"], r["document_id"]) for r in rows("charge_step")})
        for k, blocks in steps.items():
            blocks.sort(key=lambda r: int(r["step_index"]))
            self.assertEqual([int(r["step_index"]) for r in blocks], list(range(1, len(blocks) + 1)), k)
            self.assertEqual(len({(r["quantity_unit"], r["reset_period"]) for r in blocks}), 1, k)
            self.assertEqual(blocks[-1]["upper_bound"], "", k)
            for a, b in zip(blocks, blocks[1:]):
                # consecutive blocks meet at one boundary that belongs to exactly one of them
                self.assertEqual(Decimal(a["upper_bound"]), Decimal(b["lower_bound"]), k)
                self.assertNotEqual(a["upper_inclusive"], b["lower_inclusive"], k)

    def test_component_repeated_in_document(self):
        c = by("charge", "charge_id")
        for i in instances("component_repeated_in_document"):
            first = c[i["charge_id"]]
            same = [x for x in rows("charge") if x["listing_id"] == first["listing_id"]
                    and x["component_label"] == first["component_label"] and x["price_basis"] == first["price_basis"]
                    and x["gst"] == first["gst"]]
            self.assertGreaterEqual(len(same), int(i["quantity"]))
            self.assertEqual(len({x["locator"] + x["charge_id"] for x in same}), len(same))

    def test_time_stated_in_daylight_time(self):
        """Times the source states in daylight time are stored as stated, flagged, and never converted."""
        sched = by("tou_schedule", "tou_schedule_id")
        daylight = {k for k, v in sched.items() if v["time_basis"] == "daylight_time"}
        self.assertEqual(daylight, {i["detail"].split(" ", 1)[0] for i in instances("time_stated_in_daylight_time")})
        self.assertEqual({sched[k]["distributor_id"] for k in daylight}, {"ausnet"})
        for w in rows("tou_window"):
            if w["tou_schedule_id"] in daylight:
                self.assertRegex(w["quote"], r"ADST|daylight", w["window_id"])
                # the stated clock time is kept: the quote shows the same start hour as the stored window
                hour = int(w["start_time"][:2]) % 12 or 12
                self.assertRegex(w["quote"].lower(), rf"(?<!\d){hour}(?!\d)", w["window_id"])

    def test_price_status_unverified(self):
        """No document asserts a regulatory status that no held source supports: every AER-hosted document is
        'unverified', and every unverified document has an instance."""
        docs = rows("source_document")
        unverified = {d["document_id"] for d in docs if d["price_status"] == "unverified"}
        self.assertEqual(unverified, {i["document_id"] for i in instances("price_status_unverified")})
        self.assertEqual(unverified, {d["document_id"] for d in docs if d["recon_side"] == "AER_HOSTED"})
        for a in rows("price_adjustment"):
            d = by("source_document", "document_id")[a["aer_document_id"]]
            if d["price_status"] == "unverified":
                self.assertNotIn("approved charges =", a["formula"])

    def test_trial_flag_needs_a_trial_tariff(self):
        """'trial' flags a tariff offered as a trial, not a rebate or note that only mentions a trial."""
        import build
        pattern = dict(build.STATUS_FLAGS)["trial"]
        for text, hit in [("Residential trial tariff", True), ("Trial: closes 30 June", True),
                          ("tariff-trial rebate applies", False), ("pre-trial review", False),
                          ("Network tariff trial rebate", False)]:
            self.assertEqual(bool(re.search(pattern, text, re.I)), hit, text)


if __name__ == "__main__":
    unittest.main()
